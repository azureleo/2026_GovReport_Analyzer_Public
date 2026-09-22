"""사람 확정 객체 목록을 고정 분모로 사용하는 시트 라우팅 평가."""

from __future__ import annotations

import csv
import json
import re
from dataclasses import asdict, dataclass
from difflib import SequenceMatcher
from pathlib import Path
from typing import Any

from openpyxl import load_workbook

import config
from utils.object_evidence_match import descriptor, match_object_evidence
from utils.object_routing import normalize_object_label, normalize_object_number


_SPLIT_RE = re.compile(r"[,|;\n]+")
_PAGE_RE = re.compile(r"\d+")
_TRUE_VALUES = {"1", "true", "y", "yes", "예", "대상", "포함"}
_FALSE_VALUES = {"0", "false", "n", "no", "아니오", "제외"}
_AUXILIARY_SHEETS = {"visual_inventory", "16_시각자료목록"}


@dataclass(frozen=True, slots=True)
class FixedRoutingItem:
    item_id: str
    page: int
    object_type: str
    number: str
    caption: str
    allowed_sheets: tuple[str, ...]
    note: str = ""
    evidence_id: str = ""
    key_values: str = ""


@dataclass(frozen=True, slots=True)
class RoutingEvaluationDetail:
    item_id: str
    page: int
    number: str
    caption: str
    allowed_sheets: tuple[str, ...]
    matched_object_ids: tuple[str, ...]
    linked_body_sheets: tuple[str, ...]
    status: str
    message: str
    match_stage: str = ""
    match_reason: str = ""

    def to_dict(self) -> dict[str, Any]:
        return asdict(self)


def _text(value: Any) -> str:
    return str(value or "").strip()


def _truth(value: Any, default: bool = True) -> bool:
    text = _text(value).casefold()
    if not text:
        return default
    if text in _TRUE_VALUES:
        return True
    if text in _FALSE_VALUES:
        return False
    return default


def _page(value: Any) -> int:
    match = _PAGE_RE.search(_text(value))
    return int(match.group(0)) if match else 0


def _sheet_aliases() -> dict[str, str]:
    aliases: dict[str, str] = {}
    for key, name in getattr(config, "SHEET_KEY_TO_NAME", {}).items():
        for value in (key, name):
            aliases[normalize_object_label(value)] = name
    return aliases


def canonical_sheet_name(value: Any) -> str:
    text = _text(value)
    if not text:
        return ""
    return _sheet_aliases().get(normalize_object_label(text), text)


def parse_sheet_set(value: Any) -> tuple[str, ...]:
    if isinstance(value, (list, tuple, set)):
        parts = [str(item) for item in value]
    else:
        parts = _SPLIT_RE.split(_text(value))
    names: list[str] = []
    for part in parts:
        name = canonical_sheet_name(part)
        if name and name not in names:
            names.append(name)
    return tuple(names)


def _row_value(row: dict[str, Any], *names: str) -> Any:
    for name in names:
        if name in row and row[name] not in (None, ""):
            return row[name]
    return None


def _rows_from_xlsx(path: Path) -> list[dict[str, Any]]:
    workbook = load_workbook(path, read_only=True, data_only=True)
    try:
        preferred = next(
            (
                name for name in ("라우팅객체", "시각요소", "원문객체")
                if name in workbook.sheetnames
            ),
            workbook.sheetnames[0],
        )
        sheet = workbook[preferred]
        headers = [_text(cell.value) for cell in sheet[1]]
        return [
            {header: value for header, value in zip(headers, values) if header}
            for values in sheet.iter_rows(min_row=2, values_only=True)
            if any(_text(value) for value in values)
        ]
    finally:
        workbook.close()


def _rows_from_path(path: Path) -> list[dict[str, Any]]:
    suffix = path.suffix.casefold()
    if suffix == ".xlsx":
        return _rows_from_xlsx(path)
    if suffix == ".jsonl":
        return [
            json.loads(line)
            for line in path.read_text(encoding="utf-8").splitlines()
            if line.strip()
        ]
    if suffix == ".json":
        payload = json.loads(path.read_text(encoding="utf-8"))
        rows = payload.get("objects") if isinstance(payload, dict) else payload
        return [row for row in (rows or []) if isinstance(row, dict)]
    if suffix == ".csv":
        with path.open("r", encoding="utf-8-sig", newline="") as stream:
            return list(csv.DictReader(stream))
    raise ValueError(f"지원하지 않는 라우팅 인벤토리 형식입니다: {path.suffix}")


def load_fixed_routing_inventory(path: str | Path) -> list[FixedRoutingItem]:
    source = Path(path).expanduser().resolve()
    items: list[FixedRoutingItem] = []
    for index, row in enumerate(_rows_from_path(source), start=1):
        expected = _truth(
            _row_value(row, "평가대상", "기대추출", "라우팅평가대상"),
            default=True,
        )
        allowed = parse_sheet_set(_row_value(row, "허용시트", "관련시트", "예상시트"))
        if not expected or not allowed:
            continue
        items.append(FixedRoutingItem(
            item_id=_text(_row_value(row, "객체ID", "요소ID", "근거ID")) or f"routing-{index}",
            page=_page(_row_value(row, "페이지", "출처페이지")),
            object_type=_text(_row_value(row, "객체유형", "요소유형", "유형")),
            number=_text(_row_value(row, "번호", "표그림번호")),
            caption=_text(_row_value(row, "캡션", "제목")),
            allowed_sheets=allowed,
            note=_text(_row_value(row, "비고", "검수메모")),
            evidence_id=_text(_row_value(row, "근거ID")),
            key_values=_text(_row_value(row, "핵심값", "대표값")),
        ))
    return items


def _output_inventory_rows(path: Path) -> list[dict[str, Any]]:
    workbook = load_workbook(path, read_only=True, data_only=True)
    try:
        if "21_원문객체인벤토리" not in workbook.sheetnames:
            return []
        sheet = workbook["21_원문객체인벤토리"]
        headers = [_text(cell.value) for cell in sheet[1]]
        return [
            {header: value for header, value in zip(headers, values) if header}
            for values in sheet.iter_rows(min_row=2, values_only=True)
            if any(_text(value) for value in values)
        ]
    finally:
        workbook.close()


def _row_ids(row: dict[str, Any]) -> set[str]:
    values = [row.get("객체ID"), row.get("근거ID")]
    values.extend(_SPLIT_RE.split(_text(row.get("중복객체ID"))))
    return {_text(value) for value in values if _text(value)}


def _body_sheets(row: dict[str, Any]) -> set[str]:
    auxiliary = {
        canonical_sheet_name(value)
        for value in (*_AUXILIARY_SHEETS, *parse_sheet_set(row.get("보조연결시트")))
    }
    return {
        name for name in parse_sheet_set(row.get("연결시트"))
        if name and name not in auxiliary
    }


def _caption_similarity(item: FixedRoutingItem, row: dict[str, Any]) -> float:
    left = normalize_object_label(item.caption)
    right = normalize_object_label(row.get("캡션"))
    if not left or not right:
        return 0.0
    if left == right:
        return 1.0
    if len(left) >= 8 and (left in right or right in left):
        return min(len(left), len(right)) / max(len(left), len(right))
    return SequenceMatcher(None, left[:500], right[:500]).ratio()


def _matches(item: FixedRoutingItem, rows: list[dict[str, Any]]):
    candidates = [
        descriptor(
            key=str(index),
            page=_page(row.get("출처페이지")),
            evidence_ids=row.get("근거ID"),
            object_ids=[row.get("객체ID"), *_SPLIT_RE.split(_text(row.get("중복객체ID")))],
            number=row.get("번호"),
            caption=row.get("캡션"),
            content=json.dumps(row, ensure_ascii=False, sort_keys=True, default=str),
        )
        for index, row in enumerate(rows)
    ]
    match = match_object_evidence(
        descriptor(
            key=item.item_id,
            page=item.page,
            evidence_ids=item.evidence_id,
            object_ids=item.item_id,
            number=item.number or item.caption,
            caption=item.caption,
            content=item.key_values or item.caption,
        ),
        candidates,
    )
    matched = [rows[int(key)] for key in match.matched_keys if key.isdigit()]
    return matched, match


def evaluate_fixed_routing_inventory(
    output_path: str | Path,
    inventory_path: str | Path,
) -> tuple[dict[str, Any], list[RoutingEvaluationDetail]]:
    items = load_fixed_routing_inventory(inventory_path)
    rows = _output_inventory_rows(Path(output_path).expanduser().resolve())
    details: list[RoutingEvaluationDetail] = []
    for item in items:
        matched, evidence_match = _matches(item, rows)
        body_sheets = sorted({sheet for row in matched for sheet in _body_sheets(row)})
        allowed = set(item.allowed_sheets)
        if evidence_match.status == "ambiguous":
            status = "복수객체_검토필요"
            message = evidence_match.reason
        elif evidence_match.status == "partial":
            status = "근거불충분"
            message = evidence_match.reason
        elif not evidence_match.confirmed:
            status = "미추출"
            message = "고정 인벤토리 객체가 출력 원문 객체 인벤토리에 없습니다."
        elif allowed.intersection(body_sheets):
            status = "일치"
            message = "허용 본문 시트 중 하나와 연결되었습니다."
        elif body_sheets:
            status = "오배치의심"
            message = "객체는 추출됐지만 허용 시트 밖의 본문 시트에 연결되었습니다."
        else:
            status = "본문미연결"
            message = "보조 시각자료 목록 외의 본문 시트 연결을 찾지 못했습니다."
        details.append(RoutingEvaluationDetail(
            item_id=item.item_id,
            page=item.page,
            number=item.number,
            caption=item.caption,
            allowed_sheets=item.allowed_sheets,
            matched_object_ids=tuple(
                sorted({_text(row.get("객체ID")) for row in matched if _text(row.get("객체ID"))})
            ),
            linked_body_sheets=tuple(body_sheets),
            status=status,
            message=message,
            match_stage=evidence_match.status,
            match_reason=evidence_match.reason,
        ))
    evaluable = len(details)
    mismatches = sum(detail.status != "일치" for detail in details)
    missing = sum(detail.status == "미추출" for detail in details)
    return {
        "routing_error_rate": round(mismatches / evaluable, 4) if evaluable else None,
        "routing_evaluable_objects": evaluable,
        "routing_mismatch_objects": mismatches,
        "routing_missing_objects": missing,
        "routing_fixed_denominator": True,
    }, details
