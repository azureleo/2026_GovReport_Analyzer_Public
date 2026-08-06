"""Deterministic fixed visual sample datasets for no-cache VLM experiments."""

from __future__ import annotations

import hashlib
import io
import json
import os
import re
from collections import Counter, defaultdict, deque
from contextlib import redirect_stdout
from copy import deepcopy
from dataclasses import asdict, dataclass
from datetime import datetime, timezone
from pathlib import Path
from typing import Any, Iterable, Sequence

import fitz
from openpyxl import load_workbook

from utils.visual_contract import (
    VISUAL_CONTRACT_VERSION,
    comparison_equal,
    normalize_visual_table_rows,
    numeric_values,
)


SCHEMA_VERSION = 1
DEFAULT_FAILURE_STATUSES = ("vision_유실",)
VALID_RESULT_STATUSES = {
    "extracted",
    "not_relevant",
    "no_data",
    "needs_review",
    "failed",
}
VALID_EXPECTED_MERGE = {
    "pending", "accept", "fix_then_merge", "duplicate", "needs_review", "reject"
}
VALID_EVIDENCE_FIXTURES = {
    "single", "missing", "multiple_ids", "multiple_objects", "object_not_extracted",
}
REVIEWED_ANNOTATION_STATUSES = {"confirmed", "확정", "reviewed", "검토완료"}


class FixedSampleError(RuntimeError):
    """Raised when a fixed sample contract is invalid or has drifted."""


@dataclass(frozen=True, slots=True)
class InventoryElement:
    sample_id: str
    page: int
    element_type: str
    title: str
    data_included: bool
    expected: bool
    related_sheet: str
    note: str


@dataclass(frozen=True, slots=True)
class FixedSampleItem:
    sample_id: str
    page: int
    element_type: str
    title: str
    data_included: bool
    expected: bool
    related_sheet: str
    note: str
    role: str
    expected_status: str
    source_status: str
    image_path: str
    image_sha256: str
    page_text_excerpt: str
    bbox_pdf: list[float] | None = None

    @classmethod
    def from_dict(cls, value: dict[str, Any]) -> "FixedSampleItem":
        allowed = set(cls.__dataclass_fields__)
        return cls(**{key: value.get(key) for key in allowed})


@dataclass(frozen=True, slots=True)
class FixedSampleDataset:
    manifest_path: Path
    payload: dict[str, Any]
    items: tuple[FixedSampleItem, ...]

    @property
    def dataset_id(self) -> str:
        return str(self.payload.get("dataset_id") or "")

    @property
    def fingerprint(self) -> str:
        return str(self.payload.get("dataset_fingerprint") or "")


@dataclass(frozen=True, slots=True)
class ValidationReport:
    valid: bool
    errors: tuple[str, ...]
    warnings: tuple[str, ...]
    dataset_id: str
    object_count: int
    source_path: Path | None

    def to_dict(self) -> dict[str, Any]:
        return {
            "valid": self.valid,
            "errors": list(self.errors),
            "warnings": list(self.warnings),
            "dataset_id": self.dataset_id,
            "object_count": self.object_count,
            "source_path": str(self.source_path) if self.source_path else None,
        }


def sha256_file(path: str | Path) -> str:
    digest = hashlib.sha256()
    with Path(path).open("rb") as stream:
        for chunk in iter(lambda: stream.read(1024 * 1024), b""):
            digest.update(chunk)
    return digest.hexdigest()


def _canonical_hash(value: Any) -> str:
    encoded = json.dumps(
        value,
        ensure_ascii=False,
        sort_keys=True,
        separators=(",", ":"),
    ).encode("utf-8")
    return hashlib.sha256(encoded).hexdigest()


def _dataset_fingerprint(
    *,
    source_sha256: str,
    inventory_sha256: str,
    audit_sha256: str,
    selection: dict[str, Any],
    render: dict[str, Any],
) -> str:
    return _canonical_hash({
        "schema_version": SCHEMA_VERSION,
        "source_sha256": source_sha256,
        "inventory_sha256": inventory_sha256,
        "audit_sha256": audit_sha256,
        "selection": selection,
        "render": render,
    })


def _text(value: Any) -> str:
    return "" if value is None else str(value).strip()


def _truthy(value: Any) -> bool:
    return _text(value).casefold() in {"y", "yes", "true", "1", "예", "참"}


def _relative(path: Path, base: Path) -> str:
    try:
        return Path(os.path.relpath(path.resolve(), base.resolve())).as_posix()
    except (OSError, ValueError):
        return str(path.resolve())


def _resolve(path_value: str, base: Path) -> Path:
    path = Path(path_value)
    return path if path.is_absolute() else (base / path).resolve()


def load_inventory(path: str | Path) -> tuple[InventoryElement, ...]:
    workbook = load_workbook(Path(path), read_only=True, data_only=True)
    try:
        if "시각요소" not in workbook.sheetnames:
            raise FixedSampleError("인벤토리에 '시각요소' 시트가 없습니다.")
        sheet = workbook["시각요소"]
        headers = [_text(cell.value) for cell in sheet[1]]
        required = {
            "요소ID",
            "페이지",
            "요소유형",
            "제목",
            "데이터포함",
            "기대추출",
            "관련시트",
            "비고",
        }
        missing = sorted(required - set(headers))
        if missing:
            raise FixedSampleError("시각 인벤토리 필수 컬럼 누락: " + ", ".join(missing))
        rows: list[InventoryElement] = []
        seen: set[str] = set()
        for row_num, values in enumerate(sheet.iter_rows(min_row=2, values_only=True), start=2):
            row = {header: value for header, value in zip(headers, values)}
            sample_id = _text(row.get("요소ID"))
            if not sample_id:
                continue
            if sample_id in seen:
                raise FixedSampleError(f"시각 인벤토리 요소ID 중복: {sample_id}")
            seen.add(sample_id)
            try:
                page = int(row.get("페이지"))
            except (TypeError, ValueError) as exc:
                raise FixedSampleError(f"시각 인벤토리 {row_num}행 페이지가 정수가 아닙니다.") from exc
            if page < 1:
                raise FixedSampleError(f"시각 인벤토리 {row_num}행 페이지는 1 이상이어야 합니다.")
            rows.append(InventoryElement(
                sample_id=sample_id,
                page=page,
                element_type=_text(row.get("요소유형")),
                title=_text(row.get("제목")),
                data_included=_truthy(row.get("데이터포함")),
                expected=_truthy(row.get("기대추출")),
                related_sheet=_text(row.get("관련시트")),
                note=_text(row.get("비고")),
            ))
        return tuple(rows)
    finally:
        workbook.close()


def load_audit_statuses(path: str | Path | None) -> dict[str, str]:
    if path is None:
        return {}
    payload = json.loads(Path(path).read_text(encoding="utf-8"))
    details = payload.get("details", []) if isinstance(payload, dict) else []
    statuses: dict[str, str] = {}
    for row in details if isinstance(details, list) else []:
        if not isinstance(row, dict):
            continue
        sample_id = _text(row.get("요소ID"))
        if sample_id:
            statuses[sample_id] = _text(row.get("상태"))
    return statuses


def _phase(page: int, low: int, high: int) -> int:
    if high <= low:
        return 1
    ratio = (page - low) / (high - low)
    return 0 if ratio < 1 / 3 else 1 if ratio < 2 / 3 else 2


def _stratified_select(
    candidates: Sequence[InventoryElement],
    limit: int,
    *,
    page_low: int,
    page_high: int,
) -> tuple[InventoryElement, ...]:
    if limit <= 0 or not candidates:
        return ()
    buckets: dict[tuple[int, str], deque[InventoryElement]] = defaultdict(deque)
    for item in sorted(candidates, key=lambda row: (row.page, row.sample_id)):
        buckets[(_phase(item.page, page_low, page_high), item.element_type)].append(item)
    keys = sorted(buckets, key=lambda key: (key[0], key[1]))
    selected: list[InventoryElement] = []
    while len(selected) < limit and keys:
        remaining: list[tuple[int, str]] = []
        for key in keys:
            bucket = buckets[key]
            if bucket and len(selected) < limit:
                selected.append(bucket.popleft())
            if bucket:
                remaining.append(key)
        keys = remaining
    return tuple(selected)


def select_fixed_sample(
    inventory: Sequence[InventoryElement],
    statuses: dict[str, str],
    *,
    failure_statuses: Sequence[str] = DEFAULT_FAILURE_STATUSES,
    positive_controls: int = 12,
    negative_controls: int = 5,
) -> tuple[tuple[InventoryElement, str, str], ...]:
    """Select failures and controls without randomness or model involvement."""
    if not inventory:
        raise FixedSampleError("시각 인벤토리가 비어 있습니다.")
    failure_set = {_text(status) for status in failure_statuses if _text(status)}
    low = min(item.page for item in inventory)
    high = max(item.page for item in inventory)
    failures = tuple(sorted(
        (
            item for item in inventory
            if item.expected and statuses.get(item.sample_id, "") in failure_set
        ),
        key=lambda item: (item.page, item.sample_id),
    ))
    failure_ids = {item.sample_id for item in failures}
    positives = [
        item for item in inventory
        if item.expected
        and item.sample_id not in failure_ids
        and (not statuses or statuses.get(item.sample_id) == "기록됨")
    ]
    negatives = [item for item in inventory if not item.expected]
    chosen_positive = _stratified_select(
        positives,
        positive_controls,
        page_low=low,
        page_high=high,
    )
    chosen_negative = _stratified_select(
        negatives,
        negative_controls,
        page_low=low,
        page_high=high,
    )
    selected: list[tuple[InventoryElement, str, str]] = []
    selected.extend((item, "failure", statuses.get(item.sample_id, "")) for item in failures)
    selected.extend((item, "positive_control", statuses.get(item.sample_id, "")) for item in chosen_positive)
    selected.extend((item, "negative_control", statuses.get(item.sample_id, "")) for item in chosen_negative)
    return tuple(selected)


def _page_excerpt(text: str, title: str, limit: int = 1800) -> str:
    compact = re.sub(r"\s+", " ", text or "").strip()
    if len(compact) <= limit:
        return compact
    needle = re.sub(r"\s+", " ", title or "").strip()
    index = compact.find(needle) if len(needle) >= 4 else -1
    if index < 0:
        return compact[:limit]
    start = max(0, index - limit // 3)
    return compact[start:start + limit]


def _write_jsonl(path: Path, rows: Iterable[dict[str, Any]]) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    with path.open("w", encoding="utf-8", newline="\n") as stream:
        for row in rows:
            stream.write(json.dumps(row, ensure_ascii=False, sort_keys=True) + "\n")


def build_fixed_sample_dataset(
    source_pdf: str | Path,
    inventory_path: str | Path,
    output_dir: str | Path,
    *,
    dataset_id: str,
    audit_path: str | Path | None = None,
    failure_statuses: Sequence[str] = DEFAULT_FAILURE_STATUSES,
    positive_controls: int = 12,
    negative_controls: int = 5,
    dpi: int = 144,
    materialize_images: bool = True,
    force: bool = False,
) -> FixedSampleDataset:
    source = Path(source_pdf).resolve()
    inventory_file = Path(inventory_path).resolve()
    audit_file = Path(audit_path).resolve() if audit_path else None
    output = Path(output_dir).resolve()
    manifest_path = output / "manifest.json"
    if manifest_path.exists() and not force:
        raise FixedSampleError(
            f"이미 고정 표본이 존재합니다: {manifest_path}. 새 ID를 쓰거나 --force를 지정하세요."
        )
    if not source.exists() or not inventory_file.exists():
        missing = source if not source.exists() else inventory_file
        raise FileNotFoundError(missing)
    if audit_file is not None and not audit_file.exists():
        raise FileNotFoundError(audit_file)
    if dpi < 72 or dpi > 600:
        raise FixedSampleError("렌더 DPI는 72~600 범위여야 합니다.")

    inventory = load_inventory(inventory_file)
    statuses = load_audit_statuses(audit_file)
    selected = select_fixed_sample(
        inventory,
        statuses,
        failure_statuses=failure_statuses,
        positive_controls=positive_controls,
        negative_controls=negative_controls,
    )
    if not selected:
        raise FixedSampleError("선택 조건을 만족하는 고정 표본이 없습니다.")

    output.mkdir(parents=True, exist_ok=True)
    images_dir = output / "images"
    if materialize_images:
        images_dir.mkdir(parents=True, exist_ok=True)

    items: list[FixedSampleItem] = []
    document = fitz.open(str(source))
    source_page_count = len(document)
    try:
        unique_pages = sorted({item.page for item, _, _ in selected})
        for page_number in unique_pages:
            if page_number > len(document):
                raise FixedSampleError(
                    f"표본 페이지 p{page_number}가 원본 {len(document)}페이지 범위를 벗어납니다."
                )
            if materialize_images:
                page = document[page_number - 1]
                pixmap = page.get_pixmap(matrix=fitz.Matrix(dpi / 72, dpi / 72), alpha=False)
                pixmap.save(str(images_dir / f"page_{page_number:04d}.png"))

        for inventory_item, role, source_status in selected:
            page = document[inventory_item.page - 1]
            image_file = images_dir / f"page_{inventory_item.page:04d}.png"
            item = FixedSampleItem(
                sample_id=inventory_item.sample_id,
                page=inventory_item.page,
                element_type=inventory_item.element_type,
                title=inventory_item.title,
                data_included=inventory_item.data_included,
                expected=inventory_item.expected,
                related_sheet=inventory_item.related_sheet,
                note=inventory_item.note,
                role=role,
                expected_status="extracted" if inventory_item.expected else "not_relevant",
                source_status=source_status,
                image_path=_relative(image_file, output) if materialize_images else "",
                image_sha256=sha256_file(image_file) if materialize_images else "",
                page_text_excerpt=_page_excerpt(
                    page.get_text("text", sort=True),
                    inventory_item.title,
                ),
                bbox_pdf=None,
            )
            items.append(item)
    finally:
        document.close()

    object_rows = [asdict(item) for item in items]
    objects_path = output / "objects.jsonl"
    _write_jsonl(objects_path, object_rows)
    inventory_snapshot = output / "inventory_snapshot.xlsx"
    inventory_snapshot.write_bytes(inventory_file.read_bytes())
    audit_snapshot: Path | None = None
    if audit_file is not None:
        audit_snapshot = output / "audit_snapshot.json"
        audit_snapshot.write_bytes(audit_file.read_bytes())
    annotation_path = output / "annotations_template.jsonl"
    _write_jsonl(annotation_path, (
        {
            "sample_id": item.sample_id,
            "review_status": "pending",
            "expected_status": item.expected_status,
            "expected_title": item.title,
            "expected_unit": None,
            "expected_sheet": item.related_sheet,
            "expected_merge": "reject" if not item.expected else "pending",
            "expected_visual_rows": [],
            "expected_merged_rows": [],
            # 이전 표본 파일과의 하위 호환용. 새 주석은 위 두 필드를 사용한다.
            "expected_rows": [],
            "fixture_evidence": "single",
            "fixture_baseline_rows": [],
            "reviewer": "",
            "review_basis": "",
            "reviewed_at": None,
            "notes": "",
        }
        for item in items
    ))

    selection_contract = {
        "failure_statuses": list(failure_statuses),
        "positive_controls": int(positive_controls),
        "negative_controls": int(negative_controls),
        "selected": [
            {
                "sample_id": item.sample_id,
                "page": item.page,
                "role": item.role,
                "expected_status": item.expected_status,
            }
            for item in items
        ],
    }
    source_hash = sha256_file(source)
    inventory_hash = sha256_file(inventory_file)
    audit_hash = sha256_file(audit_file) if audit_file else ""
    render_contract = {"dpi": dpi, "format": "png", "mode": "full_page"}
    fingerprint = _dataset_fingerprint(
        source_sha256=source_hash,
        inventory_sha256=inventory_hash,
        audit_sha256=audit_hash,
        selection=selection_contract,
        render=render_contract,
    )
    role_counts = Counter(item.role for item in items)
    expected_counts = Counter(item.expected_status for item in items)
    payload: dict[str, Any] = {
        "schema_version": SCHEMA_VERSION,
        "dataset_id": dataset_id,
        "role": "development",
        "review_status": "draft",
        "locked": True,
        "generated_at": datetime.now(timezone.utc).isoformat(),
        "dataset_fingerprint": fingerprint,
        "source": {
            "path": _relative(source, output),
            "sha256": source_hash,
            "page_count": source_page_count,
        },
        "inputs": {
            "inventory": {
                "path": "inventory_snapshot.xlsx",
                "original_path": _relative(inventory_file, output),
                "sha256": inventory_hash,
            },
            "audit": {
                "path": _relative(audit_snapshot, output) if audit_snapshot else None,
                "original_path": _relative(audit_file, output) if audit_file else None,
                "sha256": audit_hash,
            },
        },
        "selection": selection_contract,
        "render": {
            "dpi": dpi,
            "format": "png",
            "mode": "full_page",
            "materialized": materialize_images,
        },
        "objects_file": "objects.jsonl",
        "objects_sha256": sha256_file(objects_path),
        "annotations_template": "annotations_template.jsonl",
        "counts": {
            "objects": len(items),
            "unique_pages": len({item.page for item in items}),
            "roles": dict(sorted(role_counts.items())),
            "expected_statuses": dict(sorted(expected_counts.items())),
        },
        "metric_contract": {
            "primary": [
                "response_coverage",
                "expected_recall",
                "negative_specificity",
                "classification_accuracy",
                "numeric_value_coverage",
            ],
            "value_accuracy": "confirmed annotations의 expected_rows가 있을 때만 계산",
        },
    }
    manifest_path.write_text(
        json.dumps(payload, ensure_ascii=False, indent=2),
        encoding="utf-8",
    )
    return FixedSampleDataset(manifest_path, payload, tuple(items))


def load_fixed_sample_dataset(manifest_path: str | Path) -> FixedSampleDataset:
    manifest = Path(manifest_path).resolve()
    payload = json.loads(manifest.read_text(encoding="utf-8"))
    if not isinstance(payload, dict):
        raise FixedSampleError("고정 표본 manifest는 JSON 객체여야 합니다.")
    objects_path = _resolve(_text(payload.get("objects_file")), manifest.parent)
    items: list[FixedSampleItem] = []
    for line_num, line in enumerate(objects_path.read_text(encoding="utf-8").splitlines(), start=1):
        if not line.strip():
            continue
        try:
            raw = json.loads(line)
        except json.JSONDecodeError as exc:
            raise FixedSampleError(f"objects.jsonl {line_num}행 JSON 오류") from exc
        if not isinstance(raw, dict):
            raise FixedSampleError(f"objects.jsonl {line_num}행은 JSON 객체여야 합니다.")
        items.append(FixedSampleItem.from_dict(raw))
    return FixedSampleDataset(manifest, payload, tuple(items))


def validate_fixed_sample_dataset(
    manifest_path: str | Path,
    *,
    source_override: str | Path | None = None,
    require_images: bool = True,
) -> ValidationReport:
    errors: list[str] = []
    warnings: list[str] = []
    try:
        dataset = load_fixed_sample_dataset(manifest_path)
    except (OSError, json.JSONDecodeError, FixedSampleError) as exc:
        return ValidationReport(False, (str(exc),), (), "", 0, None)
    payload = dataset.payload
    if payload.get("schema_version") != SCHEMA_VERSION:
        errors.append(
            f"schema_version 불일치: {payload.get('schema_version')} != {SCHEMA_VERSION}"
        )
    source_data = payload.get("source") if isinstance(payload.get("source"), dict) else {}
    source = (
        Path(source_override).resolve()
        if source_override
        else _resolve(_text(source_data.get("path")), dataset.manifest_path.parent)
    )
    if not source.exists():
        errors.append(f"원본 PDF를 찾을 수 없습니다: {source}")
    elif sha256_file(source) != _text(source_data.get("sha256")):
        errors.append("원본 PDF SHA256이 manifest와 다릅니다.")
    inputs = payload.get("inputs") if isinstance(payload.get("inputs"), dict) else {}
    for label in ("inventory", "audit"):
        entry = inputs.get(label) if isinstance(inputs.get(label), dict) else {}
        path_value = _text(entry.get("path"))
        expected_hash = _text(entry.get("sha256"))
        if not path_value:
            if label == "inventory":
                errors.append("manifest에 inventory snapshot 경로가 없습니다.")
            continue
        input_path = _resolve(path_value, dataset.manifest_path.parent)
        if not input_path.exists():
            errors.append(f"{label} snapshot을 찾을 수 없습니다: {input_path}")
        elif expected_hash and sha256_file(input_path) != expected_hash:
            errors.append(f"{label} snapshot SHA256이 manifest와 다릅니다.")
    objects_path = _resolve(_text(payload.get("objects_file")), dataset.manifest_path.parent)
    if not objects_path.exists():
        errors.append(f"objects 파일을 찾을 수 없습니다: {objects_path}")
    elif sha256_file(objects_path) != _text(payload.get("objects_sha256")):
        errors.append("objects.jsonl SHA256이 manifest와 다릅니다.")
    expected_count = int((payload.get("counts") or {}).get("objects", -1))
    if expected_count != len(dataset.items):
        errors.append(f"표본 개수 불일치: manifest={expected_count}, objects={len(dataset.items)}")
    selection = payload.get("selection") if isinstance(payload.get("selection"), dict) else {}
    render = payload.get("render") if isinstance(payload.get("render"), dict) else {}
    render_contract = {
        "dpi": render.get("dpi"),
        "format": render.get("format"),
        "mode": render.get("mode"),
    }
    expected_fingerprint = _dataset_fingerprint(
        source_sha256=_text(source_data.get("sha256")),
        inventory_sha256=_text((inputs.get("inventory") or {}).get("sha256")),
        audit_sha256=_text((inputs.get("audit") or {}).get("sha256")),
        selection=selection,
        render=render_contract,
    )
    if expected_fingerprint != dataset.fingerprint:
        errors.append("dataset_fingerprint가 선택·렌더 계약과 일치하지 않습니다.")
    ids = [item.sample_id for item in dataset.items]
    if len(ids) != len(set(ids)):
        errors.append("표본 sample_id가 중복되었습니다.")
    for item in dataset.items:
        if item.expected_status not in {"extracted", "not_relevant"}:
            errors.append(f"{item.sample_id}: expected_status 오류")
        if not item.image_path:
            if require_images:
                errors.append(f"{item.sample_id}: 렌더 이미지가 없습니다.")
            else:
                warnings.append(f"{item.sample_id}: 렌더 이미지 미생성")
            continue
        image_path = _resolve(item.image_path, dataset.manifest_path.parent)
        if not image_path.exists():
            errors.append(f"{item.sample_id}: 이미지 파일 누락: {image_path}")
        elif item.image_sha256 and sha256_file(image_path) != item.image_sha256:
            errors.append(f"{item.sample_id}: 이미지 SHA256 불일치")
    if payload.get("review_status") != "confirmed":
        warnings.append("사람 값 정답(expected_rows)이 확정되지 않은 draft 표본입니다.")
    return ValidationReport(
        valid=not errors,
        errors=tuple(dict.fromkeys(errors)),
        warnings=tuple(dict.fromkeys(warnings)),
        dataset_id=dataset.dataset_id,
        object_count=len(dataset.items),
        source_path=source,
    )


def load_result_rows(path: str | Path) -> list[dict[str, Any]]:
    rows: list[dict[str, Any]] = []
    for line_num, line in enumerate(Path(path).read_text(encoding="utf-8").splitlines(), start=1):
        if not line.strip():
            continue
        try:
            row = json.loads(line)
        except json.JSONDecodeError as exc:
            raise FixedSampleError(f"결과 JSONL {line_num}행 오류") from exc
        if isinstance(row, dict):
            rows.append(row)
    return rows


def load_annotations(path: str | Path | None) -> list[dict[str, Any]]:
    if path is None:
        return []
    return load_result_rows(path)


def prepare_merge_annotations(
    dataset: FixedSampleDataset,
    annotations: Sequence[dict[str, Any]] = (),
) -> list[dict[str, Any]]:
    """기존 값 주석을 보존하면서 병합 A/B용 주석 계약을 생성한다."""
    by_id = {
        _text(row.get("sample_id")): dict(row)
        for row in annotations
        if isinstance(row, dict) and _text(row.get("sample_id"))
    }
    rows: list[dict[str, Any]] = []
    for item in dataset.items:
        existing = by_id.get(item.sample_id, {})
        expected_merge = _text(existing.get("expected_merge")).casefold()
        if expected_merge not in VALID_EXPECTED_MERGE:
            expected_merge = "reject" if not item.expected else "pending"
        fixture_evidence = _text(existing.get("fixture_evidence")).casefold()
        if fixture_evidence not in VALID_EVIDENCE_FIXTURES:
            fixture_evidence = "single"
        legacy_expected_rows = existing.get("expected_rows")
        expected_visual_rows = existing.get("expected_visual_rows")
        if not isinstance(expected_visual_rows, list):
            expected_visual_rows = legacy_expected_rows if isinstance(legacy_expected_rows, list) else []
        expected_merged_rows = existing.get("expected_merged_rows")
        if not isinstance(expected_merged_rows, list):
            expected_merged_rows = legacy_expected_rows if isinstance(legacy_expected_rows, list) else []
        baseline_rows = existing.get("fixture_baseline_rows")
        rows.append({
            "sample_id": item.sample_id,
            "review_status": existing.get("review_status", "pending"),
            "expected_status": existing.get("expected_status", item.expected_status),
            "expected_title": existing.get("expected_title", item.title),
            "expected_unit": existing.get("expected_unit"),
            "expected_sheet": existing.get("expected_sheet") or item.related_sheet,
            "expected_merge": expected_merge,
            "expected_visual_rows": expected_visual_rows,
            "expected_merged_rows": expected_merged_rows,
            "expected_rows": legacy_expected_rows if isinstance(legacy_expected_rows, list) else [],
            # fixture_*는 병합 계약의 실패 사례를 재현하는 입력이고 expected_* 채점값과 분리한다.
            "fixture_evidence": fixture_evidence,
            "fixture_baseline_rows": baseline_rows if isinstance(baseline_rows, list) else [],
            "reviewer": existing.get("reviewer", ""),
            "review_basis": existing.get("review_basis", ""),
            "reviewed_at": existing.get("reviewed_at"),
            "notes": existing.get("notes", ""),
        })
    return rows


def _cell_equal(expected: Any, actual: Any, field_name: str = "") -> bool:
    return comparison_equal(expected, actual, field_name)


def _annotation_value_score(
    result_by_id: dict[str, dict[str, Any]],
    annotations: Sequence[dict[str, Any]],
) -> tuple[int, int, int]:
    matched_cells = 0
    expected_cells = 0
    confirmed_objects = 0
    for annotation in annotations:
        if _text(annotation.get("review_status")).casefold() not in REVIEWED_ANNOTATION_STATUSES:
            continue
        expected_rows = annotation.get("expected_visual_rows")
        if not isinstance(expected_rows, list):
            expected_rows = annotation.get("expected_rows")
        if not isinstance(expected_rows, list) or not expected_rows:
            continue
        confirmed_objects += 1
        result = result_by_id.get(_text(annotation.get("sample_id")), {})
        raw_actual_rows = result.get("table") if isinstance(result.get("table"), list) else []
        raw_actual_rows = normalize_visual_table_rows(
            raw_actual_rows,
            chart_type=_text(result.get("object_type")),
            title=_text(result.get("title")),
        )
        actual_rows: list[dict[str, Any]] = []
        for raw_row in raw_actual_rows:
            if not isinstance(raw_row, dict):
                continue
            fields = raw_row.get("fields") if isinstance(raw_row.get("fields"), dict) else {}
            actual_rows.append({**fields, **raw_row})
        used_actual: set[int] = set()
        for expected_row in expected_rows:
            if not isinstance(expected_row, dict):
                continue
            cells = {
                str(key): value for key, value in expected_row.items()
                if value is not None and _text(value) != ""
            }
            expected_cells += len(cells)
            if not cells:
                continue
            candidates = [
                (index, row)
                for index, row in enumerate(actual_rows)
                if index not in used_actual
            ]
            if not candidates:
                continue
            best_index, best = max(
                candidates,
                key=lambda pair: sum(
                    _cell_equal(value, pair[1].get(key), key)
                    for key, value in cells.items()
                ),
            )
            used_actual.add(best_index)
            matched_cells += sum(
                _cell_equal(value, best.get(key), key) for key, value in cells.items()
            )
    return matched_cells, expected_cells, confirmed_objects


def _annotation_numeric_score(
    result_by_id: dict[str, dict[str, Any]],
    annotations: Sequence[dict[str, Any]],
) -> tuple[int, int, int]:
    """Compare explicit numeric values independently from labels and row layout."""
    matched = 0
    expected_count = 0
    actual_count = 0
    for annotation in annotations:
        if _text(annotation.get("review_status")).casefold() not in REVIEWED_ANNOTATION_STATUSES:
            continue
        expected_rows = annotation.get("expected_visual_rows")
        if not isinstance(expected_rows, list):
            expected_rows = annotation.get("expected_rows")
        if not isinstance(expected_rows, list) or not expected_rows:
            continue
        result = result_by_id.get(_text(annotation.get("sample_id")), {})
        actual_rows = result.get("table") if isinstance(result.get("table"), list) else []
        actual_rows = normalize_visual_table_rows(
            actual_rows,
            chart_type=_text(result.get("object_type")),
            title=_text(result.get("title")),
        )
        expected_numbers = numeric_values(expected_rows)
        actual_numbers = numeric_values(actual_rows)
        expected_count += len(expected_numbers)
        actual_count += len(actual_numbers)
        used: set[int] = set()
        for expected in expected_numbers:
            for index, actual in enumerate(actual_numbers):
                if index in used or not _cell_equal(expected, actual, "값"):
                    continue
                used.add(index)
                matched += 1
                break
    return matched, expected_count, actual_count


def _annotation_status_score(
    result_by_id: dict[str, dict[str, Any]],
    annotations: Sequence[dict[str, Any]],
) -> tuple[int, int]:
    """원본 대조를 마친 주석의 상태 정답과 모델 판정을 비교한다."""
    matched = 0
    expected = 0
    for annotation in annotations:
        if _text(annotation.get("review_status")).casefold() not in REVIEWED_ANNOTATION_STATUSES:
            continue
        expected_status = _text(annotation.get("expected_status")).casefold()
        if expected_status not in VALID_RESULT_STATUSES - {"failed"}:
            continue
        expected += 1
        result = result_by_id.get(_text(annotation.get("sample_id")), {})
        actual_status = _text(result.get("status")).casefold() or "failed"
        if actual_status not in VALID_RESULT_STATUSES:
            actual_status = "needs_review"
        matched += actual_status == expected_status
    return matched, expected


def _has_numeric_value(row: dict[str, Any]) -> bool:
    table = row.get("table")
    if not isinstance(table, list):
        return False
    for item in table:
        if not isinstance(item, dict):
            continue
        value = item.get("값")
        if isinstance(value, (int, float)) and not isinstance(value, bool):
            return True
        values = item.get("values") or item.get("연도별")
        if isinstance(values, dict) and any(
            isinstance(cell, (int, float)) and not isinstance(cell, bool)
            for cell in values.values()
        ):
            return True
    return False


def evaluate_fixed_sample_results(
    dataset: FixedSampleDataset,
    result_rows: Sequence[dict[str, Any]],
    annotations: Sequence[dict[str, Any]] = (),
) -> dict[str, Any]:
    by_id: dict[str, dict[str, Any]] = {}
    duplicate_ids: list[str] = []
    extra_ids: list[str] = []
    expected_ids = {item.sample_id for item in dataset.items}
    for row in result_rows:
        sample_id = _text(row.get("sample_id"))
        if not sample_id:
            continue
        if sample_id not in expected_ids:
            extra_ids.append(sample_id)
            continue
        if sample_id in by_id:
            duplicate_ids.append(sample_id)
            continue
        by_id[sample_id] = row

    positives = [item for item in dataset.items if item.expected]
    negatives = [item for item in dataset.items if not item.expected]
    data_positives = [item for item in positives if item.data_included]
    received = 0
    positive_hits = 0
    negative_hits = 0
    numeric_hits = 0
    statuses: Counter[str] = Counter()
    missing_ids: list[str] = []
    for item in dataset.items:
        row = by_id.get(item.sample_id)
        status = _text(row.get("status")) if row else "failed"
        if status not in VALID_RESULT_STATUSES:
            status = "needs_review"
        statuses[status] += 1
        if row and bool(row.get("response_received", True)) and status != "failed":
            received += 1
        else:
            missing_ids.append(item.sample_id)
        if item.expected and status == "extracted":
            positive_hits += 1
        if not item.expected and status == "not_relevant":
            negative_hits += 1
        if item.data_included and item.expected and row and status == "extracted" and _has_numeric_value(row):
            numeric_hits += 1

    total = len(dataset.items)
    correct = positive_hits + negative_hits
    matched_cells, expected_cells, confirmed_objects = _annotation_value_score(
        by_id,
        annotations,
    )
    reviewed_status_matches, reviewed_status_expected = _annotation_status_score(
        by_id,
        annotations,
    )
    numeric_matches, numeric_expected, numeric_actual = _annotation_numeric_score(
        by_id,
        annotations,
    )
    value_accuracy = round(matched_cells / expected_cells, 4) if expected_cells else None
    return {
        "dataset_id": dataset.dataset_id,
        "dataset_fingerprint": dataset.fingerprint,
        "metrics": {
            "sample_count": total,
            "response_coverage": round(received / total, 4) if total else 0.0,
            "expected_recall": round(positive_hits / len(positives), 4) if positives else None,
            "negative_specificity": round(negative_hits / len(negatives), 4) if negatives else None,
            "classification_accuracy": round(correct / total, 4) if total else 0.0,
            "reviewed_classification_accuracy": (
                round(reviewed_status_matches / reviewed_status_expected, 4)
                if reviewed_status_expected else None
            ),
            "numeric_value_coverage": round(numeric_hits / len(data_positives), 4) if data_positives else None,
            "structured_cell_accuracy": value_accuracy,
            "numeric_value_recall": (
                round(numeric_matches / numeric_expected, 4) if numeric_expected else None
            ),
            "numeric_value_precision": (
                round(numeric_matches / numeric_actual, 4) if numeric_actual else None
            ),
            "value_accuracy": value_accuracy,
            "value_accuracy_note": (
                f"원본 대조 완료 객체 {confirmed_objects}개, 기대 셀 {expected_cells}개 기준"
                if expected_cells
                else "검토 완료된 expected_visual_rows가 없어 계산하지 않음"
            ),
        },
        "counts": {
            "positives": len(positives),
            "positive_hits": positive_hits,
            "negatives": len(negatives),
            "negative_hits": negative_hits,
            "data_positives": len(data_positives),
            "numeric_hits": numeric_hits,
            "confirmed_annotation_objects": confirmed_objects,
            "reviewed_status_objects": reviewed_status_expected,
            "reviewed_status_matches": reviewed_status_matches,
            "value_cells_matched": matched_cells,
            "value_cells_expected": expected_cells,
            "numeric_values_matched": numeric_matches,
            "numeric_values_expected": numeric_expected,
            "numeric_values_actual": numeric_actual,
            "statuses": dict(sorted(statuses.items())),
        },
        "missing_ids": missing_ids,
        "duplicate_ids": sorted(set(duplicate_ids)),
        "extra_ids": sorted(set(extra_ids)),
    }


def _merge_sheet_key(value: Any) -> str:
    import config

    text = _text(value)
    if text in config.SHEET_KEY_TO_NAME:
        return text
    by_name = {name: key for key, name in config.SHEET_KEY_TO_NAME.items()}
    if text in by_name:
        return by_name[text]
    legacy = {
        "vehicle": "regional_conditions",
        "energy": "regional_conditions",
        "ghg": "emissions_regional",
        "forecast": "emissions_forecast",
        "target": "reduction_targets",
        "strategy": "mitigation_projects",
    }
    return legacy.get(text, "")


def _fixed_result_analysis(item: FixedSampleItem, result: dict[str, Any]) -> dict[str, Any]:
    """정답 주석을 보지 않고 고정 표본 결과를 운영 ImageAgent 입력으로 변환한다."""
    from agents.image_agent import ImageAgent

    target_sheet = _merge_sheet_key(result.get("target_sheet"))
    if not target_sheet:
        target_sheet = _merge_sheet_key(item.related_sheet)
    inference_payload = {
        "target_sheet": target_sheet,
        "title": result.get("title") or item.title,
        "summary": result.get("summary", ""),
        "unit": result.get("unit"),
    }
    if not target_sheet:
        target_sheet = ImageAgent()._infer_target_sheet(inference_payload)

    object_type = _text(result.get("object_type") or item.element_type)
    if "표" in object_type:
        chart_type = "표"
    elif "막대" in object_type:
        chart_type = "막대"
    else:
        chart_type = "꺾은선"

    normalized_rows: list[dict[str, Any]] = []
    raw_rows = result.get("table") if isinstance(result.get("table"), list) else []
    for raw in raw_rows:
        if not isinstance(raw, dict):
            continue
        row = dict(raw)
        fields = dict(row.get("fields")) if isinstance(row.get("fields"), dict) else {}
        for key in ("항목", "연도", "단위"):
            if row.get(key) is not None and fields.get(key) is None:
                fields[key] = row.get(key)
        item_name = row.get("항목")
        if target_sheet == "regional_conditions":
            fields.setdefault("지표명", item_name)
        elif target_sheet == "emissions_regional":
            fields.setdefault("부문", item_name)
        elif target_sheet == "emissions_management":
            fields.setdefault("관리부문", item_name)
        elif target_sheet == "emissions_forecast":
            fields.setdefault("부문", item_name)
        elif target_sheet == "reduction_targets":
            fields.setdefault("부문", item_name)
            fields.setdefault("목표연도", row.get("연도"))
        elif target_sheet == "financial_plan":
            fields.setdefault("부문", item_name)
        row["fields"] = {key: value for key, value in fields.items() if value is not None}
        normalized_rows.append(row)
    normalized_rows = normalize_visual_table_rows(
        normalized_rows,
        chart_type=chart_type,
        title=_text(result.get("title") or item.title),
    )

    summary = _text(result.get("summary"))
    if _text(result.get("scope")).casefold() == "external_reference":
        summary = f"참고자료/해외사례 {summary}".strip()
    return {
        "contract_version": result.get("contract_version") or VISUAL_CONTRACT_VERSION,
        "type": "chart_table",
        "target_sheet": target_sheet,
        "chart_type": chart_type,
        "title": result.get("title") or item.title,
        "summary": summary,
        "unit": result.get("unit"),
        "confidence": result.get("confidence") or "low",
        "page_number": item.page,
        "municipality": "",
        "table": normalized_rows,
    }


def _evidence_fixture(
    item: FixedSampleItem,
    result: dict[str, Any],
    annotation: dict[str, Any],
) -> tuple[list[str], list[dict[str, Any]]]:
    fixture = _text(annotation.get("fixture_evidence")).casefold() or "single"
    if fixture not in VALID_EVIDENCE_FIXTURES:
        raise FixedSampleError(f"{item.sample_id}: 지원하지 않는 fixture_evidence={fixture}")
    evidence_id = f"ev-fixed-{item.sample_id}"
    status = _text(result.get("status")) or "needs_review"
    final_status = "extracted" if status == "extracted" else status

    def obj(object_id: str, eid: str, object_status: str = final_status) -> dict[str, Any]:
        return {
            "object_id": object_id,
            "object_type": "chart",
            "page_number": item.page,
            "sequence": 1,
            "metadata": {
                "evidence_id": eid,
                "final_status": object_status,
            },
        }

    if fixture == "missing":
        return [], [obj(f"fixed-{item.sample_id}", evidence_id)]
    if fixture == "multiple_ids":
        second = f"{evidence_id}-secondary"
        return [evidence_id, second], [
            obj(f"fixed-{item.sample_id}-1", evidence_id),
            obj(f"fixed-{item.sample_id}-2", second),
        ]
    if fixture == "multiple_objects":
        return [evidence_id], [
            obj(f"fixed-{item.sample_id}-1", evidence_id),
            obj(f"fixed-{item.sample_id}-2", evidence_id),
        ]
    if fixture == "object_not_extracted":
        return [evidence_id], [obj(f"fixed-{item.sample_id}", evidence_id, "needs_review")]
    return [evidence_id], [obj(f"fixed-{item.sample_id}", evidence_id)]


def _baseline_rows_by_sheet(
    annotation: dict[str, Any],
    fallback_sheet: str,
) -> dict[str, list[dict[str, Any]]]:
    grouped: dict[str, list[dict[str, Any]]] = defaultdict(list)
    rows = annotation.get("fixture_baseline_rows")
    for raw in rows if isinstance(rows, list) else []:
        if not isinstance(raw, dict):
            continue
        row = dict(raw)
        sheet_key = _merge_sheet_key(
            row.pop("_sheet", None)
            or row.pop("target_sheet", None)
            or row.pop("대상시트", None)
        ) or fallback_sheet
        if not sheet_key:
            continue
        row.setdefault("데이터상태", "reported")
        grouped[sheet_key].append(row)
    return dict(grouped)


def _merge_decision(observations: Sequence[dict[str, Any]], result_status: str) -> str:
    statuses = {_text(row.get("병합상태")) for row in observations if isinstance(row, dict)}
    if "fix_then_merge" in statuses:
        return "fix_then_merge"
    if "accept" in statuses or "merged" in statuses or any(
        row.get("반영여부") == "반영" for row in observations
    ):
        return "accept"
    if statuses & {"duplicate", "verified_duplicate", "duplicate_visual"}:
        return "duplicate"
    if "reject" in statuses:
        return "reject"
    if "needs_review" in statuses or "candidate" in statuses:
        return "needs_review"
    if result_status in {"not_relevant", "no_data", "failed"} or not observations:
        return "reject"
    return "needs_review"


def _run_merge_mode(
    item: FixedSampleItem,
    result: dict[str, Any],
    annotation: dict[str, Any],
    *,
    mode: str,
    municipality: str,
) -> dict[str, Any]:
    import config
    from agents.image_agent import ImageAgent
    from agents.organizer_agent import OrganizerAgent
    from utils.evidence_merge import build_evidence_catalog

    analysis = _fixed_result_analysis(item, result)
    analysis["municipality"] = municipality
    target_sheet = _merge_sheet_key(analysis.get("target_sheet"))
    evidence_ids, document_objects = _evidence_fixture(item, result, annotation)
    analysis["source_evidence_ids"] = evidence_ids
    raw: dict[str, Any] = {"municipality_name": municipality}
    for sheet_key, rows in _baseline_rows_by_sheet(annotation, target_sheet).items():
        raw[sheet_key] = deepcopy(rows)

    image_agent = ImageAgent()
    previous_labeled = getattr(config, "VISUAL_MERGE_LABELED_ENABLED", True)
    previous_evidence = getattr(config, "VISUAL_EVIDENCE_MERGE_ENABLED", True)
    try:
        config.VISUAL_MERGE_LABELED_ENABLED = True
        config.VISUAL_EVIDENCE_MERGE_ENABLED = True
        if mode == "evidence":
            catalog = build_evidence_catalog(document_objects)
            raw = image_agent._merge_image_results(
                raw,
                [analysis],
                municipality,
                evidence_catalog=catalog,
            )
            raw["document_objects"] = deepcopy(document_objects)
            raw["object_triage"] = []
        else:
            raw = image_agent._merge_image_results(raw, [analysis], municipality)
        # 표본 A/B 보고서에는 Organizer의 진행 로그를 섞지 않는다.
        with redirect_stdout(io.StringIO()):
            cleaned = OrganizerAgent().organize(raw)
    finally:
        config.VISUAL_MERGE_LABELED_ENABLED = previous_labeled
        config.VISUAL_EVIDENCE_MERGE_ENABLED = previous_evidence

    observations = [
        deepcopy(row) for row in cleaned.get("chart_observations", [])
        if isinstance(row, dict)
    ]
    merged_rows: list[dict[str, Any]] = []
    # 기존 경로는 target_sheet와 다른 시트에 직접 쓰는 오라우팅도 가능하므로
    # 대상 시트만 보지 않고 전체 운영 시트의 실제 visual_only 출력을 수집한다.
    for sheet_key in getattr(config, "EXTRACTION_SHEETS", []):
        rows = cleaned.get(sheet_key, [])
        for row in rows if isinstance(rows, list) else []:
            if not isinstance(row, dict) or _text(row.get("데이터상태")) != "visual_only":
                continue
            merged = deepcopy(row)
            merged["_sheet"] = sheet_key
            merged_rows.append(merged)
    return {
        "sample_id": item.sample_id,
        "page": item.page,
        "mode": mode,
        "target_sheet": target_sheet,
        "result_status": _text(result.get("status")) or "failed",
        "scope": _text(result.get("scope")),
        "fixture_evidence": _text(annotation.get("fixture_evidence")) or "single",
        "decision": _merge_decision(observations, _text(result.get("status"))),
        "merged_rows": merged_rows,
        "observations": observations,
    }


def _rows_match(expected: dict[str, Any], actual: dict[str, Any]) -> bool:
    cells = {
        str(key): value for key, value in expected.items()
        if value is not None and _text(value) != ""
    }
    return bool(cells) and all(
        _cell_equal(value, actual.get(key), key) for key, value in cells.items()
    )


def _match_rows(
    expected_rows: Sequence[dict[str, Any]],
    actual_rows: Sequence[dict[str, Any]],
) -> int:
    used: set[int] = set()
    matched = 0
    for expected in expected_rows:
        if not isinstance(expected, dict):
            continue
        for index, actual in enumerate(actual_rows):
            if index in used or not isinstance(actual, dict):
                continue
            if _rows_match(expected, actual):
                used.add(index)
                matched += 1
                break
    return matched


def _ratio(numerator: int, denominator: int) -> float | None:
    return round(numerator / denominator, 4) if denominator else None


def _merge_mode_metrics(details: Sequence[dict[str, Any]]) -> dict[str, Any]:
    confirmed = [row for row in details if row.get("scored")]
    observations = [
        obs for row in details for obs in row.get("observations", []) if isinstance(obs, dict)
    ]
    merged_rows = [
        item for row in details for item in row.get("merged_rows", []) if isinstance(item, dict)
    ]
    expected_rows_total = 0
    matched_rows_total = 0
    false_merge_rows = 0
    scored_merged_rows = 0
    decision_matches = 0
    sheet_evaluable = 0
    sheet_matches = 0
    merged_sheet_rows = 0
    merged_sheet_matches = 0
    for row in confirmed:
        expected_merge = row.get("expected_merge")
        expected_rows = [
            item for item in row.get("expected_merged_rows", [])
            if isinstance(item, dict)
        ]
        actual_rows = [item for item in row.get("merged_rows", []) if isinstance(item, dict)]
        if row.get("decision") == expected_merge:
            decision_matches += 1
        expected_sheet = _merge_sheet_key(row.get("expected_sheet"))
        if expected_sheet:
            sheet_evaluable += 1
            if _merge_sheet_key(row.get("target_sheet")) == expected_sheet:
                sheet_matches += 1
            if expected_merge == "accept":
                merged_sheet_rows += len(actual_rows)
                merged_sheet_matches += sum(
                    _merge_sheet_key(actual.get("_sheet")) == expected_sheet
                    for actual in actual_rows
                )
        if expected_merge == "accept" and expected_rows:
            matched = _match_rows(expected_rows, actual_rows)
            expected_rows_total += len(expected_rows)
            matched_rows_total += matched
            scored_merged_rows += len(actual_rows)
            false_merge_rows += max(0, len(actual_rows) - matched)
        elif expected_merge != "accept":
            scored_merged_rows += len(actual_rows)
            false_merge_rows += len(actual_rows)

    exact = sum(_text(row.get("근거매칭상태")) == "exact" for row in observations)
    needs_review = [row for row in observations if _text(row.get("병합상태")) == "needs_review"]
    reason_count = sum(bool(_text(row.get("병합차단사유"))) for row in needs_review)

    multiple_cases = [
        row for row in details
        if row.get("fixture_evidence") in {"multiple_ids", "multiple_objects"}
    ]
    multiple_isolated = sum(
        not row.get("merged_rows") and row.get("decision") == "needs_review"
        for row in multiple_cases
    )
    null_cases = [
        row for row in details
        if row.get("observations")
        and all(obs.get("값") is None for obs in row.get("observations", []))
    ]
    null_isolated = sum(
        not row.get("merged_rows") and row.get("decision") == "needs_review"
        for row in null_cases
    )
    conflict_cases = [
        row for row in details
        if any("텍스트-시각" in _text(obs.get("병합차단사유")) for obs in row.get("observations", []))
    ]
    conflict_isolated = sum(not row.get("merged_rows") for row in conflict_cases)

    return {
        "sample_count": len(details),
        "confirmed_samples": len(confirmed),
        "decision_accuracy": _ratio(decision_matches, len(confirmed)),
        "target_sheet_accuracy": _ratio(sheet_matches, sheet_evaluable),
        "merged_sheet_accuracy": _ratio(merged_sheet_matches, merged_sheet_rows),
        "auto_merge_rows": len(merged_rows),
        "scored_auto_merge_rows": scored_merged_rows,
        "matched_merged_rows": matched_rows_total,
        "false_merge_rows": false_merge_rows,
        "auto_merge_precision": _ratio(
            scored_merged_rows - false_merge_rows,
            scored_merged_rows,
        ),
        "auto_merge_recall": _ratio(matched_rows_total, expected_rows_total),
        "exact_evidence_match_rate": _ratio(exact, len(observations)),
        "needs_review_candidates": len(needs_review),
        "needs_review_reason_coverage": _ratio(reason_count, len(needs_review)),
        "multiple_evidence_isolation_rate": _ratio(multiple_isolated, len(multiple_cases)),
        "null_isolation_rate": _ratio(null_isolated, len(null_cases)),
        "text_conflict_isolation_rate": _ratio(conflict_isolated, len(conflict_cases)),
        "unscored_samples": len(details) - len(confirmed),
    }


def evaluate_fixed_sample_merge_ab(
    dataset: FixedSampleDataset,
    result_rows: Sequence[dict[str, Any]],
    annotations: Sequence[dict[str, Any]],
    *,
    municipality: str = "서울특별시",
) -> dict[str, Any]:
    """동일 VLM 결과에 legacy/evidence 병합을 적용해 정책 차이만 비교한다."""
    result_by_id = {
        _text(row.get("sample_id")): dict(row)
        for row in result_rows
        if isinstance(row, dict) and _text(row.get("sample_id"))
    }
    prepared_annotations = prepare_merge_annotations(dataset, annotations)
    annotation_by_id = {row["sample_id"]: row for row in prepared_annotations}
    details_by_mode: dict[str, list[dict[str, Any]]] = {"legacy": [], "evidence": []}
    for item in dataset.items:
        result = result_by_id.get(item.sample_id, {
            "sample_id": item.sample_id,
            "status": "failed",
            "table": [],
            "confidence": "low",
        })
        annotation = annotation_by_id[item.sample_id]
        expected_merge = _text(annotation.get("expected_merge")).casefold()
        review_status = _text(annotation.get("review_status")).casefold()
        scored = (
            review_status in {"confirmed", "확정", "reviewed", "검토완료"}
            and expected_merge in VALID_EXPECTED_MERGE - {"pending"}
        )
        for mode in ("legacy", "evidence"):
            detail = _run_merge_mode(
                item,
                result,
                annotation,
                mode=mode,
                municipality=municipality,
            )
            detail.update({
                "review_status": annotation.get("review_status"),
                "expected_sheet": annotation.get("expected_sheet"),
                "expected_merge": expected_merge,
                "expected_visual_rows": annotation.get("expected_visual_rows", []),
                "expected_merged_rows": annotation.get("expected_merged_rows", []),
                # 이전 소비자 호환용 별칭. 병합 채점은 expected_merged_rows만 사용한다.
                "expected_rows": annotation.get("expected_merged_rows", []),
                "scored": scored,
            })
            details_by_mode[mode].append(detail)

    metrics = {
        mode: _merge_mode_metrics(details)
        for mode, details in details_by_mode.items()
    }
    legacy = metrics["legacy"]
    evidence = metrics["evidence"]

    def delta(key: str) -> float | None:
        left = legacy.get(key)
        right = evidence.get(key)
        if isinstance(left, (int, float)) and isinstance(right, (int, float)):
            return round(float(right) - float(left), 4)
        return None

    return {
        "dataset_id": dataset.dataset_id,
        "dataset_fingerprint": dataset.fingerprint,
        "metrics": metrics,
        "comparison": {
            "auto_merge_precision_delta": delta("auto_merge_precision"),
            "auto_merge_recall_delta": delta("auto_merge_recall"),
            "false_merge_reduction": legacy["false_merge_rows"] - evidence["false_merge_rows"],
            "auto_merge_row_delta": evidence["auto_merge_rows"] - legacy["auto_merge_rows"],
        },
        "details": details_by_mode,
        "annotation_status": {
            "total": len(prepared_annotations),
            "confirmed": sum(row.get("scored", False) for row in details_by_mode["evidence"]),
            "pending": sum(not row.get("scored", False) for row in details_by_mode["evidence"]),
        },
    }
