"""Frozen-input recovery helpers. No model, PDF, or gold-answer dependencies."""
from __future__ import annotations

from collections import Counter, defaultdict
from copy import deepcopy
import hashlib
import json
import math
from pathlib import Path
import re

VERSION = "vision-recovery/1"


def canonical(value):
    return json.dumps(value, ensure_ascii=False, sort_keys=True, allow_nan=False, default=str)


def digest(value):
    return hashlib.sha256(canonical(value).encode("utf-8")).hexdigest()


def file_hash(path):
    return hashlib.sha256(Path(path).read_bytes()).hexdigest()


def read_json(path):
    return json.loads(Path(path).read_text(encoding="utf-8-sig"))


def write_json(path, value):
    Path(path).write_text(json.dumps(value, ensure_ascii=False, indent=2, allow_nan=False,
                                    default=str) + "\n", encoding="utf-8")


def evidence_ids(row):
    values = row.get("근거ID목록") or [row.get("근거ID", "")]
    if isinstance(values, str):
        values = [values]
    return sorted({str(v) for v in values if v})


def observation_identity(row):
    """Use the same source/metric identity as the production merge gate."""
    from utils.visual_fact_identity import fact_identity
    return canonical(fact_identity(row))


def observation_payload(row):
    # Include meaning/wording/unit differences, not just numeric differences.
    return {k: row.get(k) for k in ("항목", "원문값", "원문단위", "값", "단위", "판독필드")}


def merge_fresh(original, fresh, allowed_page, allowed_evidence):
    """Append new identities only; changed existing readings are quarantined."""
    result = deepcopy(original)
    existing = defaultdict(list)
    for row in result:
        existing[observation_identity(row)].append(row)
    audit = []
    for row in fresh:
        if row.get("페이지") != allowed_page or not set(evidence_ids(row)) <= set(allowed_evidence) or not evidence_ids(row):
            audit.append({"status": "out_of_scope", "candidate": row})
            continue
        identity = observation_identity(row)
        matches = existing.get(identity, [])
        if matches:
            same = any(observation_payload(x) == observation_payload(row) for x in matches)
            audit.append({"status": "unchanged" if same else "conflict_held",
                          "difference_kind": "none" if same else (
                              "metadata_or_unit_difference" if any(
                                  x.get("원문값") == row.get("원문값") for x in matches)
                              else "value_difference"),
                          "existing": matches, "candidate": row})
        else:
            copied = deepcopy(row)
            copied["_recovery_id"] = "retry-" + digest(row)[:20]
            result.append(copied)
            existing[identity].append(copied)
            audit.append({"status": "added", "recovery_id": copied["_recovery_id"]})
    return result, audit


def record_response_success(raw, task):
    """Update only the exact successful source object on a working copy.

    Call only after observations_from_response has validated the response.
    Failed tasks and unrelated objects must never be promoted.
    """
    matches = [obj for obj in raw.get("document_objects", [])
               if obj.get("object_id") in task["object_ids"]]
    if (len(matches) != 1 or len(task["object_ids"]) != 1 or len(task["evidence_ids"]) != 1
            or matches[0].get("page_number") != task["page"]
            or matches[0].get("object_type") not in {"image", "chart", "table", "diagram", "infographic"}
            or (matches[0].get("metadata") or {}).get("evidence_id") != task["evidence_ids"][0]
            or matches[0].get("bbox") != task.get("bbox")):
        raise ValueError("Successful task does not identify its original source object")
    obj = matches[0]
    previous = deepcopy(obj.get("metadata", {}))
    obj.setdefault("metadata", {}).update(
        final_status="extracted", ocr_status="completed",
        recovery_response_verified=True,
        terminal_reason="Validated successful Vision response")
    triage_changes = []
    for row in raw.get("object_triage", []):
        if (row.get("object_id") == obj["object_id"]
                and row.get("evidence_id") == task["evidence_ids"][0]):
            triage_changes.append(deepcopy(row))
            row.update(final_status="extracted", ocr_status="completed",
                       terminal_reason="Validated successful Vision response")
    return {"task": task.get("id"), "object_id": obj["object_id"],
            "before": previous, "after": deepcopy(obj["metadata"]),
            "previous_triage": triage_changes}


def scoped_region(fields, municipality):
    """Resolve only an explicitly evidenced row/object scope; never a filename."""
    from utils.reading_context_facts import scoped_municipality
    return scoped_municipality(fields, municipality)


def prepare_bindings(raw):
    """Conservative, opt-in refinements. Never read gold values or repair readings."""
    result = deepcopy(raw)
    events = []
    for i, row in enumerate(result.get("chart_observations", []), 1):
        row.setdefault("_recovery_id", f"original-{i:06}")
        f = row.get("판독필드") or {}
        target = row.get("대상시트")
        named_project = bool(f.get("사업명") or f.get("감축사업명") or f.get("관리번호"))
        legend = row.get("범례목록") or []
        non_project = target == "mitigation_projects" and not named_project and (
            row.get("항목") in legend or f.get("구조역할") in {"기타", "단계", "지표"}
            or f.get("정보유형") == "성과집계")
        if non_project:
            row["자동병합정책"] = "block_structured_visual"
            events.append({"record_id": row["_recovery_id"], "action": "non_project_held",
                           "reason": "Explicit legend/structural item is not a named project"})
        region = scoped_region(f, result.get("municipality_name"))
        if region and region != result.get("municipality_name"):
            events.append({"record_id": row["_recovery_id"], "action": "scoped_region",
                           "region": region, "evidence": f["_reading"]["객체문맥"]})
    return result, events


def build_retry_queue(model, raw, comparisons):
    """Gold determines *where* to review, never supplies answers to the call."""
    observations = raw.get("chart_observations", [])
    objects = raw.get("document_objects", [])
    by_page = defaultdict(dict)
    for obj in objects:
        if obj.get("object_type") not in {"image", "chart", "table", "diagram", "infographic"}:
            continue  # A full-page text proxy is not an additional Vision object.
        ev = (obj.get("metadata") or {}).get("evidence_id")
        page = obj.get("page_number")
        if ev and isinstance(page, int):
            by_page[page].setdefault(ev, obj)
    queue = {}
    held = []
    for fact in comparisons:
        if fact.get("model") != model:
            continue
        reading_problem = not fact.get("present") or not fact.get("complete")
        if not reading_problem:
            continue
        page = fact["page"]
        if len(by_page[page]) != 1:
            held.append({"gold_id": fact["id"], "page": page,
                         "reason": "full_page_retry_requires_one_physical_object"})
            continue
        # Prefer the exact object already associated with this evaluation fact.
        match = re.match(r"[A-Za-z](\d+)", fact.get("pred_id", ""))
        obs = observations[int(match[1])-1] if match and 0 < int(match[1]) <= len(observations) else None
        evs = evidence_ids(obs) if obs else list(by_page[page])
        for ev in evs:
            obj = by_page[page].get(ev)
            if not obj:
                held.append({"gold_id": fact["id"], "page": page, "reason": "source_object_missing"})
                continue
            qid = f"{model}:p{page}:{ev}"
            if qid not in queue:
                queue[qid] = {"id": qid, "model": model, "page": page, "evidence_ids": [ev],
                              "object_ids": [obj["object_id"]], "bbox": obj.get("bbox"),
                              "fact_ids": [], "reasons": [], "render": "full_page_context"}
            q = queue[qid]
            q["fact_ids"].append(fact["id"])
            reason = "missing_or_timeout" if not fact.get("present") else "reading_fields_mismatch"
            if reason not in q["reasons"]:
                q["reasons"].append(reason)
        if not evs:
            held.append({"gold_id": fact["id"], "page": page, "reason": "page_objects_missing"})
    return sorted(queue.values(), key=lambda x: (x["page"], x["id"])), held


def projected_cell(value):
    if value == "":
        return None
    if isinstance(value, bool):
        return "Y" if value else "N"
    if isinstance(value, (dict, list)):
        return json.dumps(value, ensure_ascii=False, sort_keys=True)
    if isinstance(value, float) and value.is_integer():
        return int(value)
    return value


def verify_saved_workbook(path, data, schema):
    """Independent multiset check; evidence IDs are object IDs, not unique row IDs."""
    import openpyxl
    from utils.excel_writer import HEADER_KEY_ALIASES, _stable_output_rows, _headers_for_output
    wb = openpyxl.load_workbook(path, read_only=True, data_only=False)
    issues, trace = [], []
    try:
        for key in schema.EXTRACTION_SHEETS:
            name = schema.SHEET_KEY_TO_NAME[key]
            if name not in wb:
                issues.append({"kind": "missing_sheet", "sheet": name})
                continue
            expected_headers = _headers_for_output(name, schema.EXCEL_HEADERS[name])
            actual_headers = [c.value for c in next(wb[name].iter_rows())]
            if expected_headers != actual_headers:
                issues.append({"kind": "header_mismatch", "sheet": name})
                continue
            expected = Counter()
            for row in _stable_output_rows(name, data.get(key, [])):
                vals = [projected_cell(row.get(h) if row.get(h) is not None else
                                      row.get(HEADER_KEY_ALIASES.get(h, h))) for h in expected_headers]
                expected[canonical(vals)] += 1
            actual = Counter()
            for cells in wb[name].iter_rows(min_row=2):
                vals = [c.value for c in cells]
                if not any(v is not None for v in vals):
                    continue
                actual[canonical(vals)] += 1
                for c, h in zip(cells, actual_headers):
                    if c.data_type in ("e", "f"):
                        issues.append({"kind": "formula_or_error", "sheet": name, "cell": c.coordinate})
                trace.append({"sheet_key": key, "sheet": name, "row": cells[0].row,
                              "cells": {h: {"address": c.coordinate, "value": c.value}
                                        for h, c in zip(actual_headers, cells)}})
            if actual != expected:
                issues.append({"kind": "row_multiset_mismatch", "sheet": name,
                               "missing": list((expected-actual).elements()),
                               "extra": list((actual-expected).elements())})
    finally:
        wb.close()
    return {"passed": not issues, "issues": issues, "trace": trace}


def validate_budget(budget):
    for name in ("max_calls", "call_timeout_seconds", "total_seconds"):
        value = budget.get(name)
        if isinstance(value, bool) or not isinstance(value, (int, float)) or not math.isfinite(value) or value <= 0:
            raise ValueError(f"Positive finite budget required: {name}")
    if not isinstance(budget["max_calls"], int):
        raise ValueError("max_calls must be an integer")
