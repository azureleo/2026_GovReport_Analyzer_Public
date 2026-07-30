"""
구조화된 데이터를 엑셀 파일로 작성하는 유틸리티

carbon_guideline.md 기반 16개 시트 구조의 Excel 파일을 생성합니다.
"""

from datetime import datetime
import json
from pathlib import Path

import openpyxl
from openpyxl.cell.cell import ILLEGAL_CHARACTERS_RE
from openpyxl.styles import (
    Alignment, Font, PatternFill, Border, Side,
)
from openpyxl.utils import get_column_letter

import config
from utils.run_state import merge_rows_stably


COLOR_HEADER_BG = "2E75B6"
COLOR_HEADER_FONT = "FFFFFF"
COLOR_ALT_ROW = "F2F2F2"

# 엑셀 헤더 라벨과 데이터 dict 키가 다른 경우의 별칭 매핑.
# 예: 헤더는 "감축률(%)"로 표기하지만 추출·정제 단계의 dict 키는 "감축률"이다.
# 별칭이 없으면 header == key 로 간주한다.
HEADER_KEY_ALIASES = {
    "감축률(%)": "감축률",
}


def _headers_for_output(sheet_name: str, headers: list[str]) -> list[str]:
    """출처페이지·데이터상태 opt-out을 엑셀 쓰기 직전에 적용한다."""
    output = list(headers)
    if sheet_name[:2].isdigit() and int(sheet_name[:2]) <= 15:
        if not getattr(config, "PROVENANCE_ENABLED", True):
            output = [header for header in output if header != "출처페이지"]
        if not getattr(config, "DATA_STATUS_ENABLED", True):
            output = [header for header in output if header != "데이터상태"]
    return output


def _thin_border() -> Border:
    side = Side(style="thin", color="AAAAAA")
    return Border(left=side, right=side, top=side, bottom=side)


def _header_fill() -> PatternFill:
    return PatternFill("solid", fgColor=COLOR_HEADER_BG)


def _alt_fill() -> PatternFill:
    return PatternFill("solid", fgColor=COLOR_ALT_ROW)


def _apply_header_row(ws, headers: list, row: int = 1):
    for col_idx, header in enumerate(headers, start=1):
        cell = ws.cell(row=row, column=col_idx, value=header)
        cell.font = Font(bold=True, color=COLOR_HEADER_FONT, name="맑은 고딕", size=10)
        cell.fill = _header_fill()
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        cell.border = _thin_border()


def _auto_column_widths(ws, min_width: float = 10, max_width: float = 45):
    for col in ws.columns:
        max_len = 0
        col_letter = get_column_letter(col[0].column)
        for cell in col:
            try:
                cell_len = len(str(cell.value)) if cell.value is not None else 0
                max_len = max(max_len, cell_len)
            except Exception:
                pass
        ws.column_dimensions[col_letter].width = max(min(max_len * 1.3 + 2, max_width), min_width)


def _fallback_output_path(output_path: Path) -> Path:
    timestamp = datetime.now().strftime("%Y%m%d_%H%M%S")
    suffix = output_path.suffix or ".xlsx"
    return output_path.with_name(f"{output_path.stem}_{timestamp}{suffix}")


def _sanitize_cell_value(value):
    if isinstance(value, (dict, list)):
        return json.dumps(value, ensure_ascii=False)
    if isinstance(value, str):
        return ILLEGAL_CHARACTERS_RE.sub("", value)
    return value


def _write_cell(ws, row: int, column: int, value):
    sanitized = _sanitize_cell_value(value)
    cell = ws.cell(row=row, column=column, value=sanitized)
    if isinstance(sanitized, str) and sanitized.startswith(("=", "+", "@")):
        cell.data_type = "s"
    return cell


def _write_generic_sheet(ws, headers: list[str], data: list[dict]):
    """범용 시트 작성: 헤더 기반으로 데이터를 행에 씀"""
    _apply_header_row(ws, headers)
    ws.freeze_panes = "B2"

    for row_idx, row_data in enumerate(data, start=2):
        for col_idx, header in enumerate(headers, start=1):
            val = row_data.get(header, None)
            if val is None and header in HEADER_KEY_ALIASES:
                val = row_data.get(HEADER_KEY_ALIASES[header], None)
            if isinstance(val, bool):
                val = "Y" if val else "N"
            cell = _write_cell(ws, row_idx, col_idx, val)
            cell.border = _thin_border()
            cell.alignment = Alignment(
                horizontal="center" if col_idx <= 4 else "left",
                vertical="center",
                wrap_text=True,
            )
            if row_idx % 2 == 0:
                cell.fill = _alt_fill()
            if isinstance(val, (int, float)) and not isinstance(val, bool):
                cell.number_format = "#,##0.##"

    _auto_column_widths(ws)


def _stable_output_rows(sheet_name: str, data: list[dict]) -> list[dict]:
    """병렬 완료 순서와 무관하게 같은 행 집합을 같은 순서로 기록한다."""
    data_key = next(
        (key for key, configured_name in config.SHEET_KEY_TO_NAME.items() if configured_name == sheet_name),
        sheet_name,
    )
    rows, _ = merge_rows_stably(data_key, [], data)
    return rows


def write_excel(extracted_data: dict, output_path: str | Path) -> Path:
    """
    추출·정제된 데이터 딕셔너리를 받아 16개 시트 엑셀 파일을 생성합니다.

    Args:
        extracted_data: 시트키 → list[dict] 매핑
        output_path: 저장할 엑셀 파일 경로

    Returns:
        저장된 파일의 Path 객체
    """
    output_path = Path(output_path)

    wb = openpyxl.Workbook()
    del wb[wb.sheetnames[0]]

    optional_sheets = getattr(config, "OPTIONAL_EXCEL_SHEETS", set())

    for sheet_name, headers in config.EXCEL_HEADERS.items():
        if sheet_name == "90_코드북" and not getattr(config, "CODEBOOK_SHEET_ENABLED", True):
            continue

        # config.SHEET_KEY_TO_NAME 역매핑으로 데이터 키 찾기
        data_key = None
        for key, name in config.SHEET_KEY_TO_NAME.items():
            if name == sheet_name:
                data_key = key
                break

        data = extracted_data.get(data_key, []) if data_key else []
        if not isinstance(data, list):
            data = []

        # 선택 시트(보조검수후보·병합로그 등)는 데이터가 있을 때만 생성한다.
        # 하이브리드 검수를 끄면 항상 비므로 빈 헤더 시트를 만들지 않는다.
        if sheet_name in optional_sheets and not data:
            continue

        ws = wb.create_sheet(sheet_name)
        _write_generic_sheet(
            ws,
            _headers_for_output(sheet_name, headers),
            _stable_output_rows(sheet_name, data),
        )

    try:
        wb.save(str(output_path))
        return output_path
    except (PermissionError, OSError):
        fallback_path = _fallback_output_path(output_path)
        wb.save(str(fallback_path))
        print(f"원본 경로가 잠겨 있어 대체 경로에 저장했습니다: {fallback_path}")
        return fallback_path
