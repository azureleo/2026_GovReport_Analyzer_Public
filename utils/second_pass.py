"""확정 누락만 재판독하는 독립 실행 흐름. 원본/기존 판독/모델 설정은 변경하지 않는다."""
from __future__ import annotations

from collections import Counter
from contextlib import contextmanager
from copy import deepcopy
from datetime import datetime, timezone
import hashlib
import json
import math
import os
from pathlib import Path
import re
import tempfile
import time
import uuid

from utils.source_coverage import RowIndex, canonical, compare, itemize, number, persist, read_excel, sha256_file, token, unit

VERSION = "confirmed-second-pass-v2-context"
ROOT = Path(__file__).resolve().parents[1]
DECISIONS = {"final_cell_missing", "final_budget_cell_missing"}


def read_json(path):
    return json.loads(Path(path).read_text(encoding="utf-8-sig"))


def write_json(path, data):
    path = Path(path)
    path.parent.mkdir(parents=True, exist_ok=True)
    fd, temporary = tempfile.mkstemp(prefix=path.name + ".", suffix=".tmp", dir=path.parent)
    try:
        with os.fdopen(fd, "w", encoding="utf-8") as stream:
            json.dump(data, stream, ensure_ascii=False, indent=2, default=str)
        os.replace(temporary, path)
    finally:
        if Path(temporary).exists():
            Path(temporary).unlink()


def digest(data):
    return hashlib.sha256(canonical(data).encode()).hexdigest()


def label(value):
    return token(re.sub(r"\\+n", " ", str(value or "")))


def runtime_hash():
    paths = [ROOT / "config.py", ROOT / "main.py"]
    for package in ("agents", "utils"):
        paths.extend(sorted((ROOT / package).rglob("*.py")))
    return digest({str(path.relative_to(ROOT)): sha256_file(path) for path in paths})


def require(condition, message):
    if not condition:
        raise ValueError(message)


def _read_workbook(path):
    import openpyxl
    book = openpyxl.load_workbook(path, read_only=True, data_only=False)
    sheets = {}
    try:
        for sheet in book:
            iterator = sheet.iter_rows()
            headers = [cell.value for cell in next(iterator, ())]
            require(all(isinstance(h, str) and h for h in headers) and len(set(headers)) == len(headers),
                    "빈/중복 헤더는 지원하지 않습니다: " + sheet.title)
            values = [[cell.value for cell in row] for row in iterator]
            require(not any(any(isinstance(v, str) and v.startswith("=") for v in row) for row in values),
                    "수식이 있는 기준 Excel은 재작성으로 보존할 수 없어 중단합니다.")
            sheets[sheet.title] = {"headers": headers, "values": values}
    finally:
        book.close()
    return sheets


def verify_inputs(prepared):
    require(all(sha256_file(p) == value for p, value in prepared["plan"]["input_hashes"].items()),
            "실행 중 입력 파일 변경")


def verify_schema(book):
    import config
    require(set(book) == set(config.SHEET_KEY_TO_NAME.values()), "기준 Excel의 시트 구성이 현재 스키마와 다릅니다.")
    for display in config.SHEET_KEY_TO_NAME.values():
        require(book[display]["headers"] == config.EXCEL_HEADERS[display],
                "기준 Excel/현재 쓰기 스키마 차이: " + display)
    display = config.SHEET_KEY_TO_NAME["document_meta"]
    cities = {dict(zip(book[display]["headers"], row)).get("지자체명") for row in book[display]["values"]}
    require(len(cities - {None, ""}) == 1, "기준 지자체를 유일하게 확인할 수 없습니다.")


def preserve_baseline_cell_types(baseline_path, generated_path):
    """재작성 시 값이 없는 inlineStr 셀까지 원래 자료형을 보존한다."""
    import openpyxl
    original = openpyxl.load_workbook(baseline_path, read_only=True, data_only=False)
    generated = openpyxl.load_workbook(generated_path, data_only=False)
    changed = 0
    try:
        for sheet in original:
            for row in sheet:
                for old in row:
                    if not hasattr(old, "coordinate"):
                        continue
                    new = generated[sheet.title][old.coordinate]
                    require(new.value == old.value, "기존 셀 값 변경: " + sheet.title + "!" + old.coordinate)
                    if new.data_type != old.data_type:
                        require(new.value is None and old.value is None, "기존 비어있지 않은 셀의 자료형 변경")
                        new.data_type = old.data_type
                        changed += 1
        if changed:
            generated.save(generated_path)
    finally:
        generated.close()
        original.close()
    return changed


def prepare(target_path, pdf_path, *, max_calls=4, max_seconds=300, call_timeout=90):
    """입력 바인딩/검토 근거/현재 셀 부재를 검증한다. 모델·dotenv를 호출하지 않는다."""
    target_path, pdf_path = Path(target_path).resolve(), Path(pdf_path).resolve()
    require(max_calls >= 1 and max_seconds > 0 and call_timeout > 0, "호출/시간 상한은 양수여야 합니다.")
    manifest = read_json(target_path)
    require(manifest.get("selection_status") == "source_review_confirmed_missing" and
            manifest.get("second_pass_executed") is False, "확정 대상 준비 목록 형식이 아닙니다.")
    targets = manifest.get("targets")
    require(isinstance(targets, list) and targets and manifest.get("target_count") == len(targets), "대상 수 불일치")
    ids = [row["candidate_id"] for row in targets]
    require(len(ids) == len(set(ids)), "중복 대상 ID")
    summary_path = Path(manifest["source_audit"]).resolve()
    summary = read_json(summary_path)
    provenance = summary["provenance"]
    excel_path, audit_path = Path(provenance["source_excel_path"]).resolve(), Path(provenance["source_pipeline_audit_path"]).resolve()
    journal_path = Path(summary["missing_candidates_path"]).resolve()
    input_files = [target_path, pdf_path, summary_path, excel_path, audit_path, journal_path]
    require(sha256_file(pdf_path) == summary["document_sha256"], "PDF 해시 불일치")
    require(sha256_file(excel_path) == provenance["excel_sha256"] == manifest["source_excel_sha256"], "Excel 해시 불일치")
    require(sha256_file(audit_path) == provenance["pipeline_audit_sha256"], "원본 감사 파일 해시 불일치")
    require(sha256_file(journal_path) == summary["missing_candidates_sha256"], "후보 원장 해시 불일치")
    journal = [json.loads(line) for line in journal_path.read_text(encoding="utf-8-sig").splitlines() if line.strip()]
    by_id = {row["candidate_id"]: row for row in journal}
    require(len(by_id) == len(journal), "원장 ID 중복")
    reviews = {}
    for filename in manifest.get("review_files", []):
        path = Path(filename).resolve()
        input_files.append(path)
        for row in read_json(path)["records"]:
            if row.get("candidate_id") not in ids:
                continue
            require(row.get("review_verdict") in DECISIONS, "확정 대상에 미승인/상충 검토가 있습니다.")
            candidate = row["candidate"]
            require(candidate["document_sha256"] == summary["document_sha256"] and
                    candidate["provenance"]["excel_sha256"] == provenance["excel_sha256"], "다른 실행의 검토 자료")
            current = by_id.get(row["candidate_id"])
            # 과거 검토는 감사 규칙 버전이 다를 수 있다. 설명/보조 메타데이터가 아닌
            # 사실과 위치, 원본 실행 해시는 엄격하게 일치해야 한다.
            require(current is not None and row["candidate_id"] == candidate["candidate_id"] and
                    all(candidate.get(k) == current.get(k) for k in ("sheet", "field", "identity", "raw_value", "raw_unit")) and
                    all(candidate["source"].get(k) == current["source"].get(k) for k in
                        ("object_id", "evidence_id", "page", "bbox", "row", "column")) and
                    candidate["provenance"]["pipeline_audit_sha256"] == provenance["pipeline_audit_sha256"],
                    "검토 레코드/원장 대응 불일치")
            reviews[row["candidate_id"]] = {"review_file": str(path), "review_sha256": sha256_file(path),
                "sample_id": row["sample_id"], "verdict": row["review_verdict"],
                "review_rule_version":candidate.get("schema_version"), "review_binding":"fact_position_and_baseline_hashes"}
    parent = read_json(audit_path)
    objects = parent.get("ledgers", {}).get("document_objects", [])
    objects_by_id = {row["object_id"]: row for row in objects}
    require(len(objects_by_id) == len(objects), "원본 객체 ID 중복")
    canonical_facts = {row["candidate_id"]: row for row in itemize(objects, summary["document_sha256"])[0]}
    baseline_index = RowIndex(read_excel(excel_path))
    for row in targets:
        cid = row["candidate_id"]
        require(cid in reviews and row == by_id.get(cid), "원문 검토/후보 원장 대응 불일치: " + cid)
        require(row["status"] == "missing_candidate" and row.get("second_pass_status") == "not_requested", "비누락/기처리 대상")
        require(row.get("sheet") == "financial_plan" and row.get("field") == "예산액", "첫 버전은 확정 예산 항목만 지원합니다.")
        require(number(row.get("raw_value")) is not None and unit(row.get("raw_unit"))[0] == "krw", "확정 금액/통화 단위가 불명확합니다.")
        fact = canonical_facts.get(cid)
        require(fact is not None and all(row[k] == fact[k] for k in ["identity", "raw_value", "raw_unit", "source"]), "원문 객체 사실 불일치")
        require(baseline_index.match(row)["status"] == "absent", "현재 기준 Excel에 이미 있거나 충돌하는 대상: " + cid)
    producer_manifest_path = excel_path.with_name(excel_path.stem + "_run_manifest.json")
    producer = read_json(producer_manifest_path)
    input_files.append(producer_manifest_path)
    require(producer["input_sha256"] == summary["document_sha256"] and producer["output_sha256"] == provenance["excel_sha256"], "기준 실행 매니페스트 불일치")
    provider = producer["execution_info"]["text_backend"]
    model = producer["execution_info"]["text_model"]
    require(provider == "codex" and bool(model), "첫 버전은 기준 실행의 Codex 모델만 지원합니다. 자동 대체하지 않습니다.")
    first_seconds = (datetime.fromisoformat(producer["completed_at"]) - datetime.fromisoformat(producer["created_at"])).total_seconds()
    require(first_seconds > 0, "기준 실행 시간 불명확")
    budget = min(float(max_seconds), first_seconds * .20)
    import fitz
    groups = {}
    with fitz.open(pdf_path) as document:
        for row in targets:
            source = row["source"]
            obj = objects_by_id.get(source["object_id"])
            require(obj is not None and obj["page_number"] == source["page"], "원문 객체 페이지 불일치")
            require(type(source["page"]) is int and 1 <= source["page"] <= len(document), "원문 페이지 범위 오류")
            page = document[source["page"] - 1]
            bounds = page.rect * page.derotation_matrix
            bbox = source.get("bbox")
            require(isinstance(bbox, list) and len(bbox) == 4 and all(type(v) in (int, float) and math.isfinite(v) for v in bbox), "표 좌표 없음/오류")
            box = fitz.Rect(bbox)
            require(not box.is_empty and bounds.contains(box), "표 좌표가 페이지 범위를 벗어납니다.")
            key = digest([summary["document_sha256"], source["object_id"], source["page"], bbox])[:24]
            group = groups.setdefault(key, {"group_id": key, "object_id": source["object_id"], "page": source["page"],
                "bbox": bbox, "context_bbox": bbox[:], "caption": source.get("caption", ""), "candidate_ids": []})
            group["candidate_ids"].append(row["candidate_id"])
            for evidence in source.get("unit_evidence", []):
                if evidence.get("page") == source["page"] and evidence.get("bbox"):
                    unit_box = fitz.Rect(evidence["bbox"])
                    require(bounds.contains(unit_box), "명시 단위 좌표가 페이지 범위를 벗어납니다.")
                    group["context_bbox"] = list(fitz.Rect(group["context_bbox"]) | unit_box)
    hashes = {str(p): sha256_file(p) for p in input_files}
    public = {"version": VERSION, "input_hashes": hashes, "provider": provider, "model": model,
        "reasoning_override": False, "target_count": len(targets), "group_count": len(groups),
        "groups": list(groups.values()), "budget": {"max_calls": int(max_calls), "max_seconds": budget,
            "per_call_seconds": float(call_timeout), "baseline_manifest_seconds": first_seconds, "fraction_ceiling": .20,
            "scope": "cumulative_reading_loop_includes_render_response_checks_excludes_downstream_finalization"},
        "selection_evidence": {cid: reviews[cid] for cid in ids}, "runtime_sha256": runtime_hash(),
        "expected_values_sent_to_model": False, "automatic_retry": False, "model_accuracy_evaluated": False}
    # 운영 상한 확대는 같은 체크포인트에서 명시적으로 가능하되 지문과 누적 사용량은 유지.
    public["fingerprint"] = digest({k:v for k,v in public.items() if k != "budget"})
    baseline = _read_workbook(excel_path)
    verify_schema(baseline)  # 호출 비용이 발생하기 전에 저장 호환성을 점검한다.
    return {"plan": public, "targets": {r["candidate_id"]: r for r in targets}, "objects": objects,
        "parent_audit": parent, "pdf_path": pdf_path, "excel_path": excel_path,
        "baseline_book": baseline}


def save_plan(prepared, output_dir):
    directory = Path(output_dir).resolve()
    require(directory not in {prepared["excel_path"].parent, prepared["pdf_path"].parent}, "전용 새 출력 폴더를 지정하세요.")
    path = directory / "plan.json"
    if directory.exists() and not path.exists():
        require(not any(directory.iterdir()), "이미 다른 파일이 있는 출력 폴더입니다.")
    if path.exists():
        require(read_json(path)["fingerprint"] == prepared["plan"]["fingerprint"], "다른 입력/구현의 계획이 존재합니다.")
    directory.mkdir(parents=True, exist_ok=True)
    write_json(path, prepared["plan"])
    return directory


def prompt_for(prepared, group):
    requests = []
    for cid in group["candidate_ids"]:
        c = prepared["targets"][cid]
        requests.append({"candidate_id": cid, "사업명": re.sub(r"\\+n", " ", c["identity"]["name"]),
                         "연도": c["identity"]["year"], "재원구분": c["identity"]["context"].get("재원구분")})
    return ("첨부한 원본 PDF 표 이미지만 판독하세요. 외부 조회·파일 수정·다른 실행/검토 자료 열람을 금지합니다. "
        "이름/연도/재원에 대응하는 예산액과 원문 단위를 그대로 읽고 합계와 연간 값, 국비/시비를 구분하세요. "
        "이전 결과나 예상 금액은 제공하지 않습니다. 보이지 않으면 추정하지 말고 unreadable로 반환하세요. "
        "계획구분/부문은 추론하지 마세요. 각 candidate_id당 결과 하나만 반환하세요.\n"
        + json.dumps({"object_id": group["object_id"], "출처페이지": group["page"], "caption": group["caption"],
                      "requests": requests}, ensure_ascii=False)
        + '\nJSON만 반환: {"results":[{"candidate_id":"...","status":"read|unreadable|not_found",'
          '"사업명":"원문","연도":2024,"재원구분":"계|국비|시비|구비|기타",'
          '"예산액":123,"예산단위":"원문 단위","출처페이지":1,"객체ID":"...",'
          '"근거ID":"원문 객체 근거 ID 또는 객체 ID","evidence_text":"대응하는 원문 행/머리글 인용"}]}')


def render_group(prepared, group, directory):
    import fitz
    path = directory / "images" / (group["group_id"] + ".png")
    path.parent.mkdir(exist_ok=True)
    with fitz.open(prepared["pdf_path"]) as doc:
        page = doc[group["page"] - 1]
        bounds = page.rect * page.derotation_matrix
        box = fitz.Rect(group["context_bbox"])
        # 표 위 단위/제목과 왼쪽 병합 사업명까지 포함한다.
        box = fitz.Rect(box.x0 - 18, box.y0 - 72, box.x1 + 18, box.y1 + 24) & bounds
        page.get_pixmap(matrix=fitz.Matrix(2, 2), clip=box * page.rotation_matrix).save(path)
    return path


def validate_response(prepared, group, text):
    """모델 금액을 원문 사실·검토와 대조한다. 원문 예상 금액은 모델 입력에 사용하지 않는다."""
    require(isinstance(text, str) and len(text) <= 2_000_000, "응답 형식/크기 오류")
    parsed = json.loads(text)
    require(isinstance(parsed, dict) and isinstance(parsed.get("results"), list), "results 배열 누락")
    buckets = {}
    for row in parsed["results"]:
        if isinstance(row, dict) and row.get("candidate_id") in group["candidate_ids"]:
            buckets.setdefault(row["candidate_id"], []).append(row)
    result = {}
    for cid in group["candidate_ids"]:
        c = prepared["targets"][cid]
        responses = buckets.get(cid, [])
        record = {"candidate_id": cid, "status": "held", "reason": "not_returned"}
        result[cid] = record
        if not responses:
            continue
        if len({canonical(r) for r in responses}) != 1:
            record["reason"], record["responses"] = "conflicting_duplicate_response", responses
            continue
        row = responses[0]
        record["response"] = row
        if row.get("status") != "read":
            record["reason"] = "model_" + str(row.get("status", "unknown"))
            continue
        context = c["identity"]["context"]
        funding = lambda v: {"계":"합계","총계":"합계"}.get(label(v), label(v))
        checks = [(label(row.get("사업명")) == label(c["identity"]["name"]), "business_mismatch"),
            (type(row.get("연도")) is int and row["연도"] == c["identity"]["year"], "year_mismatch"),
            (funding(row.get("재원구분")) == funding(context.get("재원구분")), "funding_mismatch"),
            (type(row.get("출처페이지")) is int and row["출처페이지"] == group["page"], "page_mismatch"),
            (row.get("객체ID") == group["object_id"] and row.get("근거ID") in c["source"]["aliases"], "evidence_mismatch"),
            (number(row.get("예산액")) is not None, "invalid_amount"),
            (bool(row.get("예산단위")) and unit(row["예산단위"]) == unit(c["raw_unit"]), "unit_or_scale_mismatch")]
        invalid = next((reason for ok, reason in checks if not ok), None)
        if invalid:
            record["reason"] = invalid
            continue
        if number(row["예산액"]) != number(c["raw_value"]):
            record["reason"] = "source_value_conflict"
            continue
        quote = str(row.get("evidence_text", "")).replace(",", "")
        amount = format(number(row["예산액"]), "f").rstrip("0").rstrip(".") if "." in str(number(row["예산액"])) else str(number(row["예산액"]))
        if not re.search(r"(?<![\d.])" + re.escape(amount) + r"(?![\d.])", quote):
            record["reason"] = "amount_quote_missing"
            continue
        record.update(status="validated_reading", reason="source_fact_and_review_agree")
    return result


@contextmanager
def run_lock(directory):
    path = directory / ".run.lock"
    fd = os.open(path, os.O_CREAT | os.O_EXCL | os.O_WRONLY)
    try:
        os.write(fd, str(os.getpid()).encode())
        os.close(fd)
        yield
    finally:
        path.unlink(missing_ok=True)


def _downstream(prepared, records, directory, iteration, *, fixture=False):
    import config
    from agents.organizer_agent import OrganizerAgent
    from agents.excel_agent import ExcelAgent
    from utils.document_objects import DocumentObject
    from utils.integration_audit import persist_audit, source_validation, object_inventory
    from utils.pdf_reader import PDFContent, PageContent
    from utils.semantic_routing import validate_and_reclassify
    from utils.second_pass_context import ContextResolver, summary as context_summary
    baseline, data = prepared["baseline_book"], {}
    for key, display in config.SHEET_KEY_TO_NAME.items():
        require(display in baseline and baseline[display]["headers"] == config.EXCEL_HEADERS[display], "기준 Excel/현재 쓰기 스키마 차이: " + display)
        data[key] = [dict(zip(baseline[display]["headers"], values)) for values in baseline[display]["values"]]
    municipalities = {row.get("지자체명") for row in data.get("document_meta", []) if row.get("지자체명")}
    require(len(municipalities) == 1, "기준 지자체를 유일하게 확인할 수 없습니다.")
    city = next(iter(municipalities))
    data["municipality_name"] = city
    resolver = ContextResolver(prepared["objects"],data,city)
    additions = {"municipality_name": city, "financial_plan": []}
    order = []
    for cid, record in records.items():
        if record["status"] != "validated_reading":
            continue
        c, response = prepared["targets"][cid], record["response"]
        row = {"지자체명": city, "사업명": c["identity"]["name"], "연도": response["연도"],
            "재원구분": {"계":"합계","총계":"합계"}.get(response["재원구분"], response["재원구분"]),
            "예산액": float(number(response["예산액"])), "예산단위": response["예산단위"],
            "출처페이지": c["source"]["page"], "객체ID": c["source"]["object_id"],
            "근거ID": c["source"]["evidence_id"], "추출원천": "second_pass", "추출엔진": "fixture_not_model" if fixture else "codex_cli",
            "_second_pass_candidate_id": cid,
            "검수상태": "2차_근거검증", "출처계보": [{"candidate_id": cid, "object_id": c["source"]["object_id"],
                "evidence_id": c["source"]["evidence_id"], "page": c["source"]["page"], "bbox": c["source"]["bbox"]}]}
        for key in ("계획구분", "부문"):
            if c["identity"]["context"].get(key):
                row[key] = c["identity"]["context"][key]
        # 같은 사업·지자체·출처 페이지에 명시된 유일한 문맥만 재사용한다.
        for key in ("계획구분", "부문"):
            values = {r.get(key) for r in data["financial_plan"] if label(r.get("사업명")) == label(row["사업명"])
                      and r.get("지자체명") == city and r.get(key)
                      and str(c["source"]["page"]) in re.findall(r"\d+", str(r.get("출처페이지", "")))}
            if not row.get(key) and len(values) == 1:
                row[key] = next(iter(values))
        resolution = resolver.resolve(c)
        for key, decision in resolution["fields"].items():
            previous = row.get(key)
            if previous:
                # 최초 연결에서 이미 확보한 값은 보존하고 새 충돌은 별도로 표시한다.
                decision["previous_value"] = previous
                if decision["status"] == "held":
                    decision.update(status="preserved_with_conflict",value=previous)
                else:
                    decision.update(status="preserved",value=previous)
            elif decision["status"] == "resolved":
                row[key] = decision["value"]
            if row.get(key):
                for proof in decision["evidence"]:
                    if proof["value"] == decision["value"]:
                        row["출처계보"].append({"context_field":key,"context_value":row[key],**proof})
        record["context_resolution"] = resolution
        record["unfilled_context"] = [key for key in ("계획구분", "부문") if not row.get(key)]
        additions["financial_plan"].append(row)
        order.append(cid)
    import fitz
    pages = sorted({c["source"]["page"] for c in prepared["targets"].values()})
    with fitz.open(prepared["pdf_path"]) as document:
        content = PDFContent(len(document), [PageContent(p, document[p-1].get_text(), [], []) for p in pages], "")
    clean = OrganizerAgent().organize(additions)
    objects = [DocumentObject.from_dict(o) for o in prepared["objects"] if o["page_number"] in pages]
    semantic = validate_and_reclassify(clean, content, auto_reclassify=False, document_objects=objects)
    checks = source_validation(clean, content)
    cleaned_rows = clean.get("financial_plan", [])
    require(len(checks) == len(cleaned_rows), "후단 원문 대조 행 수 불일치")
    by_id = {}
    for row, check in zip(cleaned_rows, checks):
        cid = row.get("_second_pass_candidate_id")
        require(cid in order and cid not in by_id, "정제 후 대상 ID 유실/중복")
        by_id[cid] = (row, check)
    merged = []
    for cid in order:
        if cid not in by_id:
            records[cid].update(status="held", reason="downstream_row_removed")
            continue
        row, check = by_id[cid]
        c = prepared["targets"][cid]
        if check["status"] != "source_present" or RowIndex({"financial_plan":[row]}).match(c)["status"] != "covered":
            records[cid].update(status="held", reason="downstream_source_or_cleaning_review")
            continue
        current = RowIndex(data).match(c)
        if current["status"] == "covered":
            records[cid].update(status="duplicate", reason="already_merged")
        elif current["status"] != "absent":
            records[cid].update(status="held", reason="existing_value_or_context_conflict")
        else:
            data["financial_plan"].append(row)
            records[cid].update(status="merged", reason="validated_new_fact")
            merged.append(cid)
    data["document_objects"] = prepared["objects"]
    for issue in semantic.validation_issues(city):
        data["validation_report"].append({"지자체명":city,"대상시트":"financial_plan",
            "항목":issue.get("항목"),"사유":issue.get("문제내용"),"출처페이지":",".join(map(str,pages))})
    for cid, record in records.items():
        if record["status"] not in {"merged", "duplicate"}:
            data["validation_report"].append({"지자체명":city,"대상시트":"financial_plan","항목":"2차 판독 보류",
                "사유":cid + ": " + record["reason"],"출처페이지":prepared["targets"][cid]["source"]["page"],
                "감사파일":str(directory / f"outcome_{iteration:03}.json")})
        elif record.get("context_resolution"):
            for field, decision in record["context_resolution"]["fields"].items():
                if decision["status"] in {"held","missing","preserved_with_conflict"}:
                    data["validation_report"].append({"지자체명":city,"대상시트":"financial_plan",
                        "항목":"2차 예산 분류 검토: " + field,
                        "사유":cid + ": " + decision["reason"],"출처페이지":prepared["targets"][cid]["source"]["page"],
                        "감사파일":str(directory / f"recovered_{iteration:03}_context_resolution.json")})
    temporary = directory / f"pending_{iteration:03}.xlsx"
    final = directory / f"recovered_{iteration:03}.xlsx"
    require(not temporary.exists() and not final.exists(), "출력 Excel을 덮어쓰지 않습니다.")
    verify_inputs(prepared)
    ExcelAgent().write(data, temporary)
    restored_types = preserve_baseline_cell_types(prepared["excel_path"], temporary)
    saved_book = _read_workbook(temporary)
    for name, sheet in baseline.items():
        require(saved_book[name]["headers"] == sheet["headers"] and
                saved_book[name]["values"][:len(sheet["values"])] == sheet["values"], "기존 셀 보존 실패: " + name)
    saved_index = RowIndex(read_excel(temporary))
    for cid in merged:
        result = saved_index.match(prepared["targets"][cid])
        require(result["status"] == "covered", "최종 셀 반영 실패: " + cid)
        saved_values = saved_book[config.SHEET_KEY_TO_NAME["financial_plan"]]
        for match in result["matches"]:
            saved_row = dict(zip(saved_values["headers"],saved_values["values"][match["row"]-1]))
            for field in ("계획구분","부문"):
                require(saved_row.get(field) == by_id[cid][0].get(field), "분류 셀 반영 실패: " + cid + " " + field)
        records[cid].update(status="saved_verified", saved_cells=result["matches"])
        records[cid]["classification_saved_verified"] = True
    temporary.rename(final)
    result = {"excel_path":str(final),"excel_sha256":sha256_file(final),"pipeline_audit_path":None,
              "coverage_summary_path":None,"new_rows":len(merged),"baseline_cells_preserved":True,"records":records,
              "baseline_empty_cell_types_restored":restored_types,
              "context_summary":context_summary(records),"context_audit_path":None,
              "audit_saved":False,"warnings":[]}
    try:
        context_path = final.with_name(final.stem + "_context_resolution.json")
        write_json(context_path,{"version":VERSION,"summary":result["context_summary"],
            "records":{cid:{"unfilled_context":r.get("unfilled_context"),"resolution":r.get("context_resolution"),
                            "saved_cells":r.get("saved_cells")} for cid,r in records.items() if r.get("context_resolution")}})
        result["context_audit_path"] = str(context_path)
        journal, coverage = compare(prepared["objects"], sha256_file(prepared["pdf_path"]), read_excel(final),
            provenance={"source":"second_pass_incremental","source_excel_path":str(final),
                        "excel_sha256":sha256_file(final),"baseline_excel_sha256":sha256_file(prepared["excel_path"])})
        coverage_saved = persist(final.with_suffix(""), journal, coverage)
        result["coverage_summary_path"] = coverage_saved["coverage_summary_path"]
        new_data = {"municipality_name":city,"financial_plan":[by_id[cid][0] for cid in merged]}
        context = {"input_sha256":sha256_file(prepared["pdf_path"]), "source_coverage":coverage_saved,
            "second_pass":{"plan_fingerprint":prepared["plan"]["fingerprint"],"baseline_excel":str(prepared["excel_path"]),
                           "fixture":fixture,"original_raw_cleaned_snapshots":"unknown","semantic":semantic.metrics(),"records":records,
                           "source_validation":{cid:check for cid,(_,check) in by_id.items()},
                           "context_summary":result["context_summary"],"context_audit_path":str(context_path),
                           "response_reuse":prepared.get("response_reuse"),
                           "new_row_object_inventory":object_inventory(objects,new_data,{})}}
        # 기존 내부 원장도 복사해 보존한다. 현재 실행에서 덮어쓰는 것은 객체 목록뿐이다.
        for key, value in prepared["parent_audit"].get("ledgers", {}).items():
            if key in config.INTERNAL_OBJECT_KEYS and key not in data:
                data[key] = deepcopy(value)
        audit = persist_audit(final, prepared["pdf_path"], data, context)
        result.update(pipeline_audit_path=str(audit), audit_saved=True)
    except Exception as error:
        result["warnings"].append("Excel 저장·셀 검증 완료. 별도 감사 저장 실패: " + str(error))
    return result


def execute(prepared, output_dir, *, resume=False, reader=None, fixture=False, clock=time.monotonic, progress=None):
    execution_started = time.perf_counter()
    directory = save_plan(prepared, output_dir)
    require(not (directory / "replay_source.json").exists(), "응답 재평가 전용 폴더입니다. 실제 호출 상태로 전환하지 않습니다.")
    plan = prepared["plan"]
    with run_lock(directory):
        path = directory / "state.json"
        if path.exists():
            require(resume, "상태 파일이 존재합니다. 중복 호출 방지를 위해 --resume을 지정하세요.")
            state = read_json(path)
            require(state["fingerprint"] == plan["fingerprint"] and state["fixture"] == fixture, "체크포인트 입력/구현/모드 불일치")
        else:
            state = {"fingerprint":plan["fingerprint"],"fixture":fixture,"attempts":{},"spent_seconds":0.,"exports":[]}
        require(reader is not None, "판독 경계가 제공되지 않았습니다.")
        verify_inputs(prepared)
        prior_attempts = len(state["attempts"])
        history = state.setdefault("budget_history", [])
        if not history or history[-1]["budget"] != plan["budget"]:
            history.append({"at":datetime.now(timezone.utc).isoformat(),"budget":deepcopy(plan["budget"])})
        write_json(path, state)
        started = clock()
        records, reused = {}, 0
        budget = plan["budget"]
        emit = progress or (lambda message: None)
        try:
            for position, group in enumerate(plan["groups"], 1):
                gid = group["group_id"]
                prefix = f"[2차 판독 {position}/{plan['group_count']}] p{group['page']} {gid}: "
                record = state["attempts"].get(gid)
                response_path = directory / "responses" / (gid + ".json")
                if record and record.get("status") == "saved":
                    require(response_path.is_file() and sha256_file(response_path) == record["sha256"], "응답 체크포인트 손상")
                    response = read_json(response_path)
                    require(response["fingerprint"] == plan["fingerprint"] and response["group_id"] == gid, "응답 체크포인트 불일치")
                    text = response["text"]
                    reused += 1
                    emit(prefix + "저장 응답 재사용")
                elif record:
                    emit(prefix + "기시도 표 보류; 자동 재호출 없음")
                    for cid in group["candidate_ids"]:
                        records[cid] = {"candidate_id":cid,"status":"held","reason":record.get("reason", "attempt_interrupted_no_auto_retry")}
                    continue
                else:
                    remaining = budget["max_seconds"] - state["spent_seconds"] - (clock() - started)
                    if len(state["attempts"]) >= budget["max_calls"] or remaining <= .1:
                        emit(prefix + "호출/시간 상한으로 미실행")
                        for cid in group["candidate_ids"]:
                            records[cid] = {"candidate_id":cid,"status":"deferred","reason":"budget_exhausted"}
                        continue
                    image = render_group(prepared, group, directory)
                    remaining = budget["max_seconds"] - state["spent_seconds"] - (clock() - started)
                    if remaining <= .1:
                        emit(prefix + "이미지 준비 후 시간 상한 도달; 미실행")
                        for cid in group["candidate_ids"]:
                            records[cid] = {"candidate_id":cid,"status":"deferred","reason":"budget_exhausted"}
                        continue
                    prompt = prompt_for(prepared, group)
                    timeout = min(budget["per_call_seconds"], remaining)
                    state["attempts"][gid] = {"status":"attempting","attempt_id":uuid.uuid4().hex,
                        "started_at":datetime.now(timezone.utc).isoformat(), "timeout_seconds":timeout,
                        "image_sha256":sha256_file(image), "prompt_sha256":hashlib.sha256(prompt.encode()).hexdigest()}
                    write_json(path, state)  # 호출 전에 횟수를 확정. 중단돼도 자동 재호출하지 않는다.
                    emit(prefix + ("fixture 경계 실행" if fixture else "실제 판독 호출") +
                         f" (제한 {min(budget['per_call_seconds'], remaining):.1f}초)")
                    try:
                        text = reader(prompt, image, timeout, group)
                        write_json(response_path, {"fingerprint":plan["fingerprint"],"group_id":gid,"text":text})
                        state["attempts"][gid].update(status="saved",sha256=sha256_file(response_path))
                    except Exception as error:
                        emit(prefix + "실패; 나머지 성공 응답 보존: " + type(error).__name__)
                        state["attempts"][gid].update(status="failed",reason=type(error).__name__ + ": " + str(error)[-1600:])
                        write_json(path, state)
                        for cid in group["candidate_ids"]:
                            records[cid] = {"candidate_id":cid,"status":"held","reason":state["attempts"][gid]["reason"]}
                        continue
                    write_json(path, state)
                try:
                    records.update(validate_response(prepared, group, text))
                except (ValueError, TypeError, KeyError) as error:
                    for cid in group["candidate_ids"]:
                        records[cid] = {"candidate_id":cid,"status":"held","reason":"invalid_response: " + str(error)}
        finally:
            state["spent_seconds"] += max(0., clock() - started)
            write_json(path, state)
        verify_inputs(prepared)
        iteration = int(state.get("export_sequence", 0)) + 1
        # 최종 저장 중 중단돼도 다음 재개에서 기존 산출물을 덮어쓰지 않는다.
        state["export_sequence"] = iteration
        write_json(path, state)
        emit("[2차 후단] 신규 후보 정제·의미검증·원문 대조·병합·Excel 저장/재조회")
        result = _downstream(prepared, records, directory, iteration, fixture=fixture)
        result.update(version=VERSION, fixture=fixture, model_accuracy_evaluated=False,
            model_requests=0 if fixture else len(state["attempts"]) - prior_attempts,
            cumulative_model_requests=0 if fixture else len(state["attempts"]),
            simulated_requests=len(state["attempts"]) if fixture else 0, reused_responses=reused,
            budget=budget, spent_reading_seconds=state["spent_seconds"],
            elapsed_execution_seconds_before_outcome=time.perf_counter() - execution_started,
            original_files_modified=False, provider=plan["provider"], model=plan["model"])
        verify_inputs(prepared)
        result["status_counts"] = dict(Counter(r["status"] for r in records.values()))
        result["exit_code"] = 0 if result["audit_saved"] and all(r["status"] in {"saved_verified","duplicate"} for r in records.values()) else 2
        outcome = directory / f"outcome_{iteration:03}.json"
        write_json(outcome, result)
        state["exports"].append({"outcome":str(outcome),"excel_path":result["excel_path"]})
        write_json(path, state)
        return result


def replay_saved(prepared, source_run, output_dir, *, progress=None):
    """입력과 실제 전송 계약을 검증한 과거 응답만 새 후단으로 재평가한다. 호출 없음."""
    source_run = Path(source_run).resolve()
    require(not Path(output_dir).resolve().is_relative_to(source_run), "기존 실행 폴더 안에는 저장하지 않습니다. 별도 출력 폴더를 지정하세요.")
    source_plan, source_state = read_json(source_run/"plan.json"), read_json(source_run/"state.json")
    plan = prepared["plan"]
    require(source_plan.get("version") in {"confirmed-second-pass-v1","confirmed-second-pass-v2-context"}, "지원하지 않는 응답 원본 버전")
    require(source_plan["fingerprint"] == digest({k:v for k,v in source_plan.items() if k not in {"budget","fingerprint"}}), "원본 계획 지문 검증 실패")
    require(source_plan["input_hashes"] == plan["input_hashes"], "재사용 입력 바인딩 불일치")
    require(source_state["fingerprint"] == source_plan["fingerprint"], "원본 상태/계획 지문 불일치")
    require(source_plan["provider"] == plan["provider"] and source_plan["model"] == plan["model"] and
            source_plan.get("reasoning_override") is False, "원본 모델 계약 불일치")
    require(source_plan["groups"] == plan["groups"], "원본 판독 표/대상/좌표 계약 불일치")
    verify_inputs(prepared)
    imported = {str(source_run/filename):sha256_file(source_run/filename) for filename in ("plan.json","state.json")}
    records, reused = {}, 0
    emit = progress or (lambda message:None)
    started = time.perf_counter()
    for group in plan["groups"]:
        gid = group["group_id"]
        attempt = source_state["attempts"].get(gid)
        if not attempt or attempt.get("status") != "saved":
            for cid in group["candidate_ids"]:
                records[cid] = {"candidate_id":cid,"status":"held" if attempt else "deferred",
                    "reason":"source_attempt_unusable_no_automatic_call" if attempt else "source_not_attempted_no_automatic_call"}
            continue
        path, image = source_run/"responses"/(gid+".json"), source_run/"images"/(gid+".png")
        require(sha256_file(path) == attempt["sha256"] and sha256_file(image) == attempt["image_sha256"], "원본 응답/이미지 해시 불일치")
        require(hashlib.sha256(prompt_for(prepared,group).encode()).hexdigest() == attempt["prompt_sha256"], "판독 프롬프트가 변경돼 응답을 자동 재사용할 수 없습니다.")
        response = read_json(path)
        require(response["fingerprint"] == source_plan["fingerprint"] and response["group_id"] == gid, "응답/원본 계획 불일치")
        imported[str(path)], imported[str(image)] = sha256_file(path), sha256_file(image)
        try:
            checked = validate_response(prepared,group,response["text"])
        except (ValueError,TypeError,KeyError) as error:
            checked = {cid:{"candidate_id":cid,"status":"held","reason":"invalid_saved_response: "+str(error)} for cid in group["candidate_ids"]}
        for record in checked.values():
            record["reading_reused_from"] = {"path":str(path),"sha256":imported[str(path)],"source_fingerprint":source_plan["fingerprint"]}
        records.update(checked)
        reused += 1
        emit(f"[응답 재평가] p{group['page']} {gid}: 해시·프롬프트 일치, 기존 응답 사용")
    directory = save_plan(prepared,output_dir)
    with run_lock(directory):
        require(not (directory/"replay_source.json").exists(), "재평가 출력이 이미 있습니다. 새 폴더를 지정하세요.")
        binding = {"source_run":str(source_run),"source_hashes":imported,"source_fingerprint":source_plan["fingerprint"],
                   "current_fingerprint":plan["fingerprint"],"new_model_requests":0,"reason":"explicit_saved_response_downstream_replay"}
        write_json(directory/"replay_source.json",binding)
        prepared = {**prepared,"response_reuse":binding}
        result = _downstream(prepared,records,directory,1,fixture=source_state["fixture"])
        require(all(sha256_file(p) == value for p,value in imported.items()), "재평가 중 원본 실행 파일 변경")
        verify_inputs(prepared)
        result.update(version=VERSION,execution_mode="saved_response_replay",fixture=source_state["fixture"],
            model_requests=0,cumulative_model_requests=0,reused_responses=reused,model_accuracy_evaluated=False,
            original_files_modified=False,provider=plan["provider"],model=plan["model"],response_reuse=binding,
            elapsed_execution_seconds_before_outcome=time.perf_counter()-started)
        result["status_counts"] = dict(Counter(r["status"] for r in records.values()))
        result["exit_code"] = 0 if result["audit_saved"] and all(r["status"] in {"saved_verified","duplicate"} for r in records.values()) else 2
        write_json(directory/"outcome_001.json",result)
        return result
