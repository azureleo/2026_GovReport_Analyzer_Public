"""원문 → 추출 → 정제 → 저장 셀의 제한적 대조. 모델/추출기를 호출하지 않는다.

정답 채점기가 아니다. 명시적인 사업·과제 식별자와 연도·수치·단위가 있는
표만 비교하며, 해석할 수 없는 표를 누락으로 단정하지 않는다.
"""
from __future__ import annotations

from collections import Counter, defaultdict
from copy import deepcopy
from decimal import Decimal, InvalidOperation
import hashlib
import json
import os
from pathlib import Path
import re
import tempfile
import time
import unicodedata

VERSION = "source-coverage-v3"
# 기존 사실은 동일 ID로 규칙 전후를 대조한다. 분류 버전과 사실 ID 버전을 분리.
FACT_ID_VERSION = "source-coverage-v1"
SHEETS = {
    "quantitative_reductions": "10_정량감축량",
    "financial_plan": "11_재정투자계획",
    "foundation_measures": "12_대응기반강화",
}
NAME_FIELDS = ("사업명", "과제명", "세부사업명", "세부과제명")
CONTEXT_FIELDS = ("계획구분", "대응기반영역", "부문", "재원구분", "모니터링인자")
VALUE_FIELDS = ("활동량", "예상감축량", "예산액", "과제명")
ABSENT = {"", "-", "—", "–", "미정", "미표기", "없음", "n/a"}


def canonical(value):
    return json.dumps(value, ensure_ascii=False, sort_keys=True, separators=(",", ":"), default=str)


def sha256_file(path):
    h = hashlib.sha256()
    with Path(path).open("rb") as stream:
        for block in iter(lambda: stream.read(1024 * 1024), b""):
            h.update(block)
    return h.hexdigest()


def token(value):
    return re.sub(r"\s+", "", unicodedata.normalize("NFKC", str(value or "").replace("\\n", " "))).casefold()


def number(value):
    if isinstance(value, bool) or value is None:
        return None
    text = unicodedata.normalize("NFKC", str(value)).strip().replace(",", "").replace("−", "-")
    if not re.fullmatch(r"[+-]?(?:\d+(?:\.\d*)?|\.\d+)", text):
        return None
    try:
        result = Decimal(text)
        return result if result.is_finite() else None
    except InvalidOperation:
        return None


def unit(value):
    """차원과 배율. CO2와 CO2eq, 화폐와 활동량은 서로 바꾸지 않는다."""
    text = token(value).replace("㎡", "m2").replace("m²", "m2").replace("제곱미터", "m2")
    text = text.replace("co₂", "co2").replace("co2-eq", "co2eq").replace("co2e", "co2eq") if "co2eq" not in text else text
    aliases = {
        "m2": ("area", "1"), "ha": ("area", "10000"),
        "kw": ("power", "1"), "mw": ("power", "1000"),
        "그루": ("trees", "1"), "천그루": ("trees", "1000"),
        "원": ("krw", "1"), "천원": ("krw", "1000"), "백만원": ("krw", "1000000"),
        "억원": ("krw", "100000000"),
        "tco2eq": ("co2eq", "1"), "톤co2eq": ("co2eq", "1"),
        "천톤co2eq": ("co2eq", "1000"), "천tco2eq": ("co2eq", "1000"),
        "tco2": ("co2", "1"), "톤co2": ("co2", "1"), "천톤co2": ("co2", "1000"),
    }
    dimension, factor = aliases.get(text, ("literal:" + text, "1"))
    return dimension, Decimal(factor)


def pages(value):
    if isinstance(value, (list, tuple, set)):
        return sorted({n for v in value for n in pages(v)})
    return sorted({int(v) for v in re.findall(r"\d+", str(value or "")) if int(v) > 0})


def ids(row):
    result = set()
    for key in ("object_id", "evidence_id", "객체ID", "근거ID", "source_object_ids",
                "source_evidence_ids", "original_object_ids", "original_evidence_ids",
                "original_object_id", "original_evidence_id", "alias_object_ids", "근거ID목록"):
        value = row.get(key)
        if isinstance(value, (list, tuple, set)):
            result.update(str(v) for v in value if v)
        elif value:
            result.add(str(value))
    return result


def snapshot(data):
    """감사 대상 세 시트와 계층 연결용 연차별 행을 보존. 업무 입력은 변경하지 않는다."""
    return {key: deepcopy(data.get(key, [])) for key in (*SHEETS, "annual_implementation")}


def source_snapshot(objects):
    """모델/OCR 이전 원문 객체의 비교 범위만 보존. 이미지 본체/HTML은 복제하지 않는다."""
    selected = []
    for original in objects:
        obj = original.to_dict() if hasattr(original, "to_dict") else original
        is_unit_text = obj.get("object_type") == "text" and bool(re.search(r"단위\s*[:：]", obj.get("text") or ""))
        if not (obj.get("rows") or is_unit_text or obj.get("object_type") in {"image", "chart", "graph"} or
                int(obj.get("page_number") or 0) <= 2):
            continue
        copied = {key: deepcopy(obj.get(key)) for key in (
            "object_id", "object_type", "page_number", "sequence", "rows", "caption", "number",
            "section", "nearby_text", "bbox")}
        copied["text"] = obj.get("text", "") if is_unit_text or obj.get("rows") or int(obj.get("page_number") or 0) <= 2 else ""
        copied["metadata"] = {key: deepcopy(value) for key, value in (obj.get("metadata") or {}).items()
            if key in {"engine", "evidence_id", "physical_object_id", "original_object_ids", "source_object_ids",
                       "source_evidence_ids", "original_evidence_ids", "alias_object_ids", "index_reference", "is_index_reference",
                       "header_row_count", "table_parse_error", "parse_errors", "unit", "source_kind", "render_proxy", *CONTEXT_FIELDS}}
        selected.append(copied)
    return selected


def _header_field(header):
    h = token(str(header).split(" / ")[-1])
    aliases = {
        "사업명": "사업명", "세부사업명": "사업명", "과제명": "과제명", "세부과제명": "과제명",
        "관리번호": "관리번호", "사업코드": "관리번호", "과제id": "과제ID",
        "연도": "연도", "년": "연도", "활동량": "활동량", "목표물량": "활동량",
        "예상감축량": "예상감축량", "감축량": "예상감축량", "온실가스감축량": "예상감축량",
        "예산액": "예산액", "예산": "예산액", "투자액": "예산액", "사업비": "예산액",
        "단위": "단위", "활동단위": "활동단위", "예산단위": "예산단위",
        "부서": "주관부서", "담당부서": "주관부서", "주관부서": "주관부서",
        **{token(v): v for v in CONTEXT_FIELDS},
    }
    # 단위가 괄호로 명시된 단일 머리글만 허용. 복수 상위 머리글의 의미는 추정하지 않는다.
    return aliases.get(re.sub(r"\([^()]*\)", "", h), "")


def _embedded_unit(text):
    matches = re.findall(r"단위\s*[:：]\s*([^()\n]+)", str(text))
    if not matches:
        matches = re.findall(r"\(([^()\n]+)\)", str(text))
    units = set()
    for value in matches:
        cleaned = value.strip().rstrip(")").strip()
        if unit(cleaned)[0].split(":")[0] in {"area", "power", "trees", "krw", "co2", "co2eq"} or token(cleaned) in {"개소", "세대", "대", "건", "개", "%", "kwh", "mwh", "tj"}:
            units.add(cleaned)
    return next(iter(units)) if len(units) == 1 else ""


def _year_scope(objects):
    """명시된 문서 계획기간이 유일할 때만 ’24 등 머리글의 세기를 해소한다.

    기간에서 연도별 사실을 생성하는 기능은 아니다. 실제 연도 열만 해석한다.
    """
    scopes = {}
    for obj in objects:
        text = " ".join(str(obj.get(k, "")) for k in ("text", "nearby_text"))
        for match in re.finditer(r"기본계획\s*\(\s*((?:19|20)\d{2})\s*[-~–]\s*((?:19|20)\d{2})\s*\)", text):
            start, end = map(int, match.groups())
            if start <= end and end - start < 100:
                scopes.setdefault((start, end), {"start": start, "end": end, "object_id": obj.get("object_id"),
                    "page": obj.get("page_number"), "snippet": match.group(0)})
    if len(scopes) == 1:
        return next(iter(scopes.values()))
    return {"ambiguous": True, "ranges": list(scopes.values())} if scopes else None


def _header_years(headers, scope):
    result = {}
    for c, header in enumerate(headers):
        label = header.split(" / ")[-1].strip()
        if match := re.fullmatch(r"((?:19|20)\d{2})\s*년?", label):
            result[c] = int(match.group(1))
        elif scope and not scope.get("ambiguous") and (match := re.fullmatch(r"[’'‘ʼ](\d{2})\s*년?", label)):
            possible = [year for year in range(scope["start"], scope["end"] + 1) if year % 100 == int(match.group(1))]
            if len(possible) == 1:
                result[c] = possible[0]
    return result


def _context_token(key, value):
    value = token(value)
    return {"계": "합계", "총계": "합계"}.get(value, value) if key == "재원구분" else value


def _caption_project(caption):
    text = re.sub(r"^\s*\[?\s*표\s*[\dⅠⅡⅢⅣⅤⅥⅦⅧⅨⅩ]+(?:[-–]\d+)*\s*\]?\s*", "", str(caption or ""))
    return re.sub(r"\s*(?:온실가스\s*감축효과|투자계획|추진계획|성과\s*목표)\s*$", "", text).strip()


def _subject_token(value):
    raw_unit = _embedded_unit(value)
    if raw_unit:
        value = re.sub(r"\(([^()]*)\)\s*$", lambda m: "" if token(m.group(1)) == token(raw_unit) else m.group(0), str(value))
    return token(value)


def _vehicle_qualifier_base(name):
    """차량 동력원 열거만 제한적으로 분리. 지역·단계·세부 사업 괄호는 유지."""
    text = token(name)
    groups = list(re.finditer(r"\(([^()]*)\)", text))
    if len(groups) != 1:
        return ""
    fuels = set(re.split(r"[,/·ㆍ+]|및", groups[0][1]))
    if len(fuels) < 2 or not fuels <= {"cng", "lng", "lpg", "전기", "수소", "경유", "휘발유", "하이브리드"}:
        return ""
    base = text[:groups[0].start()] + text[groups[0].end():]
    return base if re.search(r"청소차|운반차|화물차|승용차|자동차|차량|버스|택시", base) else ""


def _financial_role(caption, headers):
    """표 제목에 명시된 수입과 지출을 구분. 사업명에 '정산금'만 있다고 제외하지 않는다."""
    label = token(caption)
    income = bool(re.search(r"수익금|수입금|수익액|수입액|매전수익|판매수익", label))
    spending = bool(re.search(r"투자|예산|사업비|지출", label)) or any(
        _header_field(h.rsplit(" / ", 1)[0]) == "예산액" for h in headers)
    return "mixed" if income and spending else "revenue" if income else "unspecified_or_spending"


def _funding_components(rows, start, name_col, year_col, raw_unit):
    """명시 합계 바로 아래 병합 재원 행을 별칭 검증 근거로만 읽는다. 업무 사실을 생성하지 않는다."""
    components = {}
    for row_index in range(start, len(rows)):
        cells = rows[row_index]
        if token(cells[name_col]):
            break
        labels = {token(v) for c, v in enumerate(cells) if c not in {name_col, year_col}} & {"국비", "시비", "구비", "도비", "지방비", "기타"}
        value = number(cells[year_col])
        if len(labels) != 1 or value is None:
            continue
        label = next(iter(labels))
        if label in components:
            return []  # 같은 재원에 둘 이상의 금액을 임의 선택하지 않음
        components[label] = {"funding": label, "raw_value": cells[year_col], "raw_unit": raw_unit,
                             "source_row": row_index + 1, "source_column": year_col + 1}
    return list(components.values())


def _unit_annotations(objects):
    """원래 파서가 읽은 명시 단위만 위치와 함께 재사용. PDF/모델 재호출 없음."""
    result = defaultdict(list)
    for obj in objects:
        text, bbox = obj.get("text") or "", obj.get("bbox")
        if obj.get("object_type") != "text" or not bbox or len(bbox) != 4 or not re.search(r"단위\s*[:：]", text):
            continue
        literal = _embedded_unit(text)
        if literal:
            result[obj.get("page_number")].append({"object_id": obj.get("object_id"), "page": obj.get("page_number"),
                "bbox": bbox, "literal": literal, "text": text})
    return result


def _spatial_unit(obj, annotations):
    bbox = obj.get("bbox")
    if not bbox or len(bbox) != 4:
        return "", []
    candidates = [a for a in annotations.get(obj.get("page_number"), []) if
        -6 <= bbox[1] - a["bbox"][3] <= 48 and
        max(bbox[0], a["bbox"][0]) < min(bbox[2], a["bbox"][2])]
    if len({token(a["literal"]) for a in candidates}) != 1:
        return "", candidates
    return candidates[0]["literal"], candidates


def itemize(objects, document_sha256):
    """기계적으로 해석 가능한 명시 표 셀/과제. 나머지는 진단으로 보존."""
    facts, diagnostics, counts = [], [], Counter()
    seen_objects, seen_facts, seen_nonbudget = set(), {}, {}
    objects = [obj.to_dict() if hasattr(obj, "to_dict") else obj for obj in objects]
    annotations = _unit_annotations(objects)
    # 표지(물리 1~2쪽)의 명시 계획기간을 우선한다. 본문에 인용된 다른 계획의
    # 기간을 문서 자체의 기간으로 오인하지 않는다. 빈 표지는 전체 유일성 조건.
    document_year_scope = _year_scope([o for o in objects if int(o.get("page_number") or 0) <= 2]) or _year_scope(objects)
    for obj in objects:
        meta = obj.get("metadata") or {}
        year_scope = _year_scope([obj]) or document_year_scope
        counts["objects_seen"] += 1
        p = int(obj.get("page_number") or 0)
        oid = str(obj.get("object_id", ""))
        source = {"object_id": oid, "evidence_id": meta.get("evidence_id", ""),
                  "aliases": sorted(ids(obj) | ids(meta)), "page": p,
                  "bbox": obj.get("bbox"), "bbox_scope": "object_not_individual_cell", "table_number": obj.get("number", ""),
                  "caption": obj.get("caption", ""), "engine": meta.get("engine", "native")}
        source["caption_project"] = _caption_project(source["caption"])
        def diagnose(reason, *, status="unassessable", **extra):
            counts[reason] += 1
            diagnostics.append({"status": status, "reason": reason, "source": source, **extra})
        if meta.get("index_reference") or meta.get("is_index_reference"):
            counts["excluded_reference"] += 1
            continue
        rows = obj.get("rows") or []
        if not rows:
            if obj.get("object_type") in {"image", "chart", "graph"}:
                if meta.get("render_proxy") or meta.get("source_kind") == "page_render":
                    diagnose("page_render_proxy", status="excluded", exclusion_scope="coverage_audit_only")
                else:
                    diagnose("image_or_caption_only")
            else:
                counts["outside_explicit_table_scope"] += 1
            continue
        physical = (p, canonical(obj.get("bbox")), canonical(rows))
        if physical in seen_objects:
            counts["duplicate_object"] += 1
        seen_objects.add(physical)
        if meta.get("table_parse_error") or meta.get("parse_errors") or not all(isinstance(r, list) for r in rows):
            diagnose("unsupported_table_structure")
            continue
        width = max(map(len, rows), default=0)
        if not width or any(len(r) != width for r in rows):
            diagnose("non_rectangular_table")
            continue
        header_count = int(meta.get("header_row_count") or 1)
        if header_count < 1 or header_count >= len(rows):
            if len(rows) == 1 and not any(_header_field(str(v)) for v in rows[0]):
                diagnose("non_tabular_diagram_or_fragment")
            else:
                diagnose("unresolved_header")
            continue
        headers = [" / ".join(dict.fromkeys(str(rows[r][c]).strip() for r in range(header_count)
                                          if str(rows[r][c]).strip())) for c in range(width)]
        fields = [_header_field(h) for h in headers]
        active_year_headers, active_year_header_row = headers, header_count
        name_columns = [c for c, f in enumerate(fields) if f in NAME_FIELDS]
        if len(name_columns) != 1:
            if _header_years(headers, year_scope) or any("성과지표" in h for h in headers):
                diagnose("caption_business_context_required" if source["caption_project"] else "business_identity_unresolved")
            else:
                counts["outside_supported_identity"] += 1
            continue
        name_col = name_columns[0]
        subject_names = {_subject_token(row[name_col]) for row in rows[header_count:] if
            token(row[name_col]) not in ABSENT | {"합계", "계", "소계", "총계"} and not _header_field(str(row[name_col]))}
        source["business_subject_count"] = len(subject_names)
        source["business_subject_names"] = sorted(subject_names)
        years = _header_years(headers, year_scope)
        counts["excluded_total_columns"] += sum(token(h) in {"계", "합계", "소계", "총계"} for h in headers)
        metric_columns = [(c, f) for c, f in enumerate(fields) if f in VALUE_FIELDS and f != "과제명"]
        foundation = fields[name_col] == "과제명" and "주관부서" in fields and not metric_columns and not years
        if not foundation and not years and not ("연도" in fields and metric_columns):
            if re.search(r"이행성과|점검실적|달성도", str(obj.get("caption", "")) + " ".join(headers)):
                diagnose("outside_supported_performance_schema", status="excluded", exclusion_scope="coverage_audit_only")
            else:
                diagnose("no_explicit_year_or_metric")
            continue
        context_text = " ".join(str(obj.get(k, "")) for k in ("caption", "section", "text"))
        explicit = {k: str(meta[k]) for k in CONTEXT_FIELDS if meta.get(k) not in (None, "")}
        # 표 단위도 명시된 표기만 사용. 주변 페이지 전체의 단위를 가져오지 않는다.
        table_unit = str(meta.get("unit") or _embedded_unit(context_text))
        unit_origin = "object_header_or_metadata"
        unit_evidence = []
        if not table_unit:
            table_unit, unit_evidence = _spatial_unit(obj, annotations)
            if table_unit:
                unit_origin = "spatial_explicit_native_unit"
        # 표 안의 단위가 우선. 투자 표에서 페이지의 명시 화폐 단위가 하나일 때만
        # 보조 문맥으로 사용하고, 그 문자열과 귀속 한계를 원장에 남긴다.
        if not table_unit and not unit_evidence and any("사업명" in h for h in headers) and (
                re.search(r"투자|예산|재정", obj.get("caption", "")) or "예산액" in fields):
            nearby_units = {_embedded_unit("단위:" + m) for m in re.findall(r"단위\s*[:：]\s*([^()\n]+)", obj.get("nearby_text", ""))}
            nearby_units.discard("")
            if len(nearby_units) == 1 and unit(next(iter(nearby_units)))[0] == "krw":
                table_unit = next(iter(nearby_units))
                unit_origin = "unique_explicit_page_currency_unit_reviewable"
        for r, cells in enumerate(rows[header_count:], header_count + 1):
            repeated_years = _header_years([str(v) for v in cells], year_scope)
            header_labels = {"사업명", "사업명(단위)", "성과지표", "성과지표(단위)", "구분", "합계", "계", "비고", "단위", "예산단위", "활동단위"}
            normal_header = all(c in repeated_years or not str(v).strip() or token(v) in header_labels for c, v in enumerate(cells))
            labelled_header = (name_col not in repeated_years and token(cells[name_col]) not in ABSENT | header_labels and
                number(cells[name_col]) is None and
                all(c == name_col or c in repeated_years or not str(v).strip() or token(v) in header_labels for c, v in enumerate(cells)) and
                (any(token(v) == "구분" for c, v in enumerate(cells) if c != name_col) or
                 all(re.match(r"[’'‘ʼ]", str(cells[c]).strip()) for c in repeated_years)))
            if len(repeated_years) >= 2 and (normal_header or labelled_header):
                years = repeated_years
                active_year_headers, active_year_header_row = list(cells), r
                counts["repeated_year_header"] += 1
                if labelled_header:
                    # 연도 위치는 보존하되 머리글을 측정값/재판독 후보로 세지 않는다.
                    for c, year in repeated_years.items():
                        diagnose("repeated_header_with_merged_business_label", status="excluded", source_row=r,
                                 source_column=c + 1, raw_value=cells[c], header_year=year,
                                 business_label=cells[name_col], raw_row=list(cells))
                continue
            name = str(cells[name_col] or "").strip()
            if token(name) in ABSENT or token(name) in {"합계", "총계", "소계", "계"}:
                numeric_columns = list(years) or [c for c, _ in metric_columns]
                if numeric_columns and all(token(cells[c]) in ABSENT for c in numeric_columns):
                    diagnose("source_values_blank_row", status="excluded", source_row=r)
                else:
                    diagnose("merged_business_identity_unresolved" if token(name) in ABSENT else "aggregate_subject_unresolved", source_row=r)
                continue
            row_fields = {field: str(cells[c]).strip() for c, field in enumerate(fields) if field and cells[c] not in (None, "")}
            context = {**explicit, **{k: row_fields[k] for k in CONTEXT_FIELDS if row_fields.get(k)}}
            name_unit = _embedded_unit(name)
            if name_unit:
                name = re.sub(r"\(([^()]*)\)\s*$", lambda m: "" if token(m.group(1)) == token(name_unit) else m.group(0), name).strip()
            tasks = []
            if foundation:
                tasks.append((name_col, "과제명", None, name, ""))
            elif years:
                if all(token(cells[c]) in ABSENT for c in years):
                    aggregate_values = [{"source_column": c + 1, "header": h, "raw_value": cells[c]}
                        for c, h in enumerate(active_year_headers) if token(h) in {"합계", "총계", "누계"} and token(cells[c]) not in ABSENT]
                    diagnose("source_annual_values_blank_total_only" if aggregate_values else "source_values_blank_row",
                             status="excluded", source_row=r, raw_row=list(cells),
                             year_headers=active_year_headers, aggregate_values=aggregate_values)
                    continue
                raw_unit = row_fields.get("활동단위") or row_fields.get("예산단위") or row_fields.get("단위") or name_unit or table_unit
                if unit(raw_unit)[0] == "krw":
                    funding = {str(cells[c]).strip() for c, h in enumerate(headers) if c != name_col and
                               (token(h) == "구분" or h.startswith("Col")) and
                               token(cells[c]) in {"계", "합계", "국비", "시비", "구비", "도비", "기타", "지방비"}}
                    if len(funding) == 1:
                        context["재원구분"] = next(iter(funding))
                parent_metrics = {_header_field(h.rsplit(" / ", 1)[0]) for c, h in enumerate(headers) if c in years and " / " in h}
                parent_metrics &= {"활동량", "예상감축량", "예산액"}
                if len(parent_metrics) > 1:
                    diagnose("mixed_wide_metrics_require_unit_mapping", source_row=r)
                    continue
                metric = next(iter(parent_metrics), "") or ("예산액" if unit(raw_unit)[0] == "krw" else "예상감축량" if unit(raw_unit)[0] in {"co2", "co2eq"} else "활동량" if row_fields.get("활동단위") or name_unit or "활동량" in context_text or "목표물량" in context_text else "")
                if not metric:
                    diagnose("explicit_unit_evidence_required" if not raw_unit else "ambiguous_wide_metric", source_row=r,
                             unit_evidence=unit_evidence)
                    continue
                tasks.extend((c, metric, year, cells[c], raw_unit) for c, year in years.items())
            else:
                year_raw = row_fields.get("연도", "")
                year_match = re.fullmatch(r"((?:19|20)\d{2})년?", year_raw)
                if not year_match:
                    diagnose("period_or_ambiguous_year", source_row=r, raw_year=year_raw)
                    continue
                for c, metric in metric_columns:
                    raw_unit = row_fields.get("예산단위" if metric == "예산액" else "활동단위" if metric == "활동량" else "단위") or _embedded_unit(headers[c]) or table_unit
                    tasks.append((c, metric, int(year_match.group(1)), cells[c], raw_unit))
            for c, metric, year, raw_value, raw_unit in tasks:
                if metric != "과제명" and token(raw_value) in ABSENT and raw_value != 0:
                    diagnose("source_value_blank", status="excluded", source_row=r, source_column=c + 1,
                             raw_value=raw_value, raw_unit=raw_unit)
                    continue
                if metric != "과제명" and (number(raw_value) is None or not raw_unit):
                    diagnose("explicit_unit_evidence_required" if not raw_unit else "source_value_unreadable", source_row=r, source_column=c + 1,
                             raw_value=raw_value, raw_unit=raw_unit)
                    continue
                dimension = unit(raw_unit)[0]
                if (metric == "활동량" and dimension == "krw") or (metric == "예산액" and dimension != "krw"):
                    diagnose("source_unit_metric_conflict", source_row=r, source_column=c + 1,
                             raw_value=raw_value, raw_unit=raw_unit, field=metric, unit_evidence=unit_evidence)
                    continue
                sheet = "foundation_measures" if metric == "과제명" else "financial_plan" if metric == "예산액" else "quantitative_reductions"
                identity = {"name": name, "raw_name": str(cells[name_col]), "code": row_fields.get("관리번호", row_fields.get("과제ID", "")),
                            "year": year, "context": context}
                fact_source = {**source, "row": r, "column": c + 1,
                               "snippet": " | ".join(str(v) for v in cells), "headers": headers,
                               "year_headers": active_year_headers, "year_header_row": active_year_header_row,
                               "year_scope_evidence": year_scope if any("’" in str(h) or "'" in str(h) for h in active_year_headers) else None,
                               "unit_origin": unit_origin, "unit_evidence": unit_evidence}
                signature = [FACT_ID_VERSION, document_sha256, p, token(name), identity["code"], year, metric,
                             {k: _context_token(k, v) for k, v in context.items()},
                             str(number(raw_value)) if metric != "과제명" else [token(name), row_fields.get("주관부서", "")], token(raw_unit)]
                fact_id = hashlib.sha256(canonical(signature).encode()).hexdigest()[:24]
                role = _financial_role(obj.get("caption"), headers) if metric == "예산액" else ""
                if role in {"revenue", "mixed"}:
                    if fact_id in seen_nonbudget:
                        existing = seen_nonbudget[fact_id]
                        existing["source"]["aliases"] = sorted(set(existing["source"]["aliases"]) | set(fact_source["aliases"]))
                        existing["source_instances"].append(fact_source)
                        continue
                    diagnose("non_investment_revenue_outside_budget_scope" if role == "revenue" else "mixed_revenue_spending_context",
                             status="excluded" if role == "revenue" else "unassessable",
                             candidate_id=fact_id, source=fact_source, source_row=r, source_column=c + 1,
                             raw_value=raw_value, raw_unit=raw_unit, financial_role=role,
                             observed_identity=identity, source_instances=[fact_source], exclusion_scope="coverage_audit_only")
                    seen_nonbudget[fact_id] = diagnostics[-1]
                    continue
                if metric == "예산액" and _vehicle_qualifier_base(name) and _context_token("재원구분", context.get("재원구분")) == "합계":
                    fact_source["funding_components"] = _funding_components(rows, r, name_col, c, raw_unit)
                if fact_id in seen_facts:
                    counts["duplicate_fact"] += 1
                    prior = facts[seen_facts[fact_id]]
                    prior["source"]["aliases"] = sorted(set(prior["source"]["aliases"]) | set(source["aliases"]))
                    prior["source_instances"].append(fact_source)
                    continue
                seen_facts[fact_id] = len(facts)
                facts.append({"candidate_id": fact_id, "sheet": sheet, "field": metric,
                              "identity": identity, "raw_value": raw_value, "raw_unit": raw_unit,
                              "normalized_value": str(number(raw_value) * unit(raw_unit)[1]) if metric != "과제명" else token(name),
                              "normalized_dimension": unit(raw_unit)[0] if metric != "과제명" else "task_existence",
                              "department": row_fields.get("주관부서", ""), "source": fact_source,
                              "source_instances": [deepcopy(fact_source)]})
    # 같은 연도·생략명에 서로 다른 동력원 집합이 대응하면 이름 생략으로 확정하지 않는다.
    vehicle_variants = defaultdict(set)
    for fact in facts:
        base = _vehicle_qualifier_base(fact["identity"]["name"]) if fact["sheet"] == "financial_plan" else ""
        if base:
            vehicle_variants[(base, fact["identity"]["year"])].add(token(fact["identity"]["name"]))
    for fact in facts:
        base = _vehicle_qualifier_base(fact["identity"]["name"]) if fact["sheet"] == "financial_plan" else ""
        if base:
            fact["source"]["vehicle_qualifier_variant_count"] = len(vehicle_variants[(base, fact["identity"]["year"])])
    # 수치/연도와 원문 사업명이 같은 활동량 표가 있을 때만 후단 계층 연결 근거로 제공.
    activity_sources = defaultdict(list)
    for fact in facts:
        if fact["field"] == "활동량":
            activity_sources[(token(fact["identity"]["name"]), fact["identity"]["year"])].append(fact)
    for fact in facts:
        if fact["field"] != "예상감축량" or not fact["source"]["caption_project"]:
            continue
        fact["related_activity_sources"] = [{"source": peer["source"], "value": peer["raw_value"], "unit": peer["raw_unit"]}
            for peer in activity_sources[(token(fact["identity"]["name"]), fact["identity"]["year"])]
            if token(peer["source"]["caption_project"]) == token(fact["source"]["caption_project"]) and
               abs(peer["source"]["page"] - fact["source"]["page"]) <= 2]
    return facts, diagnostics, dict(counts)


def read_excel(path):
    """재개방한 실제 저장 파일만 읽는다. 수식은 값으로 간주하지 않는다."""
    import openpyxl
    wb = openpyxl.load_workbook(path, read_only=True, data_only=False)
    data = {}
    try:
        # 활동량이 연차별 시트에 반영된 경우도 보존 여부를 확인한다.
        for key, display in {**SHEETS, "annual_implementation": "09_연차별이행계획"}.items():
            if display not in wb.sheetnames:
                data[key] = []
                continue
            iterator = wb[display].iter_rows()
            header_cells = next(iterator, ())
            headers = [cell.value for cell in header_cells]
            data[key] = []
            for cells in iterator:
                row = {h: cell.value for h, cell in zip(headers, cells) if h is not None}
                if not any(v is not None for v in row.values()):
                    continue
                row["_excel_cells"] = {h: f"'{display}'!{cell.coordinate}" for h, cell in zip(headers, cells) if h}
                data[key].append(row)
    finally:
        wb.close()
    return data


class RowIndex:
    def __init__(self, data):
        self.rows = defaultdict(list)
        self.by_year = defaultdict(list)
        self.by_code = defaultdict(list)
        for sheet, rows in (data or {}).items():
            if sheet not in {*SHEETS, "annual_implementation"} or not isinstance(rows, list):
                continue
            for ordinal, row in enumerate(rows, 1):
                if isinstance(row, dict):
                    name = next((row.get(k) for k in NAME_FIELDS if row.get(k)), "")
                    numeric_year = number(row.get("연도"))
                    year = str(int(numeric_year)) if numeric_year is not None and numeric_year == int(numeric_year) else str(row.get("연도") or "")
                    entry = (sheet, ordinal, row)
                    self.rows[(token(name), year)].append(entry)
                    self.by_year[year].append(entry)
                    code = row.get("관리번호") or row.get("과제ID")
                    if code:
                        self.by_code[(token(code), year)].append(entry)

    @staticmethod
    def _anchored(row, source):
        return bool(set(source.get("aliases", [])) & ids(row)) or source["page"] in pages(row.get("출처페이지"))

    @staticmethod
    def _equal_measure(value, raw_unit, expected, expected_unit):
        actual, expected = number(value), number(expected)
        if actual is None or expected is None or not raw_unit or not expected_unit:
            return False
        dimension, scale = unit(raw_unit)
        expected_dimension, expected_scale = unit(expected_unit)
        return dimension == expected_dimension and actual * scale == expected * expected_scale

    def _vehicle_funding_proof(self, fact, row, year):
        components = fact["source"].get("funding_components", [])
        # 표의 동력원 열거와 재원별 내역까지 대응할 때만 생략된 이름을 확정한다.
        if fact["source"].get("business_subject_count") != 1 or fact["source"].get("vehicle_qualifier_variant_count") != 1 or len(components) < 2:
            return []
        total = sum(number(c["raw_value"]) * unit(c["raw_unit"])[1] for c in components)
        if total != number(fact["raw_value"]) * unit(fact["raw_unit"])[1]:
            return []
        proof = []
        for component in components:
            peers = [(sheet, ordinal, peer) for sheet, ordinal, peer in self.rows.get((token(row.get("사업명")), year), []) if
                sheet == "financial_plan" and self._anchored(peer, fact["source"]) and
                _context_token("재원구분", peer.get("재원구분")) == component["funding"] and
                all(not row.get(k) or peer.get(k) == row.get(k) for k in ("계획구분", "부문", "관리번호"))]
            if not peers or any(not self._equal_measure(peer.get("예산액"), peer.get("예산단위"), component["raw_value"], component["raw_unit"]) for _, _, peer in peers):
                return []
            proof.extend({"source_component": component, "sheet": sheet, "row": ordinal,
                          "cells": peer.get("_excel_cells", {}), "evidence_ids": sorted(ids(peer))} for sheet, ordinal, peer in peers)
        return proof

    def _identity_options(self, fact):
        """정확 이름·명시 제목·재원으로 검증한 차량 생략명·원문 활동량 연결만 허용."""
        identity, source = fact["identity"], fact["source"]
        year = str(identity["year"] or "")
        options = {(sheet, ordinal): (sheet, ordinal, row, {"method": "exact_business_name", "uncertain": False})
                   for sheet, ordinal, row in self.rows.get((token(identity["name"]), year), [])}
        project = token(source.get("caption_project"))
        if fact["sheet"] == "financial_plan" and project:
            caption_options = []
            for sheet, ordinal, row in self.by_year[year]:
                if sheet != "financial_plan" or not self._anchored(row, source):
                    continue
                if token(_caption_project(row.get("사업명"))) != project:
                    continue
                explicit_subproject = token(row.get("부문")) == token(identity["name"])
                if not explicit_subproject and token(row.get("부문")) in source.get("business_subject_names", []):
                    continue  # 표에 명시된 다른 하위 사업의 행
                bound = explicit_subproject or source.get("business_subject_count") == 1
                link = {"method": "caption_and_subproject_field" if explicit_subproject else "single_subject_table_caption" if bound else "unbound_multi_subject_caption",
                        "source_caption": source.get("caption"), "subject_count": source.get("business_subject_count"),
                        "matched_subproject_field": "부문" if explicit_subproject else None, "uncertain": not bound}
                caption_options.append((sheet, ordinal, row, link))
            has_bound_caption = any(not option[3]["uncertain"] for option in caption_options)
            for sheet, ordinal, row, link in caption_options:
                if not has_bound_caption or not link["uncertain"]:
                    options.setdefault((sheet, ordinal), (sheet, ordinal, row, link))
        vehicle_base = _vehicle_qualifier_base(identity["name"]) if fact["sheet"] == "financial_plan" else ""
        if vehicle_base:
            for sheet, ordinal, row in self.rows.get((vehicle_base, year), []):
                if sheet != "financial_plan" or not self._anchored(row, source):
                    continue
                proof = self._vehicle_funding_proof(fact, row, year)
                options.setdefault((sheet, ordinal), (sheet, ordinal, row,
                    {"method": "vehicle_fuel_qualifier_omission", "source_name": identity["name"],
                     "saved_name": row.get("사업명"), "funding_proof": proof, "uncertain": not bool(proof)}))
        for peer in fact.get("related_activity_sources", []):
            for peer_sheet, peer_ordinal, annual in self.rows.get((token(identity["name"]), year), []):
                code = annual.get("관리번호")
                if peer_sheet != "annual_implementation" or not code or not self._anchored(annual, peer["source"]):
                    continue
                if not self._equal_measure(annual.get("목표물량"), annual.get("목표단위"), peer["value"], peer["unit"]):
                    continue
                for sheet, ordinal, row in self.by_code.get((token(code), year), []):
                    if sheet != "quantitative_reductions" or not self._anchored(row, source):
                        continue
                    if not self._equal_measure(row.get("활동량"), row.get("활동단위"), peer["value"], peer["unit"]):
                        continue
                    options.setdefault((sheet, ordinal), (sheet, ordinal, row,
                        {"method": "source_activity_and_annual_code", "uncertain": False,
                         "activity_source": peer["source"], "annual_sheet": peer_sheet,
                         "annual_row": peer_ordinal, "annual_cells": annual.get("_excel_cells", {}),
                         "management_code": code}))
        return list(options.values())

    def match(self, fact):
        identity = fact["identity"]
        options = self._identity_options(fact)
        unresolved_context = {key for key in CONTEXT_FIELDS if key not in identity["context"] and
            len({_context_token(key, row[key]) for _, _, row, link in options if row.get(key) and
                 link.get("matched_subproject_field") != key}) > 1}
        good, conflict, ambiguous = [], [], []
        for sheet, ordinal, row, identity_link in options:
            # 동명이과제는 명시 영역/계획/재원/인자로 분리. 코드만 같다고 매칭하지 않는다.
            if identity["code"] and row.get("관리번호", row.get("과제ID")) and token(identity["code"]) != token(row.get("관리번호", row.get("과제ID"))):
                continue
            if any(row.get(k) and _context_token(k, row[k]) != _context_token(k, v) for k, v in identity["context"].items()):
                continue
            anchored = self._anchored(row, fact["source"])
            uncertain = identity_link["uncertain"] or bool(unresolved_context) or not anchored or any(not row.get(k) for k in identity["context"])
            uncertain |= any(row.get(k) and k not in identity["context"] for k in CONTEXT_FIELDS) and not anchored
            if fact["sheet"] == "financial_plan" and row.get("재원구분") and "재원구분" not in identity["context"]:
                uncertain = True
            field = fact["field"]
            if sheet == "annual_implementation" and field == "활동량":
                field, unit_field = "목표물량", "목표단위"
            else:
                unit_field = "예산단위" if field == "예산액" else "활동단위" if field == "활동량" else "단위"
            if field not in row:
                continue
            reference = {"sheet": sheet, "row": ordinal, "field": field,
                         "cells": row.get("_excel_cells", {}), "value": row.get(field),
                         "unit": row.get(unit_field), "evidence_ids": sorted(ids(row))}
            reference["identity_link"] = identity_link
            if field == "과제명":
                if fact["department"] and token(row.get("주관부서")) != token(fact["department"]):
                    (conflict if row.get("주관부서") else ambiguous).append(reference)
                    continue
                equal = token(row.get(field)) == token(fact["raw_value"])
            else:
                actual = number(row.get(field))
                if actual is None or not row.get(unit_field):
                    ambiguous.append(reference)
                    continue
                expected_dimension, expected_scale = unit(fact["raw_unit"])
                dimension, scale = unit(row[unit_field])
                if dimension != expected_dimension:
                    reference["reason_codes"] = ["unit_dimension_mismatch"]
                    reference["expected_unit"] = fact["raw_unit"]
                    ambiguous.append(reference)
                    continue
                equal = actual * scale == number(fact["raw_value"]) * expected_scale
            if uncertain:
                ambiguous.append(reference)
            elif equal:
                good.append(reference)
            else:
                conflict.append(reference)
        # 같은 사실에 다른 수치가 병존하면 정상 반영으로 숨기지 않는다.
        status = "conflict_review" if conflict else "unit_review" if any("unit_dimension_mismatch" in r.get("reason_codes", []) for r in ambiguous) else "ambiguous_review" if ambiguous else "covered" if good else "absent"
        return {"status": status, "matches": good, "conflicts": conflict, "ambiguous": ambiguous}


def _exposure(fact, records, routes):
    matches = []
    source = fact["source"]
    for record in records:
        if record.get("sheet_key") != fact["sheet"]:
            continue
        aliases = set(source["aliases"])
        exact_object = bool(ids(record) & aliases)
        if exact_object or source["page"] in pages(record.get("page_nums")):
            matches.append({k: record[k] for k in ("batch_id", "object_id", "status", "rows", "model_requested",
                "no_call_reason", "reused", "error", "input_sha256") if k in record})
            matches[-1]["association"] = "object_association" if exact_object else "page_association_only"
    route = next((r for r in routes if r.get("sheet_key") == fact["sheet"]), {})
    no_call = "not_routed" if source["page"] in route.get("not_selected_pages", []) else "unknown_historical_record"
    return {"records": matches, "no_record_reason": "" if matches else no_call,
            "exact_cell_content_exposure": "not_verified"}


def compare(objects, document_sha256, excel_data, *, raw=None, cleaned=None,
            records=(), routes=(), quarantine=(), provenance=None):
    facts, diagnostics, counts = itemize(objects, document_sha256)
    indices = {"raw": RowIndex(raw) if raw is not None else None,
               "cleaned": RowIndex(cleaned) if cleaned is not None else None,
               "excel": RowIndex(excel_data)}
    held_rows = defaultdict(list)
    held_details = {}
    for entry in quarantine:
        if isinstance(entry, dict) and isinstance(entry.get("row"), dict):
            sheet = entry.get("sheet_key", entry.get("sheet", ""))
            held_rows[sheet].append(entry["row"])
            held_details[(sheet, len(held_rows[sheet]))] = entry
    held_index = RowIndex(held_rows)
    journal = []
    for fact in facts:
        stages = {stage: index.match(fact) if index else {"status": "unknown", "reason": "snapshot_not_available"}
                  for stage, index in indices.items()}
        final = stages["excel"]
        held = held_index.match(fact)
        held_refs = held["matches"] + held["conflicts"] + held["ambiguous"]
        quarantine_details = [deepcopy(held_details[(ref["sheet"], ref["row"])]) for ref in held_refs]
        status, reason = final["status"], "actual_saved_cell_comparison"
        if status == "absent":
            if held["status"] != "absent":
                status, reason = "cleaning_excluded", "quarantine_preserved_no_auto_retry"
            elif stages["cleaned"]["status"] == "covered":
                status, reason = "storage_review", "cleaned_fact_not_in_saved_cells"
            elif stages["raw"]["status"] == "covered":
                status, reason = "missing_candidate", "present_raw_absent_after_cleaning"
            else:
                status, reason = "missing_candidate", "explicit_source_fact_absent_from_saved_cells"
        elif status == "covered" and any(r["sheet"] != fact["sheet"] for r in final["matches"]):
            status, reason = "covered_elsewhere", "equivalent_fact_in_other_supported_sheet"
        journal.append({"schema_version": VERSION, **fact, "status": status, "reason": reason,
                        "selection_reason": "원문 표의 명시 사실을 저장 셀과 대조합니다. 정확 이름 이외에는 원문 표 제목·세부 사업 열 또는 활동량/연도/관리번호의 연결 근거를 기록합니다. 임의 유사도 매칭은 하지 않습니다.",
                        "stages": stages, "extraction": _exposure(fact, records, routes),
                        "quarantine": quarantine_details,
                        "document_sha256": document_sha256, "provenance": provenance or {},
                        "source_truth": "parser_observation_not_verified_gold",
                        "review_status": "not_required" if status == "covered" else "pending",
                        "recall_policy": "review_before_recall" if status == "missing_candidate" else "no_automatic_recall",
                        "second_pass_status": "not_requested"})
    for diagnostic in diagnostics:
        signature = [document_sha256, diagnostic]
        journal.append({"schema_version": VERSION, "candidate_id": hashlib.sha256(canonical(signature).encode()).hexdigest()[:24],
                        **diagnostic, "document_sha256": document_sha256,
                        "provenance": provenance or {},
                        "review_status": "not_required" if diagnostic["status"] == "excluded" else "pending",
                        "recall_policy": "excluded_from_recall" if diagnostic["status"] == "excluded" else "review_before_recall",
                        "second_pass_status": "not_requested"})
    summary = {"schema_version": VERSION, "mode": "audit", "model_calls_added": 0,
               "business_data_modified": False, "model_accuracy_evaluated": False,
               "document_sha256": document_sha256, "supported_sheets": SHEETS,
               "comparable_facts": len(facts), "counts": dict(Counter(r["status"] for r in journal)),
               "coverage_counts": {
                   "compared_facts": len(facts),
                   "reflected_facts": sum(r["status"] in {"covered", "covered_elsewhere"} for r in journal),
                   "missing_candidates": sum(r["status"] == "missing_candidate" for r in journal),
                   "review_or_quarantined_facts": sum(r["status"] in {"conflict_review", "ambiguous_review", "unit_review", "storage_review", "cleaning_excluded"} for r in journal),
                   "unassessable_diagnostics_not_fact_count": sum(r["status"] == "unassessable" for r in diagnostics),
                   "excluded_diagnostics_not_fact_count": sum(r["status"] == "excluded" for r in diagnostics)},
               "source_scope": counts, "provenance": provenance or {},
               "stage_availability": {"raw": raw is not None, "cleaned": cleaned is not None, "excel": True},
               "extraction_records": {"scope": "supplied_text_sheet_records_only_not_provider_billing",
                   "records": len(records),
                   "boundary_requested": sum(r.get("model_requested") is True for r in records),
                   "no_request": sum(r.get("model_requested") is False for r in records),
                   "unknown_historical": sum("model_requested" not in r for r in records),
                   "called_empty_success": sum(r.get("model_requested") is True and r.get("status") == "ok" and r.get("rows") == 0 for r in records),
                   "no_call_reasons": dict(Counter(r["no_call_reason"] for r in records if r.get("no_call_reason")))},
               "limitations": ["명시 표의 세 시트 한정; 본문·복잡 표·이미지 사실의 전수 누락률이 아님",
                               "배치 페이지 연관은 해당 셀의 실제 프롬프트 포함을 보증하지 않음",
                               "누락 후보는 재판독·정답 판정·자동 수정 대상 확정이 아님"]}
    return journal, summary


def _atomic_write(path, writer):
    path = Path(path)
    path.parent.mkdir(parents=True, exist_ok=True)
    fd, temporary = tempfile.mkstemp(prefix=path.name + ".", suffix=".tmp", dir=path.parent)
    try:
        with os.fdopen(fd, "w", encoding="utf-8", newline="\n") as stream:
            writer(stream)
        os.replace(temporary, path)
    finally:
        if Path(temporary).exists():
            Path(temporary).unlink()


def persist(output_prefix, journal, summary, *, overwrite=False):
    prefix = Path(output_prefix)
    journal_path = prefix.with_name(prefix.name + "_missing_candidates.jsonl")
    summary_path = prefix.with_name(prefix.name + "_coverage_summary.json")
    if not overwrite and (journal_path.exists() or summary_path.exists()):
        raise FileExistsError("감사 출력이 이미 있습니다. 새 output-prefix를 사용하세요.")
    result = {**summary, "missing_candidates_path": str(journal_path.resolve()),
              "coverage_summary_path": str(summary_path.resolve()), "journal_records": len(journal)}
    def write_journal(stream):
        for row in journal:
            stream.write(json.dumps(row, ensure_ascii=False, default=str) + "\n")
    _atomic_write(journal_path, write_journal)
    result["missing_candidates_sha256"] = sha256_file(journal_path)
    _atomic_write(summary_path, lambda stream: json.dump(result, stream, ensure_ascii=False, indent=2, default=str))
    return result


def audit_saved_run(input_path, excel_path, pipeline_audit, output_prefix, *, overwrite=False):
    """독립 실행용. 저장된 객체를 재사용; dotenv/모델/네트워크/문서 재판독 없음."""
    started = time.perf_counter()
    link_path = Path(output_prefix).with_name(Path(output_prefix).name + "_pipeline_audit.json")
    if link_path.resolve() == Path(pipeline_audit).resolve():
        raise ValueError("원본 pipeline_audit는 변경하지 않습니다. 새 output-prefix를 사용하세요.")
    if link_path.exists() and not overwrite:
        raise FileExistsError("감사 연결 파일이 이미 있습니다. 새 output-prefix를 사용하세요.")
    audit = json.loads(Path(pipeline_audit).read_text(encoding="utf-8-sig"))
    digest = sha256_file(input_path)
    expected = audit.get("context", {}).get("input_sha256")
    if not expected or expected != digest:
        raise ValueError("문서 SHA256 불일치 또는 누락: 다른 문서의 객체를 적용하지 않습니다.")
    saved_excel = audit.get("excel_path", "")
    expected_excel_hash = audit.get("context", {}).get("source_coverage", {}).get("excel_sha256")
    actual_excel_hash = sha256_file(excel_path)
    if expected_excel_hash and expected_excel_hash != actual_excel_hash:
        raise ValueError("Excel SHA256 불일치")
    if not expected_excel_hash and Path(saved_excel.replace('\\', '/')).name != Path(excel_path).name:
        raise ValueError("과거 감사 기록의 Excel 파일명과 불일치: 실행 대응을 먼저 확인하세요.")
    ledgers = audit.get("ledgers", {})
    capture = ledgers.get("source_coverage_snapshot") or {}
    objects = capture.get("source_objects") or ledgers.get("document_objects")
    if not isinstance(objects, list) or not objects:
        raise ValueError("저장된 document_objects가 없습니다. 자동 재판독하지 않습니다.")
    provenance = {"producer_run": Path(saved_excel.replace('\\', '/')).stem,
                  "producer_input_path": audit.get("input_path"),
                  "source_pipeline_audit_path": str(Path(pipeline_audit).resolve()),
                  "source_excel_path": str(Path(excel_path).resolve()),
                  "source_snapshot": "pre_extraction_native_objects" if capture.get("source_objects") else "legacy_post_extraction_objects",
                  "producer_coverage_provenance": audit.get("context", {}).get("source_coverage", {}).get("provenance"),
                  "auditor_rule_version": VERSION, "auditor_module_sha256": sha256_file(__file__),
                  "excel_sha256": actual_excel_hash, "pipeline_audit_sha256": sha256_file(pipeline_audit),
                  "excel_binding": "sha256" if expected_excel_hash else "legacy_filename_only_unverified_hash"}
    journal, summary = compare(objects, digest, read_excel(excel_path),
        raw=capture.get("raw"), cleaned=capture.get("cleaned"),
        records=ledgers.get("extraction_records", []), routes=capture.get("routes", []),
        quarantine=ledgers.get("_semantic_contract_quarantine", []), provenance=provenance)
    summary["audit_elapsed_seconds_before_persistence"] = round(time.perf_counter() - started, 6)
    summary["pipeline_audit_link_path"] = str(link_path.resolve())
    saved = persist(output_prefix, journal, summary, overwrite=overwrite)
    link = {"schema_version": "source-coverage-link-v1", "original_audit_unchanged": True,
            "parent_pipeline_audit": str(Path(pipeline_audit).resolve()),
            "parent_pipeline_audit_sha256": provenance["pipeline_audit_sha256"],
            "context": {"input_sha256": digest, "source_coverage": saved},
            "excel_path": str(Path(excel_path).resolve())}
    _atomic_write(link_path, lambda stream: json.dump(link, stream, ensure_ascii=False, indent=2))
    return saved
