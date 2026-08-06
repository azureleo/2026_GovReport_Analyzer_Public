"""고정 골든셋·홀드아웃 계약에 따라 파이프라인 출력을 평가한다."""

from __future__ import annotations

import hashlib
import json
import re
import time
from dataclasses import dataclass
from pathlib import Path
from typing import Any

from openpyxl import load_workbook

from scripts.audit_visual_inventory import AuditPaths, run_audit
from scripts.score_against_golden import score_workbooks
from utils.routing_benchmark import evaluate_fixed_routing_inventory


class EvaluationContractError(RuntimeError):
    """평가 데이터셋 계약 또는 파일 해시가 유효하지 않다."""


@dataclass(frozen=True, slots=True)
class EvaluationDataset:
    dataset_id: str
    role: str
    review_status: str
    locked: bool
    source_path: Path
    source_sha256: str
    golden_path: Path | None
    golden_sha256: str
    visual_inventory_path: Path | None
    visual_inventory_sha256: str
    routing_inventory_path: Path | None
    routing_inventory_sha256: str
    notes: str = ""


@dataclass(frozen=True, slots=True)
class BenchmarkEvaluationResult:
    dataset: EvaluationDataset
    metrics: dict[str, Any]
    report_json: Path
    report_markdown: Path
    artifacts: dict[str, str]


def sha256_file(path: str | Path) -> str:
    digest = hashlib.sha256()
    with Path(path).open("rb") as stream:
        for chunk in iter(lambda: stream.read(1024 * 1024), b""):
            digest.update(chunk)
    return digest.hexdigest().lower()


def _resolve_path(base: Path, value: Any) -> Path | None:
    text = str(value or "").strip()
    if not text:
        return None
    path = Path(text).expanduser()
    return path.resolve() if path.is_absolute() else (base / path).resolve()


def load_evaluation_datasets(manifest_path: str | Path) -> list[EvaluationDataset]:
    manifest = Path(manifest_path).expanduser().resolve()
    try:
        payload = json.loads(manifest.read_text(encoding="utf-8"))
    except FileNotFoundError as exc:
        raise EvaluationContractError(f"평가 매니페스트를 찾을 수 없습니다: {manifest}") from exc
    except json.JSONDecodeError as exc:
        raise EvaluationContractError(f"평가 매니페스트 JSON 오류: {exc}") from exc

    raw_datasets = payload.get("datasets")
    if not isinstance(raw_datasets, list):
        raise EvaluationContractError("평가 매니페스트의 datasets는 배열이어야 합니다.")
    datasets: list[EvaluationDataset] = []
    for index, raw in enumerate(raw_datasets, start=1):
        if not isinstance(raw, dict):
            raise EvaluationContractError(f"datasets[{index}]는 객체여야 합니다.")
        dataset_id = str(raw.get("id") or "").strip()
        role = str(raw.get("role") or "").strip().lower()
        if not dataset_id or role not in {"development", "holdout"}:
            raise EvaluationContractError(
                f"datasets[{index}]에 id와 role(development|holdout)이 필요합니다."
            )
        source_path = _resolve_path(manifest.parent, raw.get("source"))
        if source_path is None:
            raise EvaluationContractError(f"{dataset_id}: source 경로가 필요합니다.")
        datasets.append(EvaluationDataset(
            dataset_id=dataset_id,
            role=role,
            review_status=str(raw.get("review_status") or "draft").strip().lower(),
            locked=bool(raw.get("locked", False)),
            source_path=source_path,
            source_sha256=str(raw.get("source_sha256") or "").strip().lower(),
            golden_path=_resolve_path(manifest.parent, raw.get("golden")),
            golden_sha256=str(raw.get("golden_sha256") or "").strip().lower(),
            visual_inventory_path=_resolve_path(manifest.parent, raw.get("visual_inventory")),
            visual_inventory_sha256=str(
                raw.get("visual_inventory_sha256") or ""
            ).strip().lower(),
            routing_inventory_path=_resolve_path(manifest.parent, raw.get("routing_inventory")),
            routing_inventory_sha256=str(
                raw.get("routing_inventory_sha256") or ""
            ).strip().lower(),
            notes=str(raw.get("notes") or "").strip(),
        ))
    return datasets


def select_evaluation_dataset(
    datasets: list[EvaluationDataset],
    source_path: str | Path,
    *,
    dataset_id: str | None = None,
    allow_holdout: bool = False,
) -> EvaluationDataset:
    source = Path(source_path).expanduser().resolve()
    source_hash = sha256_file(source)
    if dataset_id:
        matches = [dataset for dataset in datasets if dataset.dataset_id == dataset_id]
    else:
        matches = [
            dataset for dataset in datasets
            if dataset.source_sha256 and dataset.source_sha256 == source_hash
        ]
    if not matches:
        marker = f"ID={dataset_id}" if dataset_id else f"원문 SHA256={source_hash}"
        raise EvaluationContractError(f"일치하는 평가 데이터셋이 없습니다: {marker}")
    if len(matches) > 1:
        raise EvaluationContractError("평가 데이터셋이 둘 이상 일치합니다. ID를 명시하세요.")
    dataset = matches[0]
    if dataset.role == "holdout" and not allow_holdout:
        raise EvaluationContractError(
            f"{dataset.dataset_id}는 홀드아웃입니다. 일상 튜닝과 분리하기 위해 "
            "--evaluate-holdout을 명시해야 합니다."
        )
    return dataset


def _verify_file(path: Path | None, expected_hash: str, label: str) -> None:
    if path is None:
        raise EvaluationContractError(f"{label} 경로가 등록되지 않았습니다.")
    if not path.exists():
        raise EvaluationContractError(f"{label} 파일을 찾을 수 없습니다: {path}")
    if not expected_hash:
        raise EvaluationContractError(f"{label} SHA256이 매니페스트에 없습니다.")
    actual = sha256_file(path)
    if actual != expected_hash:
        raise EvaluationContractError(
            f"{label} 해시 불일치: 기대 {expected_hash}, 실제 {actual}. "
            "파일을 조용히 수정하지 말고 새 데이터셋 버전을 등록하세요."
        )


def verify_evaluation_contract(
    dataset: EvaluationDataset,
    source_path: str | Path,
    *,
    allow_draft: bool = False,
) -> None:
    source = Path(source_path).expanduser().resolve()
    if not dataset.locked:
        raise EvaluationContractError(
            f"{dataset.dataset_id}는 locked=false입니다. 고정 평가 기준으로 사용할 수 없습니다."
        )
    _verify_file(source, dataset.source_sha256, "입력 원문")
    _verify_file(dataset.golden_path, dataset.golden_sha256, "골든셋")
    if dataset.review_status != "verified" and not allow_draft:
        raise EvaluationContractError(
            f"{dataset.dataset_id}는 사람 확정 전({dataset.review_status})입니다. "
            "테스트 목적이면 --allow-draft-evaluation을 명시하세요."
        )
    if dataset.visual_inventory_path is not None:
        _verify_file(
            dataset.visual_inventory_path,
            dataset.visual_inventory_sha256,
            "시각요소 인벤토리",
        )
    if dataset.routing_inventory_path is not None:
        _verify_file(
            dataset.routing_inventory_path,
            dataset.routing_inventory_sha256,
            "라우팅 인벤토리",
        )


def routing_metrics_from_workbook(
    output_path: str | Path,
    routing_inventory_path: str | Path | None = None,
) -> dict[str, Any]:
    if routing_inventory_path is not None:
        metrics, _details = evaluate_fixed_routing_inventory(
            output_path,
            routing_inventory_path,
        )
        return metrics
    workbook = load_workbook(output_path, read_only=True, data_only=True)
    try:
        if "21_원문객체인벤토리" not in workbook.sheetnames:
            return {
                "routing_error_rate": None,
                "routing_evaluable_objects": 0,
                "routing_mismatch_objects": 0,
                "routing_missing_objects": 0,
                "routing_fixed_denominator": False,
            }
        sheet = workbook["21_원문객체인벤토리"]
        headers = [str(cell.value or "").strip() for cell in sheet[1]]
        if "시트정합상태" not in headers:
            return {
                "routing_error_rate": None,
                "routing_evaluable_objects": 0,
                "routing_mismatch_objects": 0,
                "routing_missing_objects": 0,
                "routing_fixed_denominator": False,
            }
        status_index = headers.index("시트정합상태")
        statuses = [
            str(values[status_index] or "").strip()
            for values in sheet.iter_rows(min_row=2, values_only=True)
            if status_index < len(values)
        ]
        evaluable_statuses = {"일치", "오배치의심", "본문미연결"}
        evaluable = sum(status in evaluable_statuses for status in statuses)
        mismatches = sum(status in {"오배치의심", "본문미연결"} for status in statuses)
        return {
            "routing_error_rate": round(mismatches / evaluable, 4) if evaluable else None,
            "routing_evaluable_objects": evaluable,
            "routing_mismatch_objects": mismatches,
            "routing_missing_objects": 0,
            "routing_fixed_denominator": False,
        }
    finally:
        workbook.close()


def _visual_metrics(
    dataset: EvaluationDataset,
    output_path: Path,
    source_path: Path,
    report_dir: Path,
) -> tuple[dict[str, Any], dict[str, str]]:
    if dataset.visual_inventory_path is None:
        return {
            "object_recall": None,
            "object_expected": 0,
            "object_recorded": 0,
        }, {}
    reports = run_audit(AuditPaths(
        inventory=dataset.visual_inventory_path,
        pipeline_output=output_path,
        source_pdf=source_path,
        report_dir=report_dir / "visual",
    ))
    payload = json.loads(reports.json_path.read_text(encoding="utf-8"))
    expected = int(payload.get("expected_rows", 0) or 0)
    recorded = int(payload.get("recorded_rows", 0) or 0)
    return {
        "object_recall": round(recorded / expected, 4) if expected else None,
        "object_expected": expected,
        "object_recorded": recorded,
    }, {
        "visual_audit_json": str(reports.json_path),
        "visual_audit_markdown": str(reports.markdown),
    }


def _safe_label(value: str) -> str:
    return re.sub(r"[^0-9A-Za-z가-힣_.-]+", "_", value).strip("_") or "evaluation"


def _routing_metrics(
    dataset: EvaluationDataset,
    output_path: Path,
    report_dir: Path,
) -> tuple[dict[str, Any], dict[str, str]]:
    if dataset.routing_inventory_path is None:
        return routing_metrics_from_workbook(output_path), {}
    metrics, details = evaluate_fixed_routing_inventory(
        output_path,
        dataset.routing_inventory_path,
    )
    payload = {
        "schema_version": 1,
        "dataset_id": dataset.dataset_id,
        "routing_inventory": str(dataset.routing_inventory_path),
        "metrics": metrics,
        "objects": [detail.to_dict() for detail in details],
    }
    json_path = report_dir / "routing_fixed_inventory.json"
    markdown_path = report_dir / "routing_fixed_inventory.md"
    json_path.write_text(
        json.dumps(payload, ensure_ascii=False, sort_keys=True, indent=2),
        encoding="utf-8",
    )
    lines = [
        f"# 고정 라우팅 평가: {dataset.dataset_id}",
        "",
        "| 객체ID | 페이지 | 허용 시트 | 연결 본문 시트 | 판정 |",
        "|---|---:|---|---|---|",
    ]
    lines.extend(
        "| {id} | {page} | {allowed} | {linked} | {status} |".format(
            id=detail.item_id,
            page=detail.page,
            allowed=", ".join(detail.allowed_sheets),
            linked=", ".join(detail.linked_body_sheets) or "-",
            status=detail.status,
        )
        for detail in details
    )
    markdown_path.write_text("\n".join(lines) + "\n", encoding="utf-8")
    return metrics, {
        "routing_evaluation_json": str(json_path),
        "routing_evaluation_markdown": str(markdown_path),
    }


def _markdown(payload: dict[str, Any]) -> str:
    metrics = payload["metrics"]

    def ratio(key: str) -> str:
        value = metrics.get(key)
        return "-" if value is None else f"{float(value):.2%}"

    lines = [
        f"# 고정 평가 리포트: {payload['dataset_id']}",
        "",
        f"- 역할: `{payload['role']}`",
        f"- 골든 검토 상태: `{payload['review_status']}`",
        f"- 출력: `{payload['output']}`",
        f"- 원문: `{payload['source']}`",
        "",
        "## 독립 지표",
        "",
        "| 지표 | 결과 | 분모 |",
        "|---|---:|---:|",
        f"| 셀 정확도 | {ratio('cell_accuracy')} | {metrics.get('cell_expected', 0)}셀 |",
        f"| 객체 재현율 | {ratio('object_recall')} | {metrics.get('object_expected', 0)}개 |",
        f"| 라우팅 오류율 | {ratio('routing_error_rate')} | {metrics.get('routing_evaluable_objects', 0)}개 |",
        f"| 행 리콜 | {ratio('row_recall')} | {metrics.get('golden_rows', 0)}행 |",
        f"| 행 정밀도 | {ratio('row_precision')} | {metrics.get('output_rows', 0)}행 |",
        "",
        "셀 정확도는 골든에 값이 있으나 출력에서 비어 있는 셀도 오답으로 계산합니다. "
        "객체 재현율은 사람 확정 시각요소 인벤토리를 분모로 사용합니다. "
        "라우팅 인벤토리가 등록된 데이터셋은 미추출 객체를 포함한 고정 분모를, "
        "미등록 문서는 운영 객체의 동적 지표를 사용합니다.",
        "",
    ]
    return "\n".join(lines)


def evaluate_benchmark(
    output_path: str | Path,
    source_path: str | Path,
    *,
    manifest_path: str | Path,
    dataset_id: str | None = None,
    allow_draft: bool = False,
    allow_holdout: bool = False,
    report_dir: str | Path | None = None,
) -> BenchmarkEvaluationResult:
    output = Path(output_path).expanduser().resolve()
    source = Path(source_path).expanduser().resolve()
    manifest = Path(manifest_path).expanduser().resolve()
    datasets = load_evaluation_datasets(manifest)
    dataset = select_evaluation_dataset(
        datasets,
        source,
        dataset_id=dataset_id,
        allow_holdout=allow_holdout,
    )
    verify_evaluation_contract(dataset, source, allow_draft=allow_draft)
    if not output.exists():
        raise EvaluationContractError(f"평가할 출력 파일이 없습니다: {output}")

    target_dir = (
        Path(report_dir).expanduser().resolve()
        if report_dir is not None
        else output.parent / f"{output.stem}_evaluation"
    )
    target_dir.mkdir(parents=True, exist_ok=True)
    score = score_workbooks(
        output,
        dataset.golden_path,
        report_dir=target_dir / "golden",
        label=dataset.dataset_id,
        write_json=True,
    )
    if score.get("형식오류"):
        raise EvaluationContractError(
            "골든셋 형식 오류: " + "; ".join(str(item) for item in score["형식오류"][:5])
        )

    cell_metrics = dict(score.get("독립평가지표") or {})
    visual_metrics, visual_artifacts = _visual_metrics(
        dataset, output, source, target_dir
    )
    routing_metrics, routing_artifacts = _routing_metrics(dataset, output, target_dir)
    metrics = {**cell_metrics, **visual_metrics, **routing_metrics}
    artifacts = {
        "golden_score_json": str(score.get("리포트", {}).get("json", "")),
        "golden_score_markdown": str(score.get("리포트", {}).get("md", "")),
        **visual_artifacts,
        **routing_artifacts,
    }
    payload = {
        "schema_version": 1,
        "dataset_id": dataset.dataset_id,
        "role": dataset.role,
        "review_status": dataset.review_status,
        "locked": dataset.locked,
        "evaluated_at": time.strftime("%Y-%m-%d %H:%M:%S"),
        "source": str(source),
        "source_sha256": sha256_file(source),
        "output": str(output),
        "output_sha256": sha256_file(output),
        "golden": str(dataset.golden_path),
        "golden_sha256": dataset.golden_sha256,
        "visual_inventory": (
            str(dataset.visual_inventory_path)
            if dataset.visual_inventory_path is not None
            else None
        ),
        "routing_inventory": (
            str(dataset.routing_inventory_path)
            if dataset.routing_inventory_path is not None
            else None
        ),
        "metrics": metrics,
        "artifacts": artifacts,
        "notes": dataset.notes,
    }
    label = _safe_label(dataset.dataset_id)
    json_path = target_dir / f"benchmark_evaluation_{label}.json"
    markdown_path = target_dir / f"benchmark_evaluation_{label}.md"
    json_path.write_text(
        json.dumps(
            payload,
            ensure_ascii=False,
            sort_keys=True,
            indent=2,
            default=str,
        ),
        encoding="utf-8",
    )
    markdown_path.write_text(_markdown(payload), encoding="utf-8")
    return BenchmarkEvaluationResult(
        dataset=dataset,
        metrics=metrics,
        report_json=json_path,
        report_markdown=markdown_path,
        artifacts=artifacts,
    )
