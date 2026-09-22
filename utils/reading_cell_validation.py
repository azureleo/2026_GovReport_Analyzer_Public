"""Independent, evidence-keyed comparisons for the Stage A smoke run."""
from collections import defaultdict
import json


def equal(a, b):
    if isinstance(a, bool) or isinstance(b, bool):
        return type(a) is type(b) and a == b
    return a == b  # int/float equality is intentional; strings are never coerced.


def index_rows(data, keys):
    result = defaultdict(list)
    for key in keys:
        for row in data.get(key, []):
            result[(key, row.get('근거ID'))].append(row)
    return result


def compare_backend(before, after, keys):
    left, right = index_rows(before, keys), index_rows(after, keys)
    changes = []
    for identity in sorted(set(left) | set(right), key=str):
        a, b = left.get(identity, []), right.get(identity, [])
        if len(a) != 1 or len(b) != 1:
            changes.append(dict(sheet=identity[0], evidence_id=identity[1], kind='row_count_or_identity', before=a, after=b))
            continue
        for field in sorted(set(a[0]) | set(b[0])):
            if not equal(a[0].get(field), b[0].get(field)):
                changes.append(dict(sheet=identity[0], evidence_id=identity[1], field=field,
                                    before=a[0].get(field), after=b[0].get(field)))
    return changes


def inspect_workbook(path, data, schema):
    import openpyxl
    wb = openpyxl.load_workbook(path, read_only=True, data_only=False)
    issues, cells = [], []
    try:
        for key in schema['EXTRACTION_SHEETS']:
            name = schema['SHEET_KEY_TO_NAME'][key]
            if name not in wb:
                issues.append(dict(kind='missing_sheet', sheet=name))
                continue
            ws = wb[name]
            headers = [c.value for c in next(ws.iter_rows())]
            expected_headers = schema['EXCEL_HEADERS'][name]
            if headers != expected_headers:
                issues.append(dict(kind='header_mismatch', sheet=name, expected=expected_headers, actual=headers))
            if '근거ID' not in headers:
                issues.append(dict(kind='missing_evidence_column', sheet=name))
                continue
            expected = index_rows(data, [key])
            seen = defaultdict(int)
            for row in ws.iter_rows(min_row=2):
                if not any(c.value is not None for c in row):
                    continue
                ev = row[headers.index('근거ID')].value
                identity = (key, ev)
                seen[identity] += 1
                matches = expected.get(identity, [])
                if len(matches) != 1:
                    issues.append(dict(kind='unexpected_or_ambiguous_row', sheet=name, row=row[0].row, evidence_id=ev))
                    continue
                source = matches[0]
                for field in source:
                    if field not in headers and not (field == '감축률' and '감축률(%)' in headers) and source[field] is not None:
                        issues.append(dict(kind='unrepresented_field', sheet=name, evidence_id=ev, field=field))
                for header, cell in zip(headers, row):
                    value = source.get(header)
                    if value is None and header == '감축률(%)':
                        value = source.get('감축률')
                    if isinstance(value, bool):
                        value = 'Y' if value else 'N'
                    elif isinstance(value, (list, dict)):
                        value = json.dumps(value, ensure_ascii=False, sort_keys=True)
                    elif value == '':
                        value = None
                    ok = equal(value, cell.value) and cell.data_type not in ('e', 'f')
                    entry = dict(sheet=name, cell=cell.coordinate, evidence_id=ev, field=header,
                                 expected=value, actual=cell.value, stored_type=cell.data_type, match=ok)
                    cells.append(entry)
                    if not ok:
                        issues.append(dict(kind='cell_mismatch', **entry))
            for identity in expected:
                if seen[identity] != 1:
                    issues.append(dict(kind='missing_or_duplicate_row', sheet=name, evidence_id=identity[1], actual_count=seen[identity]))
        for ws in wb:
            for row in ws:
                for cell in row:
                    if cell.data_type in ('e', 'f'):
                        issues.append(dict(kind='unexpected_formula_or_error', sheet=ws.title, cell=cell.coordinate))
    finally:
        wb.close()
    return dict(passed=not issues, issues=issues, cells=cells, checked_cells=len(cells))
