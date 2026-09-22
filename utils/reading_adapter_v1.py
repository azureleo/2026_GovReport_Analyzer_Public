"""Frozen-input, offline adapter. No PDF/image access and no model API.

Semantic decisions are the model-authored rules in mapping_rules_v1.json.
Backend source functions are invoked through backend_bridge, not replaced.
"""
from __future__ import annotations
import collections, copy, hashlib, json, math, re, time
from pathlib import Path
import openpyxl
from utils.reading_position import same_cell_key

ROOT = Path(__file__).resolve().parent

def canonical(v):
    return json.dumps(v, ensure_ascii=False, sort_keys=True, separators=(',', ':'), allow_nan=False)
def digest(v): return hashlib.sha256(canonical(v).encode('utf-8')).hexdigest()
def filehash(p): return hashlib.sha256(Path(p).read_bytes()).hexdigest()
def text(v):
    if v is None: return ''
    if isinstance(v, dict): return ' '.join(text(x) for x in v.values())
    if isinstance(v, list): return ' '.join(text(x) for x in v)
    return str(v)
def has(p, v): return bool(re.search(p, text(v)))
def numeric(v): return isinstance(v, (int,float)) and not isinstance(v,bool) and math.isfinite(v)
def nonempty(v): return v is not None and v != '' and v != []
def dump(p,v): Path(p).write_text(json.dumps(v,ensure_ascii=False,indent=2,allow_nan=False)+'\n',encoding='utf-8')

def load_fixed():
    start=time.perf_counter()
    path=ROOT/'provided/첨부파일/reading_only.xlsx'
    expected=json.loads((ROOT/'input_check.json').read_text(encoding='utf-8'))
    if filehash(path)!=expected['xlsx_sha256']: raise RuntimeError('Fixed XLSX hash changed')
    def decode(v):
        if isinstance(v,str) and v.startswith(('[','{','"')):
            try: return json.loads(v)
            except ValueError: pass
        return v
    wb=openpyxl.load_workbook(path,read_only=True,data_only=False,keep_links=False)
    source={}
    for key in ('objects','records','relations'):
        rows=list(wb[key].iter_rows(values_only=True)); headers=rows[0]
        source[key]=[{h:decode(v) for h,v in zip(headers,row)} for row in rows[1:] if any(v is not None for v in row)]
    rows=list(wb['run'].iter_rows(values_only=True));source['run']={r[0]:decode(r[1]) for r in rows[1:] if r[0]};wb.close()
    companion=ROOT/'provided/첨부파일/reading_only.json'
    if companion.exists():
        if filehash(companion)!=expected['companion_json_sha256']: raise RuntimeError('Companion JSON hash changed')
        other=json.loads(companion.read_text(encoding='utf-8'))
        for key,idkey in [('objects','객체ID'),('records','레코드ID'),('relations','관계ID')]:
            om={r[idkey]:r for r in other[key]}
            if len(source[key])!=len(om):raise RuntimeError('Input count mismatch')
            for row in source[key]:
                orig=om[row[idkey]]
                if any(row.get(k)!=orig.get(k) for k in set(row)|set(orig)):raise RuntimeError('XLSX/JSON mismatch')
        if source['run']!=other['run']: raise RuntimeError('Input run mismatch')
    return source, time.perf_counter()-start

def classify(r,rules):
    c=rules['classification'];ctx=r.get('문맥 구분') or {}
    semantic=' '.join(text(v) for v in [ctx.get('값의미'),r.get('항목원문'),r.get('행머리글 경로'),r.get('열머리글 경로')])
    item=r.get('항목원문');unit=r.get('단위원문');kind=r.get('레코드종류')
    a=c['C01_identifier']
    if has(a['semantic_pattern'],semantic) or has(a['item_pattern'],item) or has(a['unit_pattern'],unit):return 'identifier','C01_identifier'
    a=c['C02_qualifier']
    if has(a['semantic_pattern'],semantic) or has(a['item_pattern'],item) or has(a['unit_pattern'],unit) or (kind=='노드' and has(a['node_location_pattern'],r.get('객체 내 위치'))):return 'qualifier','C02_qualifier'
    if kind in ('노드','문자값') and has(c['C03_unreadable_text']['pattern'],[item,r.get('값원문')]):return 'unresolved','C03_unreadable_text'
    if kind in c['C04_measurement']['record_kinds']:return 'measurement','C04_measurement'
    if kind in c['C05_node']['record_kinds']:return 'node_or_relation','C05_node'
    if kind in c['C06_text_fact']['record_kinds']:return 'text_fact','C06_text_fact'
    return 'unresolved','C07_unresolved'

def scope(r,o,rules):
    ctx=r.get('문맥 구분') or {};target=rules['target']
    explicit=text([ctx.get('자료구분'),ctx.get('자료지역')])
    if has(target['reference_pattern'],explicit):return 'reference',explicit
    if has(target['own_pattern'],explicit) or ctx.get('자료지역') in target['source_aliases']:return 'own',explicit
    common=o.get('공통문맥') or {};hint=text([o.get('제목'),common.get('자료구분'),common.get('자료지역')])
    if has(target['reference_pattern'],hint):return 'reference',hint
    if has(target['own_pattern'],hint):return 'own',hint
    return 'unknown',explicit or hint

def normalize(source,rules):
    objects={o['객체ID']:o for o in source['objects']};items=[]
    for r in source['records']:
        cl,rule=classify(r,rules);sc,evidence=scope(r,objects[r['객체ID']],rules)
        items.append({'id':'N-'+r['레코드ID'],'source_record_ids':[r['레코드ID']], 'object_id':r['객체ID'],
                      'classification':cl,'classification_rule':rule,'scope':sc,'scope_evidence':evidence,
                      'original':copy.deepcopy(r),'attributes':{'year':r.get('연도'),'period':r.get('기간원문'),'unit':r.get('단위원문')},
                      'metadata_links':[],'attribute_conflicts':[]})
    byid={i['source_record_ids'][0]:i for i in items};rels=source['relations'];incoming=collections.defaultdict(list)
    for rel in rels: incoming[rel['도착 레코드 또는 노드ID']].append(rel)
    def link(target,q,rule,relation_ids):
        if target==q or q not in byid or target not in byid:return
        if byid[q]['classification']!='qualifier':return
        entry={'source_record_id':q,'rule_id':rule,'relation_ids':relation_ids}
        if entry not in byid[target]['metadata_links']:byid[target]['metadata_links'].append(entry)
    cells=collections.defaultdict(list)
    for i in items:
        pos=i['original'].get('객체 내 위치')
        try:
            cell_key = same_cell_key(i['object_id'], pos)
        except ValueError as exc:
            raise ValueError(f"Invalid position for {i['source_record_ids'][0]}: {exc}") from exc
        if cell_key is not None:
            cells[cell_key].append(i)
    for cell in cells.values():
        for m in cell:
            if m['classification']=='measurement':
                for q in cell:
                    if q['classification']=='qualifier':link(m['source_record_ids'][0],q['source_record_ids'][0],'L02_same_cell',[])
    for rel in rels:
        a,b=rel['출발 레코드 또는 노드ID'],rel['도착 레코드 또는 노드ID']
        if a in byid and b in byid:
            if byid[a]['classification']=='qualifier':link(b,a,'L03_explicit_relation',[rel['관계ID']])
            if byid[b]['classification']=='qualifier':link(a,b,'L03_explicit_relation',[rel['관계ID']])
    lr=rules['linking']['L04_ancestor_period']
    for i in items:
        if i['classification']!='measurement':continue
        rid=i['source_record_ids'][0];front=[(rid,[])];seen={rid}
        for depth in range(lr['max_depth']):
            nxt=[]
            for child,path in front:
                for rel in incoming[child]:
                    parent=rel['출발 레코드 또는 노드ID']
                    if parent in seen or not has(lr['containment_pattern'],rel['관계원문']):continue
                    seen.add(parent);newpath=path+[rel['관계ID']];nxt.append((parent,newpath))
                    for qrel in rels:
                        if qrel['출발 레코드 또는 노드ID']==parent and has(lr['period_pattern'],qrel['관계원문']):
                            link(rid,qrel['도착 레코드 또는 노드ID'],'L04_ancestor_period',newpath+[qrel['관계ID']])
            front=nxt
    # Preserve linked qualifiers as typed evidence; do not overwrite embedded attributes.
    for i in items:
        alternatives=collections.defaultdict(set)
        for l in i['metadata_links']:
            qr=byid[l['source_record_id']]['original']
            role=text(qr.get('항목원문'))
            alternatives[role].add(canonical(qr.get('값원문')))
        i['attribute_conflicts']=[{'role':k,'alternatives':[json.loads(x) for x in sorted(v)]} for k,v in alternatives.items() if len(v)>1]
    return items

def make_candidates(source,items,rules,schema):
    objects={o['객체ID']:o for o in source['objects']};candidates=[];reasons={}
    for i in items:
        rid=i['source_record_ids'][0];r=i['original'];o=objects[i['object_id']];ctx=r.get('문맥 구분') or {};common=o.get('공통문맥') or {};cl=i['classification']
        if i['scope']!='own':
            if i['scope']=='unknown':reasons[rid]='서울 자체/참고자료 범위가 고정 입력에서 명확하지 않음'
            continue
        if i['attribute_conflicts']:reasons[rid]='명시 연결된 설명 정보가 서로 충돌함';continue
        if cl=='measurement' and (not numeric(r.get('숫자값')) or r.get('값상태')!='명시값'):
            reasons[rid]='명시 숫자 라벨 없음 또는 판독불가; 원본 null을 유지';continue
        row=None;key=None;rule=None
        if cl=='measurement':
            mr=rules['mapping']['M02_regional'];s=text([r.get('행머리글 경로'),r.get('열머리글 경로'),ctx.get('값의미'),o.get('제목')])
            category=next((a['category'] for a in mr['category_rules'] if has(a['pattern'],s)),None)
            if category and not has(mr['exclude_role_pattern'],s):
                labels=[text(v) for v in (r.get('행머리글 경로') or [])+(r.get('열머리글 경로') or []) if not re.fullmatch(r'\d{4}년?',text(v))]
                row={'지자체명':rules['target']['municipality'],'지표범주':category,'지표세부범주':ctx.get('값의미'),'지표명':' / '.join(labels),'연도':r.get('연도'),'값':r.get('숫자값'),'단위':r.get('단위원문'),'출처':common.get('출처')}
                key=mr['sheet'];rule='M02_regional'
        if row is None and cl=='text_fact':
            mr=rules['mapping']['M01_consultation']
            if has(mr['context_pattern'],common) and has(mr['activity_pattern'],o.get('제목')):
                key=mr['sheet'];rule='M01_consultation';row={'지자체명':rules['target']['municipality'],'개요유형':mr['overview_type'],'항목명':o.get('제목')+' / '+text(r.get('항목원문')),'항목값':r.get('값원문')}
            else:
                mr=rules['mapping']['M13_governance'];org=re.search(mr['organization_pattern'],text(o.get('제목')))
                if org and has(mr['role_context_pattern'],o.get('제목')):
                    key=mr['sheet'];rule='M13_governance';row={'지자체명':rules['target']['municipality'],'거버넌스기구':org.group(0),'역할':r.get('값원문'),'절차단계':' / '.join(r.get('행머리글 경로') or [])}
        if row is None and isinstance(ctx.get('업무필드'),dict) and ctx.get('대상시트') in schema['EXTRACTION_SHEETS']:
            key=ctx['대상시트'];rule='M_direct_contract';heads=schema['EXCEL_HEADERS'][schema['SHEET_KEY_TO_NAME'][key]]
            row={k:v for k,v in ctx['업무필드'].items() if k in heads}
            # A direct-contract candidate must also retain the original quantitative cell.
            if cl=='measurement' and r.get('숫자값') not in [v for v in row.values() if numeric(v)]:row=None
        if row is None:
            if cl in ('measurement','unresolved','qualifier','identifier'):reasons[rid]='고정 입력에 업무 계약의 필수 식별자·명시 목표 성격 또는 안전한 매핑 규칙이 없음'
            continue
        missing=[f for f in rules['required_fields'][key] if not nonempty(row.get(f))]
        if missing:reasons[rid]='필수 필드 부족: '+', '.join(missing);continue
        allowed=rules['units']['allowed'].get(key)
        unit=row.get('단위',row.get('예산단위',row.get('목표단위')))
        if cl=='measurement' and not nonempty(unit):reasons[rid]='원문 단위 없음';continue
        if allowed and unit not in allowed:reasons[rid]='출력 계약 단위 불일치; 환산 금지';continue
        row.update({'근거ID':rid,'출처페이지':o.get('PDF 실제 페이지'),'derivation_type':'normalized','데이터상태':'고정 판독 연결; 원문 재검증 미수행'})
        candidates.append({'sheet_key':key,'sheet_name':schema['SHEET_KEY_TO_NAME'][key],'source_record_ids':[rid], 'values':row,'field_sources':{k:[rid] for k,v in row.items() if nonempty(v)},'rule_id':rule,'object_id':i['object_id']})
    # Protect the legacy regional key before the backend can coalesce distinct cells.
    groups=collections.defaultdict(list)
    for c in candidates:
        r=c['values'];k=c['sheet_key']
        fields=['지자체명','지표범주','지표명','연도'] if k=='regional_conditions' else [x for x in r if x not in ('근거ID','출처페이지','derivation_type','데이터상태')]
        groups[(k,canonical([r.get(f) for f in fields]))].append(c)
    byid={i['source_record_ids'][0]:i for i in items};safe=[];duplicates=[]
    for group in groups.values():
        if len(group)==1:safe+=group;continue
        def signature(c):
            r=byid[c['source_record_ids'][0]]['original']
            return canonical([r.get(k) for k in ('객체ID','행머리글 경로','열머리글 경로','연도','기간원문','단위원문','문맥 구분','값원문','숫자값')])
        if len({signature(c) for c in group})==1:
            c=copy.deepcopy(group[0]);ids=[rid for g in group for rid in g['source_record_ids']];c['source_record_ids']=ids;c['values']['근거ID']=';'.join(ids)
            c['field_sources']={k:ids for k in c['field_sources']};safe.append(c);duplicates.append({'rule':'D01','source_record_ids':ids,'action':'동일 객체·동일 의미·동일 값 중복 통합'})
        else:
            ids=[rid for g in group for rid in g['source_record_ids']]
            for rid in ids:reasons[rid]='D02: 기존 후단 중복 키가 서로 다른 객체·범위·값을 구분하지 못함'
            duplicates.append({'rule':'D02','source_record_ids':ids,'action':'후보 전부 보류','candidates':group})
    return safe,reasons,duplicates

def apply_backend(candidates,items,rules,schema,bridge):
    payload={k:[] for k in schema['EXTRACTION_SHEETS']};payload['municipality_name']=rules['target']['municipality']
    for c in candidates:payload[c['sheet_key']].append(copy.deepcopy(c['values']))
    cleaned,backend_info=bridge.organize(payload)
    original={c['values']['근거ID']:c for c in candidates};retained=[];changes=[];reasons={};seen=set()
    for k in schema['EXTRACTION_SHEETS']:
        for row in cleaned.get(k,[]):
            evidence=row.get('근거ID');c=original.get(evidence)
            if c is None:
                changes.append({'type':'source_id_lost','returned_row':row,'sheet':k});continue
            issues=[]
            for f,v in row.items():
                before=c['values'].get(f)
                if numeric(v) and (not numeric(before) or v!=before):issues.append(f+': 숫자 변경 또는 새 수치')
                if f in ('단위','예산단위','목표단위','활동단위') and v!=before:issues.append(f+': 원문 단위 변경')
            for f in rules['required_fields'][k]:
                if not nonempty(row.get(f)):issues.append(f+': 필수 필드 없음')
                elif row.get(f)!=c['values'].get(f):issues.append(f+': 필수 식별/값 필드 변경')
            deltas=[{'field':f,'before':c['values'].get(f),'after':row.get(f)} for f in set(c['values'])|set(row) if c['values'].get(f)!=row.get(f)]
            if deltas:changes.append({'source_record_ids':c['source_record_ids'],'before':c['values'],'after':row,'differences':deltas,'held':bool(issues)})
            if issues:
                for rid in c['source_record_ids']:reasons[rid]='G01: '+'; '.join(issues)
            else:
                out=copy.deepcopy(c);out['values']=row;out['output_row_id']='OUT-'+digest([k,c['source_record_ids'],row])[:20];out['excel_row']=None;retained.append(out);seen.add(evidence)
    for evidence,c in original.items():
        if evidence not in seen:
            for rid in c['source_record_ids']:reasons.setdefault(rid,'기존 정리/의미 검증에서 반환되지 않음; 반환·격리 원장 보존')
    return retained,reasons,changes,backend_info,payload,cleaned

def finalize(source,items,business,reasons,duplicates,changes,schema):
    bysource={rid:c for c in business for rid in c['source_record_ids']}
    metadata_used=collections.defaultdict(list)
    for i in items:
        for l in i['metadata_links']:metadata_used[l['source_record_id']].append(i['id'])
    trace=[];holds=[];references=[]
    for i in items:
        rid=i['source_record_ids'][0];r=i['original'];c=bysource.get(rid);reason=reasons.get(rid);rules=[i['classification_rule'],'L01_embedded']
        if i['scope']=='reference':outcome='참고자료로 보존';references.append(copy.deepcopy(i));reason=None
        elif c:
            outcome='중복 통합' if len(c['source_record_ids'])>1 and rid!=c['source_record_ids'][0] else '업무 데이터에 연결';rules.append(c['rule_id']);reason=None
        elif i['classification']=='qualifier' and metadata_used[rid]:outcome='다른 레코드의 메타데이터로 연결';rules+=['L03_explicit_relation' if any(l['rule_id']=='L03_explicit_relation' and l['source_record_id']==rid for x in items for l in x['metadata_links']) else 'L02_same_cell/L04_ancestor_period'];reason=None
        elif i['classification'] in ('node_or_relation','text_fact') and not reason:outcome='정성·관계 자료로 보존';rules.append('M_preserve_nodes')
        else:outcome='연결 보류';reason=reason or '안전한 연결 근거 없음';rules.append('M_hold')
        row={'원본 레코드ID':rid,'객체ID':i['object_id'],'분류':i['classification'],'연결된 중간 데이터ID':[i['id']]+metadata_used[rid], '적용 규칙ID':rules,'대상 시트':c['sheet_name'] if c else None,'출력 행ID':c['output_row_id'] if c else None,'대상 열':list(c['field_sources']) if c else [],'처리 결과':outcome,'근거':{'원문':r.get('근거 문구'),'문맥':r.get('문맥 근거'),'범위':i['scope_evidence']},'보류 사유':reason,'Excel 행':c['excel_row'] if c else None}
        trace.append(row)
        if outcome=='연결 보류':holds.append({'원본 레코드ID':rid,'객체ID':i['object_id'],'분류':i['classification'],'보류 사유':reason,'원본 레코드':r,'중간 데이터ID':i['id']})
    return {'schema_version':'1.0','source':copy.deepcopy(source),'items':items,'separated':{cl:[i['id'] for i in items if i['classification']==cl] for cl in ('measurement','qualifier','identifier','text_fact','node_or_relation','unresolved')},'relations':copy.deepcopy(source['relations']),'business_rows':business,'trace':trace,'holds':holds,'references':references,'duplicate_events':duplicates,'backend_changes':changes,'limits':['원문 재판독 없음','의미 규칙은 모델 작성; 레코드 실행은 결정론적 코드','사실 정확도 자체 평가 없음']}

def write_workbook(business,schema,bridge,path):
    payload={k:[] for k in schema['EXTRACTION_SHEETS']}
    for row in business:payload[row['sheet_key']].append(row['values'])
    path,info=bridge.write_excel(payload,path)
    tick=time.perf_counter();wb=openpyxl.load_workbook(path,keep_links=False)
    issues=[]
    for key in schema['EXTRACTION_SHEETS']:
        name=schema['SHEET_KEY_TO_NAME'][key];ws=wb[name];heads=[c.value for c in ws[1]];evcol=heads.index('근거ID')
        expected={r['values']['근거ID']:r for r in business if r['sheet_key']==key};found=set()
        for idx,row in enumerate(ws.iter_rows(min_row=2,values_only=True),2):
            if not any(v is not None for v in row):continue
            ev=row[evcol];c=expected.get(ev)
            if c is None:issues.append({'sheet':name,'row':idx,'issue':'unexpected writer row'});continue
            found.add(ev);c['excel_row']=idx
            for h,v in zip(heads,row):
                original=c['values'].get(h,c['values'].get('감축률') if h=='감축률(%)' else None)
                if isinstance(original,bool):original='Y' if original else 'N'
                elif isinstance(original,(dict,list)):original=json.dumps(original,ensure_ascii=False,sort_keys=True)
                if original=='':original=None
                if v!=original:issues.append({'sheet':name,'row':idx,'field':h,'source_ids':c['source_record_ids'],'expected':original,'actual':v})
        for ev in set(expected)-found:issues.append({'sheet':name,'evidence':ev,'issue':'writer omitted row'})
    wb.close();return path,info,issues,time.perf_counter()-tick

def append_helpers(path,norm,run):
    tick=time.perf_counter();wb=openpyxl.load_workbook(path,keep_links=False)
    tables={'추적로그':norm['trace'],'연결보류':norm['holds'],'참고자료':[{'원본 레코드ID':i['source_record_ids'][0],'객체ID':i['object_id'],'분류':i['classification'],'범위 근거':i['scope_evidence'],'원본 레코드':i['original'],'메타데이터 연결':i['metadata_links']} for i in norm['references']], '실행정보':[{'항목':k,'내용':v} for k,v in run.items()]}
    for name,rows in tables.items():
        if name in wb:raise RuntimeError('Refusing to overwrite helper sheet')
        ws=wb.create_sheet(name);heads=list(dict.fromkeys(k for row in rows for k in row)) or ['기록']
        ws.append(heads)
        for row in rows:
            values=[canonical(row.get(h)) if isinstance(row.get(h),(dict,list)) else row.get(h) for h in heads];ws.append(values)
            for cell in ws[ws.max_row]:
                if isinstance(cell.value,str):cell.data_type='s'
        ws.freeze_panes='A2';ws.auto_filter.ref=ws.dimensions
    wb.save(path);wb.close();return time.perf_counter()-tick
