import unittest
from unittest.mock import patch
import openpyxl
from utils.reading_cell_validation import compare_backend, inspect_workbook


class CellValidationTests(unittest.TestCase):
    def fixture(self):
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = 'test'
        ws.append(['근거ID', '값'])
        ws.append(['R2', 20])  # Deliberately sorted differently from input.
        ws.append(['R1', 10])
        schema = dict(EXTRACTION_SHEETS=['key'], SHEET_KEY_TO_NAME={'key':'test'}, EXCEL_HEADERS={'test':['근거ID','값']})
        data = {'key':[{'근거ID':'R1','값':10},{'근거ID':'R2','값':20}]}
        return wb, schema, data

    def test_order_independent(self):
        wb,s,d = self.fixture()
        with patch('openpyxl.load_workbook', return_value=wb):
            self.assertTrue(inspect_workbook('unused',d,s)['passed'])

    def test_numeric_string_is_failure(self):
        wb,s,d = self.fixture()
        wb.active['B3']='10'
        with patch('openpyxl.load_workbook', return_value=wb):
            self.assertFalse(inspect_workbook('unused',d,s)['passed'])

    def test_missing_duplicate_and_formula(self):
        wb,s,d = self.fixture()
        wb.active['A3']='R2'
        wb.active['B3']='=10'
        with patch('openpyxl.load_workbook', return_value=wb):
            result=inspect_workbook('unused',d,s)
        self.assertFalse(result['passed'])
        self.assertTrue(any(i['kind']=='missing_or_duplicate_row' for i in result['issues']))

    def test_backend_changes_not_hidden(self):
        before={'k':[{'근거ID':'R','값':1,'연도':2024}]}
        after={'k':[{'근거ID':'R','값':2,'연도':2024}]}
        self.assertEqual(compare_backend(before,after,['k'])[0]['field'],'값')
        self.assertTrue(compare_backend(before,{'k':[]},['k']))


if __name__ == '__main__':
    unittest.main()
