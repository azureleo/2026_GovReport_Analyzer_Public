import config
import unittest
from scripts.golden_score_contract import 계약헤더
from utils.excel_writer import _headers_for_output, _write_generic_sheet
from openpyxl import Workbook


class RawSectorTests(unittest.TestCase):
    def test_raw_sector_columns_are_appended_and_not_scored(self):
        self.assertEqual(len(config.EXTRACTION_SHEETS), 16)
        for sheet, field in config.RAW_SECTOR_EXCEL_COLUMNS.items():
            self.assertEqual(config.EXCEL_HEADERS[sheet][-1], field)
            self.assertEqual(config.EXCEL_HEADERS[sheet].count(field), 1)
            self.assertNotIn(field, 계약헤더(sheet))


    def test_raw_sector_writer_preserves_original_and_blank(self):
        for sheet, field in config.RAW_SECTOR_EXCEL_COLUMNS.items():
            wb = Workbook()
            ws = wb.active
            headers = _headers_for_output(sheet, config.EXCEL_HEADERS[sheet])
            raw = '에너지(연료 공급량 기준)'
            _write_generic_sheet(ws, headers, [{field: raw}, {}])
            col = headers.index(field) + 1
            self.assertEqual(ws.cell(2, col).value, raw)
            self.assertIsNone(ws.cell(3, col).value)
            wb.close()
