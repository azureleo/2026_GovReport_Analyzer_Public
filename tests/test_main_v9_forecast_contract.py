"""Offline regression of selectively integrated keys, guards, replay and scorers."""
from copy import deepcopy
from itertools import permutations
import json
from pathlib import Path

import config
import pytest
from agents.organizer_agent import OrganizerAgent, deduplicate_forecast_rows
from scripts.golden_score_contract import 행, 필수채점필드
from scripts.golden_score_matching import 매칭하기
from utils.document_objects import DocumentObject
from utils.pdf_reader import PDFContent, PageContent
from utils.run_state import RunState, merge_rows_stably, stable_entity_id
from utils.semantic_routing import validate_and_reclassify


def forecast(**overrides):
    row = {"지자체명": "테스트시", "전망범위": "지역전체", "시나리오": "BAU",
           "부문": "건물", "세부부문": "가정", "연도": 2030, "전망값": 100,
           "단위": "천톤CO2eq", "출처페이지": 10}
    row.update(overrides)
    return row


def clean(rows, sheet="emissions_forecast"):
    return OrganizerAgent().organize({"municipality_name": "테스트시", sheet: rows})


def test_scope_same_value_stays_separate():
    result = clean([forecast(), forecast(전망범위="관리 권한 내")])
    assert {row["전망범위"] for row in result["emissions_forecast"]} == {"지역전체", "관리권한"}


def test_unique_blank_absorption_order_independent_and_evidence_union():
    rows = [forecast(근거ID="a"), forecast(전망범위="", 근거ID="b", 출처페이지=11),
            forecast(전망범위="관리권한", 전망값=80)]
    results = [deduplicate_forecast_rows(deepcopy(list(order))) for order in permutations(rows)]
    assert all(result == results[0] for result in results)
    regional = next(row for row in results[0] if row["전망범위"] == "지역전체")
    assert regional["근거ID목록"] == ["a", "b"]
    assert regional["출처페이지"] == "10,11"
    assert len(results[0]) == 2


def test_ambiguous_blank_scope_is_kept_and_reported():
    rows = [forecast(), forecast(전망범위="관리권한"), forecast(전망범위="")]
    result = clean(rows)
    assert len(result["emissions_forecast"]) == 3
    blank = next(row for row in result["emissions_forecast"] if row["전망범위"] == "")
    assert blank["_forecast_scope_review"]["reason"] == "ambiguous_blank_scope"
    assert any("ambiguous_blank_scope" == row["항목"] for row in result["validation_report"])


@pytest.mark.parametrize("field,value", [("시나리오", "정책반영"), ("세부부문", "상업"), ("단위", "톤CO2eq")])
def test_forecasts_separate_scenario_subsector_unit(field, value):
    result = clean([forecast(), forecast(**{field: value}), forecast(전망범위="", **{field: value})])
    assert len(result["emissions_forecast"]) == 2


def test_scope_not_guessed_from_locality_and_old_checkpoint_blank():
    legacy = forecast()
    legacy.pop("전망범위")
    rows = clean([legacy, forecast(전망범위="서울시 산정")])["emissions_forecast"]
    assert all(row["전망범위"] == "" for row in rows)
    assert any(row.get("_forecast_scope_review", {}).get("original") == "서울시 산정" for row in rows)


@pytest.mark.parametrize("scope", ["관리권한 외", "관리권한 제외", "비관리권한", "관리권한 없음"])
def test_scope_negative_or_exclusion_labels_are_not_management_scope(scope):
    rows = clean([forecast(전망범위=scope)])["emissions_forecast"]
    assert rows[0]["전망범위"] == ""
    assert rows[0]["_forecast_scope_review"]["original"] == scope


def test_axis_original_and_reason_are_preserved_and_grounded_refill_only():
    rows = [dict(관리부문="건물", 세부부문="가정", 직간접구분="에너지", 연도=2030, 배출량=1),
            dict(관리부문="건물", 세부부문="가정(전력)", 직간접구분="부문합계", 연도=2030, 배출량=2),
            dict(관리부문="건물", 세부부문="공공", 직간접구분="direct", 연도=2030, 배출량=3)]
    result = clean(rows, "emissions_management")
    by_sector = {row["세부부문"]: row for row in result["emissions_management"]}
    assert by_sector["가정"]["직간접구분"] == ""
    assert by_sector["가정"]["_axis_original"] == {"직간접구분": "에너지"}
    assert by_sector["가정(전력)"]["직간접구분"] == "간접"
    assert by_sector["공공"]["직간접구분"] == "직접"
    assert rows[0]["직간접구분"] == "에너지"  # caller data unchanged
    assert result["_key_blank_stats"]["emissions_management"]["blank_counts"]["직간접구분"] == 1


def test_invalid_axis_different_originals_do_not_absorb_each_other():
    rows = [dict(부문="건물", 세부부문="가정", 배출유형=value, 연도=2030, 배출량=1)
            for value in ("에너지", "부문합계")]
    assert len(clean(rows, "emissions_regional")["emissions_regional"]) == 2


def test_nonvalue_fact_conflicts_do_not_disappear_during_cleaning():
    result = clean([forecast(주요가정="현 추세"), forecast(주요가정="정책 반영")])
    assert len(result["emissions_forecast"]) == 2
    assert all(row["_merge_conflict"]["reason"] == "same_entity_different_field" for row in result["emissions_forecast"])


@pytest.mark.parametrize("field", ["주요가정", "전망방법원문"])
def test_blank_scope_does_not_hide_nonvalue_field_conflicts(field):
    rows = [forecast(**{field: "현 추세"}), forecast(전망범위="", **{field: "정책 변경"})]
    for order in permutations(rows):
        result = clean(list(order))["emissions_forecast"]
        assert len(result) == 2
        blank = next(row for row in result if not row["전망범위"])
        assert blank["_forecast_scope_review"]["reason"] == "scope_candidate_payload_conflict"


def routing_fixture(caption="온실가스 배출량 전망 BAU"):
    text = "온실가스 배출 전망\n표 4-2 " + caption
    page = PageContent(10, text, [], [])
    doc = PDFContent(10, [page], text)
    obj = DocumentObject("table10", "table", 10, 1, caption=caption,
                         section="온실가스 배출 전망", metadata={"evidence_id": "e10"})
    data = {"document_meta": [{"계획시작연도": 2024}], "emissions_forecast": [],
            "emissions_regional": [{"부문": "건물", "배출유형": "", "연도": 2030,
                                      "배출량": 100, "출처페이지": 10, "근거ID": "e10"}],
            "emissions_management": []}
    return data, doc, [obj]


def route(data, doc, objects):
    return validate_and_reclassify(data, doc, document_objects=objects, auto_reclassify=True)


@pytest.mark.parametrize("caption,scenario", [("온실가스 배출량 전망 BAU", "BAU"),
                                               ("정책 반영 온실가스 배출량 전망", "정책반영")])
def test_routing_preserves_grounded_scenario(caption, scenario):
    data, doc, objects = routing_fixture(caption)
    route(data, doc, objects)
    assert len(data["emissions_forecast"]) == 1
    assert data["emissions_forecast"][0]["시나리오"] == scenario
    assert data["emissions_forecast"][0]["전망범위"] == "지역전체"


@pytest.mark.parametrize("kind", ["generic", "valid_axis", "invalid_axis", "historical", "target"])
def test_routing_holds_unresolved_and_nonforecast_rows(kind):
    caption = "온실가스 배출량 전망" if kind == "generic" else "온실가스 배출량 전망 BAU"
    data, doc, objects = routing_fixture(caption)
    row = data["emissions_regional"][0]
    if kind == "valid_axis":
        row["배출유형"] = "직접배출"
    elif kind == "invalid_axis":
        row["_axis_original"] = {"배출유형": "에너지"}
    elif kind == "historical":
        row["연도"] = 2021
    elif kind == "target":
        row["시나리오"] = "목표값"
    route(data, doc, objects)
    assert not data["emissions_forecast"]
    assert len(data["emissions_regional"]) == 1


def test_routing_same_page_objects_and_scenarios_do_not_merge():
    data, doc, objects = routing_fixture()
    objects.append(DocumentObject("policy10", "table", 10, 2,
                                  caption="정책반영 온실가스 배출량 전망",
                                  section="온실가스 배출 전망", metadata={"evidence_id": "p10"}))
    data["emissions_regional"].append(dict(data["emissions_regional"][0], 근거ID="p10"))
    route(data, doc, objects)
    assert {row["시나리오"] for row in data["emissions_forecast"]} == {"BAU", "정책반영"}


def test_routing_known_scope_and_ambiguous_blank_order_independent():
    data, doc, objects = routing_fixture()
    data["emissions_management"] = [dict(관리부문="건물", 직간접구분="", 연도=2030,
                                        배출량=100, 출처페이지=10, 근거ID="e10")]
    data["emissions_forecast"] = [forecast(지자체명="", 전망범위="", 세부부문="")]
    for rows in (data["emissions_regional"], data["emissions_management"]):
        rows[0]["지자체명"] = ""
    route(data, doc, objects)
    assert len(data["emissions_forecast"]) == 3


def test_ledger_keeps_conflicting_sources_and_separates_full_entity_key():
    a = forecast(_추출원천="표객체")
    b = forecast(전망값=80, _추출원천="본문청크")
    rows, conflicts = merge_rows_stably("emissions_forecast", [a], [b], prefer=("표객체", "본문청크"))
    assert len(rows) == 2 and conflicts[0]["resolution"] == "both_preserved_for_review"
    for field, value in [("시나리오", "政策"), ("전망범위", "관리권한"), ("세부부문", "상업"), ("단위", "톤CO2eq")]:
        assert stable_entity_id("emissions_forecast", a) != stable_entity_id("emissions_forecast", forecast(**{field: value}))


def test_legacy_checkpoint_and_source_replay_are_idempotent():
    old = forecast()
    old.pop("전망범위")
    rows, _ = merge_rows_stably("emissions_forecast", [old], [forecast(전망범위="", 근거ID="second")])
    assert len(rows) == 1
    again, _ = merge_rows_stably("emissions_forecast", rows, [old])
    assert again == rows


def test_run_fingerprint_changes_with_external_ocr_hash(tmp_path, monkeypatch):
    monkeypatch.setattr(config, "RUN_STATE_DIR", str(tmp_path / "runs"), raising=False)
    args = dict(input_path=tmp_path / "missing.pdf", guideline_path=None,
                extraction_prompts={"forecast": "same"}, output_path=tmp_path / "out.xlsx")
    a = RunState.create(**args, execution_info={"ocr_file_sha256": "a"})
    b = RunState.create(**args, execution_info={"ocr_file_sha256": "b"})
    assert a.run_id != b.run_id
    a.register_expected("x", {"kind": "native_all"})
    a.record_batch(batch_id="x", kind="native_all", sheet_keys=["emissions_forecast"],
                   page_nums=[10], status="ok", result={"emissions_forecast": [forecast()]})
    resumed = RunState.create(**args, execution_info={"ocr_file_sha256": "a"}, resume=True)
    assert resumed.restored_result("x")["emissions_forecast"][0]["전망범위"] == "지역전체"


@pytest.mark.parametrize("setting,value", [("EXTRACTION_OBJECT_BATCHING_ENABLED", True),
                                           ("MAX_BATCH_CHARS", 777), ("MAX_BATCH_TABLE_CELLS", 12),
                                           ("TABLE_OBJECT_ROWS_PER_CHUNK", 2)])
def test_run_fingerprint_uses_actual_main_batch_settings(tmp_path, monkeypatch, setting, value):
    monkeypatch.setattr(config, "RUN_STATE_DIR", str(tmp_path / "runs"), raising=False)
    args = dict(input_path=tmp_path / "missing.pdf", guideline_path=None,
                extraction_prompts={}, output_path=tmp_path / "out.xlsx", execution_info={})
    first = RunState.create(**args)
    monkeypatch.setattr(config, setting, value, raising=False)
    second = RunState.create(**args)
    assert first.run_id != second.run_id


def scored(row, number=2):
    return 행(number, row, "텍스트표", "10")


@pytest.mark.parametrize("left_scope,right_scope", [(None, "지역전체"), ("관리권한", None),
                                                       ("지역전체", "지역전체"), (None, None)])
def test_scorer_cross_old_new_scope_contract(left_scope, right_scope):
    golden = forecast(전망범위=left_scope)
    output = forecast(전망범위=right_scope)
    assert len(매칭하기([scored(golden)], [scored(output)], "05_배출전망")[0]) == 1
    assert "전망범위" not in 필수채점필드("05_배출전망")


@pytest.mark.parametrize("field,value", [("전망범위", "관리권한"), ("시나리오", "정책반영"),
                                         ("세부부문", "상업"), ("단위", "톤CO2eq")])
def test_scorer_explicit_axis_conflicts_block_relaxed_matching(field, value):
    matches = 매칭하기([scored(forecast())], [scored(forecast(**{field: value}))], "05_배출전망")[0]
    assert not matches


def test_scorer_ambiguous_legacy_scope_does_not_consume_one_arbitrarily():
    result = 매칭하기([scored(forecast(전망범위=""))],
                     [scored(forecast()), scored(forecast(전망범위="관리권한"), 3)], "05_배출전망")
    assert not result[0] and result[1] == {0} and result[2] == {0, 1}


def test_scorer_explicit_scope_match_precedes_legacy_wildcard():
    result = 매칭하기([scored(forecast(전망범위="")), scored(forecast(), 3)],
                     [scored(forecast()), scored(forecast(전망범위="관리권한"), 3)], "05_배출전망")
    assert len(result[0]) == 2
    explicit = next(match for match in result[0] if match.골든.값["전망범위"])
    assert explicit.출력.값["전망범위"] == "지역전체"


@pytest.mark.parametrize("reverse", [False, True])
def test_unknown_scope_cannot_rescue_an_explicit_conflicting_scope(reverse):
    explicit = [scored(forecast())]
    mixed = [scored(forecast(전망범위="관리권한")), scored(forecast(전망범위=""), 3)]
    for order in permutations(mixed):
        golden, output = (list(order), explicit) if reverse else (explicit, list(order))
        matches = 매칭하기(golden, output, "05_배출전망")[0]
        assert not matches


def test_duplicate_scope_candidates_are_not_arbitrarily_consumed():
    rows = [scored(forecast()), scored(forecast(), 3)]
    for order in permutations(rows):
        assert not 매칭하기([scored(forecast())], list(order), "05_배출전망")[0]
        assert not 매칭하기(list(order), [scored(forecast())], "05_배출전망")[0]


def test_fixed_benchmark_does_not_follow_live_header_mutations(monkeypatch):
    from scripts.benchmark_golden_contract import 계약헤더
    before = 계약헤더("05_배출전망")
    monkeypatch.setitem(config.EXCEL_HEADERS, "05_배출전망", ["unexpected"])
    assert 계약헤더("05_배출전망") == before
    from scripts.benchmark_golden_matching import 매칭하기 as benchmark_match
    assert not benchmark_match([scored(forecast())], [scored(forecast(전망범위="관리권한"))], "05_배출전망")[0]


def test_existing_main_nonemission_cleaners_and_metadata_preserved():
    data = {"municipality_name": "테스트시", "financial_plan": [{"부문": "건물", "연도": "2030", "예산액": "1,000"}],
            "visual_inventory": [{"시각자료ID": "v1"}], "validation_report": [{"항목": "existing"}]}
    result = OrganizerAgent().organize(data)
    assert result["financial_plan"][0]["예산액"] == 1000
    assert result["visual_inventory"] == data["visual_inventory"]
    assert result["validation_report"] == data["validation_report"]


def test_dropped_emission_is_visible_in_quarantine(tmp_path):
    from utils.semantic_contract_guard import write_quarantine_ledger
    agent = OrganizerAgent()
    result = agent.organize({"emissions_forecast": [forecast(전망값=None)]})
    assert not result["emissions_forecast"]
    path = write_quarantine_ledger(agent.get_quarantine_records(), tmp_path / "output.xlsx")
    payload = json.loads(path.read_text(encoding="utf-8"))
    assert payload["records"][0]["row"]["부문"] == "건물"


def write_forecast_book(path, scope, *, golden=False):
    import openpyxl
    from scripts.benchmark_contract_schema import FIXED_HEADERS
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "05_배출전망"
    headers = list(FIXED_HEADERS[ws.title])
    if scope is not None:
        headers.insert(1, "전망범위")
    if golden:
        headers.extend(["골든_출처유형", "골든_출처페이지", "골든_채점제외", "골든_비고"])
    ws.append(headers)
    data = forecast(전망범위=scope)
    data.update({"골든_출처유형": "텍스트표", "골든_출처페이지": "10"})
    ws.append([data.get(field) for field in headers])
    wb.save(path)
    wb.close()


@pytest.mark.parametrize("golden_scope,output_scope", [(None, "지역전체"), ("지역전체", None),
                                                       ("지역전체", "지역전체"), ("지역전체", "관리권한")])
def test_workbook_cross_contract_both_scorers(tmp_path, golden_scope, output_scope):
    from scripts.score_against_golden import score_workbooks
    from scripts.benchmark_golden_score import score_workbooks as benchmark_score
    gold, output = tmp_path / "gold.xlsx", tmp_path / "output.xlsx"
    write_forecast_book(gold, golden_scope, golden=True)
    write_forecast_book(output, output_scope)
    for index, scorer in enumerate((score_workbooks, benchmark_score)):
        result = scorer(output, gold, report_dir=tmp_path, label=str(index), write_json=True)
        assert not result["형식오류"]
        expected = 0 if golden_scope == "지역전체" and output_scope == "관리권한" else 1
        assert result["시트별"]["05_배출전망"]["매칭수"] == expected
        assert len(result["골든SHA256"]) == 64


def test_cross_backend_and_agreed_pair_tools_remain_offline(tmp_path):
    from scripts.cross_backend_probe import probe_workbooks
    from scripts.d30_agreed_pair_accuracy import measure
    gold, left, right = (tmp_path / name for name in ("gold.xlsx", "left.xlsx", "right.xlsx"))
    write_forecast_book(gold, "지역전체", golden=True)
    write_forecast_book(left, "지역전체")
    write_forecast_book(right, "지역전체")
    result = probe_workbooks(left, right, report_dir=tmp_path, label="probe")
    assert result
    agreed = measure(left, right, gold)
    assert agreed["종합_q_필드"] == 1.0


def test_visual_fixture_hashes_and_pending_review(tmp_path):
    import fitz
    from utils.visual_regression_fixture import build_visual_regression_fixture, validate_visual_regression_fixture
    pdf = fitz.open()
    page = pdf.new_page()
    page.insert_text((72, 72), "Offline fixture")
    source = tmp_path / "source.pdf"
    pdf.save(source)
    pdf.close()
    specs = tmp_path / "cases.jsonl"
    specs.write_text(json.dumps({"case_id": "table1", "source_page": 1, "caption": "table",
                                 "expected_object_type": "table", "expected_triage": "ocr_required",
                                 "expected_status": "extracted", "expected_sheet": "emissions_forecast",
                                 "expected_render_variants": ["full", "crop"], "bbox_pdf": [50, 50, 200, 100]}), encoding="utf-8")
    manifest = build_visual_regression_fixture(source, specs, tmp_path / "fixture", dataset_id="offline-test")
    assert validate_visual_regression_fixture(manifest)["valid"]
    assert json.loads(manifest.read_text(encoding="utf-8"))["review_status"] == "pending"
    (manifest.parent / "crops" / "table1.png").write_bytes(b"damaged")
    assert not validate_visual_regression_fixture(manifest)["valid"]


def test_vision_explicit_target_row_precedes_chart_forecast_target(monkeypatch):
    from agents.image_agent import ImageAgent
    agent = ImageAgent()
    monkeypatch.setattr(agent, "_chart_merge_decision", lambda *args: (True, "high", []))
    analysis = {"type": "chart_table", "target_sheet": "emissions_forecast", "title": "BAU 전망 및 목표",
                "page_number": 10, "table": [{"항목": "건물", "연도": 2030, "값": 80, "종류": "목표"}]}
    result = agent._merge_image_results({}, [analysis], "테스트시")
    assert not result["emissions_forecast"]
    assert result["reduction_targets"][0]["목표배출량"] == 80


def test_vision_mixed_scenario_chart_title_does_not_fix_missing_scenario_to_bau(monkeypatch):
    from agents.image_agent import ImageAgent
    agent = ImageAgent()
    monkeypatch.setattr(agent, "_chart_merge_decision", lambda *args: (True, "high", []))
    analysis = {"type": "chart_table", "target_sheet": "emissions_forecast", "title": "BAU 및 정책반영 전망",
                "page_number": 10, "table": [{"항목": "건물", "연도": 2030, "값": 80, "종류": "전망"}]}
    result = agent._merge_image_results({}, [analysis], "테스트시")
    assert all(row.get("시나리오") != "BAU" for row in result["emissions_forecast"])


def test_vision_management_sheet_preserves_declared_axis(monkeypatch):
    from agents.image_agent import ImageAgent
    agent = ImageAgent()
    monkeypatch.setattr(agent, "_chart_merge_decision", lambda *args: (True, "high", []))
    analysis = {"type": "chart_table", "target_sheet": "emissions_management", "title": "관리권한 배출현황",
                "page_number": 10, "table": [{"항목": "건물", "관리부문": "건물", "세부부문": "가정",
                                              "직간접구분": "간접", "연도": 2021, "값": 80, "종류": "현황"}]}
    result = agent._merge_image_results({}, [analysis], "테스트시")
    assert not result["emissions_regional"]
    assert result["emissions_management"][0]["직간접구분"] == "간접"
    assert result["emissions_management"][0]["관리부문"] == "건물"


@pytest.mark.parametrize("kind,sheet,value_field", [("전망", "emissions_forecast", "전망값"),
                                                   ("목표", "reduction_targets", "목표배출량")])
def test_vision_legacy_series_preserves_declared_semantic_type(kind, sheet, value_field):
    from agents.image_agent import ImageAgent
    analysis = {"type": "그래프", "page_number": 10,
                "ghg_data": [{"label": "건물", "data_type": kind, "전망범위": "관리권한",
                              "시나리오": "정책반영", "values": {"2030": 80}}]}
    result = ImageAgent()._merge_image_results({}, [analysis], "테스트시")
    assert not result["emissions_regional"]
    assert result[sheet][0][value_field] == 80
    if kind == "전망":
        assert result[sheet][0]["전망범위"] == "관리권한"
        assert result[sheet][0]["시나리오"] == "정책반영"


def test_scope_absorption_regenerates_serialized_lineage_for_final_excel(tmp_path):
    from utils.excel_writer import write_excel
    import openpyxl
    native_record = {"engine": "pymupdf", "object_id": "native-table", "evidence_id": "native-e"}
    ocr_record = {"engine": "unlimited_ocr", "object_id": "ocr-table", "evidence_id": "ocr-e"}
    result = clean([
        forecast(_extraction_sources=[native_record], 출처계보=json.dumps([native_record]), source_engines=["pymupdf"]),
        forecast(전망범위="", 출처계보=json.dumps([ocr_record]), source_engines=["unlimited_ocr"],
                 ocr_source_file="ocr.jsonl", ocr_source_line=7),
    ])
    assert len(result["emissions_forecast"]) == 1
    merged = result["emissions_forecast"][0]
    assert json.loads(merged["출처계보"]) == merged["_extraction_sources"]
    assert {record["engine"] for record in merged["_extraction_sources"]} == {"pymupdf", "unlimited_ocr"}
    assert merged["source_engines"] == ["pymupdf", "unlimited_ocr"]
    path = write_excel(result, tmp_path / "lineage.xlsx")
    book = openpyxl.load_workbook(path, data_only=True)
    sheet = book["05_배출전망"]
    headers = [cell.value for cell in sheet[1]]
    serialized = sheet.cell(2, headers.index("출처계보") + 1).value
    assert {record["engine"] for record in json.loads(serialized)} == {"pymupdf", "unlimited_ocr"}
    book.close()


def test_source_engines_and_ocr_file_line_are_provenance_not_fact_conflicts():
    records = [{"engine": "pymupdf"}, {"engine": "unlimited_ocr"}]
    rows = [forecast(source_engines=["pymupdf"], _extraction_sources=records[:1]),
            forecast(source_engines=["unlimited_ocr"], ocr_source_file="input.jsonl", ocr_source_line=2,
                     _extraction_sources=records[1:])]
    result = clean(rows)
    assert len(result["emissions_forecast"]) == 1
    row = result["emissions_forecast"][0]
    assert row["source_engines"] == ["pymupdf", "unlimited_ocr"]
    assert len(json.loads(row["출처계보"])) == 2
    assert not row.get("_merge_conflict")
