"""v7-3 06_감축목표 사업명힌트 내부 처리 회귀 테스트."""

from openpyxl import load_workbook

from agents.organizer_agent import OrganizerAgent
from utils.excel_writer import write_excel


def _target_row(**overrides) -> dict:
    row = {
        "목표수준": "세부사업",
        "목표범위": "관리권한",
        "부문": "건물",
        "기준연도": 2018,
        "목표연도": 2030,
        "목표감축량": 100,
        "출처페이지": 10,
    }
    row.update(overrides)
    return row


def _organize_sheet(rows: list[dict]) -> list[dict]:
    return OrganizerAgent().organize_sheet("reduction_targets", rows, "테스트시")


def test_사업명힌트를_내부필드로_이관하고_엑셀에는_출력하지_않는다(tmp_path) -> None:
    cleaned = _organize_sheet([_target_row(사업명힌트=" 공공건물 개선 ")])

    assert cleaned[0]["_사업명힌트"] == "공공건물 개선"
    assert "사업명힌트" not in cleaned[0]

    output_path = write_excel({"reduction_targets": cleaned}, tmp_path / "targets.xlsx")
    workbook = load_workbook(output_path, read_only=True)
    headers = [cell.value for cell in next(workbook["06_감축목표"].iter_rows())]
    workbook.close()

    assert "사업명힌트" not in headers
    assert "_사업명힌트" not in headers


def test_세부사업의_사업명힌트가_다르면_dedup하지_않는다() -> None:
    cleaned = _organize_sheet([
        _target_row(사업명힌트="공공건물 개선"),
        _target_row(사업명힌트="민간건물 개선", 출처페이지=11),
    ])

    assert len(cleaned) == 2
    assert {row["_사업명힌트"] for row in cleaned} == {"공공건물 개선", "민간건물 개선"}


def test_세부사업의_사업명힌트는_정규화값으로_dedup한다() -> None:
    cleaned = _organize_sheet([
        _target_row(사업명힌트="공공 건물·개선"),
        _target_row(사업명힌트="공공건물ㆍ개선", 출처페이지=11),
    ])

    assert len(cleaned) == 1
    assert cleaned[0]["출처페이지"] == "10,11"


def test_기준연도_빈칸_흡수도_서로다른_사업명힌트를_병합하지_않는다() -> None:
    cleaned = _organize_sheet([
        _target_row(사업명힌트="공공건물 개선", 기준연도=None),
        _target_row(사업명힌트="민간건물 개선", 출처페이지=11),
    ])

    assert len(cleaned) == 2


def test_힌트없는_세부사업은_현행키로_dedup한다() -> None:
    cleaned = _organize_sheet([
        _target_row(),
        _target_row(출처페이지=11, 단위="tCO2eq"),
    ])

    assert len(cleaned) == 1
    assert cleaned[0]["출처페이지"] == "10,11"
    assert cleaned[0]["단위"] == "tCO2eq"


def test_S1_재태깅으로_세부사업이_된_힌트없는_행은_현행키와_호환된다() -> None:
    cleaned = OrganizerAgent().organize({
        "municipality_name": "테스트시",
        "reduction_targets": [
            _target_row(목표수준="부문", 근거ID="ev-card"),
            _target_row(목표수준="부문", 출처페이지=11, 근거ID="ev-card"),
        ],
        "mitigation_projects": [
            {"사업명": "공공건물 개선", "출처페이지": "10,11", "근거ID": "ev-card"},
        ],
    })["reduction_targets"]

    assert len(cleaned) == 1
    assert cleaned[0]["목표수준"] == "세부사업"
    assert "_사업명힌트" not in cleaned[0]
    assert cleaned[0]["출처페이지"] == "10,11"
