from __future__ import annotations

import json
from pathlib import Path

import openpyxl
import pytest

import config
from scripts.golden_score_contract import 계약헤더
from utils.benchmark_evaluation import (
    EvaluationContractError,
    evaluate_benchmark,
    load_evaluation_datasets,
    select_evaluation_dataset,
    sha256_file,
)


def _write_output(path: Path) -> None:
    workbook = openpyxl.Workbook()
    sheet = workbook.active
    sheet.title = "03_배출현황_지역"
    headers = config.EXCEL_HEADERS["03_배출현황_지역"]
    sheet.append(headers)
    row = {
        "지자체명": "서울특별시",
        "배출유형": "직접배출",
        "부문": "건물",
        "세부부문": "가정",
        "연도": 2021,
        "배출량": 100,
        "단위": "천톤CO2eq",
        "출처페이지": 1,
    }
    sheet.append([row.get(header) for header in headers])

    inventory = workbook.create_sheet("21_원문객체인벤토리")
    inventory_headers = config.EXCEL_HEADERS["21_원문객체인벤토리"]
    inventory.append(inventory_headers)
    for status in ("일치", "오배치의심"):
        values = {
            "객체ID": f"object-{status}",
            "시트정합상태": status,
        }
        inventory.append([values.get(header) for header in inventory_headers])
    workbook.save(path)


def _write_golden(path: Path) -> None:
    workbook = openpyxl.Workbook()
    sheet = workbook.active
    sheet.title = "03_배출현황_지역"
    headers = [*계약헤더("03_배출현황_지역"), "골든_출처유형", "골든_출처페이지", "골든_채점제외", "골든_비고"]
    sheet.append(headers)
    row = {
        "지자체명": "서울특별시",
        "배출유형": "직접배출",
        "부문": "건물",
        "세부부문": "가정",
        "연도": 2021,
        "배출량": 100,
        "단위": "천톤CO2eq",
        "골든_출처유형": "텍스트표",
        "골든_출처페이지": 1,
    }
    sheet.append([row.get(header) for header in headers])
    workbook.save(path)


def _manifest(path: Path, source: Path, golden: Path, *, role: str = "development") -> Path:
    payload = {
        "schema_version": 1,
        "datasets": [{
            "id": f"sample-{role}",
            "role": role,
            "review_status": "verified",
            "locked": True,
            "source": source.name,
            "source_sha256": sha256_file(source),
            "golden": golden.name,
            "golden_sha256": sha256_file(golden),
            "visual_inventory": None,
            "visual_inventory_sha256": "",
        }],
    }
    path.write_text(json.dumps(payload, ensure_ascii=False), encoding="utf-8")
    return path


def test_fixed_evaluation_reports_three_independent_metrics(tmp_path: Path) -> None:
    source = tmp_path / "source.pdf"
    source.write_bytes(b"fixed source")
    output = tmp_path / "output.xlsx"
    golden = tmp_path / "golden.xlsx"
    _write_output(output)
    _write_golden(golden)
    manifest = _manifest(tmp_path / "manifest.json", source, golden)

    result = evaluate_benchmark(
        output,
        source,
        manifest_path=manifest,
        report_dir=tmp_path / "reports",
    )

    assert result.metrics["cell_accuracy"] == 1.0
    assert result.metrics["object_recall"] is None
    assert result.metrics["routing_error_rate"] == 0.5
    assert result.report_json.exists()
    assert result.report_markdown.exists()


def test_holdout_requires_explicit_opt_in(tmp_path: Path) -> None:
    source = tmp_path / "source.pdf"
    golden = tmp_path / "golden.xlsx"
    source.write_bytes(b"holdout source")
    _write_golden(golden)
    manifest = _manifest(
        tmp_path / "manifest.json",
        source,
        golden,
        role="holdout",
    )
    datasets = load_evaluation_datasets(manifest)

    with pytest.raises(EvaluationContractError, match="홀드아웃"):
        select_evaluation_dataset(datasets, source)


def test_cell_accuracy_counts_values_from_unmatched_golden_rows(tmp_path: Path) -> None:
    source = tmp_path / "source.pdf"
    source.write_bytes(b"fixed source")
    output = tmp_path / "output.xlsx"
    golden = tmp_path / "golden.xlsx"
    _write_output(output)
    _write_golden(golden)

    workbook = openpyxl.load_workbook(golden)
    sheet = workbook["03_배출현황_지역"]
    headers = [cell.value for cell in sheet[1]]
    missing_row = {
        "지자체명": "서울특별시",
        "배출유형": "직접배출",
        "부문": "수송",
        "세부부문": "도로",
        "연도": 2022,
        "배출량": 50,
        "단위": "천톤CO2eq",
        "골든_출처유형": "텍스트표",
        "골든_출처페이지": 2,
    }
    sheet.append([missing_row.get(header) for header in headers])
    workbook.save(golden)
    manifest = _manifest(tmp_path / "manifest.json", source, golden)

    result = evaluate_benchmark(
        output,
        source,
        manifest_path=manifest,
        report_dir=tmp_path / "reports",
    )

    assert result.metrics["cell_matches"] < result.metrics["cell_expected"]
    assert result.metrics["cell_accuracy"] == 0.5
