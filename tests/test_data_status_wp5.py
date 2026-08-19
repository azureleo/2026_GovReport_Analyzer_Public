from __future__ import annotations

import tempfile
from pathlib import Path

from openpyxl import load_workbook

import config
from agents.image_agent import ImageAgent
from agents.organizer_agent import OrganizerAgent
from utils import excel_writer


def test_data_status_marks_visual_gapfill_calculated_and_conflicting_rows() -> None:
    # Given: 이미지 유래, 빈칸보완, 산식계산, 중복충돌 행이 함께 들어오면
    raw = {
        "municipality_name": "서울특별시",
        "regional_conditions": [
            {"지표범주": "교통", "지표명": "통행량", "연도": 2020, "값": 10, "출처": "이미지 p.12"},
        ],
        "plan_overview": [
            {"개요유형": "절차", "항목명": "공청회", "항목값": "개최", "보완출처": "gap_fill"},
        ],
        "reduction_targets": [
            {"목표수준": "총괄", "목표범위": "GIR", "부문": "합계", "기준배출량": 100, "목표배출량": 60, "목표연도": 2030},
        ],
        "emissions_management": [
            {"관리부문": "건물", "세부부문": "공공", "직간접구분": "직접", "연도": 2030, "배출량": 10},
            {"관리부문": "건물", "세부부문": "공공", "직간접구분": "직접", "연도": 2030, "배출량": 20},
        ],
    }

    # When: organizer가 행 단위 상태를 부여하면
    cleaned = OrganizerAgent().organize(raw)

    # Then: 상태 우선순위에 맞는 데이터상태가 남는다.
    assert cleaned["regional_conditions"][0]["데이터상태"] == "visual_only"
    assert cleaned["plan_overview"][0]["데이터상태"] == "gap_fill"
    assert cleaned["reduction_targets"][0]["데이터상태"] == "reported"
    assert cleaned["reduction_targets"][0]["감축률계산값"] == 40.0
    assert cleaned["reduction_targets"][0]["derivation_type"] == "inferred"
    assert cleaned["emissions_management"][0]["데이터상태"] == "conflicting"


def test_data_status_correction_keeps_conflicting_priority() -> None:
    # Given: 감축률 교정 대상인 두 중복 행의 수치가 충돌하면
    raw = {
        "municipality_name": "서울특별시",
        "reduction_targets": [
            {"목표수준": "총괄", "목표범위": "GIR", "부문": "합계", "기준배출량": 100, "목표배출량": 60, "감축률": 10, "목표연도": 2030},
            {"목표수준": "총괄", "목표범위": "GIR", "부문": "합계", "기준배출량": 120, "목표배출량": 70, "감축률": 20, "목표연도": 2030},
        ],
    }

    # When: dedup 충돌 후 감축률 재계산을 수행하면
    cleaned = OrganizerAgent().organize(raw)

    # Then: 원문 충돌 상태가 유지된다.
    assert cleaned["reduction_targets"][0]["데이터상태"] == "conflicting"


def test_image_merged_rows_keep_visual_only_for_projects_and_financial() -> None:
    # Given: 이미지 병합 경로가 08/11 시트에 행을 추가하면
    text_results = {"municipality_name": "서울특별시"}
    analyses = [
        {
            "type": "chart_table",
            "target_sheet": "mitigation_projects",
            "chart_type": "표",
            "title": "감축사업",
            "unit": "건",
            "page_number": 11,
            "municipality": "서울특별시",
            "confidence": "high",
            "table": [{"연도": 2030, "항목": "건물", "값": 1, "fields": {"감축사업명": "건물 효율화"}}],
        },
        {
            "type": "chart_table",
            "target_sheet": "financial_plan",
            "chart_type": "표",
            "title": "재정 투자",
            "unit": "백만원",
            "page_number": 12,
            "municipality": "서울특별시",
            "confidence": "high",
            "table": [{"연도": 2030, "항목": "건물", "값": 100, "단위": "백만원"}],
        },
    ]

    # When: 실제 이미지 병합 후 organizer를 통과시키면
    merged = ImageAgent()._merge_image_results(text_results, analyses, "서울특별시")
    cleaned = OrganizerAgent().organize(merged)

    # Then: 두 시트의 이미지 유래 행은 reported로 강등되지 않고 visual_only를 유지한다.
    assert cleaned["mitigation_projects"][0]["데이터상태"] == "visual_only"
    assert cleaned["financial_plan"][0]["데이터상태"] == "visual_only"


def test_data_status_header_is_after_provenance_and_opt_out_restores_legacy_columns() -> None:
    # Given: 데이터상태와 출처페이지를 가진 엑셀 데이터
    previous_status = getattr(config, "DATA_STATUS_ENABLED", True)
    previous_provenance = getattr(config, "PROVENANCE_ENABLED", True)
    data = {
        "document_meta": [
            {"지자체명": "서울특별시", "계획명": "계획", "출처페이지": "12", "데이터상태": "reported"}
        ]
    }

    try:
        with tempfile.TemporaryDirectory() as tmp:
            config.DATA_STATUS_ENABLED = True
            config.PROVENANCE_ENABLED = True
            path = Path(tmp) / "enabled.xlsx"
            excel_writer.write_excel(data, path)
            wb = load_workbook(path)
            headers = [cell.value for cell in wb["00_문서메타"][1]]
            assert headers[-2:] == ["출처페이지", "데이터상태"]

        with tempfile.TemporaryDirectory() as tmp:
            config.DATA_STATUS_ENABLED = False
            config.PROVENANCE_ENABLED = False
            path = Path(tmp) / "disabled.xlsx"
            excel_writer.write_excel(data, path)
            wb = load_workbook(path)
            headers = [cell.value for cell in wb["00_문서메타"][1]]
            assert "출처페이지" not in headers
            assert "데이터상태" not in headers
    finally:
        config.DATA_STATUS_ENABLED = previous_status
        config.PROVENANCE_ENABLED = previous_provenance
