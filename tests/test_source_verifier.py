from pathlib import Path

import fitz
from openpyxl import load_workbook

from utils.excel_writer import write_excel
from utils.pdf_reader import PDFContent, PageContent
from utils.source_verifier import (
    SourceVerificationReport,
    _numeric_variants,
    assess_quality,
    build_source_object_inventory,
    create_marked_pdf,
    verify_final_data,
)


def test_numeric_variants_keep_deterministic_order():
    assert _numeric_variants("2023") == ["2023", "2,023"]
    assert _numeric_variants("2,023") == ["2,023", "2023"]


def _document(*texts: str) -> PDFContent:
    pages = [
        PageContent(page_number=index, text=text, tables=[], images=[])
        for index, text in enumerate(texts, start=1)
    ]
    return PDFContent(total_pages=len(pages), pages=pages, full_text="\n".join(texts))


def test_source_verifier_uses_global_fallback_and_keeps_unconfirmed_rows():
    document = _document(
        "문서 표지",
        "전환 부문 태양광 보급 사업을 2030년까지 추진한다.",
    )
    final_data = {
        "municipality_name": "서울특별시",
        "mitigation_projects": [
            {"지자체명": "서울특별시", "부문": "전환", "사업명": "태양광 보급 사업", "출처페이지": "1"},
            {"지자체명": "서울특별시", "부문": "수송", "사업명": "해저 도시 조성", "출처페이지": "1"},
        ],
    }

    report = verify_final_data(final_data, document, page_radius=0, global_search=True)

    assert report.total_rows == 2
    assert report.rows[0].status == "확인"
    assert report.rows[0].matched_page == 2
    assert "다른 페이지" in report.rows[0].message
    assert report.rows[1].status == "미확인"


def test_quality_score_uses_grounding_fill_provenance_and_is_capped():
    final_data = {
        "document_meta": [{"지자체명": "서울", "계획명": "기본계획", "계획시작연도": 2024, "계획종료연도": 2033}],
        "emissions_regional": [{"배출유형": "직접배출", "부문": "합계", "연도": 2021, "배출량": 10, "단위": "천톤CO2eq"}],
        "reduction_targets": [{"목표수준": "총괄", "기준연도": 2018, "목표연도": 2030, "목표배출량": 80}],
        "mitigation_projects": [{"사업명": "태양광", "부문": "전환"}],
        "financial_plan": [{"사업명": "태양광", "연도": 2030, "예산액": 100, "예산단위": "백만원"}],
        "validation_report": [],
    }
    verification = SourceVerificationReport(
        rows=[], by_sheet={}, total_rows=5, verifiable_rows=5, confirmed_rows=5,
        weak_rows=0, unconfirmed_rows=0, skipped_rows=0, rows_with_provenance=5,
        grounding_ratio=1.0, provenance_ratio=1.0,
    )

    assessment = assess_quality(final_data, verification)

    assert assessment.score == 95.0
    assert assessment.metrics["source_inventory_connected"] is False
    assert any("점수 상한" in issue for issue in assessment.issues)


def test_quality_score_penalizes_extraction_ledger_failures():
    final_data = {
        "document_meta": [{"지자체명": "서울", "계획명": "기본계획", "계획시작연도": 2024, "계획종료연도": 2033}],
        "emissions_regional": [{"배출유형": "직접배출", "부문": "합계", "연도": 2021, "배출량": 10, "단위": "천톤CO2eq"}],
        "reduction_targets": [{"목표수준": "총괄", "기준연도": 2018, "목표연도": 2030, "목표배출량": 80}],
        "mitigation_projects": [{"사업명": "태양광", "부문": "전환"}],
        "financial_plan": [{"사업명": "태양광", "연도": 2030, "예산액": 100, "예산단위": "백만원"}],
        "validation_report": [],
        "pipeline_metrics": {"extraction_batches_total": 10, "extraction_batches_ok": 5},
    }
    verification = SourceVerificationReport(
        rows=[], by_sheet={}, total_rows=5, verifiable_rows=5, confirmed_rows=5,
        weak_rows=0, unconfirmed_rows=0, skipped_rows=0, rows_with_provenance=5,
        grounding_ratio=1.0, provenance_ratio=1.0,
    )

    assessment = assess_quality(final_data, verification)

    assert assessment.score == 92.5
    assert assessment.metrics["extraction_success_ratio"] == 0.5
    assert assessment.metrics["components"]["추출성공"] == 7.5
    assert any("추출 배치 성공률" in issue for issue in assessment.issues)


def test_source_object_inventory_connects_table_and_unlocks_full_score():
    table = """
    <table><tr><th>사업명</th><th>예산액</th></tr>
    <tr><td>태양광 보급</td><td>100</td></tr></table>
    """
    document = PDFContent(
        total_pages=1,
        pages=[PageContent(
            page_number=1,
            text="표 5-1 태양광 보급 재정계획",
            tables=[table],
            images=[],
        )],
        full_text="표 5-1 태양광 보급 재정계획",
    )
    final_data = {
        "document_meta": [{"지자체명": "서울", "계획명": "기본계획", "계획시작연도": 2024, "계획종료연도": 2033}],
        "emissions_regional": [{"배출유형": "직접배출", "부문": "합계", "연도": 2021, "배출량": 10, "단위": "천톤CO2eq"}],
        "reduction_targets": [{"목표수준": "총괄", "기준연도": 2018, "목표연도": 2030, "목표배출량": 80}],
        "mitigation_projects": [{"사업명": "태양광 보급", "부문": "전환", "출처페이지": 1}],
        "financial_plan": [{"사업명": "태양광 보급", "연도": 2030, "예산액": 100, "예산단위": "백만원", "출처페이지": 1}],
        "validation_report": [],
    }
    verification = SourceVerificationReport(
        rows=[], by_sheet={}, total_rows=5, verifiable_rows=5, confirmed_rows=5,
        weak_rows=0, unconfirmed_rows=0, skipped_rows=0, rows_with_provenance=5,
        grounding_ratio=1.0, provenance_ratio=1.0,
    )

    inventory = build_source_object_inventory(final_data, document)
    assessment = assess_quality(final_data, verification, inventory)

    assert inventory.total_objects == 1
    assert inventory.confirmed_objects == 1
    assert inventory.coverage_ratio == 1.0
    assert assessment.score == 100.0
    assert assessment.metrics["source_inventory_connected"] is True


def test_source_object_inventory_marks_unlinked_table():
    document = PDFContent(
        total_pages=1,
        pages=[PageContent(
            page_number=1,
            text="표 2-12 연도별 에너지 소비량",
            tables=["<table><tr><th>연도</th><th>전력</th></tr><tr><td>2021</td><td>10</td></tr></table>"],
            images=[],
        )],
        full_text="표 2-12 연도별 에너지 소비량",
    )

    inventory = build_source_object_inventory({}, document)

    assert inventory.unconfirmed_objects == 1
    assert inventory.coverage_ratio == 0.0


def test_source_verification_rows_are_written_to_optional_sheet(tmp_path: Path):
    data = {
        "source_verification": [{
            "지자체명": "서울", "대상시트": "08_감축사업목록", "원본행번호": 2,
            "상태": "확인", "신뢰도": "high", "출처페이지": "2", "확인페이지": 2,
            "점수": 5.0, "검색어": "태양광", "일치검색어": "태양광",
            "행요약": "사업명=태양광", "검수메시지": "확인",
        }],
    }
    output = write_excel(data, tmp_path / "verified.xlsx")

    workbook = load_workbook(output, read_only=True)
    assert "20_원문대조" in workbook.sheetnames
    assert workbook["20_원문대조"].cell(2, 4).value == "확인"


def test_source_object_inventory_rows_are_written_to_optional_sheet(tmp_path: Path):
    data = {
        "source_object_inventory": [{
            "지자체명": "서울", "객체ID": "p1_table_1", "객체유형": "표",
            "출처페이지": 1, "번호": "표 1-1", "캡션": "표 1-1 계획",
            "섹션": "계획", "행수": 2, "열수": 2, "연결상태": "확인",
            "완전성점수": 1.0, "연결시트": "08_감축사업목록", "연결행수": 1,
            "검수메시지": "확인",
        }],
    }
    output = write_excel(data, tmp_path / "inventory.xlsx")

    workbook = load_workbook(output, read_only=True)
    assert "21_원문객체인벤토리" in workbook.sheetnames
    assert workbook["21_원문객체인벤토리"].cell(2, 10).value == "확인"


def test_marked_pdf_is_created_for_exact_matched_term(tmp_path: Path):
    source = tmp_path / "source.pdf"
    document = fitz.open()
    page = document.new_page()
    page.insert_text((72, 72), "Solar Project 2030")
    document.save(source)
    document.close()

    report = verify_final_data(
        {"mitigation_projects": [{"사업명": "Solar Project", "부문": "energy", "출처페이지": "1"}]},
        _document("Solar Project 2030"),
    )
    output, count = create_marked_pdf(source, report, tmp_path / "marked.pdf")

    assert output is not None and output.exists()
    assert count >= 1
