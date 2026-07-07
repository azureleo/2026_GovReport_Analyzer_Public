from __future__ import annotations

from pathlib import Path

from openpyxl import Workbook

from scripts.reclean_offline import analyze_workbook, render_markdown


def _save_workbook(path: Path) -> None:
    wb = Workbook()
    ws = wb.active
    ws.title = "03_배출현황_지역"
    ws.append(["지자체명", "배출유형", "부문", "세부부문", "연도", "배출량", "단위"])
    ws.append(["강원특별자치도", "", "에너지", "연료연소", 2020, 10, "천톤"])
    ws.append(["강원특별자치도", "직접배출", "에너지", "연료연소", 2020, 10, None])

    ws = wb.create_sheet("04_배출현황_관리권한")
    ws.append(["지자체명", "관리부문", "세부부문", "직간접구분", "연도", "배출량", "단위"])
    ws.append(["강원특별자치도", "건물", "공공", "", 2030, 7, "천톤"])
    ws.append(["강원특별자치도", "건물", "공공", "간접", 2030, 7, None])

    ws = wb.create_sheet("19_검증리포트")
    ws.append(["지자체명", "심각도", "영역", "항목", "문제내용", "권장조치"])
    ws.append(["강원특별자치도", "경고", "중복제거", "중복 키 값 충돌(테스트)", "수정 전 충돌", "확인"])
    wb.save(path)


def test_reclean_offline_reports_conflict_blank_key_and_row_deltas(tmp_path: Path) -> None:
    # Given: 기존 출력 xlsx에 중복제거 이슈와 03/04 빈 키 행이 있으면
    xlsx = tmp_path / "강원_결과_v6.xlsx"
    _save_workbook(xlsx)

    # When: 오프라인 재정제 분석을 실행하면
    report = analyze_workbook(xlsx)

    # Then: 재추출 없이 organizer 재정제 전후 감소 수치가 계산된다.
    assert report.municipality == "강원특별자치도"
    assert report.conflicts_before == 1
    assert report.conflicts_after == 0
    assert report.sheet_metrics["03_배출현황_지역"].blank_key_rows_before == 1
    assert report.sheet_metrics["03_배출현황_지역"].blank_key_rows_after == 0
    assert report.sheet_metrics["03_배출현황_지역"].rows_before == 2
    assert report.sheet_metrics["03_배출현황_지역"].rows_after == 1
    assert report.sheet_metrics["04_배출현황_관리권한"].blank_key_rows_before == 1
    assert report.sheet_metrics["04_배출현황_관리권한"].blank_key_rows_after == 0


def test_reclean_offline_markdown_contains_gate_table(tmp_path: Path) -> None:
    # Given: 오프라인 재정제 분석 결과가 있으면
    xlsx = tmp_path / "강원_결과_v6.xlsx"
    _save_workbook(xlsx)
    report = analyze_workbook(xlsx)

    # When: 마크다운 리포트를 렌더링하면
    markdown = render_markdown([report])

    # Then: 최종 게이트에 붙일 수 있는 감소 표가 생성된다.
    assert "| 파일 | 지자체 | 지표 | 수정 전 | 수정 후 | 감소 |" in markdown
    assert "중복제거 충돌" in markdown
    assert "03_배출현황_지역 빈 키 행" in markdown
