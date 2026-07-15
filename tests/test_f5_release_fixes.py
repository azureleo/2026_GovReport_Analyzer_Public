from __future__ import annotations

import subprocess
import sys
from pathlib import Path

import pytest
from openpyxl import load_workbook

import config
import main
from agents.supervisor import Supervisor
from utils import llm_client
from utils.pdf_reader import PDFContent, PageContent


def test_retries_zero_is_argparse_error(monkeypatch, tmp_path, capsys) -> None:
    # Given: 사용자가 전체 파이프라인 재시도 횟수로 0을 입력하면
    input_path = tmp_path / "input.pdf"
    input_path.write_bytes(b"%PDF-1.4\n")
    monkeypatch.setattr(sys, "argv", ["main.py", str(input_path), "-r", "0"])

    # When / Then: argparse 단계에서 한국어 오류로 종료한다.
    with pytest.raises(SystemExit) as exc:
        main.main()
    assert exc.value.code == 2
    assert "1 이상이어야 합니다" in capsys.readouterr().err


def test_stdio_reconfigure_uses_replace_errors(monkeypatch) -> None:
    # Given: cp949 리다이렉트처럼 일부 문자를 인코딩하지 못하는 표준 스트림
    calls: dict[str, list[dict]] = {"stdout": [], "stderr": []}

    class FakeStream:
        def __init__(self, key: str):
            self.key = key

        def reconfigure(self, **kwargs) -> None:
            calls[self.key].append(kwargs)

    monkeypatch.setattr(sys, "stdout", FakeStream("stdout"))
    monkeypatch.setattr(sys, "stderr", FakeStream("stderr"))

    # When: CLI 시작부 인코딩 오류 정책을 적용하면
    main._configure_stdio_errors()

    # Then: 인코딩은 바꾸지 않고 오류 문자만 대체한다.
    assert calls == {"stdout": [{"errors": "replace"}], "stderr": [{"errors": "replace"}]}


def _patch_minimal_supervisor_pipeline(monkeypatch, pdf: PDFContent) -> None:
    monkeypatch.setattr(config, "GAP_FILL_ENABLED", False)
    monkeypatch.setattr(config, "HYBRID_REVIEW_ENABLED", False)
    monkeypatch.setattr(config, "SHEET_CLOSED_LOOP_ENABLED", False)
    monkeypatch.setattr(config, "MIN_DOCUMENT_TEXT_CHARS", 500)
    monkeypatch.setattr("agents.supervisor.extract_pdf", lambda input_path, render_graph_pages=True: pdf)
    monkeypatch.setattr("agents.supervisor.GuidelineAgent.get_all_prompts", lambda self: {"document_meta": ""})
    monkeypatch.setattr("agents.supervisor.GuidelineAgent.report", lambda self: "[가이드라인] 테스트")
    monkeypatch.setattr(
        "agents.extractor_agent.ExtractorAgent.extract",
        lambda self, pages, full_text, extraction_prompts: {
            "municipality_name": "서울특별시",
            "document_meta": [{"지자체명": "서울특별시", "계획명": "계획"}],
        },
    )


def test_textless_document_warning_is_logged_and_written_to_validation_report(monkeypatch, tmp_path, capsys) -> None:
    # Given: 전체 텍스트 레이어가 임계값보다 작은 스캔본 의심 PDF
    pdf = PDFContent(total_pages=1, pages=[PageContent(1, "  ", [], [])], full_text="  ")
    _patch_minimal_supervisor_pipeline(monkeypatch, pdf)
    monkeypatch.setattr(
        llm_client,
        "call_text",
        lambda prompt, system="", max_retries=1, stage=None: (
            '{"assessment": "ok", "quality_level": "보통", "key_issues": []}'
        ),
    )

    # When: 파이프라인을 끝까지 실행하면
    output_path = tmp_path / "scan.xlsx"
    result_path = Supervisor().run(
        input_path=tmp_path / "scan.pdf",
        output_path=output_path,
        max_pipeline_retries=1,
        include_images=False,
    )

    # Then: 경고가 출력되고 19_검증리포트 정보 항목으로 남는다.
    captured = capsys.readouterr().out
    assert result_path == output_path
    assert "텍스트 레이어가 거의 없습니다" in captured
    workbook = load_workbook(result_path)
    rows = list(workbook["19_검증리포트"].iter_rows(values_only=True))
    assert any(row[1] == "정보" and row[3] == "텍스트 레이어 부족" for row in rows[1:])


def test_final_quality_review_quota_after_excel_save_exits_successfully(monkeypatch, tmp_path, capsys) -> None:
    # Given: 엑셀 저장은 끝났지만 저장 이후 LLM 최종 검수에서 quota가 발생하면
    text = "서울특별시 탄소중립 기본계획 " * 40
    pdf = PDFContent(total_pages=1, pages=[PageContent(1, text, [], [])], full_text=text)
    _patch_minimal_supervisor_pipeline(monkeypatch, pdf)
    monkeypatch.setattr(
        Supervisor,
        "_llm_quality_review",
        lambda self, final_data: (_ for _ in ()).throw(llm_client.LLMQuotaExceededError("quota")),
    )

    # When: Supervisor 파이프라인을 실행하면
    output_path = tmp_path / "quota-after-save.xlsx"
    result_path = Supervisor().run(
        input_path=tmp_path / "input.pdf",
        output_path=output_path,
        max_pipeline_retries=1,
        include_images=False,
    )

    # Then: 결과 파일을 보존하고 정상 경로로 반환한다.
    captured = capsys.readouterr().out
    assert result_path == output_path
    assert output_path.exists()
    assert "한도 초과로 최종 검수는 생략했습니다" in captured


def test_requirements_include_pytest_dependency() -> None:
    # Given / When: 신규 설치자가 requirements.txt만 설치하면
    body = Path("requirements.txt").read_text(encoding="utf-8")

    # Then: 문서에 있는 pytest 실행 명령이 동작할 dev 도구가 포함된다.
    assert "pytest>=8.0.0" in body


def test_local_agent_timeout_help_and_env_example_match_300() -> None:
    # Given / When: CLI help와 .env.example의 로컬 에이전트 timeout 기본값을 확인하면
    completed = subprocess.run(
        [sys.executable, "main.py", "--help"],
        text=True,
        capture_output=True,
        check=False,
    )
    env_body = Path(".env.example").read_text(encoding="utf-8")

    # Then: config.py의 실제 기본값 300초와 일치한다.
    assert completed.returncode == 0
    assert "기본값: 300" in completed.stdout
    assert "LOCAL_AGENT_TIMEOUT=300" in env_body
