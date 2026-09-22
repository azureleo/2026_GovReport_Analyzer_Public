"""M3-1 교차 백엔드 무골든 품질 프로브 테스트."""
from __future__ import annotations

import json
from pathlib import Path

import openpyxl

import config
from scripts import cross_backend_probe


def _시트쓰기(workbook, sheet_name: str, rows: list[dict]) -> None:
    ws = workbook.create_sheet(sheet_name)
    headers = list(config.EXCEL_HEADERS[sheet_name])
    ws.append(headers)
    for row in rows:
        ws.append([row.get(header) for header in headers])


def _워크북(path: Path, sheets: dict[str, list[dict]]) -> Path:
    workbook = openpyxl.Workbook()
    del workbook[workbook.sheetnames[0]]
    for sheet_name, rows in sheets.items():
        _시트쓰기(workbook, sheet_name, rows)
    workbook.save(path)
    return path


def _메타() -> list[dict]:
    return [{"지자체명": "테스트시", "계획시작연도": 2020, "계획종료연도": 2030}]


def test_L3_L2_L1과_md_json을_결정론적으로_생성한다(tmp_path: Path) -> None:
    a = _워크북(
        tmp_path / "A.xlsx",
        {
            "00_문서메타": _메타(),
            "03_배출현황_지역": [
                {"배출유형": "직접배출", "부문": "건물", "세부부문": "가정", "연도": 2020, "배출량": 100, "단위": "천톤CO2eq", "데이터상태": "reported"},
                {"배출유형": "직접배출", "부문": "수송", "세부부문": "도로", "연도": 2021, "배출량": 200, "단위": "천톤CO2eq", "데이터상태": "reported"},
                {"배출유형": "직접배출", "부문": "비표준", "세부부문": "기타", "연도": 2000, "배출량": 300, "단위": "천톤CO2eq", "데이터상태": "invalid"},
                {"배출유형": "직접배출", "부문": "산업", "세부부문": "", "연도": 2023, "배출량": 400, "단위": "천톤CO2eq", "데이터상태": "reported"},
            ],
            "19_검증리포트": [
                {"심각도": "경고", "영역": "배출현황_지역", "항목": "중복 키 값 충돌(테스트)"},
                {"심각도": "경고", "영역": "배출현황_지역", "항목": "원장 호출실패"},
            ],
        },
    )
    b = _워크북(
        tmp_path / "B.xlsx",
        {
            "00_문서메타": _메타(),
            "03_배출현황_지역": [
                {"배출유형": "직접배출", "부문": "건물", "세부부문": "가정", "연도": 2020, "배출량": 100.4, "단위": "천톤CO2eq", "데이터상태": "reported"},
                {"배출유형": "직접배출", "부문": "수송", "세부부문": "도로", "연도": 2021, "배출량": 220, "단위": "천톤CO2eq", "데이터상태": "reported"},
                {"배출유형": "직접배출", "부문": "폐기물", "세부부문": "매립", "연도": 2022, "배출량": 50, "단위": "천톤CO2eq", "데이터상태": "reported"},
            ],
        },
    )

    result = cross_backend_probe.probe_workbooks(
        a, b, report_dir=tmp_path / "report", label="A_vs_B"
    )

    sheet = result["시트별"]["03_배출현황_지역"]
    assert sheet["매칭수"] == 2
    assert sheet["키합의율"] == 0.6667
    assert sheet["값합의율"] == 0.5
    assert sheet["A단독행수"] == 2
    assert sheet["B단독행수"] == 1
    assert sheet["L1"]["A"]["밀도"] == 0.5
    assert sheet["L1"]["B"]["밀도"] == 0.0
    assert sheet["L2"]["A"]["일차키채움률"] == 1.0
    assert sheet["L2"]["A"]["옵셔널필드"] == {
        "대상필드": ["세부부문"],
        "채움수": 3,
        "전체수": 4,
        "채움률": 0.75,
    }
    assert sheet["L2"]["B"]["옵셔널필드"]["채움률"] == 1.0
    assert sheet["L2"]["A"]["표준코드"]["매핑률"] == 0.75
    assert sheet["L2"]["B"]["표준코드"]["매핑률"] == 1.0
    assert sheet["L2"]["A"]["연도"]["정합률"] == 0.75
    assert sheet["L2"]["B"]["연도"]["정합률"] == 1.0
    assert result["L1_자기감사"]["A"]["심각도별"]["경고"] == 2
    assert result["L1_자기감사"]["A"]["유형별"]["충돌"] == 1
    assert result["L1_자기감사"]["A"]["유형별"]["원장"] == 1
    assert result["L1_자기감사"]["B"]["유형별"] == {
        "충돌": 0,
        "검산": 0,
        "스케일": 0,
        "원장": 0,
        "빈 시트": 0,
    }
    md_path = Path(result["리포트"]["md"])
    json_path = Path(result["리포트"]["json"])
    assert md_path.exists()
    assert json_path.exists()
    assert "# 무골든 교차 백엔드 품질 리포트 (A_vs_B)" in md_path.read_text(encoding="utf-8")
    assert "A옵셔널채움률 | B옵셔널채움률" in md_path.read_text(encoding="utf-8")
    payload = json.loads(json_path.read_text(encoding="utf-8"))
    assert payload["시트별"]["03_배출현황_지역"]["값합의율"] == 0.5
    assert payload["시트별"]["03_배출현황_지역"]["L2"]["A"]["옵셔널필드"]["채움률"] == 0.75
    assert payload["LLM호출수"] == 0


def test_L2는_지표세부범주를_필수키에서_제외하고_별도_채움률로_병기한다(
    tmp_path: Path,
) -> None:
    rows = [
        {"지표범주": "자연환경", "지표세부범주": "", "지표명": "평균기온", "연도": 2020, "값": 13},
        {"지표범주": "인문사회", "지표세부범주": "인구", "지표명": "총인구", "연도": 2020, "값": 100},
    ]
    a = _워크북(tmp_path / "A.xlsx", {"02_지역여건": rows})
    b = _워크북(tmp_path / "B.xlsx", {"02_지역여건": rows})

    result = cross_backend_probe.probe_workbooks(
        a, b, report_dir=tmp_path / "report"
    )

    l2 = result["시트별"]["02_지역여건"]["L2"]["A"]
    assert l2["일차키채움률"] == 1.0
    assert l2["옵셔널필드"] == {
        "대상필드": ["지표세부범주"],
        "채움수": 1,
        "전체수": 2,
        "채움률": 0.5,
    }


def test_서술형은_보수티어의_키합의만_계산한다(
    tmp_path: Path, monkeypatch
) -> None:
    rows_a = [
        {"개요유형": "목적", "항목명": "계획 목적", "항목값": "탄소중립"},
        {"개요유형": "배경", "항목명": "계획 수립 배경", "항목값": "기후위기"},
    ]
    rows_b = [
        {"개요유형": "목적", "항목명": "계획 목적", "항목값": "탄소중립"},
        {"개요유형": "배경", "항목명": "계획 수립 배경 및 필요성", "항목값": "기후위기"},
    ]
    a = _워크북(tmp_path / "A.xlsx", {"01_계획개요": rows_a})
    b = _워크북(tmp_path / "B.xlsx", {"01_계획개요": rows_b})
    original = cross_backend_probe.매칭하기
    calls: list[bool] = []

    def spy(golden_rows, output_rows, sheet_name, 보수티어만=False):
        calls.append(보수티어만)
        return original(golden_rows, output_rows, sheet_name, 보수티어만=보수티어만)

    monkeypatch.setattr(cross_backend_probe, "매칭하기", spy)

    result = cross_backend_probe.probe_workbooks(a, b, report_dir=tmp_path / "report")

    sheet = result["시트별"]["01_계획개요"]
    assert sheet["매칭수"] == 1
    assert sheet["키합의율"] == 0.5
    assert sheet["값합의율"] is None
    assert sheet["L2"]["A"]["연도"]["정합률"] is None
    assert all(calls)


def test_세부부문이_달라도_값이_같으면_완화티어로_매칭한다(tmp_path: Path) -> None:
    a = _워크북(
        tmp_path / "A.xlsx",
        {
            "03_배출현황_지역": [
                {"배출유형": "직접배출", "부문": "건물", "세부부문": "가정", "연도": 2020, "배출량": 100, "단위": "톤"}
            ]
        },
    )
    b = _워크북(
        tmp_path / "B.xlsx",
        {
            "03_배출현황_지역": [
                {"배출유형": "직접배출", "부문": "건물", "세부부문": "상업", "연도": 2020, "배출량": 100, "단위": "톤"}
            ]
        },
    )

    result = cross_backend_probe.probe_workbooks(a, b, report_dir=tmp_path / "report")

    sheet = result["시트별"]["03_배출현황_지역"]
    assert sheet["엄격매칭수"] == 0
    assert sheet["완화매칭수"] == 1
    assert sheet["키합의율"] == 1.0
    assert sheet["값합의율"] == 1.0


def test_L4는_값합의율_하위_3시트에서_각_10개까지만_고른다() -> None:
    sheet_results = {
        "03_배출현황_지역": {
            "A유효행수": 12,
            "B유효행수": 12,
            "값합의율": 0.2,
            "표본후보키": [{"구분": "값불일치", "키": f"k{i}"} for i in range(12)],
        },
        "04_배출현황_관리권한": {
            "A유효행수": 1,
            "B유효행수": 1,
            "값합의율": 0.1,
            "표본후보키": [{"구분": "값불일치", "키": "k"}],
        },
        "05_배출전망": {
            "A유효행수": 1,
            "B유효행수": 1,
            "값합의율": 0.3,
            "표본후보키": [{"구분": "값불일치", "키": "k"}],
        },
        "06_감축목표": {
            "A유효행수": 1,
            "B유효행수": 1,
            "값합의율": 0.4,
            "표본후보키": [{"구분": "값불일치", "키": "k"}],
        },
    }

    samples = cross_backend_probe._L4_표본(sheet_results)

    assert [row["시트"] for row in samples] == [
        "04_배출현황_관리권한",
        "03_배출현황_지역",
        "05_배출전망",
    ]
    assert len(samples[1]["표본키"]) == 10


def test_cli는_낮은_합의율을_실패로_판정하지_않는다(tmp_path: Path) -> None:
    a = _워크북(
        tmp_path / "A.xlsx",
        {"08_감축사업목록": [{"관리번호": "A-1", "사업명": "사업 A"}]},
    )
    b = _워크북(
        tmp_path / "B.xlsx",
        {"08_감축사업목록": [{"관리번호": "B-1", "사업명": "사업 B"}]},
    )

    exit_code = cross_backend_probe.main(
        [str(a), str(b), "--report-dir", str(tmp_path / "report"), "--label", "low"]
    )

    assert exit_code == 0


def test_cli는_형식오류일_때만_비정상_종료한다(tmp_path: Path) -> None:
    invalid = tmp_path / "invalid.xlsx"
    invalid.write_text("not an xlsx", encoding="utf-8")

    assert cross_backend_probe.main([str(invalid), str(invalid)]) == 2
