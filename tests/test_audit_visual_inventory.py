from __future__ import annotations

import json
import subprocess
import sys
from io import BytesIO
from pathlib import Path

import fitz
import pytest
from openpyxl import Workbook, load_workbook
from PIL import Image, ImageDraw

from agents import image_agent
from scripts import audit_visual_inventory as audit


INVENTORY_HEADERS = [
    "요소ID", "페이지", "요소유형", "제목", "데이터포함", "기대추출", "관련시트", "비고",
]


def _item(element_id: str, page: int, element_type: str = "그래프", expected: bool = True) -> audit.InventoryItem:
    return audit.InventoryItem(
        element_id=element_id,
        page=page,
        element_type=element_type,
        title=f"제목 {element_id}",
        data_included=True,
        expected=expected,
        related_sheet="03_배출현황_지역",
        note="",
    )


def _page(
    page: int,
    image_count: int,
    passed_count: int,
    score: int | None,
    reference_filtered: int = 0,
    drawing_count: int = 0,
) -> audit.PageEvidence:
    return audit.PageEvidence(
        page=page,
        image_count=image_count,
        reduced_image_count=image_count,
        triage_passed_count=passed_count,
        reference_filtered_count=reference_filtered,
        top_score=score,
        top_reasons=("fixture",),
        drawing_count=drawing_count,
    )


def _save_inventory(path: Path, rows: list[list[str | int]]) -> None:
    wb = Workbook()
    ws = wb.active
    ws.title = "시각요소"
    ws.append(INVENTORY_HEADERS)
    for row in rows:
        ws.append(row)
    wb.save(path)


def _save_pipeline_output(path: Path, rows: list[list[str]]) -> None:
    wb = Workbook()
    ws = wb.active
    ws.title = "16_시각자료목록"
    ws.append(["지자체명", "시각자료ID", "캡션", "유형", "데이터포함여부", "추출값요약", "디지타이징필요", "관련시트"])
    for row in rows:
        ws.append(row)
    wb.save(path)


def _make_chart_png() -> bytes:
    image = Image.new("RGB", (420, 300), "white")
    draw = ImageDraw.Draw(image)
    for idx, height in enumerate([80, 150, 210, 120]):
        x0 = 60 + idx * 80
        draw.rectangle((x0, 250 - height, x0 + 45, 250), fill="steelblue", outline="black")
    for y in range(60, 260, 40):
        draw.line((40, y, 390, y), fill="black")
    draw.text((50, 20), "배출량 그래프 2020 2030", fill="black")
    buf = BytesIO()
    image.save(buf, format="PNG")
    return buf.getvalue()


def _make_visual_pdf(path: Path) -> None:
    doc = fitz.open()
    page1 = doc.new_page(width=595, height=842)
    page1.insert_text((72, 72), "[그림 1-1] 온실가스 배출량 그래프\n배출량 감축목표 현황")
    page1.insert_image(fitz.Rect(72, 120, 492, 420), stream=_make_chart_png())
    page2 = doc.new_page(width=595, height=842)
    page2.insert_text((72, 72), "[그림 2-1] 벡터 차트 후보")
    for idx in range(70):
        x = 72 + (idx % 10) * 18
        y = 130 + (idx // 10) * 14
        page2.draw_rect(fitz.Rect(x, y, x + 12, y + 8), color=(0, 0, 0), fill=None)
    doc.save(path)
    doc.close()


def test_funnel_statuses_when_each_stage_loses_or_records_element() -> None:
    # Given: 페이지별 이미지 확보, 참고자료 제외, triage 탈락, vision 유실, 기록 상태가 준비되면
    items = [_item("E001", 1), _item("E002", 2), _item("E003", 3), _item("E004", 4), _item("E005", 5)]
    pages = {
        1: _page(1, image_count=0, passed_count=0, score=None, drawing_count=42),
        2: _page(2, image_count=1, passed_count=0, score=-14, reference_filtered=1),
        3: _page(3, image_count=1, passed_count=0, score=2),
        4: _page(4, image_count=1, passed_count=1, score=7),
        5: _page(5, image_count=1, passed_count=1, score=7),
    }
    records = [audit.VisualRecord("V5-001", 5, "캡션", "그래프", "Y", "N", "03_배출현황_지역")]

    # When: 요소별 funnel 상태를 판정하면
    results = audit.classify_inventory(items, pages, records)

    # Then: 각 유실 단계가 한국어 상태값으로 보존된다.
    assert [row.status for row in results] == [
        "이미지_미추출", "참고자료_제외", "triage_탈락", "vision_유실", "기록됨",
    ]


def test_parse_visual_page_id_when_id_has_page_or_only_sequence() -> None:
    # Given: 16 시트의 시각자료ID 값이 페이지 포함형과 순번형으로 섞이면
    # When / Then: V{페이지}-{idx}만 페이지 번호를 돌려준다.
    assert audit.parse_visual_page_id("V161-003") == 161
    assert audit.parse_visual_page_id(" V7-1 ") == 7
    assert audit.parse_visual_page_id("V003") is None
    assert audit.parse_visual_page_id("") is None


def test_same_page_partial_record_when_two_expected_elements_share_one_output_row() -> None:
    # Given: 같은 페이지에 기대추출 요소 2개와 16 시트 행 1개만 있으면
    items = [_item("E001", 10), _item("E002", 10)]
    pages = {10: _page(10, image_count=2, passed_count=2, score=8)}
    records = [audit.VisualRecord("V10-001", 10, "캡션", "그래프", "Y", "Y", "06_감축목표")]

    # When: 페이지 단위로 보수적 매칭을 수행하면
    results = audit.classify_inventory(items, pages, records)

    # Then: 부족분만 동일페이지_부분기록으로 표시된다.
    assert [row.status for row in results] == ["기록됨", "동일페이지_부분기록"]


def test_threshold_sensitivity_when_scores_cross_min_score() -> None:
    # Given: triage 최고 점수 4, 5, 6인 기대추출 요소 3개가 있으면
    items = [_item("E001", 1), _item("E002", 2), _item("E003", 3)]
    pages = {page: _page(page, image_count=1, passed_count=1, score=score) for page, score in [(1, 4), (2, 5), (3, 6)]}
    audits = audit.classify_inventory(items, pages, [])

    # When: 임계값 민감도를 집계하면
    sensitivity = audit.calculate_sensitivity(audits, pages)

    # Then: MIN_SCORE=5에서는 2개, 3에서는 3개가 triage 단계 리콜에 들어간다.
    assert sensitivity.image_thresholds[5].passed_elements == 2
    assert sensitivity.image_thresholds[3].passed_elements == 3
    assert sensitivity.image_thresholds[5].passed_images == 2


def test_inventory_format_error_when_required_column_missing_or_type_is_unknown(tmp_path: Path) -> None:
    # Given: 필수 컬럼이 빠진 인벤토리와 요소유형 오타가 있는 인벤토리가 있으면
    missing = tmp_path / "missing.xlsx"
    wb = Workbook()
    ws = wb.active
    ws.title = "시각요소"
    ws.append(["페이지", "요소유형", "데이터포함", "기대추출"])
    ws.append([1, "그래프", "Y", "Y"])
    wb.save(missing)

    typo = tmp_path / "typo.xlsx"
    _save_inventory(typo, [["E001", 1, "사진", "제목", "Y", "Y", "", ""]])

    # When / Then: 한국어 형식 오류로 중단된다.
    with pytest.raises(audit.InventoryFormatError, match="필수 컬럼"):
        audit.load_inventory(missing)
    with pytest.raises(audit.InventoryFormatError, match="요소유형"):
        audit.load_inventory(typo)


def test_dump_mode_smoke_when_tiny_pdf_has_image_and_vector_page(tmp_path: Path) -> None:
    # Given: 이미지 1장과 벡터 도형 다수 페이지가 있는 초소형 PDF가 있으면
    source = tmp_path / "visual.pdf"
    out = tmp_path / "draft.xlsx"
    _make_visual_pdf(source)

    # When: dump 모드가 초안 xlsx를 만들면
    row_count = audit.dump_inventory(source, out)

    # Then: 안내 문구와 자동 컬럼이 채워진 시각요소 시트가 생성된다.
    assert row_count >= 1
    wb = load_workbook(out)
    assert "안내" in wb.sheetnames
    assert "시각요소" in wb.sheetnames
    assert "통째로 놓친 요소" in str(wb["안내"]["A1"].value)
    headers = [cell.value for cell in wb["시각요소"][1]]
    assert "자동_트리아지점수" in headers
    assert "자동_이미지크기" in headers
    assert "자동_캡션후보" in headers
    assert "자동_벡터드로잉수" in headers


def test_audit_cli_writes_markdown_and_json_report_for_fake_inventory(tmp_path: Path) -> None:
    # Given: 초소형 PDF, 사람 인벤토리 1행, 16_시각자료목록 1행이 있으면
    source = tmp_path / "visual.pdf"
    inventory = tmp_path / "inventory.xlsx"
    pipeline = tmp_path / "pipeline.xlsx"
    report_dir = tmp_path / "reports"
    _make_visual_pdf(source)
    _save_inventory(inventory, [["E001", 1, "그래프", "[그림 1-1]", "Y", "Y", "03_배출현황_지역", ""]])
    _save_pipeline_output(pipeline, [["서울특별시", "V1-001", "[그림 1-1]", "그래프", "Y", "요약", "N", "03_배출현황_지역"]])

    # When: 실제 CLI audit 서브커맨드를 실행하면
    completed = subprocess.run(
        [
            sys.executable,
            "scripts/audit_visual_inventory.py",
            "audit",
            str(inventory),
            str(pipeline),
            str(source),
            "--report-dir",
            str(report_dir),
        ],
        text=True,
        capture_output=True,
        check=False,
    )

    # Then: 리콜 요약이 담긴 markdown/json 리포트가 생성된다.
    assert completed.returncode == 0, completed.stderr + completed.stdout
    report_paths = json.loads(completed.stdout.strip().splitlines()[-1])
    md_path = Path(report_paths["markdown"])
    json_path = Path(report_paths["json"])
    assert md_path.exists()
    assert json_path.exists()
    assert "리콜" in md_path.read_text(encoding="utf-8")


def test_image_agent_private_imports_are_same_objects() -> None:
    # Given: 감사 도구가 image_agent의 기존 private helper를 재사용하면
    # When / Then: 별도 래핑이나 동작 변경 없이 동일 객체를 가리킨다.
    assert audit._triage_image is image_agent._triage_image
    assert audit._coverage_reduce_images is image_agent._coverage_reduce_images
    assert audit._is_relevant_image is image_agent._is_relevant_image
