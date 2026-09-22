"""골든셋 대조 채점기 회귀 테스트."""
from __future__ import annotations

import json
import subprocess
import sys
from pathlib import Path

import openpyxl
import pytest

import config
from agents import organizer_agent
from scripts.golden_score_contract import 계약헤더
from scripts.score_against_golden import score_workbooks

골든열 = [
    "골든_출처유형", "골든_출처페이지", "골든_채점제외", "골든_비고",
    "골든_산출유형",
]


def _계약열(sheet_name: str) -> list[str]:
    return 계약헤더(sheet_name)


def _시트_쓰기(wb, sheet_name: str, rows: list[dict], *, golden: bool = False) -> None:
    ws = wb.create_sheet(sheet_name)
    headers = _계약열(sheet_name)
    if golden and sheet_name != "00_문서메타":
        headers = [*headers, *골든열]
    if golden and sheet_name == "00_문서메타":
        headers = [*headers, "골든_비고"]
    if not golden:
        headers = list(config.EXCEL_HEADERS[sheet_name])
    ws.append(headers)
    for row in rows:
        ws.append([row.get(header) for header in headers])


def _워크북(path: Path, sheets: dict[str, list[dict]], *, golden: bool = False) -> Path:
    wb = openpyxl.Workbook()
    del wb[wb.sheetnames[0]]
    for sheet_name, rows in sheets.items():
        _시트_쓰기(wb, sheet_name, rows, golden=golden)
    wb.save(path)
    return path


def _점수(tmp_path: Path, output_sheets: dict[str, list[dict]], golden_sheets: dict[str, list[dict]]) -> dict:
    output = _워크북(tmp_path / "출력.xlsx", output_sheets)
    golden = _워크북(tmp_path / "골든.xlsx", golden_sheets, golden=True)
    return score_workbooks(output, golden, report_dir=tmp_path / "리포트", label="테스트", write_json=True)


def test_organizer_공개_별칭은_기존_private_함수와_같은_객체다() -> None:
    """기존 정규화 함수의 공개 별칭은 본문을 바꾸지 않고 같은 객체를 가리킨다."""
    assert organizer_agent.dedup_key_text is organizer_agent._dedup_key_text
    assert organizer_agent.normalize_project_id is organizer_agent._normalize_project_id
    assert organizer_agent.normalize_direct_indirect_type is organizer_agent._normalize_direct_indirect_type
    assert organizer_agent.to_float is organizer_agent._to_float
    assert organizer_agent.normalize_provenance_pages is organizer_agent._normalize_provenance_pages


def test_산출유형별_reported와_calculated를_분리_평가한다(tmp_path: Path) -> None:
    result = _점수(
        tmp_path,
        {
            "03_배출현황_지역": [
                {
                    "배출유형": "직접배출", "부문": "건물", "세부부문": "전기",
                    "연도": 2020, "배출량": 100, "단위": "톤",
                    "derivation_type": "explicit",
                },
                {
                    "배출유형": "직접배출", "부문": "건물", "세부부문": "전기",
                    "연도": 2021, "배출량": 90, "단위": "톤",
                    "derivation_type": "calculated",
                },
            ]
        },
        {
            "03_배출현황_지역": [
                {
                    "배출유형": "직접배출", "부문": "건물", "세부부문": "전기",
                    "연도": 2020, "배출량": 100, "단위": "톤",
                    "골든_출처유형": "텍스트표", "골든_출처페이지": "1",
                    "골든_산출유형": "reported",
                },
                {
                    "배출유형": "직접배출", "부문": "건물", "세부부문": "전기",
                    "연도": 2021, "배출량": 90, "단위": "톤",
                    "골든_출처유형": "텍스트표", "골든_출처페이지": "2",
                    "골든_산출유형": "calculated",
                },
            ]
        },
    )

    assert result["산출유형별_전체"]["reported"]["리콜"] == 1.0
    assert result["산출유형별_전체"]["reported"]["산출유형일치율"] == 1.0
    assert result["산출유형별_전체"]["calculated"]["리콜"] == 1.0
    assert result["산출유형별_전체"]["calculated"]["값일치율"] == 1.0


def test_시트03_엄격키_매칭으로_리콜과_정밀도를_계산한다(tmp_path: Path) -> None:
    result = _점수(
        tmp_path,
        {
            "03_배출현황_지역": [
                {"배출유형": "직접배출", "부문": "건물", "세부부문": "전기", "연도": 2020, "배출량": 100, "단위": "천 tCO2eq", "출처페이지": "10"},
                {"배출유형": "직접배출", "부문": "산업", "세부부문": "공정", "연도": 2020, "배출량": 50, "단위": "천 tCO2eq", "출처페이지": "12"},
            ]
        },
        {
            "03_배출현황_지역": [
                {"배출유형": "직접배출", "부문": "건물", "세부부문": "전기", "연도": 2020, "배출량": 100, "단위": "천 tCO2eq", "골든_출처유형": "본문텍스트", "골든_출처페이지": "10"},
                {"배출유형": "직접배출", "부문": "수송", "세부부문": "도로", "연도": 2020, "배출량": 200, "단위": "천 tCO2eq", "골든_출처유형": "그래프", "골든_출처페이지": "11"},
            ]
        },
    )

    sheet = result["시트별"]["03_배출현황_지역"]
    assert sheet["엄격매칭수"] == 1
    assert sheet["완화매칭수"] == 0
    assert sheet["리콜"] == 0.5
    assert sheet["정밀도"] == 0.5
    assert sheet["값일치율"] == 1.0
    assert result["출처유형별_전체"]["그래프"]["리콜"] == 0.0
    assert result["출처유형별_전체"]["시각 유래"]["리콜"] == 0.0
    assert result["출처페이지통계"]["교집합비율"] == 1.0


def test_정규화키와_완화키로_행을_매칭한다(tmp_path: Path) -> None:
    result = _점수(
        tmp_path,
        {
            "08_감축사업목록": [
                {"관리번호": "R1", "사업명": "태양광 보급"},
            ],
            "03_배출현황_지역": [
                {"배출유형": "직접배출", "부문": "건물", "세부부문": "", "연도": 2020, "배출량": 10, "단위": "톤", "출처페이지": "5"},
            ],
        },
        {
            "08_감축사업목록": [
                {"관리번호": "R-01", "사업명": "태양광\u3000보급", "골든_출처유형": "텍스트표", "골든_출처페이지": "4"},
            ],
            "03_배출현황_지역": [
                {"배출유형": "직접배출", "부문": "건물", "세부부문": "전기", "연도": 2020, "배출량": 10, "단위": "톤", "골든_출처유형": "텍스트표", "골든_출처페이지": "5"},
            ],
        },
    )

    assert result["시트별"]["08_감축사업목록"]["엄격매칭수"] == 1
    assert result["시트별"]["03_배출현황_지역"]["완화매칭수"] == 1


def test_시트04_직간접구분_표기차이는_엄격키로_매칭한다(tmp_path: Path) -> None:
    result = _점수(
        tmp_path,
        {
            "04_배출현황_관리권한": [
                {"관리부문": "건물", "세부부문": "가정", "직간접구분": "직접", "연도": 2020, "배출량": 10, "단위": "톤"},
            ]
        },
        {
            "04_배출현황_관리권한": [
                {"관리부문": "건물", "세부부문": "가정", "직간접구분": "direct", "연도": 2020, "배출량": 10, "단위": "톤", "골든_출처유형": "텍스트표", "골든_출처페이지": "7"},
            ]
        },
    )

    sheet = result["시트별"]["04_배출현황_관리권한"]
    assert sheet["엄격매칭수"] == 1
    assert sheet["완화매칭수"] == 0
    assert sheet["값일치율"] == 1.0


def test_S5_시트04_직간접구분_sink는_양쪽_키에서_흡수로_정규화한다(tmp_path: Path) -> None:
    result = _점수(
        tmp_path,
        {"04_배출현황_관리권한": [
            {"관리부문": "흡수원", "세부부문": "산림", "직간접구분": "흡수", "연도": 2020, "배출량": -10, "단위": "톤"}
        ]},
        {"04_배출현황_관리권한": [
            {"관리부문": "흡수원", "세부부문": "산림", "직간접구분": "sink", "연도": 2020, "배출량": -10, "단위": "톤", "골든_출처유형": "텍스트표"}
        ]},
    )

    assert result["시트별"]["04_배출현황_관리권한"]["엄격매칭수"] == 1


def test_시트04_완화키는_직간접구분이_다른_행을_섞지_않는다(tmp_path: Path) -> None:
    result = _점수(
        tmp_path,
        {
            "04_배출현황_관리권한": [
                {"관리부문": "건물", "세부부문": "상업", "직간접구분": "간접", "연도": 2020, "배출량": 20, "단위": "톤"},
            ]
        },
        {
            "04_배출현황_관리권한": [
                {"관리부문": "건물", "세부부문": "가정", "직간접구분": "direct", "연도": 2020, "배출량": 10, "단위": "톤", "골든_출처유형": "텍스트표", "골든_출처페이지": "8"},
            ]
        },
    )

    sheet = result["시트별"]["04_배출현황_관리권한"]
    assert sheet["매칭수"] == 0
    assert sheet["완화매칭수"] == 0
    assert sheet["리콜"] == 0.0


@pytest.mark.parametrize(
    ("sheet_name", "base"),
    [
        ("03_배출현황_지역", {"배출유형": "직접배출", "부문": "폐기물", "연도": 2020, "단위": "톤"}),
        ("04_배출현황_관리권한", {"관리부문": "폐기물", "직간접구분": "직접", "연도": 2020, "단위": "톤"}),
    ],
)
def test_V1_시트03_04_완화는_양쪽_세부부문이_달라도_값이_허용오차면_매칭한다(
    tmp_path: Path, sheet_name: str, base: dict
) -> None:
    result = _점수(
        tmp_path,
        {sheet_name: [base | {"세부부문": "매립", "배출량": 10.04}]},
        {sheet_name: [base | {"세부부문": "총계", "배출량": 10, "골든_출처유형": "텍스트표"}]},
    )

    assert result["시트별"][sheet_name]["완화매칭수"] == 1


@pytest.mark.parametrize(
    ("sheet_name", "base"),
    [
        ("03_배출현황_지역", {"배출유형": "직접배출", "부문": "폐기물", "연도": 2020, "단위": "톤"}),
        ("04_배출현황_관리권한", {"관리부문": "폐기물", "직간접구분": "직접", "연도": 2020, "단위": "톤"}),
    ],
)
def test_V1_시트03_04_완화는_양쪽_세부부문과_값이_다르면_제외한다(
    tmp_path: Path, sheet_name: str, base: dict
) -> None:
    result = _점수(
        tmp_path,
        {sheet_name: [base | {"세부부문": "매립", "배출량": 20}]},
        {sheet_name: [base | {"세부부문": "총계", "배출량": 10, "골든_출처유형": "텍스트표"}]},
    )

    assert result["시트별"][sheet_name]["매칭수"] == 0


@pytest.mark.parametrize(
    ("sheet_name", "base"),
    [
        ("03_배출현황_지역", {"배출유형": "직접배출", "부문": "폐기물", "연도": 2020, "단위": "톤"}),
        ("04_배출현황_관리권한", {"관리부문": "폐기물", "직간접구분": "직접", "연도": 2020, "단위": "톤"}),
    ],
)
def test_V1_시트03_04_완화는_양쪽_세부부문이_다르고_값을_검증할_수_없으면_제외한다(
    tmp_path: Path, sheet_name: str, base: dict
) -> None:
    result = _점수(
        tmp_path,
        {sheet_name: [base | {"세부부문": "매립", "배출량": ""}]},
        {sheet_name: [base | {"세부부문": "총계", "배출량": 10, "골든_출처유형": "텍스트표"}]},
    )

    assert result["시트별"][sheet_name]["매칭수"] == 0


@pytest.mark.parametrize(
    ("sheet_name", "base"),
    [
        ("03_배출현황_지역", {"배출유형": "직접배출", "부문": "폐기물", "연도": 2020, "배출량": 10, "단위": "톤"}),
        ("04_배출현황_관리권한", {"관리부문": "폐기물", "직간접구분": "직접", "연도": 2020, "배출량": 10, "단위": "톤"}),
    ],
)
def test_S3_시트03_04_완화는_한쪽_세부부문이_비면_기존_매칭을_유지한다(
    tmp_path: Path, sheet_name: str, base: dict
) -> None:
    result = _점수(
        tmp_path,
        {sheet_name: [base | {"세부부문": ""}]},
        {sheet_name: [base | {"세부부문": "총계", "골든_출처유형": "텍스트표"}]},
    )

    assert result["시트별"][sheet_name]["완화매칭수"] == 1


@pytest.mark.parametrize(
    ("sheet_name", "base"),
    [
        ("03_배출현황_지역", {"배출유형": "직접배출", "부문": "폐기물", "연도": 2020, "배출량": 10, "단위": "톤"}),
        ("04_배출현황_관리권한", {"관리부문": "폐기물", "직간접구분": "직접", "연도": 2020, "배출량": 10, "단위": "톤"}),
    ],
)
def test_S3_시트03_04는_세부부문이_같으면_기존_엄격매칭을_유지한다(
    tmp_path: Path, sheet_name: str, base: dict
) -> None:
    result = _점수(
        tmp_path,
        {sheet_name: [base | {"세부부문": "매립"}]},
        {sheet_name: [base | {"세부부문": "매립", "골든_출처유형": "텍스트표"}]},
    )

    assert result["시트별"][sheet_name]["엄격매칭수"] == 1


def test_V74_시트06은_목표수준이_같아도_기준연도와_값이_다르면_완화매칭을_차단한다(
    tmp_path: Path,
) -> None:
    base = {"목표수준": "총괄", "목표범위": "지역전체", "부문": "합계", "목표연도": 2030}
    result = _점수(
        tmp_path,
        {"06_감축목표": [base | {"기준연도": 2018, "목표배출량": 200}]},
        {"06_감축목표": [base | {"기준연도": 2005, "목표배출량": 100, "골든_출처유형": "텍스트표"}]},
    )

    sheet = result["시트별"]["06_감축목표"]
    assert sheet["엄격매칭수"] == 0
    assert sheet["완화매칭수"] == 0
    assert sheet["매칭수"] == 0


def test_V74_시트06은_기준연도가_달라도_공통값이_같으면_완화매칭한다(tmp_path: Path) -> None:
    base = {"목표수준": "총괄", "목표범위": "지역전체", "부문": "합계", "목표연도": 2030}
    result = _점수(
        tmp_path,
        {"06_감축목표": [base | {"기준연도": 2018, "목표배출량": 100}]},
        {"06_감축목표": [base | {"기준연도": 2005, "목표배출량": 100, "골든_출처유형": "텍스트표"}]},
    )

    sheet = result["시트별"]["06_감축목표"]
    assert sheet["엄격매칭수"] == 0
    assert sheet["완화매칭수"] == 1


@pytest.mark.parametrize(("output_base_year", "golden_base_year"), [("", 2018), (2018, ""), ("", "")])
def test_V74_시트06은_기준연도가_한쪽이나_양쪽_빈값이면_값가드를_발동하지_않는다(
    tmp_path: Path, output_base_year: int | str, golden_base_year: int | str
) -> None:
    base = {"목표수준": "총괄", "목표범위": "지역전체", "부문": "합계", "목표연도": 2030}
    result = _점수(
        tmp_path,
        {"06_감축목표": [base | {"기준연도": output_base_year, "목표배출량": 200}]},
        {"06_감축목표": [base | {"기준연도": golden_base_year, "목표배출량": 100, "골든_출처유형": "텍스트표"}]},
    )

    sheet = result["시트별"]["06_감축목표"]
    assert sheet["엄격매칭수"] == 0
    assert sheet["완화매칭수"] == 1


def test_V74_시트06은_기준연도가_양쪽_채움이고_같으면_엄격매칭한다(tmp_path: Path) -> None:
    key = {
        "목표수준": "총괄",
        "목표범위": "지역전체",
        "부문": "합계",
        "기준연도": 2018,
        "목표연도": 2030,
    }
    result = _점수(
        tmp_path,
        {"06_감축목표": [key | {"목표배출량": 200}]},
        {"06_감축목표": [key | {"목표배출량": 100, "골든_출처유형": "텍스트표"}]},
    )

    sheet = result["시트별"]["06_감축목표"]
    assert sheet["엄격매칭수"] == 1
    assert sheet["완화매칭수"] == 0


def test_V72_시트06_완화는_목표수준이_다르고_값도_다르면_제외한다(tmp_path: Path) -> None:
    base = {"목표범위": "관리권한", "부문": "폐기물", "목표연도": 2030}
    result = _점수(
        tmp_path,
        {"06_감축목표": [base | {"목표수준": "연차경로", "목표감축량": 4906, "감축률(%)": 20}]},
        {"06_감축목표": [base | {"목표수준": "부문", "목표감축량": 338, "감축률(%)": 34.6, "골든_출처유형": "텍스트표"}]},
    )

    sheet = result["시트별"]["06_감축목표"]
    assert sheet["매칭수"] == 0
    assert sheet["완화매칭수"] == 0


def test_V72_시트06_완화는_목표수준_표기가_달라도_공통값이_전부_같으면_매칭한다(
    tmp_path: Path,
) -> None:
    base = {"목표범위": "관리권한", "부문": "폐기물", "목표연도": 2030}
    result = _점수(
        tmp_path,
        {"06_감축목표": [base | {"목표수준": "부문별", "목표감축량": 338, "감축률(%)": 34.6}]},
        {"06_감축목표": [base | {"목표수준": "부문", "목표감축량": 338, "감축률(%)": 34.6, "골든_출처유형": "텍스트표"}]},
    )

    assert result["시트별"]["06_감축목표"]["완화매칭수"] == 1


def test_V72_시트06_완화는_목표수준이_같거나_한쪽이_비면_값과_무관하게_매칭한다(
    tmp_path: Path,
) -> None:
    cases = [
        ({"목표수준": "부문", "목표범위": "지역전체"}, {"목표수준": "부문", "목표범위": "관리권한"}),
        ({"목표수준": "", "목표범위": "관리권한"}, {"목표수준": "연차경로", "목표범위": "관리권한"}),
        ({"목표수준": "부문", "목표범위": "관리권한"}, {"목표수준": "", "목표범위": "관리권한"}),
    ]
    for output_key, golden_key in cases:
        common_key = {"부문": "폐기물", "기준연도": 2018, "목표연도": 2030}
        result = _점수(
            tmp_path,
            {"06_감축목표": [common_key | output_key | {"목표감축량": 4906}]},
            {"06_감축목표": [common_key | golden_key | {"목표감축량": 338, "골든_출처유형": "텍스트표"}]},
        )

        assert result["시트별"]["06_감축목표"]["완화매칭수"] == 1


def test_V72_시트06_완화는_목표수준이_다르고_공통값_필드가_없으면_제외한다(
    tmp_path: Path,
) -> None:
    base = {"목표범위": "관리권한", "부문": "폐기물", "목표연도": 2030}
    result = _점수(
        tmp_path,
        {"06_감축목표": [base | {"목표수준": "연차경로", "목표배출량": 338}]},
        {"06_감축목표": [base | {"목표수준": "부문", "목표감축량": 338, "골든_출처유형": "텍스트표"}]},
    )

    assert result["시트별"]["06_감축목표"]["매칭수"] == 0


def test_V72_시트06_완화는_목표수준이_다르면_공통값_필드가_전부_일치해야_한다(
    tmp_path: Path,
) -> None:
    base = {"목표범위": "관리권한", "부문": "폐기물", "목표연도": 2030}
    result = _점수(
        tmp_path,
        {"06_감축목표": [base | {"목표수준": "연차경로", "목표감축량": 338, "목표배출량": 200}]},
        {"06_감축목표": [base | {"목표수준": "부문", "목표감축량": 338, "목표배출량": 100, "골든_출처유형": "텍스트표"}]},
    )

    assert result["시트별"]["06_감축목표"]["매칭수"] == 0


def test_V72_시트06_가드는_시트03_04의_기존_완화_값가드를_바꾸지_않는다(tmp_path: Path) -> None:
    cases = [
        ("03_배출현황_지역", {"배출유형": "직접배출", "부문": "폐기물", "연도": 2020, "단위": "톤"}),
        ("04_배출현황_관리권한", {"관리부문": "폐기물", "직간접구분": "직접", "연도": 2020, "단위": "톤"}),
    ]
    for sheet_name, base in cases:
        result = _점수(
            tmp_path,
            {sheet_name: [base | {"세부부문": "매립", "배출량": 10.04}]},
            {sheet_name: [base | {"세부부문": "총계", "배출량": 10, "골든_출처유형": "텍스트표"}]},
        )

        assert result["시트별"][sheet_name]["완화매칭수"] == 1


def test_수치_허용오차와_불일치_상세를_구분한다(tmp_path: Path) -> None:
    result = _점수(
        tmp_path,
        {
            "03_배출현황_지역": [
                {"배출유형": "직접배출", "부문": "건물", "세부부문": "전기", "연도": 2020, "배출량": 100.4, "단위": "톤"},
                {"배출유형": "직접배출", "부문": "수송", "세부부문": "도로", "연도": 2020, "배출량": 101, "단위": "톤"},
            ]
        },
        {
            "03_배출현황_지역": [
                {"배출유형": "직접배출", "부문": "건물", "세부부문": "전기", "연도": 2020, "배출량": 100, "단위": "톤", "골든_출처유형": "텍스트표", "골든_출처페이지": "1"},
                {"배출유형": "직접배출", "부문": "수송", "세부부문": "도로", "연도": 2020, "배출량": 100, "단위": "톤", "골든_출처유형": "텍스트표", "골든_출처페이지": "2"},
            ]
        },
    )

    sheet = result["시트별"]["03_배출현황_지역"]
    assert sheet["값일치율"] == 0.75
    assert len(result["값불일치상세"]) == 1
    assert result["값불일치상세"][0]["필드"] == "배출량"


def test_S2_시트06의_천배_값은_헤드라인을_유지하고_스케일동치로_병기한다(tmp_path: Path) -> None:
    key = {"목표수준": "총괄", "목표범위": "지역전체", "부문": "합계", "목표연도": 2030}
    result = _점수(
        tmp_path,
        {"06_감축목표": [key | {"기준배출량": 52_342_000}]},
        {"06_감축목표": [key | {"기준배출량": 52_342, "골든_출처유형": "텍스트표"}]},
    )

    sheet = result["시트별"]["06_감축목표"]
    assert sheet["값일치율"] == 0.0
    assert sheet["값일치율(스케일동치 포함)"] == 1.0
    assert sheet["스케일동치쌍수"] == 1
    assert result["스케일동치쌍수"] == 1
    assert result["스케일동치쌍"][0]["배율"] == 1000
    assert len(result["값불일치상세"]) == 1
    report = Path(result["리포트"]["md"]).read_text(encoding="utf-8")
    assert "## 스케일동치 값 쌍" in report
    assert "골든=52342 / 출력=52342000 (배율=1000)" in report


def test_S2_시트06의_천배가_아닌_값은_기존_불일치로만_남는다(tmp_path: Path) -> None:
    key = {"목표수준": "총괄", "목표범위": "지역전체", "부문": "합계", "목표연도": 2030}
    result = _점수(
        tmp_path,
        {"06_감축목표": [key | {"목표배출량": 120_000}]},
        {"06_감축목표": [key | {"목표배출량": 100, "골든_출처유형": "텍스트표"}]},
    )

    sheet = result["시트별"]["06_감축목표"]
    assert sheet["값일치율"] == 0.0
    assert sheet["값일치율(스케일동치 포함)"] == 0.0
    assert sheet["스케일동치쌍수"] == 0
    assert result["스케일동치쌍"] == []

def test_출처유형_분해와_채점제외를_반영한다(tmp_path: Path) -> None:
    result = _점수(
        tmp_path,
        {
            "03_배출현황_지역": [
                {"배출유형": "직접배출", "부문": "건물", "세부부문": "전기", "연도": 2020, "배출량": 10, "단위": "톤"},
            ]
        },
        {
            "03_배출현황_지역": [
                {"배출유형": "직접배출", "부문": "건물", "세부부문": "전기", "연도": 2020, "배출량": 10, "단위": "톤", "골든_출처유형": "그래프", "골든_출처페이지": "1"},
                {"배출유형": "직접배출", "부문": "수송", "세부부문": "도로", "연도": 2020, "배출량": 20, "단위": "톤", "골든_출처유형": "그래프", "골든_출처페이지": "2"},
                {"배출유형": "직접배출", "부문": "산업", "세부부문": "공정", "연도": 2020, "배출량": 30, "단위": "톤", "골든_출처유형": "이미지", "골든_출처페이지": "3", "골든_채점제외": 1},
            ]
        },
    )

    assert result["채점제외행수"] == 1
    assert result["출처유형별_전체"]["그래프"]["리콜"] == 0.5
    assert result["출처유형별_전체"]["시각 유래"]["리콜"] == 0.5
    assert result["출처유형별_수치시트"]["시각 유래"]["리콜"] == 0.5
    assert result["시트별"]["03_배출현황_지역"]["출처유형별"]["시각 유래"] == {
        "골든행수": 2,
        "매칭수": 1,
        "리콜": 0.5,
        "값비교수": 2,
        "값일치율": 1.0,
    }


def test_없는_골든시트와_형식오류시트를_명확히_보고한다(tmp_path: Path) -> None:
    output = _워크북(tmp_path / "출력.xlsx", {"02_지역여건": [], "03_배출현황_지역": []})
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "03_배출현황_지역"
    ws.append(["03. 배출현황 제목행"])
    golden = tmp_path / "잘못된_골든.xlsx"
    wb.save(golden)

    result = score_workbooks(output, golden, report_dir=tmp_path / "리포트", label="오류", write_json=True)

    assert result["시트별"]["02_지역여건"]["상태"] == "골든 없음"
    assert result["시트별"]["03_배출현황_지역"]["상태"] == "형식 오류"
    assert "1행이 제목 행으로 추정" in result["형식오류"][0]

def test_문서메타는_필드별로_다중값_집합을_비교한다(tmp_path: Path) -> None:
    result = _점수(
        tmp_path,
        {
            "00_문서메타": [
                {"지자체명": "서울시", "기준연도": 2018, "목표연도": "2050, 2033"},
            ]
        },
        {
            "00_문서메타": [
                {"지자체명": "서울특별시", "기준연도": 2018, "목표연도": "2033,2050"},
            ]
        },
    )

    sheet = result["시트별"]["00_문서메타"]
    assert sheet["매칭수"] == 1
    assert sheet["값일치율"] == 1.0


def test_cli가_md와_json_리포트를_생성한다(tmp_path: Path) -> None:
    output = _워크북(
        tmp_path / "출력.xlsx",
        {"03_배출현황_지역": [{"배출유형": "직접배출", "부문": "건물", "세부부문": "전기", "연도": 2020, "배출량": 10, "단위": "톤"}]},
    )
    golden = _워크북(
        tmp_path / "골든.xlsx",
        {"03_배출현황_지역": [{"배출유형": "직접배출", "부문": "건물", "세부부문": "전기", "연도": 2020, "배출량": 10, "단위": "톤", "골든_출처유형": "본문텍스트", "골든_출처페이지": "7"}]},
        golden=True,
    )

    completed = subprocess.run(
        [sys.executable, "scripts/score_against_golden.py", str(output), str(golden), "--report-dir", str(tmp_path / "리포트"), "--label", "cli", "--json"],
        text=True,
        capture_output=True,
        check=False,
    )

    assert completed.returncode == 0, completed.stderr + completed.stdout
    md_path = tmp_path / "리포트" / "golden_score_cli.md"
    json_path = tmp_path / "리포트" / "golden_score_cli.json"
    assert md_path.exists()
    assert json_path.exists()
    body = md_path.read_text(encoding="utf-8")
    payload = json.loads(json_path.read_text(encoding="utf-8"))
    assert "출처유형별 분해" in body
    assert payload["시트별"]["03_배출현황_지역"]["리콜"] == 1.0
