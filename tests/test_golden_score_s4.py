from __future__ import annotations

from pathlib import Path

import openpyxl

from scripts.score_against_golden import score_workbooks
from tests.test_score_against_golden import _계약열, _워크북, _점수


def test_값일치율은_양쪽_값이_있는_쌍만_분모로_쓴다(tmp_path: Path) -> None:
    result = _점수(
        tmp_path,
        {
            "03_배출현황_지역": [
                {"배출유형": "직접배출", "부문": "건물", "세부부문": "전기", "연도": 2020, "배출량": "", "단위": "톤"},
            ]
        },
        {
            "03_배출현황_지역": [
                {"배출유형": "직접배출", "부문": "건물", "세부부문": "전기", "연도": 2020, "배출량": 100, "단위": "톤", "골든_출처유형": "텍스트표", "골든_출처페이지": "1"},
            ]
        },
    )

    sheet = result["시트별"]["03_배출현황_지역"]
    assert sheet["값일치율"] == 1.0
    assert sheet["골든만있는값"] == 1
    assert result["출처유형별_전체"]["텍스트표"]["값비교수"] == 1
    body = Path(result["리포트"]["md"]).read_text(encoding="utf-8")
    assert "값일치율 정의: 양쪽 값이 모두 존재하는 비교 쌍 중 일치한 값의 비율" in body


def test_골든시트_계약열_불일치_분기를_보고한다(tmp_path: Path) -> None:
    output = _워크북(tmp_path / "출력.xlsx", {"03_배출현황_지역": []})
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "03_배출현황_지역"
    ws.append(["엉뚱한열", "부문", "세부부문"])
    golden = tmp_path / "계약열_불일치.xlsx"
    wb.save(golden)

    result = score_workbooks(output, golden, report_dir=tmp_path / "리포트", label="계약열", write_json=True)

    assert result["시트별"]["03_배출현황_지역"]["상태"] == "형식 오류"
    assert "채점 필수 컬럼이 없습니다" in result["형식오류"][0]


def test_골든시트_골든열_누락_분기를_보고한다(tmp_path: Path) -> None:
    output = _워크북(tmp_path / "출력.xlsx", {"03_배출현황_지역": []})
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "03_배출현황_지역"
    ws.append(_계약열("03_배출현황_지역"))
    golden = tmp_path / "골든열_누락.xlsx"
    wb.save(golden)

    result = score_workbooks(output, golden, report_dir=tmp_path / "리포트", label="골든열", write_json=True)

    assert result["시트별"]["03_배출현황_지역"]["상태"] == "형식 오류"
    assert "골든_출처유형·골든_출처페이지 컬럼이 계약 컬럼 뒤에 필요합니다" in result["형식오류"][0]


def test_골든_출처페이지_공란은_경고로만_수집한다(tmp_path: Path) -> None:
    result = _점수(
        tmp_path,
        {
            "03_배출현황_지역": [
                {"배출유형": "직접배출", "부문": "건물", "세부부문": "전기", "연도": 2020, "배출량": 10, "단위": "톤"},
            ]
        },
        {
            "03_배출현황_지역": [
                {"배출유형": "직접배출", "부문": "건물", "세부부문": "전기", "연도": 2020, "배출량": 10, "단위": "톤", "골든_출처유형": "본문텍스트", "골든_출처페이지": ""},
            ]
        },
    )

    assert result["시트별"]["03_배출현황_지역"]["상태"] == "정상"
    assert result["형식오류"] == []
    assert "03_배출현황_지역 2행: 골든_출처페이지가 비어 있습니다" in result["형식경고"][0]
    body = Path(result["리포트"]["md"]).read_text(encoding="utf-8")
    assert "## 형식 경고" in body


def test_config_헤더가_확장되어도_골든은_같은_축으로_채점된다(tmp_path, monkeypatch):
    """v9 통합 회귀: 파이프라인 config.EXCEL_HEADERS에 필드가 추가되어도(v8) 골든
    헤더 검증은 채점 축의 고정 계약(필수채점필드)만 보므로 형식 오류가 나지 않는다."""
    import copy

    import config as config_module
    from scripts.golden_score_workbook import 워크북읽기

    extended = copy.deepcopy(config_module.EXCEL_HEADERS)
    extended["03_배출현황_지역"] = [*extended["03_배출현황_지역"], "가상의새필드", "또다른새필드"]
    monkeypatch.setattr(config_module, "EXCEL_HEADERS", extended)

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "03_배출현황_지역"
    ws.append([
        "지자체명", "인벤토리출처", "배출범위", "배출유형", "부문", "세부부문",
        "연도", "배출량", "단위", "흡수원여부",
        "골든_출처유형", "골든_출처페이지", "골든_채점제외", "골든_비고",
    ])
    ws.append([
        "서울특별시", "GIR", "직접배출", "직접배출", "에너지", "연료연소",
        2020, 100, "천톤CO2eq", False, "텍스트표", 10, None, None,
    ])
    golden = tmp_path / "확장config_골든.xlsx"
    wb.save(golden)

    sheets = 워크북읽기(golden, golden=True)
    sheet = sheets["03_배출현황_지역"]
    assert sheet.상태 == "정상"
    assert len(sheet.행들) == 1
