import subprocess
import sys
import tempfile
import textwrap
import unittest
import openpyxl
from pathlib import Path


class ABValidationHarnessTests(unittest.TestCase):
    def test_run_ab_validation_writes_markdown_report_with_dummy_command(self):
        # Given: main.py 대신 최소 xlsx를 쓰는 더미 추출 명령
        with tempfile.TemporaryDirectory() as tmp:
            tmpdir = Path(tmp)
            source = tmpdir / "sample.pdf"
            source.write_text("dummy", encoding="utf-8")
            dummy = tmpdir / "dummy_extract.py"
            dummy.write_text(
                textwrap.dedent(
                    """
                    import argparse
                    from pathlib import Path
                    import openpyxl

                    parser = argparse.ArgumentParser()
                    parser.add_argument('source')
                    parser.add_argument('--output', '-o', required=True)
                    args = parser.parse_args()
                    wb = openpyxl.Workbook()
                    ws = wb.active
                    ws.title = '00_문서메타'
                    ws.append(['지자체명', '계획명'])
                    ws.append(['서울특별시', '계획'])
                    Path(args.output).parent.mkdir(parents=True, exist_ok=True)
                    wb.save(args.output)
                    print('dummy ok')
                    """
                ),
                encoding="utf-8",
            )

            # When: A/B 하네스를 실행하면
            completed = subprocess.run(
                [
                    sys.executable,
                    "scripts/run_ab_validation.py",
                    str(source),
                    "--command",
                    f"{sys.executable} {dummy}",
                    "--output-dir",
                    str(tmpdir / "out"),
                ],
                text=True,
                capture_output=True,
                check=False,
            )

            # Then: markdown 리포트 경로를 출력하고 판정 체크리스트를 포함한다.
            self.assertEqual(completed.returncode, 0, completed.stderr + completed.stdout)
            report = Path(completed.stdout.strip())
            self.assertTrue(report.exists())
            body = report.read_text(encoding="utf-8")
            self.assertIn("판정 체크리스트", body)
            self.assertIn("시트별 행 수", body)

    def test_run_ab_validation_with_golden_adds_score_sections(self):
        with tempfile.TemporaryDirectory() as tmp:
            tmpdir = Path(tmp)
            source = tmpdir / "sample.pdf"
            source.write_text("dummy", encoding="utf-8")
            dummy = tmpdir / "dummy_extract.py"
            dummy.write_text(
                textwrap.dedent(
                    """
                    import argparse
                    from pathlib import Path
                    import openpyxl

                    parser = argparse.ArgumentParser()
                    parser.add_argument('source')
                    parser.add_argument('--output', '-o', required=True)
                    args = parser.parse_args()
                    wb = openpyxl.Workbook()
                    ws = wb.active
                    ws.title = '00_문서메타'
                    ws.append(['지자체명', '지자체유형', '계획명', '발간일', '발간기관', '계획시작연도', '계획종료연도', '기준연도', '목표연도', '법적근거', '점검보고서여부'])
                    ws.append(['서울특별시', '특별시', '계획', '', '', 2022, 2033, 2018, '2033,2050', '', 'N'])
                    Path(args.output).parent.mkdir(parents=True, exist_ok=True)
                    wb.save(args.output)
                    print('dummy ok')
                    """
                ),
                encoding="utf-8",
            )
            golden = tmpdir / "golden.xlsx"
            wb = openpyxl.Workbook()
            ws = wb.active
            ws.title = '00_문서메타'
            ws.append(['지자체명', '지자체유형', '계획명', '발간일', '발간기관', '계획시작연도', '계획종료연도', '기준연도', '목표연도', '법적근거', '점검보고서여부', '골든_비고'])
            ws.append(['서울특별시', '특별시', '계획', '', '', 2022, 2033, 2018, '2050,2033', '', 'N', ''])
            wb.save(golden)

            completed = subprocess.run(
                [
                    sys.executable,
                    "scripts/run_ab_validation.py",
                    str(source),
                    "--command",
                    f"{sys.executable} {dummy}",
                    "--output-dir",
                    str(tmpdir / "out"),
                    "--golden",
                    str(golden),
                ],
                text=True,
                capture_output=True,
                check=False,
            )

            self.assertEqual(completed.returncode, 0, completed.stderr + completed.stdout)
            report = Path(completed.stdout.strip())
            body = report.read_text(encoding="utf-8")
            self.assertIn("골든셋 채점 — 스니펫", body)
            self.assertIn("골든셋 채점 — 구조", body)
            self.assertIn("golden_score_스니펫", body)


if __name__ == "__main__":
    unittest.main()
