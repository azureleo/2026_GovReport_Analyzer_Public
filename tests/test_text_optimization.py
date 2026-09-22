from copy import deepcopy
from pathlib import Path
from unittest.mock import Mock
import json

import pytest
import config
from agents import extractor_agent as ea
from agents.extractor_agent import ExtractorAgent, BatchParseError
from scripts.preflight_text import compare_plans
from scripts.run_text_optimization_ab import arm_settings, _arm_environment
from scripts.run_text_optimization_ab import compare_contracts, compare_golden_results
from tests.test_vision_toc import index_page
from utils.pdf_reader import PageContent
from utils.text_optimization import TextPolicy, compact_page, policy_snapshot, plan_summary


@pytest.fixture(autouse=True)
def safe_config(monkeypatch):
    for name, value in dict(TEXT_ROUTING_MODE='audit', TEXT_INPUT_MODE='audit', FULL_DOCUMENT_SCAN=False,
                            EXTRACTION_SHEET_CLUSTERING=False, TEXT_CLUSTER_MEMBER_RETRIES=1,
                            EXTRACTION_SPLIT_ON_FAILURE=True, TEXT_WORKERS=1).items():
        monkeypatch.setattr(config, name, value)
    monkeypatch.setattr(ea.llm_client, 'call_text', Mock(side_effect=AssertionError('unexpected model call')))
    monkeypatch.setattr(ea.llm_client, 'call_text_json', Mock(side_effect=AssertionError('unexpected model call')))


def table_page():
    html = '<table><tr><td>건물</td><td>2024</td><td>120</td></tr></table>'
    return PageContent(32, '표 1 예산\n단위: 억원\n건물\n2024\n120\n주: 연간', [html], [],
                       text_blocks=[dict(block_id=7, text='건물\n2024\n120', bbox=[20, 40, 80, 80])],
                       table_records=[dict(table_index=1, html=html, bbox=[10, 30, 100, 90], strategy='lines_strict')])


def test_bbox_exact_duplicate_only_preserves_source_and_context():
    page = table_page()
    original = deepcopy(page)
    compact, removed = compact_page(page)
    assert page == original
    assert compact.tables == page.tables
    assert compact.text == '표 1 예산\n단위: 억원\n\n주: 연간'
    assert removed[0]['block_id'] == 7


@pytest.mark.parametrize('case', ['outside', 'substring_number', 'no_html', 'no_bbox', 'repeated', 'unit', 'text_strategy', 'header'])
def test_uncertain_duplicate_keeps_text(case):
    page = table_page()
    if case == 'outside':
        page.text_blocks[0]['bbox'] = [0, 0, 50, 50]
    if case == 'substring_number':
        page.text = page.text.replace('120', '12')
        page.text_blocks[0]['text'] = page.text_blocks[0]['text'].replace('120', '12')
    if case == 'no_html':
        page.tables = []
    if case == 'no_bbox':
        page.table_records[0]['bbox'] = None
    if case == 'repeated':
        page.text += '\n' + page.text_blocks[0]['text']
    if case == 'unit':
        page.text_blocks[0]['text'] = '단위: 억원\n건물\n2024\n120'
    if case == 'text_strategy':
        page.table_records[0]['strategy'] = 'text'
    if case == 'header':
        html = page.tables[0].replace('td>', 'th>')
        page.tables = [html]
        page.table_records[0]['html'] = html
    result, removed = compact_page(page)
    assert not removed and result.text == page.text


@pytest.mark.parametrize('mode', ['off', 'audit', 'optimize'])
def test_toc_route_and_payload_modes(monkeypatch, mode):
    monkeypatch.setattr(config, 'TEXT_ROUTING_MODE', mode)
    monkeypatch.setattr(config, 'TEXT_INPUT_MODE', mode)
    toc, body = index_page(), table_page()
    policy = TextPolicy([toc, body])
    result = policy.route({'financial_plan': [toc, body], 'document_meta': [toc, body]})
    assert [p.page_number for p in result['financial_plan']] == ([32] if mode == 'optimize' else [41, 32])
    assert result['financial_plan'][-1].text == (compact_page(body)[0].text if mode == 'optimize' else body.text)
    if mode != 'off':
        assert policy.audit()['pages'][0]['toc']['entries'][0]['physical_page_resolved'] is False


def test_mixed_toc_kept(monkeypatch):
    monkeypatch.setattr(config, 'TEXT_ROUTING_MODE', 'optimize')
    page = index_page()
    page.tables = ['<table><tr><td>실제 배출량 100</td></tr></table>']
    assert TextPolicy([page]).route({'emissions_regional': [page]})['emissions_regional']


def test_small_non_index_fact_is_not_discarded(monkeypatch):
    monkeypatch.setattr(config, 'TEXT_ROUTING_MODE', 'optimize')
    page = index_page()
    page.text += '\n2024년 감축량 15톤'
    assert TextPolicy([page]).route({'emissions_regional': [page]})['emissions_regional']


def test_invalid_policy_fails():
    with pytest.MonkeyPatch.context() as m:
        m.setattr(config, 'TEXT_INPUT_MODE', 'typo')
        with pytest.raises(ValueError):
            policy_snapshot()


def test_cluster_requests_only_routed_members(monkeypatch):
    a, b, c = [PageContent(n, f'내용 {n}', [], []) for n in (1, 2, 3)]
    monkeypatch.setattr(config, 'EXTRACTION_SHEET_CLUSTERS', [['document_meta', 'plan_overview']])
    agent = ExtractorAgent()
    routed = {'document_meta': [a, b], 'plan_overview': [b, c]}
    per = agent._plan_sheet_tasks(routed, {}, 10)
    tasks = agent._plan_cluster_tasks(routed, '', {}, 10)
    assert plan_summary(per)['sheet_page_pairs'] == plan_summary(tasks)['sheet_page_pairs']
    assert [(agent._task_sheet_keys(t), t['page_nums']) for t in tasks] == [
        (['document_meta'], [1]), (['document_meta', 'plan_overview'], [2]), (['plan_overview'], [3])]


def task():
    page = PageContent(2, '실제 데이터', [], [])
    return dict(members=['document_meta', 'plan_overview'], pages=[page], page_nums=[2],
                batch_text='=== 페이지 2 ===\n실제 데이터', batch_num=1, batch_total=1)


def test_partial_cluster_preserves_success_and_retries_only_missing(monkeypatch):
    responses = Mock(side_effect=[({'document_meta': [{'문서명': '계획'}]}, True),
                                 ({'plan_overview': [{'목적': '감축'}]}, True)])
    monkeypatch.setattr(ea.llm_client, 'call_text_json', responses)
    agent = ExtractorAgent()
    work = task()
    result = agent._run_cluster_task(work, '서울특별시', {})
    assert result['document_meta'][0]['문서명'] == '계획'
    assert result['plan_overview'][0]['목적'] == '감축'
    assert responses.call_count == 2
    assert agent.call_audit[1]['sheets'] == ['plan_overview']
    assert work['effective_status'] == 'ok'
    assert agent.task_audit[1]['parent_trace_id'] == agent.task_audit[0]['trace_id']


@pytest.mark.parametrize('bad', [None, {}, ['invalid'], {'x': 1}])
def test_failed_member_remains_partial_without_runstate(monkeypatch, bad):
    responses = Mock(side_effect=[({'document_meta': [], 'plan_overview': bad}, True), ({}, True)])
    monkeypatch.setattr(ea.llm_client, 'call_text_json', responses)
    agent, work = ExtractorAgent(), task()
    result = agent._run_cluster_task(work, '', {})
    assert result['document_meta'] == []
    assert work['effective_status'] == 'partial'
    assert responses.call_count == 2  # bounded, no page recursion for a failed member
    assert agent.task_audit[0]['status'] == 'partial'


def test_partial_list_keeps_valid_rows(monkeypatch):
    monkeypatch.setattr(config, 'TEXT_CLUSTER_MEMBER_RETRIES', 0)
    monkeypatch.setattr(ea.llm_client, 'call_text_json', Mock(return_value=({'document_meta': [{'문서명': '보존'}, 3], 'plan_overview': []}, True)))
    agent, work = ExtractorAgent(), task()
    result = agent._run_cluster_task(work, '', {})
    assert result['document_meta'][0]['문서명'] == '보존'
    assert work['effective_status'] == 'partial'


def test_nested_partial_propagates_without_checkpoint(monkeypatch):
    monkeypatch.setattr(config, 'TEXT_CLUSTER_MEMBER_RETRIES', 0)
    monkeypatch.setattr(ea.llm_client, 'call_text_json', Mock(return_value=({'document_meta': []}, True)))
    agent, work = ExtractorAgent(), task()
    child = agent._split_child_task(work, work['pages'], 1)
    agent._run_cluster_children(work, [child], '', {}, reason='test')
    assert work['effective_status'] == 'partial'


def test_partial_checkpoint_survives_interrupt_and_resumes_only_missing(monkeypatch, tmp_path):
    from tests.test_run_state import _state
    state = _state(tmp_path, monkeypatch)
    responses = Mock(side_effect=[({'document_meta': [], 'plan_overview': None}, True), KeyboardInterrupt()])
    monkeypatch.setattr(ea.llm_client, 'call_text_json', responses)
    work = task()
    with pytest.raises(KeyboardInterrupt):
        ExtractorAgent(run_state=state)._run_cluster_task(work, '', {})
    record = state.record_for(work['batch_id'])
    assert record['status'] == 'partial'
    assert record['member_statuses'] == {'document_meta': 'ok', 'plan_overview': 'partial'}
    # A failed retry must preserve even a valid empty member, not only nonempty rows.
    state.record_batch(batch_id=work['batch_id'], kind='cluster', sheet_keys=work['members'],
                       page_nums=[2], status='call_fail', error='interrupted retry')
    assert state.record_for(work['batch_id'])['member_statuses']['document_meta'] == 'ok'
    resumed = _state(tmp_path, monkeypatch, resume=True)
    call = Mock(return_value=({'plan_overview': [{'목적': '보존'}]}, True))
    monkeypatch.setattr(ea.llm_client, 'call_text_json', call)
    agent = ExtractorAgent(run_state=resumed)
    result = agent._run_cluster_task(task(), '', {})
    assert call.call_count == 1
    assert agent.call_audit[0]['sheets'] == ['plan_overview']
    assert result['document_meta'] == []
    assert result['plan_overview'][0]['목적'] == '보존'


def test_policy_changes_checkpoint_fingerprint(monkeypatch, tmp_path):
    from tests.test_run_state import _state
    a = _state(tmp_path, monkeypatch)
    monkeypatch.setattr(config, 'TEXT_INPUT_MODE', 'optimize')
    b = _state(tmp_path, monkeypatch)
    assert a.run_id != b.run_id


def test_preflight_exact_modes_and_no_model_calls(monkeypatch):
    monkeypatch.setattr(config, 'DOCUMENT_ROUTE_MAX_PAGES', {})
    monkeypatch.setattr(config, 'DOCUMENT_ROUTE_FRONT_BACK_PAGES', {})
    result = compare_plans([index_page(), table_page()])
    assert all(result['checks'].values())
    assert result['model_calls'] == 0
    assert result['arms']['optimized']['input_chars'] < result['arms']['baseline']['input_chars']


@pytest.mark.parametrize('full_scan', [False, True])
def test_extract_and_preflight_share_queue(monkeypatch, full_scan):
    monkeypatch.setattr(config, 'FULL_DOCUMENT_SCAN', full_scan)
    monkeypatch.setattr(config, 'TEXT_ROUTING_MODE', 'optimize')
    pages = [index_page(), table_page()]
    agent = ExtractorAgent()
    plan = agent.plan_tasks(pages)
    monkeypatch.setattr(agent, '_extract_municipality_name', lambda _: '서울특별시')
    monkeypatch.setattr(agent, '_run_sheet_task', lambda *_: [])
    monkeypatch.setattr(ea, 'parallel_map_collect', lambda fn, ts, **_: [(fn(t), None) for t in ts])
    agent.extract(pages, '', {})
    assert len(agent.call_plan) == len(plan)
    assert all(41 not in t['pages'] for t in agent.call_plan)


def test_ab_axes_separate_and_force_cold_isolation(monkeypatch, tmp_path):
    for key in ('EXTRACTION_RESUME', 'EXTRACTION_RETRY_FAILED_ONLY', 'SHEET_CLOSED_LOOP_ENABLED', 'GAP_FILL_ENABLED'):
        monkeypatch.setenv(key, '1')
    a = _arm_environment('baseline', tmp_path, clustering=False, use_cache=False)
    b = _arm_environment('clustered', tmp_path, clustering=True, use_cache=False)
    assert {k for k in a if a[k] != b[k]} == {'EXTRACTION_SHEET_CLUSTERING', 'RUN_STATE_DIR'}
    for key in ('EXTRACTION_RESUME', 'EXTRACTION_RETRY_FAILED_ONLY', 'SHEET_CLOSED_LOOP_ENABLED', 'GAP_FILL_ENABLED', 'LLM_CACHE_ENABLED'):
        assert a[key] == b[key] == '0'
    assert arm_settings('inputs', True)['EXTRACTION_SHEET_CLUSTERING'] == '0'


def test_live_comparison_rejects_model_drift():
    a = {key: 'hash' for key in ('input_sha256', 'implementation_sha256', 'prompt_sha256', 'guideline_sha256')}
    a['config'] = dict(sheet_clustering=False, text_model='same')
    b = deepcopy(a)
    b['config']['sheet_clustering'] = True
    assert compare_contracts(a, b, 'clustering')['comparable']
    b['config']['text_model'] = 'different'
    assert not compare_contracts(a, b, 'clustering')['comparable']


def test_gold_comparison_retains_individual_losses():
    a = {'독립평가지표': dict(cell_expected=100, golden_rows=10), '형식오류': [],
         '미매칭상세': {'골든': [{'시트': 'x', '행번호': 3}]}}
    b = deepcopy(a)
    b['미매칭상세']['골든'] = [{'시트': 'x', '행번호': 7}]
    result = compare_golden_results(a, b)
    assert result['comparable']
    assert result['previously_matched_rows_lost'] == [{'시트': 'x', '행번호': 7}]
    b['독립평가지표']['cell_expected'] = 99
    assert not compare_golden_results(a, b)['comparable']


def test_supervisor_saves_text_audit_and_workbook(monkeypatch, tmp_path):
    from agents import supervisor as sm
    from utils.pdf_reader import PDFContent
    import openpyxl
    page = PageContent(1, '서울특별시 탄소중립 기본계획\n계획 수립 목적은 온실가스 감축이다.', [], [])
    monkeypatch.setattr(config, 'GAP_FILL_ENABLED', False)
    monkeypatch.setattr(config, 'HYBRID_REVIEW_ENABLED', False)
    monkeypatch.setattr(config, 'SHEET_CLOSED_LOOP_ENABLED', False)
    monkeypatch.setattr(config, 'EXTRACTION_SHEET_CLUSTERING', True)
    monkeypatch.setattr(config, 'RUN_STATE_DIR', str(tmp_path / 'states'))
    monkeypatch.setattr(sm, 'extract_pdf', lambda *a, **k: PDFContent(1, [page], page.text))
    monkeypatch.setattr(sm.GuidelineAgent, 'get_all_prompts', lambda *a: {})
    monkeypatch.setattr(ea, '_route_pages_by_sheet', lambda _: {'document_meta': [page], 'plan_overview': [page]})
    monkeypatch.setattr(ea.llm_client, 'call_text', lambda *a, **k: '{"municipality_name":"서울특별시"}')
    monkeypatch.setattr(ea.llm_client, 'call_text_json', Mock(return_value=({'document_meta': [], 'plan_overview': []}, True)))
    target = sm.Supervisor().run(input_path=tmp_path / 'source.pdf', output_path=tmp_path / 'result.xlsx',
                                max_pipeline_retries=1, include_images=False)
    audit = json.loads(target.with_name(target.stem + '_text_audit.json').read_text(encoding='utf-8'))
    assert audit['attempts'][0]['calls'][0]['sheets'] == ['document_meta', 'plan_overview']
    assert audit['attempts'][0]['tasks'][0]['status'] == 'ok'
    assert set(audit['final_rows_by_sheet']) == set(config.EXTRACTION_SHEETS)
    wb = openpyxl.load_workbook(target, read_only=True)
    try:
        assert '00_문서메타' in wb.sheetnames
    finally:
        wb.close()
