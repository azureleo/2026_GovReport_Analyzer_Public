from __future__ import annotations

import json
import subprocess
import sys
import textwrap
from io import BytesIO
from pathlib import Path

import fitz
from openpyxl import Workbook
from PIL import Image, ImageDraw

from scripts import benchmark_vision as bench
from scripts.vision_benchmark_core import Candidate, CommandStats, has_openai_key
from scripts.vision_benchmark_report import CandidateReportInput, ReportConfig, write_reports


def _chart_png() -> bytes:
    image = Image.new("RGB", (420, 260), "white")
    draw = ImageDraw.Draw(image)
    for idx, height in enumerate((80, 140, 190)):
        x0 = 60 + idx * 90
        draw.rectangle((x0, 230 - height, x0 + 50, 230), fill="steelblue", outline="black")
    draw.text((50, 20), "서울 배출량 그래프 2020 2030", fill="black")
    buf = BytesIO()
    image.save(buf, format="PNG")
    return buf.getvalue()


def _visual_pdf(path: Path) -> None:
    doc = fitz.open()
    page = doc.new_page(width=595, height=842)
    page.insert_text((72, 72), "[그림 1-1] 서울 배출량 그래프")
    page.insert_image(fitz.Rect(72, 120, 492, 400), stream=_chart_png())
    doc.save(path)
    doc.close()


def _inventory(path: Path, rows: list[list[str | int]] | None = None) -> None:
    wb = Workbook()
    ws = wb.active
    ws.title = "시각요소"
    ws.append(["요소ID", "페이지", "요소유형", "제목", "데이터포함", "기대추출", "관련시트", "비고", "근거ID"])
    for row in rows or [["E001", 1, "그래프", "[그림 1-1]", "Y", "Y", "03_배출현황_지역", "", "ev-e001"]]:
        ws.append(row)
    wb.save(path)


def _golden(path: Path) -> None:
    wb = Workbook()
    ws = wb.active
    ws.title = "03_배출현황_지역"
    ws.append([
        "지자체명", "인벤토리출처", "배출범위", "배출유형", "부문", "세부부문",
        "연도", "배출량", "단위", "흡수원여부", "골든_출처유형", "골든_출처페이지", "골든_비고",
    ])
    ws.append(["서울특별시", "", "", "직접배출", "건물", "전기", 2020, 10, "톤", "N", "그래프", "1", ""])
    wb.save(path)


def _workbook(path: Path, rows: list[list[str | int]]) -> None:
    wb = Workbook()
    ws = wb.active
    ws.title = "16_시각자료목록"
    ws.append(["지자체명", "시각자료ID", "캡션", "유형", "데이터포함여부", "추출값요약", "디지타이징필요", "관련시트"])
    for row in rows:
        ws.append(row)
    wb.save(path)


def _audit_json(path: Path, rows: list[tuple[str, int, str, str]]) -> None:
    path.write_text(
        json.dumps({"details": [
            {"요소ID": element_id, "페이지": page, "요소유형": kind, "상태": status}
            for element_id, page, kind, status in rows
        ]}, ensure_ascii=False),
        encoding="utf-8",
    )


def _dummy_extractor(path: Path) -> None:
    path.write_text(textwrap.dedent(
        """
        from __future__ import annotations
        import argparse, os
        from pathlib import Path
        from openpyxl import Workbook

        parser = argparse.ArgumentParser()
        parser.add_argument("source")
        parser.add_argument("--output", "-o", required=True)
        args = parser.parse_args()
        out = Path(args.output)
        out.parent.mkdir(parents=True, exist_ok=True)
        wb = Workbook()
        ws = wb.active
        ws.title = "03_배출현황_지역"
        ws.append(["지자체명", "인벤토리출처", "배출범위", "배출유형", "부문", "세부부문", "연도", "배출량", "단위", "흡수원여부", "출처페이지", "데이터상태"])
        ws.append(["서울특별시", "", "", "직접배출", "건물", "전기", 2020, 10, "톤", "N", "1", "visual_only"])
        visual = wb.create_sheet("16_시각자료목록")
        visual.append(["지자체명", "시각자료ID", "캡션", "유형", "데이터포함여부", "추출값요약", "디지타이징필요", "관련시트", "근거ID"])
        visual.append(["서울특별시", "V1-001", "[그림 1-1]", "그래프", "Y", "10톤", "N", "03_배출현황_지역", "ev-e001"])
        meta = wb.create_sheet("_env")
        meta.append(["key", "value"])
        for key in ["LLM_PROVIDER", "GEMINI_MODEL", "LLM_CACHE_DIR", "STAGE_PROVIDER_VISION", "STAGE_MODEL_VISION", "STAGE_PROVIDER_EXTRACTION"]:
            meta.append([key, os.environ.get(key, "")])
        wb.save(out)
        print("  - LLM 캐시: hit 12, miss 0, write 0, disabled 0")
        print("  - LLM 호출: 0회(실패 0, 재시도 0, 타임아웃 0), quota 대기 누적 0.0초")
        """), encoding="utf-8")


def test_variant_env_keeps_text_backend_gemini_and_only_changes_vision_stage() -> None:
    candidate = bench.resolve_candidate("flashlite")
    base_env = {
        "LLM_PROVIDER": "codex",
        "GEMINI_MODEL": "gemini-2.5-pro",
        "STAGE_PROVIDER_EXTRACTION": "claude",
        "STAGE_MODEL_EXTRACTION": "danger",
    }

    env = bench.build_variant_env(base_env, candidate, Path(".cache/llm_responses_gemini_v5"))

    assert env["LLM_PROVIDER"] == "gemini"
    assert env["GEMINI_MODEL"] == "gemini-2.5-flash-lite"
    assert env["LLM_CACHE_DIR"] == ".cache/llm_responses_gemini_v5"
    assert env["STAGE_PROVIDER_VISION"] == "gemini"
    assert env["STAGE_MODEL_VISION"] == "gemini-2.5-flash-lite"
    assert "STAGE_PROVIDER_EXTRACTION" not in env
    assert "STAGE_MODEL_EXTRACTION" not in env


def test_openai_candidate_skips_when_key_is_absent_and_runs_when_key_exists(tmp_path: Path) -> None:
    source = tmp_path / "서울.pdf"
    inventory = tmp_path / "inventory.xlsx"
    source.write_text("pdf-placeholder", encoding="utf-8")
    _inventory(inventory)
    config = bench.BenchmarkConfig(
        source=source,
        inventory=inventory,
        golden=None,
        output_dir=tmp_path / "out",
        cache_dir=tmp_path / "cache",
        command=f"{sys.executable} -c 'raise SystemExit(99)'",
        candidates=(Candidate("openai", "openai", "gpt-5.4-mini"),),
        skip_run=(),
        sample_size=15,
        baseline_recall=0.707,
        cache_only=False,
        force_run=False,
        env_file=tmp_path / ".env",
    )

    skipped = bench._run_candidate(config, config.candidates[0])

    assert skipped.skip_reason == "OPENAI_API_KEY 부재"
    (tmp_path / ".env").write_text("OPENAI_API_KEY=sk-test\n", encoding="utf-8")
    assert has_openai_key({}, tmp_path / ".env") is True


def test_cache_only_guard_patches_llm_client_direct_cache_binding(tmp_path: Path) -> None:
    wrapper = bench._cache_guard_script(tmp_path)

    body = wrapper.read_text(encoding="utf-8")
    assert "llm_cache.cached_response = cached_response" in body
    assert "llm_client.cached_response = cached_response" in body


def test_report_renders_elapsed_seconds_and_skip_reason(tmp_path: Path) -> None:
    inventory = tmp_path / "inventory.xlsx"
    output = tmp_path / "out.xlsx"
    audit_json = tmp_path / "audit.json"
    log = tmp_path / "pipeline.log"
    _inventory(inventory)
    _workbook(output, [["서울", "V1-001", "[그림 1-1]", "그래프", "Y", "10톤", "N", ""]])
    _audit_json(audit_json, [("E001", 1, "그래프", "기록됨")])
    log.write_text("  - LLM 캐시: hit 1, miss 0, write 0, disabled 0\n", encoding="utf-8")

    report = write_reports(
        ReportConfig(tmp_path / "서울.pdf", inventory, None, tmp_path, tmp_path / "cache", 15, 0.707),
        (
            CandidateReportInput(
                candidate="flashlite",
                vision_provider="gemini",
                vision_model="gemini-2.5-flash-lite",
                output_xlsx=output,
                pipeline_log=log,
                run=CommandStats(0, log, 1.25, 1, 0, 0, 0),
                audit_json=audit_json,
                score_json=None,
                workbook_metrics={"visual_rows": 1, "data_included_ratio": 1.0, "auto_trusted_ratio": 1.0},
                environment={},
                skip_reason="",
                reused=False,
            ),
            CandidateReportInput(
                candidate="openai",
                vision_provider="openai",
                vision_model="gpt-5.4-mini",
                output_xlsx=tmp_path / "missing.xlsx",
                pipeline_log=tmp_path / "openai.log",
                run=CommandStats(0, tmp_path / "openai.log", 0.0, None, None, None, None),
                audit_json=None,
                score_json=None,
                workbook_metrics={},
                environment={},
                skip_reason="OPENAI_API_KEY 부재",
                reused=False,
            ),
        ),
    )

    body = Path(str(report["markdown"])).read_text(encoding="utf-8")
    payload = json.loads(Path(str(report["json"])).read_text(encoding="utf-8"))
    assert "실행 시간(초)" in body
    assert "1.25" in body
    assert "OPENAI_API_KEY 부재" in body
    assert payload["candidates"][0]["run_stats"]["elapsed_seconds"] == 1.25


def test_benchmark_cli_writes_actual_sample_rows_from_visual_sheet(tmp_path: Path) -> None:
    source = tmp_path / "서울.pdf"
    inventory = tmp_path / "inventory.xlsx"
    golden = tmp_path / "golden.xlsx"
    dummy = tmp_path / "dummy_extract.py"
    _visual_pdf(source)
    _inventory(inventory)
    _golden(golden)
    _dummy_extractor(dummy)

    completed = subprocess.run(
        [
            sys.executable,
            "scripts/benchmark_vision.py",
            str(source),
            "--inventory",
            str(inventory),
            "--golden",
            str(golden),
            "--candidates",
            "flashlite",
            "--command",
            f"{sys.executable} {dummy}",
            "--output-dir",
            str(tmp_path / "out"),
            "--sample-size",
            "5",
        ],
        text=True,
        capture_output=True,
        check=False,
    )

    assert completed.returncode == 0, completed.stderr + completed.stdout
    paths = json.loads(completed.stdout.strip().splitlines()[-1])
    md_path = Path(paths["markdown"])
    json_path = Path(paths["json"])
    sample_path = Path(paths["value_sample_markdown"])
    body = md_path.read_text(encoding="utf-8")
    sample_body = sample_path.read_text(encoding="utf-8")
    payload = json.loads(json_path.read_text(encoding="utf-8"))
    assert "후보별 비교표" in body
    assert "실행 시간(초)" in body
    assert "채점·집계 LLM 호출 0" in body
    assert "E001" in sample_body
    assert "10톤" in sample_body
    row = payload["candidates"][0]
    assert row["candidate"] == "flashlite"
    assert row["inventory_audit"]["recall"] == 1.0
    assert row["run_stats"]["llm_total_calls"] == 0
    assert row["run_stats"]["elapsed_seconds"] >= 0
    assert row["workbook_metrics"]["visual_rows"] == 1
