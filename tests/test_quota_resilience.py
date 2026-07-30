import subprocess
import tempfile
import unittest
from pathlib import Path
from unittest.mock import patch

from openpyxl import load_workbook

import config
from agents.extractor_agent import ExtractorAgent
from agents.supervisor import Supervisor
from utils import llm_client
from utils.parallel import parallel_map_collect
from utils.pdf_reader import PDFContent, PageContent


class QuotaResilienceTests(unittest.TestCase):
    def setUp(self):
        self._saved_config = {
            name: getattr(config, name)
            for name in (
                "GAP_FILL_ENABLED",
                "LLM_QUOTA_WAIT_ENABLED",
                "LLM_QUOTA_WAIT_POLL_SECONDS",
                "LLM_QUOTA_WAIT_MAX_SECONDS",
                "PARALLEL_PROCESSING_ENABLED",
                "TEXT_WORKERS",
                "LOCAL_AGENT_TIMEOUT_RETRIES",
                "LOCAL_AGENT_TIMEOUT_RETRY_DELAY_SECONDS",
            )
        }
        llm_client.reset_llm_stats()

    def tearDown(self):
        for name, value in self._saved_config.items():
            setattr(config, name, value)
        llm_client.reset_llm_stats()

    def test_retry_local_call_resumes_after_one_quota_wait(self):
        # Given: 첫 호출만 quota를 내고 다음 호출은 성공하는 로컬 호출
        config.LLM_QUOTA_WAIT_ENABLED = True
        config.LLM_QUOTA_WAIT_POLL_SECONDS = 1
        config.LLM_QUOTA_WAIT_MAX_SECONDS = 3
        calls = {"count": 0}

        def flaky_call():
            calls["count"] += 1
            if calls["count"] == 1:
                raise llm_client.LLMQuotaExceededError("quota")
            return '{"ok": true}'

        # When: quota 대기 중 실제 sleep은 막고 재시도하면
        with patch.object(llm_client.time, "sleep", return_value=None):
            result = llm_client._retry_local_call(flaky_call, max_retries=1, label="quota-test")

        # Then: 일반 재시도 예산을 소모하지 않고 성공하며 대기 통계를 남긴다.
        self.assertEqual(result, '{"ok": true}')
        self.assertEqual(calls["count"], 2)
        self.assertGreater(llm_client.get_llm_stats()["quota_wait_seconds"], 0)

    def test_retry_local_call_raises_quota_type_when_wait_cap_is_exceeded(self):
        # Given: quota가 계속 발생하고 누적 대기 상한이 작은 로컬 호출
        config.LLM_QUOTA_WAIT_ENABLED = True
        config.LLM_QUOTA_WAIT_POLL_SECONDS = 2
        config.LLM_QUOTA_WAIT_MAX_SECONDS = 3

        def quota_call():
            raise llm_client.LLMQuotaExceededError("quota")

        # When / Then: 상한 초과 시 LLMQuotaExceededError 타입 그대로 전파한다.
        with patch.object(llm_client.time, "sleep", return_value=None):
            with self.assertRaises(llm_client.LLMQuotaExceededError):
                llm_client._retry_local_call(quota_call, max_retries=1, label="quota-test")

    def test_timeout_is_bounded_and_not_treated_as_quota(self):
        # Given: 로컬 호출이 계속 타임아웃되고 추가 재시도는 1회로 제한되어 있으면
        config.LOCAL_AGENT_TIMEOUT_RETRIES = 1
        config.LOCAL_AGENT_TIMEOUT_RETRY_DELAY_SECONDS = 0
        calls = {"count": 0}

        def timeout_call():
            calls["count"] += 1
            raise subprocess.TimeoutExpired(["codex"], 300)

        # When/Then: quota 대기 루프에 들어가지 않고 총 2회 뒤 전용 오류로 종료한다.
        with self.assertRaises(llm_client.LLMTimeoutError):
            llm_client._retry_local_call(timeout_call, max_retries=3, label="timeout-test")

        self.assertEqual(calls["count"], 2)
        stats = llm_client.get_llm_stats()
        self.assertEqual(stats["timeouts"], 2)
        self.assertEqual(stats["quota_wait_seconds"], 0)

    def test_parallel_map_collect_isolates_quota_and_preserves_finished_results(self):
        # Given: 두 번째 항목부터 quota가 발생하는 병렬 배치
        config.PARALLEL_PROCESSING_ENABLED = True

        def worker(value):
            if value == 2:
                raise llm_client.LLMQuotaExceededError("quota")
            return value * 10

        # When: 수집형 병렬 실행을 수행하면
        results = parallel_map_collect(worker, [1, 2, 3, 4], workers=2)

        # Then: raise 없이 입력 길이를 보존하고 quota 항목들을 실패로 격리한다.
        self.assertEqual(len(results), 4)
        self.assertEqual(results[0], (10, None))
        self.assertIsNone(results[1][0])
        self.assertIsInstance(results[1][1], llm_client.LLMQuotaExceededError)
        for result, error in results[2:]:
            if result is None:
                self.assertIsInstance(error, llm_client.LLMQuotaExceededError)
            else:
                self.assertIsNone(error)

    def test_extract_municipality_name_uses_regex_fallback_on_quota(self):
        # Given: 지자체명 LLM 호출이 quota를 내는 본문
        agent = ExtractorAgent()
        text = "서울특별시 탄소중립 녹색성장 기본계획"

        # When: 지자체명을 추출하면
        with patch.object(
            llm_client,
            "call_text",
            side_effect=llm_client.LLMQuotaExceededError("quota"),
        ):
            name = agent._extract_municipality_name(text)

        # Then: quota를 전파하지 않고 정규식 fallback 결과를 반환한다.
        self.assertEqual(name, "서울특별시")

    def test_supervisor_writes_validation_report_when_extraction_batch_hits_quota(self):
        # Given: 두 번째 추출 배치만 quota로 실패하는 최소 PDF 파이프라인
        config.GAP_FILL_ENABLED = False
        config.PARALLEL_PROCESSING_ENABLED = False
        config.TEXT_WORKERS = 1
        pages = [
            PageContent(
                page_number=page,
                text="온실가스 배출량 부문별 현황",
                tables=[],
                images=[],
            )
            for page in range(1, 17)
        ]
        pdf = PDFContent(total_pages=len(pages), pages=pages, full_text="서울특별시 탄소중립 기본계획")

        def fake_call_text(prompt, system="", max_retries=1):
            if "완성도" in prompt:
                return '{"assessment": "규칙 기반 확인", "quality_level": "보통", "key_issues": []}'
            return '{"municipality_name": "서울특별시"}'

        first_batch = {
            "emissions_regional": [
                {
                    "지자체명": "서울특별시",
                    "부문": "수송",
                    "연도": 2030,
                    "배출량": 1,
                    "단위": "천톤CO2eq",
                    "출처페이지": "1",
                }
            ]
        }

        with tempfile.TemporaryDirectory() as tmp:
            output_path = Path(tmp) / "quota.xlsx"
            with patch("agents.supervisor.extract_pdf", return_value=pdf), \
                 patch("agents.supervisor.GuidelineAgent.get_all_prompts", return_value={"emissions_regional": ""}), \
                 patch("agents.supervisor.GuidelineAgent.report", return_value="[가이드라인] 테스트"), \
                 patch(
                     "agents.extractor_agent._route_pages_by_sheet",
                     return_value={"emissions_regional": pages},
                 ), \
                 patch.object(llm_client, "call_text", side_effect=fake_call_text), \
                 patch.object(
                     llm_client,
                     "call_text_json",
                     side_effect=[
                         (first_batch, True),
                         llm_client.LLMQuotaExceededError("quota"),
                     ],
                 ):
                # When: Supervisor 파이프라인을 실행하면
                result_path = Supervisor().run(
                    input_path=Path(tmp) / "input.pdf",
                    output_path=output_path,
                    max_pipeline_retries=1,
                    include_images=False,
                )

            # Then: 엑셀이 생성되고 검증리포트에 원장 호출실패가 남는다.
            self.assertTrue(result_path.exists())
            workbook = load_workbook(result_path)
            self.assertIn("19_검증리포트", workbook.sheetnames)
            rows = list(workbook["19_검증리포트"].iter_rows(values_only=True))
            self.assertTrue(any("원장 호출실패" in row for row in rows[1:]))


if __name__ == "__main__":
    unittest.main()
