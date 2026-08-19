from __future__ import annotations

import tempfile
from pathlib import Path

from openpyxl import load_workbook

import config
from agents.organizer_agent import OrganizerAgent
from utils import excel_writer
from utils.reference_data import (
    load_appendix3_units,
    load_appendix4_projects,
    match_reference_name,
)
from scripts.build_reference_data import _parse_appendix4


def test_reference_csv_counts_and_boundaries() -> None:
    # Given/When: 커밋된 부록 CSV를 읽으면
    appendix3 = load_appendix3_units()
    appendix4 = load_appendix4_projects()

    # Then: 명세의 고정 행 수와 경계가 유지된다.
    assert len(appendix3) == 114
    assert appendix3[0]["번호"] == "1-1"
    assert appendix3[-1]["번호"] == "8-13"
    assert len(appendix4) == 507
    by_number = {int(row["연번"]): row for row in appendix4}
    assert by_number[1]["부문"] == "건물"
    assert by_number[143]["부문"] == "건물"
    assert by_number[144]["부문"] == "농축산"
    assert by_number[507]["부문"] == "이행기반"


def test_reference_name_match_confidence_levels() -> None:
    # Given: 작은 후보 사전
    candidates = [
        {"name": "태양광 발전"},
        {"name": "전기차 충전 인프라 확충"},
    ]

    # When/Then: 완전/포함/유사/무매칭 규칙이 순서대로 동작한다.
    assert match_reference_name("태양광 발전", candidates, name_key="name", threshold=0.55)["confidence"] == "high"
    assert match_reference_name("공공 태양광 발전 보급", candidates, name_key="name", threshold=0.55)["confidence"] == "medium"
    low = match_reference_name("전기차 충전시설 확대", candidates, name_key="name", threshold=0.2)
    assert low is not None and low["confidence"] == "low"
    assert match_reference_name("해양 생태 축제", candidates, name_key="name", threshold=0.9) is None


def test_reference_name_match_unifies_middle_dot_variants() -> None:
    # Given: 부록 사업명과 추출 사업명의 가운뎃점 표기가 다르면
    candidates = [{"name": "신재생에너지·보급"}]

    # When: 참조 사전 매칭을 수행하면
    match = match_reference_name("신재생에너지ㆍ보급", candidates, name_key="name", threshold=0.55)

    # Then: dedup 키와 같은 정규화 규칙으로 완전일치 high가 된다.
    assert match is not None
    assert match["confidence"] == "high"


def test_appendix4_matching_fills_blank_sector_and_reports_mismatch() -> None:
    # Given: 부록4 사업명과 일치하는 빈 부문 행, 부문이 다른 행
    raw = {
        "municipality_name": "서울특별시",
        "mitigation_projects": [
            {"부문": "", "사업명": "도시재생사업 연계 주택에너지 효율개선"},
            {"부문": "수송", "사업명": "물 수요관리(재이용) 확대"},
        ],
    }

    # When: organizer가 부록4 매칭을 수행하면
    cleaned = OrganizerAgent().organize(raw)

    # Then: 빈 부문은 채우고, 상충 부문은 덮지 않고 리포트한다.
    first, second = cleaned["mitigation_projects"]
    assert first["부문"] == "건물"
    assert first["표준사업연번"] == "1"
    assert first["표준사업명"] == "도시재생사업 연계 주택에너지 효율개선"
    assert first["매칭신뢰도"] == "high"
    assert second["부문"] == "수송"
    assert second["표준사업연번"] == "2"
    assert any("부문 불일치" in issue["항목"] for issue in cleaned["validation_report"])


def test_appendix3_unit_is_reference_only_and_value_difference_is_reported() -> None:
    # Given: 원단위가 없는 행과, 원문에 원단위가 명시된 행
    raw = {
        "municipality_name": "서울특별시",
        "quantitative_reductions": [
            {"사업명": "태양광 발전", "모니터링인자": "발전량", "활동량": 1000, "예상감축량": 1},
            {"사업명": "태양광 발전", "모니터링인자": "발전량", "감축원단위ID": "1-1", "감축원단위값": 1.0, "활동량": 1000, "예상감축량": 1},
        ],
    }

    # When: organizer가 부록3 매칭을 수행하면
    cleaned = OrganizerAgent().organize(raw)

    # Then: 부록3 값은 자동 채우지 않고, 명시된 값의 차이만 검증한다.
    first = cleaned["quantitative_reductions"][0]
    assert not first.get("감축원단위ID")
    assert first.get("감축원단위값") is None
    assert any("감축원단위 외부 참조 후보" in issue["항목"] for issue in cleaned["validation_report"])
    assert any("감축원단위값 차이" in issue["항목"] for issue in cleaned["validation_report"])


def test_codebook_sheet_is_optional_and_opt_out_removes_sheet() -> None:
    # Given: organizer가 코드북 행을 포함한 엑셀 데이터를 만든다.
    previous = getattr(config, "CODEBOOK_SHEET_ENABLED", True)
    config.CODEBOOK_SHEET_ENABLED = True
    try:
        cleaned = OrganizerAgent().organize({"municipality_name": "서울특별시"})
        assert any(row["코드유형"] == "전망방법" for row in cleaned["codebook"])
        with tempfile.TemporaryDirectory() as tmp:
            path = Path(tmp) / "with_codebook.xlsx"
            excel_writer.write_excel(cleaned, path)
            wb = load_workbook(path)
            assert "90_코드북" in wb.sheetnames

        # When: opt-out 플래그를 끄면
        config.CODEBOOK_SHEET_ENABLED = False
        with tempfile.TemporaryDirectory() as tmp:
            path = Path(tmp) / "without_codebook.xlsx"
            excel_writer.write_excel(cleaned, path)
            wb = load_workbook(path)
            assert "90_코드북" not in wb.sheetnames
    finally:
        config.CODEBOOK_SHEET_ENABLED = previous


def test_codebook_sheet_includes_execution_info_stamp() -> None:
    # Given: 실행 메타데이터가 포함된 organizer 입력
    previous = getattr(config, "CODEBOOK_SHEET_ENABLED", True)
    config.CODEBOOK_SHEET_ENABLED = True
    try:
        cleaned = OrganizerAgent().organize({
            "municipality_name": "서울특별시",
            "execution_info": {
                "git_commit": "abc1234",
                "text_backend": "gemini",
                "text_model": "gemini-2.5-flash-lite",
                "vision_backend": "gemini",
                "vision_model": "gemini-2.5-pro",
                "run_started_at": "2026-07-08T00:00:00+09:00",
                "input_file": "서울특별시_탄소중립계획.pdf",
                "pipeline_version": "v5.3",
            },
        })

        # When: 엑셀로 저장한 뒤 코드북 시트를 다시 읽으면
        with tempfile.TemporaryDirectory() as tmp:
            path = Path(tmp) / "with_execution_info.xlsx"
            excel_writer.write_excel(cleaned, path)
            wb = load_workbook(path, data_only=True)
            ws = wb["90_코드북"]
            headers = [cell.value for cell in ws[1]]
            rows = [dict(zip(headers, values)) for values in ws.iter_rows(min_row=2, values_only=True)]

        # Then: 실행정보 행이 기존 코드북 계약 컬럼 안에 행 추가로 기록된다.
        execution_rows = {row["코드"]: row for row in rows if row["코드유형"] == "실행정보"}
        assert execution_rows["git_commit"]["정의"] == "abc1234"
        assert execution_rows["text_model"]["정의"] == "gemini-2.5-flash-lite"
        assert execution_rows["vision_model"]["정의"] == "gemini-2.5-pro"
        assert execution_rows["input_file"]["정의"] == "서울특별시_탄소중립계획.pdf"
        assert execution_rows["pipeline_version"]["정의"] == "v5.3"
    finally:
        config.CODEBOOK_SHEET_ENABLED = previous


def test_appendix4_parser_reports_skipped_rows(capsys) -> None:
    # Given: 부록4 표 안에 연번이 깨진 행이 있으면
    text = "\n".join([
        "| 연번 | 부문 | 사업명 | 비고 |",
        "|---|---|---|---|",
        "| 1 | 건물 | 도시재생사업 연계 주택에너지 효율개선 | |",
        "| X | 건물 | 깨진 행 | |",
    ])

    # When: 부록4를 파싱하면
    rows = _parse_appendix4(text)

    # Then: 정상 행은 유지하고 실패 행은 stderr에 보고한다.
    assert rows == [{"연번": "1", "부문": "건물", "사업명": "도시재생사업 연계 주택에너지 효율개선"}]
    assert "부록4 행 건너뜀" in capsys.readouterr().err
