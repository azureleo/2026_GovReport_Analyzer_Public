"""Exact text root-queue comparison, no model calls and no workbook changes."""
from __future__ import annotations

import argparse
from contextlib import ExitStack
from dataclasses import asdict
import json
from pathlib import Path
import sys
from unittest.mock import patch

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT))

import config
from agents.extractor_agent import ExtractorAgent
from utils import llm_client
from utils.pdf_reader import extract_pdf
from utils.text_optimization import digest, plan_summary


def compare_plans(pages):
    results = {}
    before = digest([asdict(p) for p in pages])
    for name, routing, inputs, clustering in (
        ("baseline", "off", "off", False), ("audit", "audit", "audit", False),
        ("optimized", "optimize", "optimize", False), ("clustered", "optimize", "optimize", True),
    ):
        with ExitStack() as stack:
            for key, value in {"TEXT_ROUTING_MODE": routing, "TEXT_INPUT_MODE": inputs,
                               "EXTRACTION_SHEET_CLUSTERING": clustering}.items():
                stack.enter_context(patch.object(config, key, value))
            agent = ExtractorAgent()
            tasks = agent.plan_tasks(pages)
            results[name] = {**plan_summary(tasks), "audit": agent.text_policy.audit(),
                             "oversized_root_tasks": sum(len(t['batch_text']) > config.EXTRACTION_MAX_BATCH_CHARS for t in tasks),
                             "tasks": [{"sheets": agent._task_sheet_keys(t), "pages": t['page_nums'],
                                        "input_chars": len(t['batch_text']), "input_sha256": digest(t['batch_text'])} for t in tasks]}
    base = set(map(tuple, results['baseline']['sheet_page_pairs']))
    opt = set(map(tuple, results['optimized']['sheet_page_pairs']))
    cluster = set(map(tuple, results['clustered']['sheet_page_pairs']))
    toc = {row['page'] for row in results['optimized']['audit']['pages'] if row['text_index_only']}
    checks = {
        "source_unchanged": before == digest([asdict(p) for p in pages]),
        "audit_queue_identical": results['baseline']['task_sha256'] == results['audit']['task_sha256'],
        "body_assignments_preserved": not any(page not in toc for _, page in base - opt),
        "clustering_assignment_parity": opt == cluster,
        "no_added_assignments": not opt - base,
    }
    return {"arms": results, "checks": checks, "removed_pairs": sorted(base - opt), "model_calls": 0,
            "source_payload_sha256": before, "accuracy_measured": False,
            "limits": "Root tasks exclude recursive/transport retries and municipality call. Input characters are not tokens."}


def golden_coverage(pages, plans, path):
    import openpyxl
    from scripts.verify_routing_coverage import _harvest_anchors, _norm, _NAME_TO_KEY
    text = {p.page_number: _norm(p.text + '\n' + '\n'.join(p.tables)) for p in pages}
    index_pages = {r['page'] for r in plans['arms']['optimized']['audit']['pages'] if r['text_index_only']}
    pairs = {name: set(map(tuple, plans['arms'][name]['sheet_page_pairs'])) for name in ('baseline', 'optimized')}
    result = {}
    wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    try:
        for name, key in _NAME_TO_KEY.items():
            if name not in wb.sheetnames:
                continue
            count = 0
            baseline = candidate = 0
            lost = []
            index_only = []
            for anchor in _harvest_anchors(wb[name]):
                hit = {p for p, body in text.items() if _norm(anchor) in body}
                if not hit:
                    continue
                body_hit = hit - index_pages
                if not body_hit:
                    index_only.append(anchor)
                    continue
                count += 1
                b = any((key, p) in pairs['baseline'] for p in body_hit)
                c = any((key, p) in pairs['optimized'] for p in body_hit)
                baseline += b
                candidate += c
                if b and not c:
                    lost.append({"anchor": anchor, "pages": sorted(body_hit)})
            result[key] = {"located_body_anchors": count, "baseline_covered": baseline,
                           "candidate_covered": candidate, "lost": lost, "index_only_anchors": index_only}
    finally:
        wb.close()
    measured = any(r['located_body_anchors'] for r in result.values())
    return {"sheets": result, "passed": not any(r['lost'] for r in result.values()) if measured else None,
            "evaluated": measured,
            "note": "Source-string routing retention only. Unlocated/index-only anchors excluded and disclosed; not cell accuracy."}


def main():
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument('input', type=Path)
    parser.add_argument('--output', type=Path, required=True)
    parser.add_argument('--golden', type=Path)
    args = parser.parse_args()
    if not args.input.is_file() or (args.golden and not args.golden.is_file()):
        parser.error('입력 파일을 찾을 수 없습니다.')
    if args.output.exists():
        parser.error('기존 결과 보호: 새로운 --output 경로를 지정하세요.')
    def blocked(*args, **kwargs):
        raise AssertionError('Text preflight must not invoke models')
    with ExitStack() as stack:
        for name in ('call_text', 'call_text_json', 'call_vision', 'call_vision_json'):
            if hasattr(llm_client, name):
                stack.enter_context(patch.object(llm_client, name, blocked))
        source = extract_pdf(args.input, render_graph_pages=False)
        result = compare_plans(source.pages)
        from utils.run_state import sha256_file, _implementation_hash
        result.update(input_sha256=sha256_file(args.input), implementation_sha256=_implementation_hash())
        if args.golden:
            result['golden_source_coverage'] = golden_coverage(source.pages, result, args.golden)
    args.output.parent.mkdir(parents=True, exist_ok=True)
    with args.output.open('x', encoding='utf-8') as stream:
        json.dump(result, stream, ensure_ascii=False, indent=2)
    print(json.dumps({"checks": result['checks'], "arms": {k: {f: v[f] for f in ('root_tasks', 'input_chars', 'oversized_root_tasks')} for k, v in result['arms'].items()}, "golden_passed": result.get('golden_source_coverage', {}).get('passed'), "output": str(args.output)}, ensure_ascii=False, indent=2))
    return 0 if all(result['checks'].values()) and result.get('golden_source_coverage', {}).get('passed', True) else 2


if __name__ == '__main__':
    raise SystemExit(main())
