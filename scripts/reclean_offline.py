# /// script
# requires-python = ">=3.10"
# dependencies = ["openpyxl"]
# ///
# ─── How to run ───
# .venv/bin/python scripts/reclean_offline.py output/강원_결과_v6.xlsx output/경기_결과_v6.xlsx
# .venv/bin/python scripts/reclean_offline.py --json output/강원_결과_v6.xlsx
"""기존 출력 xlsx를 LLM 없이 organizer 재정제에 다시 통과시키는 게이트 도구."""

from __future__ import annotations

import argparse
import json
import sys
from dataclasses import dataclass
from pathlib import Path
from typing import Sequence

from openpyxl import load_workbook

_PROJECT_ROOT = Path(__file__).resolve().parents[1]
if str(_PROJECT_ROOT) not in sys.path:
    sys.path.insert(0, str(_PROJECT_ROOT))

import config
from agents.organizer_agent import OrganizerAgent

CellValue = str | int | float | bool | None

_TARGET_SHEET_KEYS = ("emissions_regional", "emissions_management")
_BLANK_KEY_FIELDS = {
    "emissions_regional": ("배출유형", "부문", "세부부문", "연도"),
    "emissions_management": ("관리부문", "세부부문", "직간접구분", "연도"),
}


@dataclass(frozen=True, slots=True)
class SheetMetrics:
    rows_before: int
    rows_after: int
    blank_key_rows_before: int
    blank_key_rows_after: int


@dataclass(frozen=True, slots=True)
class WorkbookReport:
    path: Path
    municipality: str
    conflicts_before: int
    conflicts_after: int
    sheet_metrics: dict[str, SheetMetrics]


def analyze_workbook(path: Path) -> WorkbookReport:
    raw_data = _load_excel_data(path)
    municipality = _municipality_name(raw_data, path)
    raw_data["municipality_name"] = municipality
    cleaned = OrganizerAgent().organize(raw_data)

    sheet_metrics: dict[str, SheetMetrics] = {}
    for sheet_key in _TARGET_SHEET_KEYS:
        sheet_name = config.SHEET_KEY_TO_NAME[sheet_key]
        before_rows = raw_data.get(sheet_key, [])
        after_rows = cleaned.get(sheet_key, [])
        sheet_metrics[sheet_name] = SheetMetrics(
            rows_before=len(before_rows),
            rows_after=len(after_rows),
            blank_key_rows_before=_blank_key_rows(sheet_key, before_rows),
            blank_key_rows_after=_blank_key_rows(sheet_key, after_rows),
        )

    return WorkbookReport(
        path=path,
        municipality=municipality,
        conflicts_before=_dedup_conflict_count(raw_data.get("validation_report", [])),
        conflicts_after=_dedup_conflict_count(cleaned.get("validation_report", [])),
        sheet_metrics=sheet_metrics,
    )


def render_markdown(reports: Sequence[WorkbookReport]) -> str:
    lines = [
        "# 오프라인 재정제 게이트 리포트",
        "",
        "| 파일 | 지자체 | 지표 | 수정 전 | 수정 후 | 감소 |",
        "|---|---|---|---:|---:|---:|",
    ]
    for report in reports:
        filename = report.path.name
        lines.append(_metric_row(filename, report.municipality, "중복제거 충돌", report.conflicts_before, report.conflicts_after))
        for sheet_name, metrics in report.sheet_metrics.items():
            lines.append(
                _metric_row(
                    filename,
                    report.municipality,
                    f"{sheet_name} 빈 키 행",
                    metrics.blank_key_rows_before,
                    metrics.blank_key_rows_after,
                )
            )
            lines.append(_metric_row(filename, report.municipality, f"{sheet_name} 행 수", metrics.rows_before, metrics.rows_after))
    return "\n".join(lines) + "\n"


def main(argv: Sequence[str] | None = None) -> int:
    parser = argparse.ArgumentParser(description="기존 출력 xlsx를 organizer로 오프라인 재정제해 전후 지표를 비교합니다.")
    parser.add_argument("workbooks", nargs="+", type=Path)
    parser.add_argument("--json", action="store_true", help="마크다운 대신 JSON으로 출력")
    args = parser.parse_args(argv)

    reports = [analyze_workbook(path) for path in args.workbooks]
    if args.json:
        print(json.dumps([_report_dict(report) for report in reports], ensure_ascii=False, indent=2))
    else:
        print(render_markdown(reports), end="")
    return 0


def _metric_row(filename: str, municipality: str, metric: str, before: int, after: int) -> str:
    return f"| {filename} | {municipality} | {metric} | {before} | {after} | {before - after} |"


def _report_dict(report: WorkbookReport) -> dict[str, str | int | dict[str, dict[str, int]]]:
    return {
        "path": str(report.path),
        "municipality": report.municipality,
        "conflicts_before": report.conflicts_before,
        "conflicts_after": report.conflicts_after,
        "sheet_metrics": {
            sheet_name: {
                "rows_before": metrics.rows_before,
                "rows_after": metrics.rows_after,
                "blank_key_rows_before": metrics.blank_key_rows_before,
                "blank_key_rows_after": metrics.blank_key_rows_after,
            }
            for sheet_name, metrics in report.sheet_metrics.items()
        },
    }


def _load_excel_data(path: Path) -> dict[str, list[dict[str, CellValue]]]:
    workbook = load_workbook(path, read_only=True, data_only=True)
    try:
        data: dict[str, list[dict[str, CellValue]]] = {}
        for sheet_key, sheet_name in config.SHEET_KEY_TO_NAME.items():
            if sheet_name not in workbook.sheetnames:
                continue
            data[sheet_key] = _sheet_rows(workbook[sheet_name])
        return data
    finally:
        workbook.close()


def _sheet_rows(worksheet) -> list[dict[str, CellValue]]:
    rows = worksheet.iter_rows(values_only=True)
    headers = [_header(value) for value in next(rows, ())]
    records: list[dict[str, CellValue]] = []
    for values in rows:
        record = {
            header: _cell_value(value)
            for header, value in zip(headers, values)
            if header
        }
        if any(value not in (None, "") for value in record.values()):
            records.append(record)
    return records


def _header(value) -> str:
    return "" if value is None else str(value).strip()


def _cell_value(value) -> CellValue:
    if value is None or isinstance(value, (str, int, float, bool)):
        return value
    return str(value)


def _municipality_name(data: dict[str, list[dict[str, CellValue]]], path: Path) -> str:
    for rows in data.values():
        for row in rows:
            value = row.get("지자체명")
            if value not in (None, ""):
                return str(value)
    return path.stem


def _blank_key_rows(sheet_key: str, rows: list[dict[str, CellValue]]) -> int:
    fields = _BLANK_KEY_FIELDS[sheet_key]
    return sum(
        1 for row in rows
        if any(str(row.get(field, "") or "").strip() == "" for field in fields)
    )


def _dedup_conflict_count(rows: list[dict[str, CellValue]]) -> int:
    return sum(
        1 for row in rows
        if "중복제거" in str(row.get("영역", "") or "") or "중복 키 값 충돌" in str(row.get("항목", "") or "")
    )


if __name__ == "__main__":
    raise SystemExit(main())
