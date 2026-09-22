# /// script
# requires-python = ">=3.11"
# dependencies = []
# ///
# ─── How to run ───
# .venv/bin/python scripts/audit_visual_inventory.py dump 서울특별시_탄소중립계획.pdf --out data/golden/시각요소_초안.xlsx
# .venv/bin/python scripts/audit_visual_inventory.py audit data/golden/서울_시각요소_인벤토리_v1.xlsx output/서울.xlsx 서울특별시_탄소중립계획.pdf --report-dir output
"""시각 요소 인벤토리 리콜 감사 도구."""
# noqa: SIZE_OK — M2-1 명세가 신규 production 코드를 이 단일 CLI 스크립트로 제한한다.
from __future__ import annotations

import argparse
import json
import re
import sys
import time
from collections import Counter, defaultdict
from collections.abc import Iterable, Sequence
from dataclasses import dataclass
from pathlib import Path

ROOT = Path(__file__).resolve().parents[1]
if str(ROOT) not in sys.path:
    sys.path.insert(0, str(ROOT))

import fitz
from openpyxl import Workbook, load_workbook

import config
from agents.image_agent import (  # 기존 동작을 바꾸지 않고 감사 입력으로만 재사용한다.
    _apply_reference_context_policy,
    _coverage_reduce_images,
    _is_relevant_image,
    _triage_image,
)
from utils.pdf_reader import PageContent, extract_pdf
from utils.object_evidence_match import descriptor, match_object_evidence

INVENTORY_HEADERS = ["요소ID", "페이지", "요소유형", "제목", "데이터포함", "기대추출", "관련시트", "비고"]
AUTO_HEADERS = [
    "자동_트리아지점수", "자동_이미지크기", "자동_캡션후보", "자동_벡터드로잉수",
    "자동_풀렌더여부", "자동_관련성통과", "자동_트리아지통과및사유",
]
VALID_ELEMENT_TYPES = {"그래프", "이미지표", "이미지", "지도·사진", "장식"}
STATUS_ORDER = [
    "이미지_미추출", "triage_탈락", "참고자료_제외", "vision_유실",
    "근거불충분", "복수객체_검토필요", "기록됨",
]
DRAFT_NOTICE = (
    "주의: 이 초안은 PDF 파서와 이미지 triage가 본 후보만 나열합니다. "
    "파이프라인이 통째로 놓친 요소는 초안에 없으므로 사람이 원문 PDF를 넘기며 누락 요소를 행으로 추가해야 최종 인벤토리가 됩니다."
)
_VISUAL_ID_RE = re.compile(r"^V(?P<page>\d+)-\d+")
_CAPTION_RE = re.compile(r"^\s*\[?\s*(그림|표)\s*[\d\-]+[^\n]*", re.MULTILINE)
CellValue = str | int | float | bool | None
WorkbookRow = dict[str, CellValue]
JsonScalar = str | int | float | bool | None
JsonValue = JsonScalar | list["JsonValue"] | dict[str, "JsonValue"]


class InventoryFormatError(Exception):
    """사람 인벤토리 xlsx 형식 오류."""


@dataclass(frozen=True, slots=True)
class InventoryItem:
    element_id: str
    page: int
    element_type: str
    title: str
    data_included: bool
    expected: bool
    related_sheet: str
    note: str
    evidence_id: str = ""
    object_id: str = ""
    key_values: str = ""


@dataclass(frozen=True, slots=True)
class VisualRecord:
    visual_id: str
    page: int | None
    caption: str
    visual_type: str
    data_included: str
    digitizing_needed: str
    related_sheet: str
    summary: str = ""
    evidence_id: str = ""
    object_id: str = ""


@dataclass(frozen=True, slots=True)
class PageEvidence:
    page: int
    image_count: int
    reduced_image_count: int
    triage_passed_count: int
    reference_filtered_count: int
    top_score: int | None
    top_reasons: tuple[str, ...]
    drawing_count: int
    triage_scores: tuple[int, ...] = ()
    reference_flagged_count: int = 0


@dataclass(frozen=True, slots=True)
class ElementAudit:
    item: InventoryItem
    status: str
    triage_score: int | None
    triage_reasons: tuple[str, ...]
    drawing_count: int
    matched_records: tuple[VisualRecord, ...]
    match_stage: str = ""
    match_reason: str = ""


@dataclass(frozen=True, slots=True)
class ThresholdSensitivity:
    threshold: int
    passed_elements: int
    total_expected: int
    recall: float
    passed_images: int


@dataclass(frozen=True, slots=True)
class SensitivityReport:
    image_thresholds: dict[int, ThresholdSensitivity]
    vector_thresholds: dict[int, int]


@dataclass(frozen=True, slots=True)
class AuditPaths:
    inventory: Path
    pipeline_output: Path
    source_pdf: Path
    report_dir: Path


@dataclass(frozen=True, slots=True)
class ReportPaths:
    markdown: Path
    json_path: Path


def parse_visual_page_id(value: str) -> int | None:
    match = _VISUAL_ID_RE.match(value.strip())
    return int(match.group("page")) if match else None


def _cell_text(value: CellValue) -> str:
    if value is None:
        return ""
    return str(value).strip()


def _yn(value: CellValue, field: str, row_num: int) -> bool:
    text = _cell_text(value).upper()
    if text == "Y":
        return True
    if text == "N":
        return False
    raise InventoryFormatError(f"{row_num}행 {field} 값은 Y/N이어야 합니다: {text or '빈값'}")


def _page_int(value: CellValue, row_num: int) -> int:
    try:
        page = int(value) if not isinstance(value, str) else int(value.strip())
    except (TypeError, ValueError) as exc:
        raise InventoryFormatError(f"{row_num}행 페이지 값이 정수가 아닙니다: {_cell_text(value)}") from exc
    if page < 1:
        raise InventoryFormatError(f"{row_num}행 페이지 값은 1 이상이어야 합니다: {page}")
    return page


def _sheet_rows(path: Path, sheet_name: str) -> list[WorkbookRow]:
    wb = load_workbook(path, data_only=True)
    if sheet_name not in wb.sheetnames:
        return []
    ws = wb[sheet_name]
    headers = [_cell_text(cell.value) for cell in ws[1]]
    rows: list[WorkbookRow] = []
    for values in ws.iter_rows(min_row=2, values_only=True):
        row = {header: value for header, value in zip(headers, values) if header}
        if any(_cell_text(value) for value in row.values()):
            rows.append(row)
    return rows


def load_inventory(path: Path) -> list[InventoryItem]:
    wb = load_workbook(path, data_only=True)
    if "시각요소" not in wb.sheetnames:
        raise InventoryFormatError("인벤토리에 '시각요소' 시트가 없습니다.")
    ws = wb["시각요소"]
    headers = [_cell_text(cell.value) for cell in ws[1]]
    missing = [header for header in INVENTORY_HEADERS if header not in headers]
    if missing:
        raise InventoryFormatError("필수 컬럼 누락: " + ", ".join(missing))
    items: list[InventoryItem] = []
    for row_num, values in enumerate(ws.iter_rows(min_row=2, values_only=True), start=2):
        row = {header: value for header, value in zip(headers, values)}
        if not any(_cell_text(value) for value in row.values()):
            continue
        element_type = _cell_text(row.get("요소유형"))
        if element_type not in VALID_ELEMENT_TYPES:
            raise InventoryFormatError(f"{row_num}행 요소유형 오류: {element_type or '빈값'}")
        element_id = _cell_text(row.get("요소ID"))
        if not element_id:
            raise InventoryFormatError(f"{row_num}행 요소ID가 비어 있습니다.")
        items.append(InventoryItem(
            element_id=element_id,
            page=_page_int(row.get("페이지"), row_num),
            element_type=element_type,
            title=_cell_text(row.get("제목")),
            data_included=_yn(row.get("데이터포함"), "데이터포함", row_num),
            expected=_yn(row.get("기대추출"), "기대추출", row_num),
            related_sheet=_cell_text(row.get("관련시트")),
            note=_cell_text(row.get("비고")),
            evidence_id=_cell_text(row.get("근거ID")),
            object_id=_cell_text(row.get("객체ID")),
            key_values=_cell_text(row.get("핵심값")),
        ))
    return items


def load_visual_records(path: Path) -> list[VisualRecord]:
    rows = _sheet_rows(path, "16_시각자료목록")
    records: list[VisualRecord] = []
    for row in rows:
        visual_id = _cell_text(row.get("시각자료ID"))
        records.append(VisualRecord(
            visual_id=visual_id,
            page=parse_visual_page_id(visual_id),
            caption=_cell_text(row.get("캡션")),
            visual_type=_cell_text(row.get("유형")),
            data_included=_cell_text(row.get("데이터포함여부")),
            digitizing_needed=_cell_text(row.get("디지타이징필요")),
            related_sheet=_cell_text(row.get("관련시트")),
            summary=" | ".join(filter(None, (
                _cell_text(row.get("추출값요약")),
                _cell_text(row.get("원문값")),
                _cell_text(row.get("원문단위")),
            ))),
            evidence_id=_cell_text(row.get("근거ID")),
            object_id=_cell_text(row.get("근거객체ID")),
        ))
    return records


def _drawing_counts(pdf_path: Path) -> dict[int, int]:
    counts: dict[int, int] = {}
    doc = fitz.open(str(pdf_path))
    try:
        for idx, page in enumerate(doc, start=1):
            try:
                counts[idx] = len(page.get_drawings())
            except RuntimeError:
                counts[idx] = 0
    finally:
        doc.close()
    return counts


def _caption_candidates(text: str) -> str:
    return " | ".join(match.group(0).strip() for match in _CAPTION_RE.finditer(text or ""))


def _limited_reasons_with_reference(reasons: Sequence[str], limit: int) -> tuple[str, ...]:
    selected = list(reasons[:limit])
    for reason in reasons:
        if reason.startswith("reference_context:") and reason not in selected:
            selected.append(reason)
    return tuple(selected)


def _triage_page_evidence(pdf_path: Path, municipality: str) -> dict[int, PageEvidence]:
    pdf = extract_pdf(pdf_path)
    drawings = _drawing_counts(pdf_path)
    raw_images = [(page, image) for page in pdf.pages for image in page.images if _is_relevant_image(image)]
    reduced_images = _coverage_reduce_images(raw_images)
    raw_counts = Counter(page.page_number for page, _ in raw_images)
    reduced_counts = Counter(page.page_number for page, _ in reduced_images)
    passed_counts: Counter[int] = Counter()
    reference_flagged_counts: Counter[int] = Counter()
    reference_counts: Counter[int] = Counter()
    scores: dict[int, list[int]] = defaultdict(list)
    reasons: dict[int, list[str]] = defaultdict(list)
    for page, image in reduced_images:
        item = _triage_image(page, image)
        score = int(item["score"])
        is_reference, was_filtered = _apply_reference_context_policy(
            page, image, item, municipality
        )
        reference_flagged_counts[page.page_number] += int(is_reference)
        reference_counts[page.page_number] += int(was_filtered)
        passed = bool(item["passed"])
        reason_list = [str(reason) for reason in item["reasons"]]
        if passed:
            passed_counts[page.page_number] += 1
        scores[page.page_number].append(score)
        reasons[page.page_number].extend(reason_list)
    evidence: dict[int, PageEvidence] = {}
    for page in pdf.pages:
        page_scores = tuple(scores.get(page.page_number, []))
        top_score = max(page_scores) if page_scores else None
        evidence[page.page_number] = PageEvidence(
            page=page.page_number,
            image_count=raw_counts[page.page_number],
            reduced_image_count=reduced_counts[page.page_number],
            triage_passed_count=passed_counts[page.page_number],
            reference_filtered_count=reference_counts[page.page_number],
            top_score=top_score,
            top_reasons=_limited_reasons_with_reference(tuple(reasons.get(page.page_number, [])), 6),
            drawing_count=drawings.get(page.page_number, 0),
            triage_scores=page_scores,
            reference_flagged_count=reference_flagged_counts[page.page_number],
        )
    return evidence


def _records_by_page(records: Sequence[VisualRecord]) -> dict[int, list[VisualRecord]]:
    by_page: dict[int, list[VisualRecord]] = defaultdict(list)
    for record in records:
        if record.page is not None:
            by_page[record.page].append(record)
    return by_page


def classify_inventory(
    items: Sequence[InventoryItem],
    pages: dict[int, PageEvidence],
    records: Sequence[VisualRecord],
) -> list[ElementAudit]:
    records_by_page = _records_by_page(records)
    record_by_key = {record.visual_id: record for record in records}
    record_descriptors = [
        descriptor(
            key=record.visual_id,
            page=record.page,
            evidence_ids=record.evidence_id,
            object_ids=record.object_id,
            number="",
            caption=record.caption,
            content=" | ".join(filter(None, (record.caption, record.summary))),
        )
        for record in records
    ]
    evidence_to_records: dict[str, set[str]] = defaultdict(set)
    for record in records:
        if record.evidence_id:
            evidence_to_records[record.evidence_id].add(record.visual_id)
    ambiguous_evidence_ids = {
        evidence_id
        for evidence_id, visual_ids in evidence_to_records.items()
        if len(visual_ids) > 1
    }
    results: list[ElementAudit] = []
    for item in items:
        page = pages.get(item.page, PageEvidence(item.page, 0, 0, 0, 0, None, (), 0))
        page_records = tuple(records_by_page.get(item.page, []))
        match = match_object_evidence(
            descriptor(
                key=item.element_id,
                page=item.page,
                evidence_ids=item.evidence_id,
                object_ids=item.object_id,
                number="",
                caption=item.title,
                content=item.key_values or item.title,
            ),
            record_descriptors,
            ambiguous_evidence_ids=ambiguous_evidence_ids,
        )
        matched = tuple(
            record_by_key[key] for key in match.matched_keys if key in record_by_key
        )
        if not item.expected:
            status = "기대제외"
        else:
            if page.image_count == 0 or page.reduced_image_count == 0:
                status = "이미지_미추출"
            elif page.triage_passed_count == 0 and page.reference_filtered_count > 0:
                status = "참고자료_제외"
            elif page.triage_passed_count == 0:
                status = "triage_탈락"
            elif not page_records:
                status = "vision_유실"
            elif match.confirmed:
                status = "기록됨"
            elif match.status == "partial":
                status = "근거불충분"
            elif match.status == "ambiguous":
                status = "복수객체_검토필요"
            else:
                status = "vision_유실"
        results.append(ElementAudit(
            item,
            status,
            page.top_score,
            page.top_reasons,
            page.drawing_count,
            matched,
            match_stage=match.status,
            match_reason=match.reason,
        ))
    return results


def calculate_sensitivity(audits: Sequence[ElementAudit], pages: dict[int, PageEvidence]) -> SensitivityReport:
    expected = [audit for audit in audits if audit.item.expected]
    image_thresholds: dict[int, ThresholdSensitivity] = {}
    for threshold in (3, 4, 5, 6):
        passed_elements = sum(1 for audit in expected if audit.triage_score is not None and audit.triage_score >= threshold)
        passed_images = sum(
            sum(1 for score in (page.triage_scores or (() if page.top_score is None else (page.top_score,))) if score >= threshold)
            for page in pages.values()
        )
        total = len(expected)
        image_thresholds[threshold] = ThresholdSensitivity(
            threshold=threshold,
            passed_elements=passed_elements,
            total_expected=total,
            recall=passed_elements / total if total else 0.0,
            passed_images=passed_images,
        )
    missing = [audit for audit in expected if audit.status == "이미지_미추출"]
    vector_thresholds = {threshold: sum(1 for audit in missing if audit.drawing_count >= threshold) for threshold in (30, 45, 60)}
    return SensitivityReport(image_thresholds=image_thresholds, vector_thresholds=vector_thresholds)


def _append_rows(ws, headers: Sequence[str], rows: Iterable[WorkbookRow]) -> int:
    ws.append(list(headers))
    count = 0
    for row in rows:
        ws.append([row.get(header, "") for header in headers])
        count += 1
    return count


def _draft_base_row(page: PageContent, title: str, caption_text: str, drawing_count: int) -> WorkbookRow:
    return {
        "요소ID": "",
        "페이지": page.page_number,
        "요소유형": "",
        "제목": title,
        "데이터포함": "",
        "기대추출": "",
        "관련시트": "",
        "비고": "",
        "자동_캡션후보": caption_text,
        "자동_벡터드로잉수": drawing_count,
    }


def _image_draft_row(page: PageContent, image: WorkbookRow, caption_text: str, drawing_count: int) -> WorkbookRow:
    triage = _triage_image(page, image)
    reasons = ",".join(str(reason) for reason in triage["reasons"][:4])
    passed_text = "Y" if bool(triage["passed"]) else "N"
    row = _draft_base_row(page, caption_text or str(image.get("caption", "")), caption_text, drawing_count)
    row.update({
        "자동_트리아지점수": int(triage["score"]),
        "자동_이미지크기": f"{image.get('width', '')}x{image.get('height', '')}",
        "자동_풀렌더여부": "Y" if "full render" in str(image.get("caption", "")).lower() else "N",
        "자동_관련성통과": "Y" if _is_relevant_image(image) else "N",
        "자동_트리아지통과및사유": f"{passed_text}:{reasons}",
    })
    return row


def _page_draft_row(page: PageContent, caption_text: str, drawing_count: int) -> WorkbookRow:
    row = _draft_base_row(page, caption_text, caption_text, drawing_count)
    row.update({
        "자동_트리아지점수": "",
        "자동_이미지크기": "",
        "자동_풀렌더여부": "",
        "자동_관련성통과": "",
        "자동_트리아지통과및사유": "",
    })
    return row


def dump_inventory(source_pdf: Path, out: Path) -> int:
    pdf = extract_pdf(source_pdf)
    drawings = _drawing_counts(source_pdf)
    rows: list[WorkbookRow] = []
    for page in pdf.pages:
        caption_text = _caption_candidates(page.text)
        drawing_count = drawings.get(page.page_number, 0)
        for image in page.images:
            rows.append(_image_draft_row(page, image, caption_text, drawing_count))
        if not page.images and (caption_text or drawing_count >= 30):
            rows.append(_page_draft_row(page, caption_text, drawing_count))
    out.parent.mkdir(parents=True, exist_ok=True)
    wb = Workbook()
    info = wb.active
    info.title = "안내"
    info["A1"] = DRAFT_NOTICE
    ws = wb.create_sheet("시각요소")
    count = _append_rows(ws, [*INVENTORY_HEADERS, *AUTO_HEADERS], rows)
    wb.save(out)
    return count


def _municipality(path: Path) -> str:
    for sheet_name in ("16_시각자료목록", "00_문서메타"):
        rows = _sheet_rows(path, sheet_name)
        for row in rows:
            value = _cell_text(row.get("지자체명"))
            if value:
                return value
    return ""


def _status_counts(audits: Sequence[ElementAudit]) -> Counter[str]:
    return Counter(audit.status for audit in audits if audit.item.expected)


def _summary_lines(audits: Sequence[ElementAudit]) -> list[str]:
    expected = [audit for audit in audits if audit.item.expected]
    recorded = sum(1 for audit in expected if audit.status == "기록됨")
    recall = recorded / len(expected) if expected else 0.0
    metrics = _p5_metrics(audits)
    counts = _status_counts(audits)
    status_text = ", ".join(f"{status} {counts.get(status, 0)}" for status in STATUS_ORDER)
    return [
        f"- 기대추출 요소 수: {len(expected)}",
        f"- 기록됨 수: {recorded}",
        f"- 리콜: {recall:.3f}",
        f"- Triage 누락률: {metrics['triage_miss_rate']:.3f} "
        f"({metrics['triage_missed_objects']}/{metrics['triage_evaluable_objects']})",
        f"- 판독 재현율: {metrics['vision_recall']:.3f}",
        f"- 근거 미결정률: {metrics['evidence_unresolved_rate']:.3f}",
        f"- 상태 분포: {status_text}",
    ]


def _type_breakdown(audits: Sequence[ElementAudit]) -> list[str]:
    lines = ["| 요소유형 | 기대추출 | 기록됨 | 리콜 | 상태분포 |", "|---|---:|---:|---:|---|"]
    for element_type in sorted({audit.item.element_type for audit in audits if audit.item.expected}):
        bucket = [audit for audit in audits if audit.item.expected and audit.item.element_type == element_type]
        recorded = sum(1 for audit in bucket if audit.status == "기록됨")
        counts = Counter(audit.status for audit in bucket)
        dist = ", ".join(f"{status}:{counts.get(status, 0)}" for status in STATUS_ORDER if counts.get(status, 0))
        lines.append(f"| {element_type} | {len(bucket)} | {recorded} | {recorded / len(bucket) if bucket else 0:.3f} | {dist} |")
    return lines


def _display_reasons(reasons: Sequence[str]) -> str:
    return ",".join(_limited_reasons_with_reference(reasons, 4))


def _detail_lines(audits: Sequence[ElementAudit]) -> list[str]:
    lines = [
        "| 요소ID | 페이지 | 유형 | 상태 | 근거매칭 | triage 점수·사유 | 벡터 drawing 수 | 16시트 매칭 |",
        "|---|---:|---|---|---|---|---:|---|",
    ]
    for audit in audits:
        if not audit.item.expected:
            continue
        records = "; ".join(
            f"{r.visual_id}/{r.data_included}/{r.digitizing_needed}/{r.related_sheet}" for r in audit.matched_records
        ) or ""
        score = "" if audit.triage_score is None else str(audit.triage_score)
        reasons = _display_reasons(audit.triage_reasons)
        lines.append(
            f"| {audit.item.element_id} | {audit.item.page} | {audit.item.element_type} | {audit.status} | "
            f"{audit.match_stage}: {audit.match_reason} | {score} {reasons} | "
            f"{audit.drawing_count} | {records} |"
        )
    return lines


def _sensitivity_lines(sensitivity: SensitivityReport) -> list[str]:
    lines = [
        "### IMAGE_TRIAGE_MIN_SCORE 민감도",
        "",
        "| 기준 | triage 단계 통과 요소 | 요소 리콜 | triage 통과 이미지 수 |",
        "|---:|---:|---:|---:|",
    ]
    for threshold, row in sensitivity.image_thresholds.items():
        lines.append(f"| {threshold} | {row.passed_elements}/{row.total_expected} | {row.recall:.3f} | {row.passed_images} |")
    lines.extend(["", "### VECTOR_RENDER_MIN_DRAWINGS 민감도", "", "| 기준 | 이미지_미추출 중 렌더 대상 요소 |", "|---:|---:|"])
    for threshold, count in sensitivity.vector_thresholds.items():
        lines.append(f"| {threshold} | {count} |")
    return lines


def _false_positive_stats(audits: Sequence[ElementAudit], records: Sequence[VisualRecord]) -> dict[str, int]:
    excluded = [audit for audit in audits if not audit.item.expected]
    over = sum(bool(audit.matched_records) for audit in excluded)
    return {
        "excluded_rows": len(excluded),
        "excluded_recorded_pages": over,
        "excluded_matched_objects": over,
    }


def _false_positive_line(audits: Sequence[ElementAudit], records: Sequence[VisualRecord]) -> str:
    stats = _false_positive_stats(audits, records)
    return (
        f"- 기대추출=N 요소 {stats['excluded_rows']}개 중 "
        f"근거가 실제 일치한 수: {stats['excluded_matched_objects']}"
    )


def _p5_metrics(audits: Sequence[ElementAudit]) -> dict[str, float | int]:
    expected = [audit for audit in audits if audit.item.expected]
    excluded = [audit for audit in audits if not audit.item.expected]
    triage_evaluable = [
        audit for audit in expected
        if audit.status not in {"이미지_미추출"}
    ]
    triage_missed = [
        audit for audit in triage_evaluable
        if audit.status in {"triage_탈락", "참고자료_제외"}
    ]
    vision_evaluable = [
        audit for audit in expected
        if audit.status not in {"이미지_미추출", "triage_탈락", "참고자료_제외"}
    ]
    recorded = sum(audit.status == "기록됨" for audit in expected)
    vision_recorded = sum(audit.status == "기록됨" for audit in vision_evaluable)
    unresolved = sum(
        audit.status in {"근거불충분", "복수객체_검토필요"}
        for audit in expected
    )
    true_negative = sum(not audit.matched_records for audit in excluded)
    return {
        "fixed_denominator": True,
        "expected_objects": len(expected),
        "recorded_objects": recorded,
        "object_recall": recorded / len(expected) if expected else 0.0,
        "triage_evaluable_objects": len(triage_evaluable),
        "triage_missed_objects": len(triage_missed),
        "triage_miss_rate": (
            len(triage_missed) / len(triage_evaluable) if triage_evaluable else 0.0
        ),
        "vision_evaluable_objects": len(vision_evaluable),
        "vision_recorded_objects": vision_recorded,
        "vision_recall": (
            vision_recorded / len(vision_evaluable) if vision_evaluable else 0.0
        ),
        "evidence_unresolved_objects": unresolved,
        "evidence_unresolved_rate": unresolved / len(expected) if expected else 0.0,
        "negative_objects": len(excluded),
        "negative_specificity": true_negative / len(excluded) if excluded else 0.0,
    }


def _type_breakdown_payload(audits: Sequence[ElementAudit]) -> dict[str, JsonValue]:
    payload: dict[str, JsonValue] = {}
    for element_type in sorted({audit.item.element_type for audit in audits if audit.item.expected}):
        bucket = [audit for audit in audits if audit.item.expected and audit.item.element_type == element_type]
        recorded = sum(1 for audit in bucket if audit.status == "기록됨")
        counts = Counter(audit.status for audit in bucket)
        payload[element_type] = {
            "expected": len(bucket),
            "recorded": recorded,
            "recall": recorded / len(bucket) if bucket else 0.0,
            "status_counts": {status: counts.get(status, 0) for status in STATUS_ORDER},
        }
    return payload


def _sensitivity_payload(sensitivity: SensitivityReport) -> dict[str, JsonValue]:
    return {
        "image_thresholds": {
            str(threshold): {
                "passed_elements": row.passed_elements,
                "total_expected": row.total_expected,
                "recall": row.recall,
                "passed_images": row.passed_images,
            }
            for threshold, row in sensitivity.image_thresholds.items()
        },
        "vector_thresholds": {str(threshold): count for threshold, count in sensitivity.vector_thresholds.items()},
    }


def _report_json(
    audits: Sequence[ElementAudit],
    sensitivity: SensitivityReport,
    paths: AuditPaths,
    records: Sequence[VisualRecord],
) -> dict[str, JsonValue]:
    counts = _status_counts(audits)
    p5_metrics = _p5_metrics(audits)
    return {
        "source_pdf": str(paths.source_pdf),
        "inventory": str(paths.inventory),
        "pipeline_output": str(paths.pipeline_output),
        "inventory_rows": len(audits),
        "expected_rows": sum(1 for audit in audits if audit.item.expected),
        "recorded_rows": sum(1 for audit in audits if audit.item.expected and audit.status == "기록됨"),
        "metrics": p5_metrics,
        "status_counts": {status: counts.get(status, 0) for status in STATUS_ORDER},
        "sensitivity": _sensitivity_payload(sensitivity),
        "type_breakdown": _type_breakdown_payload(audits),
        "false_positive_stats": _false_positive_stats(audits, records),
        "format_errors": [],
        "details": [
            {
                "요소ID": audit.item.element_id,
                "페이지": audit.item.page,
                "요소유형": audit.item.element_type,
                "상태": audit.status,
                "근거매칭단계": audit.match_stage,
                "근거매칭사유": audit.match_reason,
                "일치시각자료ID": [record.visual_id for record in audit.matched_records],
                "triage_score": audit.triage_score,
                "drawing_count": audit.drawing_count,
            }
            for audit in audits
        ],
    }


def _write_reports(audits: Sequence[ElementAudit], records: Sequence[VisualRecord], sensitivity: SensitivityReport, paths: AuditPaths) -> ReportPaths:
    stamp = time.strftime("%Y%m%d_%H%M%S")
    paths.report_dir.mkdir(parents=True, exist_ok=True)
    md_path = paths.report_dir / f"visual_inventory_audit_{stamp}.md"
    json_path = paths.report_dir / f"visual_inventory_audit_{stamp}.json"
    body = [
        f"# 시각 요소 인벤토리 리콜 감사 리포트 ({stamp})",
        "",
        "## 1. 요약",
        *_summary_lines(audits),
        "",
        "## 2. 요소유형별 분해",
        *_type_breakdown(audits),
        "",
        "## 3. 요소별 상세 표",
        *_detail_lines(audits),
        "",
        "## 4. 임계 민감도 시뮬레이션",
        *_sensitivity_lines(sensitivity),
        "",
        "## 5. 기대추출=N 참고 통계",
        _false_positive_line(audits, records),
        "",
        "## 6. 실행 정보",
        f"- 입력 PDF: `{paths.source_pdf}`",
        f"- 사람 인벤토리: `{paths.inventory}`",
        f"- 파이프라인 출력: `{paths.pipeline_output}`",
        f"- 인벤토리 행 수: {len(audits)}",
        "- 형식 오류: 없음",
        "- 주의: 이 도구는 LLM/vision 백엔드를 호출하지 않고 PDF 파싱, Pillow triage, xlsx 대조만 수행했습니다.",
        "",
    ]
    md_path.write_text("\n".join(body), encoding="utf-8")
    json_path.write_text(json.dumps(_report_json(audits, sensitivity, paths, records), ensure_ascii=False, indent=2), encoding="utf-8")
    return ReportPaths(markdown=md_path, json_path=json_path)


def run_audit(paths: AuditPaths) -> ReportPaths:
    inventory = load_inventory(paths.inventory)
    records = load_visual_records(paths.pipeline_output)
    evidence = _triage_page_evidence(paths.source_pdf, _municipality(paths.pipeline_output))
    audits = classify_inventory(inventory, evidence, records)
    sensitivity = calculate_sensitivity(audits, evidence)
    return _write_reports(audits, records, sensitivity, paths)


def _parser() -> argparse.ArgumentParser:
    parser = argparse.ArgumentParser(description="시각 요소 인벤토리 리콜 감사 도구")
    sub = parser.add_subparsers(dest="command", required=True)
    dump = sub.add_parser("dump", help="PDF 파싱만으로 사람 인벤토리 초안 xlsx 생성")
    dump.add_argument("source_pdf")
    dump.add_argument("--out", default="data/golden/시각요소_초안.xlsx")
    audit = sub.add_parser("audit", help="사람 인벤토리와 16_시각자료목록 대조 리포트 생성")
    audit.add_argument("inventory")
    audit.add_argument("pipeline_output")
    audit.add_argument("source_pdf")
    audit.add_argument("--report-dir", default="output")
    return parser


def main(argv: Sequence[str] | None = None) -> int:
    parser = _parser()
    args = parser.parse_args(argv)
    try:
        if args.command == "dump":
            out = Path(args.out)
            rows = dump_inventory(Path(args.source_pdf), out)
            print(json.dumps({"xlsx": str(out), "rows": rows}, ensure_ascii=False))
            return 0
        paths = AuditPaths(Path(args.inventory), Path(args.pipeline_output), Path(args.source_pdf), Path(args.report_dir))
        reports = run_audit(paths)
        print(json.dumps({"markdown": str(reports.markdown), "json": str(reports.json_path)}, ensure_ascii=False))
        return 0
    except InventoryFormatError as exc:
        print(f"인벤토리 형식 오류: {exc}", file=sys.stderr)
        return 2
    except FileNotFoundError as exc:
        print(f"입력 파일을 찾을 수 없습니다: {exc.filename}", file=sys.stderr)
        return 2


if __name__ == "__main__":
    raise SystemExit(main())
