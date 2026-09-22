# /// script
# requires-python = ">=3.11"
# dependencies = []
# ///
"""독립 백엔드 산출물 2개를 LLM 호출 없이 교차 진단한다."""
from __future__ import annotations

import argparse
import json
import sys
import time
from pathlib import Path
from typing import Any
from zipfile import BadZipFile

import openpyxl
from openpyxl.utils.exceptions import InvalidFileException

ROOT = Path(__file__).resolve().parents[1]
if str(ROOT) not in sys.path:
    sys.path.insert(0, str(ROOT))

import config  # noqa: E402
from scripts.golden_score_contract import (  # noqa: E402
    계약헤더,
    경로라벨,
    데이터시트,
    dedup_key_text,
    문자열,
    값있음,
    값필드,
    키,
    텍스트값필드,
    to_float,
    행,
)
from scripts.golden_score_matching import _상대오차일치, 매칭하기  # noqa: E402
from utils.reference_data import build_codebook_rows  # noqa: E402


_분석시트 = [name for name in 데이터시트 if name != "00_문서메타"]
_코드필드 = {
    "부문": "표준부문",
    "관리부문": "표준부문",
    "목표수준": "목표수준",
    "달성여부": "달성여부",
    "사업유형": "사업유형",
    "전망방법코드": "전망방법",
    "데이터상태": "데이터상태",
}
_감사접두 = {
    "충돌": ("중복 키 값 충돌", "충돌"),
    "검산": ("감축률", "재정", "기준배출량≠", "정량감축량 합≠", "검산"),
    "스케일": ("단위 스케일", "기준배출량 단위 스케일", "스케일"),
    "원장": ("원장",),
    "빈 시트": ("빈 시트",),
}
_옵셔널키필드 = ("세부부문", "지표세부범주")


class 프로브형식오류(ValueError):
    """입력 워크북이 비교 가능한 출력 계약을 충족하지 않을 때 발생한다."""


def _비율(분자: int, 분모: int) -> float | None:
    return round(분자 / 분모, 4) if 분모 else None


def _행목록(ws) -> tuple[list[str], list[행]]:
    headers = [문자열(cell.value) for cell in ws[1]]
    rows: list[행] = []
    for row_number, values in enumerate(
        ws.iter_rows(min_row=2, values_only=True), start=2
    ):
        if not any(값있음(value) for value in values):
            continue
        data = {
            header: values[index] if index < len(values) else None
            for index, header in enumerate(headers)
            if header
        }
        rows.append(행(row_number, data))
    return headers, rows


def _워크북읽기(path: Path) -> dict[str, Any]:
    workbook = openpyxl.load_workbook(path, data_only=True, read_only=True)
    errors: list[str] = []
    sheets: dict[str, dict[str, Any]] = {}
    present = 0
    try:
        for sheet_name in _분석시트:
            if sheet_name not in workbook.sheetnames:
                sheets[sheet_name] = {"헤더": [], "행들": []}
                continue
            present += 1
            headers, rows = _행목록(workbook[sheet_name])
            missing = [
                header for header in 계약헤더(sheet_name) if header not in headers
            ]
            if missing:
                errors.append(f"{sheet_name}: 계약 컬럼 누락={missing}")
            sheets[sheet_name] = {"헤더": headers, "행들": rows}

        meta_headers: list[str] = []
        meta_rows: list[행] = []
        if "00_문서메타" in workbook.sheetnames:
            meta_headers, meta_rows = _행목록(workbook["00_문서메타"])
            missing = [
                header
                for header in 계약헤더("00_문서메타")
                if header not in meta_headers
            ]
            if missing:
                errors.append(f"00_문서메타: 계약 컬럼 누락={missing}")

        audit_rows: list[행] = []
        if "19_검증리포트" in workbook.sheetnames:
            audit_headers, audit_rows = _행목록(workbook["19_검증리포트"])
            missing = [
                header
                for header in ("심각도", "영역", "항목")
                if header not in audit_headers
            ]
            if missing:
                errors.append(f"19_검증리포트: 필수 컬럼 누락={missing}")
    finally:
        workbook.close()

    if present == 0:
        errors.append("01~15 데이터 시트가 하나도 없습니다")
    if errors:
        raise 프로브형식오류(f"{path}: " + "; ".join(errors))
    return {"시트": sheets, "문서메타": meta_rows, "검증리포트": audit_rows}


def _연도값(value: Any) -> int | None:
    number = to_float(value)
    if number is None or not number.is_integer():
        return None
    return int(number)


def _연도밴드(meta_rows: list[행]) -> tuple[int, int] | None:
    for row in meta_rows:
        start = _연도값(row.값.get("계획시작연도"))
        end = _연도값(row.값.get("계획종료연도"))
        if start is None or end is None or start > end:
            continue
        return start - 15, end + 2
    return None


def _코드집합() -> dict[str, set[str]]:
    codes: dict[str, set[str]] = {}
    for row in build_codebook_rows():
        code_type = 문자열(row.get("코드유형"))
        if code_type in set(_코드필드.values()):
            codes.setdefault(code_type, set()).add(dedup_key_text(row.get("코드")))
    codes.setdefault("표준부문", set()).update(
        dedup_key_text(value) for value in ("합계", "총계", "전체")
    )
    return codes


def _코드매핑(rows: list[행], headers: list[str], code_sets: dict[str, set[str]]) -> dict[str, Any]:
    details: dict[str, dict[str, Any]] = {}
    total = 0
    matched = 0
    for field_name, code_type in _코드필드.items():
        if field_name not in headers:
            continue
        values = [row.값.get(field_name) for row in rows if 값있음(row.값.get(field_name))]
        field_matched = sum(
            dedup_key_text(value) in code_sets.get(code_type, set()) for value in values
        )
        details[field_name] = {
            "채움수": len(values),
            "매핑수": field_matched,
            "매핑률": _비율(field_matched, len(values)),
        }
        total += len(values)
        matched += field_matched
    return {
        "대상필드": details,
        "채움수": total,
        "매핑수": matched,
        "매핑률": _비율(matched, total),
    }


def _연도정합(rows: list[행], headers: list[str], band: tuple[int, int] | None) -> dict[str, Any]:
    year_fields = [
        header
        for header in headers
        if header == "연도"
        or header.endswith("연도")
        or header in {"기간시작", "기간종료"}
    ]
    if band is None or not year_fields:
        return {"밴드": None, "비교수": 0, "정합수": 0, "정합률": None}
    values = [
        row.값.get(field_name)
        for row in rows
        for field_name in year_fields
        if 값있음(row.값.get(field_name))
    ]
    years = [_연도값(value) for value in values]
    consistent = sum(
        year is not None and band[0] <= year <= band[1] for year in years
    )
    return {
        "밴드": [band[0], band[1]],
        "비교수": len(values),
        "정합수": consistent,
        "정합률": _비율(consistent, len(values)),
    }


def _L2(
    sheet_name: str,
    sheet: dict[str, Any],
    band: tuple[int, int] | None,
    code_sets: dict[str, set[str]],
) -> dict[str, Any]:
    rows = sheet["행들"]
    required_filled = 0
    for row in rows:
        values = dict(row.값)
        for field_name in _옵셔널키필드:
            if not 값있음(values.get(field_name)):
                values[field_name] = "__옵셔널필드__"
        required_filled += 키(행(row.번호, values), sheet_name, False) is not None
    optional_fields = [
        field_name for field_name in _옵셔널키필드 if field_name in sheet["헤더"]
    ]
    optional_total = len(rows) * len(optional_fields)
    optional_filled = sum(
        값있음(row.값.get(field_name))
        for row in rows
        for field_name in optional_fields
    )
    return {
        "일차키채움수": required_filled,
        "일차키채움률": _비율(required_filled, len(rows)),
        "옵셔널필드": {
            "대상필드": optional_fields,
            "채움수": optional_filled,
            "전체수": optional_total,
            "채움률": _비율(optional_filled, optional_total),
        },
        "표준코드": _코드매핑(rows, sheet["헤더"], code_sets),
        "연도": _연도정합(rows, sheet["헤더"], band),
    }


def _감사유형(item: Any) -> str | None:
    text = 문자열(item)
    for audit_type, prefixes in _감사접두.items():
        if text.startswith(prefixes):
            return audit_type
    return None


def _영역시트(area: Any) -> str | None:
    text = 문자열(area)
    for sheet_name in _분석시트:
        suffix = sheet_name.split("_", 1)[1]
        internal = next(
            (
                key
                for key, value in config.SHEET_KEY_TO_NAME.items()
                if value == sheet_name
            ),
            "",
        )
        if text in {sheet_name, suffix, internal}:
            return sheet_name
    return None


def _자기감사(audit_rows: list[행], sheets: dict[str, dict[str, Any]]) -> dict[str, Any]:
    total_rows = sum(len(sheets[name]["행들"]) for name in _분석시트)
    severity = {"경고": 0, "정보": 0}
    types = {name: 0 for name in _감사접두}
    by_sheet = {
        name: {"건수": 0, "유형별": {audit_type: 0 for audit_type in _감사접두}}
        for name in _분석시트
    }
    for row in audit_rows:
        severity_name = 문자열(row.값.get("심각도"))
        if severity_name in severity:
            severity[severity_name] += 1
        audit_type = _감사유형(row.값.get("항목"))
        if audit_type is None:
            continue
        types[audit_type] += 1
        sheet_name = _영역시트(row.값.get("영역"))
        if sheet_name is not None:
            by_sheet[sheet_name]["건수"] += 1
            by_sheet[sheet_name]["유형별"][audit_type] += 1
    for sheet_name, stats in by_sheet.items():
        stats["밀도"] = _비율(stats["건수"], len(sheets[sheet_name]["행들"])) or 0.0
    return {
        "데이터행수": total_rows,
        "심각도별": severity,
        "심각도별밀도": {
            name: _비율(count, total_rows) or 0.0 for name, count in severity.items()
        },
        "유형별": types,
        "유형별밀도": {
            name: _비율(count, total_rows) or 0.0 for name, count in types.items()
        },
        "시트별": by_sheet,
    }


def _값쌍일치(sheet_name: str, match: Any) -> bool | None:
    common_fields = [
        field_name
        for field_name in 값필드.get(sheet_name, [])
        if 값있음(match.골든.값.get(field_name))
        and 값있음(match.출력.값.get(field_name))
    ]
    if not common_fields:
        return None
    return all(
        dedup_key_text(match.골든.값.get(field_name))
        == dedup_key_text(match.출력.값.get(field_name))
        if field_name in 텍스트값필드
        else _상대오차일치(
            match.골든.값.get(field_name), match.출력.값.get(field_name)
        )
        for field_name in common_fields
    )


def _표본키들(
    sheet_name: str,
    matches: list[Any],
    a_rows: list[행],
    b_rows: list[행],
    remaining_a: set[int],
    remaining_b: set[int],
) -> list[dict[str, str]]:
    candidates: list[dict[str, str]] = []
    for match in matches:
        if _값쌍일치(sheet_name, match) is False:
            candidates.append({"구분": "값불일치", "키": match.키})
    for index in sorted(remaining_a):
        row = a_rows[index]
        candidates.append(
            {"구분": "A단독", "키": 키(row, sheet_name, False) or 키(row, sheet_name, True) or ""}
        )
    for index in sorted(remaining_b):
        row = b_rows[index]
        candidates.append(
            {"구분": "B단독", "키": 키(row, sheet_name, False) or 키(row, sheet_name, True) or ""}
        )
    candidates.extend({"구분": "합의", "키": match.키} for match in matches)
    result: list[dict[str, str]] = []
    seen: set[tuple[str, str]] = set()
    for candidate in candidates:
        identity = (candidate["구분"], candidate["키"])
        if not candidate["키"] or identity in seen:
            continue
        seen.add(identity)
        result.append(candidate)
    return result


def _시트비교(
    sheet_name: str,
    a_sheet: dict[str, Any],
    b_sheet: dict[str, Any],
    a_audit: dict[str, Any],
    b_audit: dict[str, Any],
    a_band: tuple[int, int] | None,
    b_band: tuple[int, int] | None,
    code_sets: dict[str, set[str]],
) -> dict[str, Any]:
    a_all = a_sheet["행들"]
    b_all = b_sheet["행들"]
    a_rows = [row for row in a_all if 키(row, sheet_name, False) or 키(row, sheet_name, True)]
    b_rows = [row for row in b_all if 키(row, sheet_name, False) or 키(row, sheet_name, True)]
    matches, remaining_a, remaining_b, semantic, character, observations = 매칭하기(
        a_rows, b_rows, sheet_name, 보수티어만=True
    )
    if semantic or character or observations:
        raise AssertionError("보수 티어 호출에서 후속 유사도 결과가 생성되었습니다")

    comparable_results = [
        result
        for match in matches
        if (result := _값쌍일치(sheet_name, match)) is not None
    ]
    agreed = sum(comparable_results)
    denominator = min(len(a_rows), len(b_rows))
    value_defined = sheet_name in 값필드
    return {
        "A행수": len(a_all),
        "B행수": len(b_all),
        "A유효행수": len(a_rows),
        "B유효행수": len(b_rows),
        "매칭수": len(matches),
        "엄격매칭수": sum(match.방식 == "엄격" for match in matches),
        "완화매칭수": sum(match.방식 == "완화" for match in matches),
        "키합의율": _비율(len(matches), denominator),
        "A단독행수": len(remaining_a),
        "A단독비율": _비율(len(remaining_a), len(a_rows)),
        "B단독행수": len(remaining_b),
        "B단독비율": _비율(len(remaining_b), len(b_rows)),
        "값비교쌍수": len(comparable_results) if value_defined else None,
        "값합의쌍수": agreed if value_defined else None,
        "값합의율": _비율(agreed, len(comparable_results)) if value_defined else None,
        "L1": {
            "A": a_audit["시트별"][sheet_name],
            "B": b_audit["시트별"][sheet_name],
        },
        "L2": {
            "A": _L2(sheet_name, a_sheet, a_band, code_sets),
            "B": _L2(sheet_name, b_sheet, b_band, code_sets),
        },
        "표본후보키": _표본키들(
            sheet_name, matches, a_rows, b_rows, remaining_a, remaining_b
        ),
    }


def _L4_표본(sheet_results: dict[str, dict[str, Any]]) -> list[dict[str, Any]]:
    candidates = [
        (sheet_name, result)
        for sheet_name, result in sheet_results.items()
        if sheet_name in 값필드 and (result["A유효행수"] or result["B유효행수"])
    ]
    candidates.sort(
        key=lambda item: (
            item[1]["값합의율"] is None,
            item[1]["값합의율"] if item[1]["값합의율"] is not None else 0.0,
            _분석시트.index(item[0]),
        )
    )
    return [
        {
            "시트": sheet_name,
            "값합의율": result["값합의율"],
            "표본키": result["표본후보키"][:10],
        }
        for sheet_name, result in candidates[:3]
    ]


def _표시(value: Any) -> str:
    if value is None:
        return "-"
    if isinstance(value, float):
        return f"{value:.4f}"
    return str(value)


def _마크다운(result: dict[str, Any]) -> str:
    lines = [
        f"# 무골든 교차 백엔드 품질 리포트 ({result['라벨']})\n\n",
        "## 실행 정보\n\n",
        f"- A 산출물: `{result['A산출물']}`\n",
        f"- B 산출물: `{result['B산출물']}`\n",
        f"- Markdown: `{result['리포트']['md']}`\n",
        f"- JSON: `{result['리포트']['json']}`\n",
        f"- 실행시각: {result['실행시각']}\n",
        "- LLM 호출: 0\n",
        "- 임계 합격/불합격 판정: 없음\n\n",
        "## 시트별 지표\n\n",
        "| 시트 | A행 | B행 | 키합의율 | 값합의율 | A단독 | B단독 | L1 A밀도 | L1 B밀도 | A키채움률 | B키채움률 | A옵셔널채움률 | B옵셔널채움률 | A코드매핑률 | B코드매핑률 | A연도정합률 | B연도정합률 |\n",
        "|---|---:|---:|---:|---:|---:|---:|---:|---:|---:|---:|---:|---:|---:|---:|---:|---:|\n",
    ]
    for sheet_name, sheet in result["시트별"].items():
        values = [
            sheet_name,
            sheet["A행수"],
            sheet["B행수"],
            sheet["키합의율"],
            sheet["값합의율"],
            sheet["A단독행수"],
            sheet["B단독행수"],
            sheet["L1"]["A"]["밀도"],
            sheet["L1"]["B"]["밀도"],
            sheet["L2"]["A"]["일차키채움률"],
            sheet["L2"]["B"]["일차키채움률"],
            sheet["L2"]["A"]["옵셔널필드"]["채움률"],
            sheet["L2"]["B"]["옵셔널필드"]["채움률"],
            sheet["L2"]["A"]["표준코드"]["매핑률"],
            sheet["L2"]["B"]["표준코드"]["매핑률"],
            sheet["L2"]["A"]["연도"]["정합률"],
            sheet["L2"]["B"]["연도"]["정합률"],
        ]
        lines.append("| " + " | ".join(_표시(value) for value in values) + " |\n")

    lines.extend(
        [
            "\n## L1 자기 감사 요약\n\n",
            "| 백엔드 | 경고 | 정보 | 충돌 | 검산 | 스케일 | 원장 | 빈 시트 | 데이터행수 |\n",
            "|---|---:|---:|---:|---:|---:|---:|---:|---:|\n",
        ]
    )
    for backend in ("A", "B"):
        audit = result["L1_자기감사"][backend]
        values = [
            backend,
            audit["심각도별"]["경고"],
            audit["심각도별"]["정보"],
            *(audit["유형별"][name] for name in _감사접두),
            audit["데이터행수"],
        ]
        lines.append("| " + " | ".join(_표시(value) for value in values) + " |\n")

    lines.append("\n## 합의율 하위 시트\n\n")
    if not result["L4_표본후보"]:
        lines.append("- 값 합의율을 계산할 수 있는 시트가 없습니다.\n")
    for sheet in result["L4_표본후보"]:
        lines.append(f"### {sheet['시트']} (값 합의율 {_표시(sheet['값합의율'])})\n\n")
        if not sheet["표본키"]:
            lines.append("- 표본 키 없음\n")
        for sample in sheet["표본키"]:
            lines.append(f"- [{sample['구분']}] `{sample['키']}`\n")
        lines.append("\n")
    return "".join(lines)


def probe_workbooks(
    a_path: str | Path,
    b_path: str | Path,
    *,
    report_dir: str | Path = "output",
    label: str | None = None,
) -> dict[str, Any]:
    a_file = Path(a_path).expanduser().resolve()
    b_file = Path(b_path).expanduser().resolve()
    a_data = _워크북읽기(a_file)
    b_data = _워크북읽기(b_file)
    code_sets = _코드집합()
    a_audit = _자기감사(a_data["검증리포트"], a_data["시트"])
    b_audit = _자기감사(b_data["검증리포트"], b_data["시트"])
    a_band = _연도밴드(a_data["문서메타"])
    b_band = _연도밴드(b_data["문서메타"])
    sheet_results = {
        sheet_name: _시트비교(
            sheet_name,
            a_data["시트"][sheet_name],
            b_data["시트"][sheet_name],
            a_audit,
            b_audit,
            a_band,
            b_band,
            code_sets,
        )
        for sheet_name in _분석시트
    }

    label_value = 경로라벨(label)
    report_base = Path(report_dir).expanduser().resolve()
    report_base.mkdir(parents=True, exist_ok=True)
    md_path = report_base / f"cross_backend_probe_{label_value}.md"
    json_path = report_base / f"cross_backend_probe_{label_value}.json"
    result: dict[str, Any] = {
        "라벨": label_value,
        "A산출물": str(a_file),
        "B산출물": str(b_file),
        "실행시각": time.strftime("%Y-%m-%d %H:%M:%S"),
        "LLM호출수": 0,
        "시트별": sheet_results,
        "L1_자기감사": {"A": a_audit, "B": b_audit},
        "L4_표본후보": _L4_표본(sheet_results),
        "리포트": {"md": str(md_path), "json": str(json_path)},
    }
    md_path.write_text(_마크다운(result), encoding="utf-8")
    json_path.write_text(
        json.dumps(result, ensure_ascii=False, indent=2, default=str), encoding="utf-8"
    )
    return result


def main(argv: list[str] | None = None) -> int:
    parser = argparse.ArgumentParser(
        description="독립 백엔드 산출물 2개의 무골든 품질 지표를 LLM 호출 없이 생성"
    )
    parser.add_argument("a", help="기준 백엔드 산출물 xlsx")
    parser.add_argument("b", help="대조 백엔드 산출물 xlsx")
    parser.add_argument("--report-dir", default="output", help="리포트 저장 디렉터리")
    parser.add_argument("--label", default=None, help="리포트 파일 라벨")
    args = parser.parse_args(argv)
    try:
        result = probe_workbooks(
            args.a, args.b, report_dir=args.report_dir, label=args.label
        )
    except (
        FileNotFoundError,
        OSError,
        InvalidFileException,
        BadZipFile,
        프로브형식오류,
    ) as exc:
        print(f"파일·형식 오류: {exc}", file=sys.stderr)
        return 2
    print(result["리포트"]["md"])
    print(result["리포트"]["json"])
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
