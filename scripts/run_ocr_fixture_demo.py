"""Persist a reproducible synthetic A/B harness demo; no server or model calls."""
from __future__ import annotations

import argparse
import io
import json
import os
from pathlib import Path
import sys
import time

PROJECT_ROOT = Path(__file__).resolve().parents[1]
EXPERIMENT_ROOT = PROJECT_ROOT / "졸업프로젝트_실험1"
for location in (PROJECT_ROOT, EXPERIMENT_ROOT):
    if str(location) not in sys.path:
        sys.path.insert(0, str(location))

import fitz
from PIL import Image, ImageDraw, ImageFont
from openpyxl import load_workbook

from uocr_experiment.baseline import extract_baseline
from uocr_experiment.evaluation import evaluate_run
from uocr_experiment.hybrid import build_hybrid
from uocr_experiment.io_utils import load_manifest, save_objects, sha256_file, write_json, write_jsonl
from uocr_experiment.models import ExperimentObject
from uocr_experiment.project_bridge import export_project_objects
from uocr_experiment.reporting import write_reports
from uocr_experiment.sampling import prepare_sample
from uocr_experiment.uocr_parser import parse_raw_directory


REFERENCE_ROWS = {
    1: [
        ["Sector", "2020", "2030", "Unit"],
        ["buildings", "100", "90", "tCO2eq"],
        ["transport", "200", "180", "tCO2eq"],
        ["industry", "300", "240", "tCO2eq"],
    ],
    2: [
        ["Sector", "2020", "2030", "Unit"],
        ["buildings", "10", "9", "tCO2eq"],
        ["transport", "20", "18", "tCO2eq"],
        ["industry", "0", "0", "tCO2eq"],
    ],
}


def _number_image(value: str) -> bytes:
    image = Image.new("RGB", (72, 22), "white")
    drawing = ImageDraw.Draw(image)
    drawing.text((4, 4), value, fill="black", font=ImageFont.load_default())
    stream = io.BytesIO()
    image.save(stream, format="PNG")
    return stream.getvalue()


def make_fixture_pdf(path: Path) -> None:
    """Raster-only value cells naturally disappear from native text extraction."""
    pdf = fitz.open()
    try:
        pdf.set_metadata({
            "title": "Synthetic offline OCR integration fixture",
            "author": "Reproducible fixture generator",
            "creationDate": "D:20261010000000Z",
            "modDate": "D:20261010000000Z",
        })
        for number, rows in REFERENCE_ROWS.items():
            page = pdf.new_page(width=600, height=400)
            page.insert_text((60, 40), f"Offline OCR fixture page {number}", fontsize=14)
            page.insert_text((60, 62), "Synthetic table; numeric cells are raster images.", fontsize=10)
            columns = (60, 200, 300, 400, 500)
            for y in (90, 120, 150, 180, 210):
                page.draw_line((60, y), (500, y), color=(0, 0, 0), width=1)
            for x in columns:
                page.draw_line((x, 90), (x, 210), color=(0, 0, 0), width=1)
            for row_index, row in enumerate(rows):
                top = 90 + row_index * 30
                for column, value in enumerate(row):
                    if row_index > 0 and column in {1, 2}:
                        page.insert_image(
                            fitz.Rect(columns[column]+6, top+4, columns[column]+78, top+26),
                            stream=_number_image(value),
                        )
                    else:
                        page.insert_text((columns[column]+6, top+20), value, fontsize=10)
            if number == 2:
                page.set_rotation(90)
        # Stable fixture bytes: timestamps and PDF IDs do not change per run.
        pdf.save(path, no_new_id=True)
    finally:
        pdf.close()


def _markdown_table(rows: list[list[str]]) -> str:
    lines = ["| " + " | ".join(rows[0]) + " |",
             "| " + " | ".join("---" for _ in rows[0]) + " |"]
    lines.extend("| " + " | ".join(row) + " |" for row in rows[1:])
    return "\n".join(lines)


def _raw_table(obj: ExperimentObject, rows: list[list[str]]) -> str:
    return "<|det|>table " + json.dumps(obj.bbox_norm) + "<|/det|>\n" + _markdown_table(rows) + "\n"


def diagnostics_demo(destination: Path, manifest_path: Path, baseline: list[ExperimentObject],
                     native_tables: dict[int, ExperimentObject]) -> dict:
    """Run parse/selection failures separately so they do not contaminate A/B."""
    diagnostic_root = destination / "diagnostic_cases"
    diagnostic_manifest = write_json(diagnostic_root / "sample_manifest.json", load_manifest(manifest_path))
    raw = diagnostic_root / "saved_raw"
    raw.mkdir(parents=True, exist_ok=True)
    table = native_tables[1]
    malformed_location = (
        "<|det|>table [0, 0, 1000, 999]<|/det|>\n"
        + _markdown_table(REFERENCE_ROWS[1]) + "\n"
    )
    bad_span = (
        "<|det|>table " + json.dumps(table.bbox_norm) + "<|/det|>\n"
        '<table><tr><th rowspan="0">Sector</th><th>Value</th></tr>'
        "<tr><td>buildings</td><td>100</td></tr></table>\n"
    )
    (raw / "page_0001.md").write_text(malformed_location + bad_span, encoding="utf-8")
    (raw / "page_9999.md").write_text("| X | Y |\n| --- | --- |\n| 1 | 2 |", encoding="utf-8")
    save_objects(diagnostic_root / "outputs/pymupdf/objects.jsonl", baseline)
    _, parsed = parse_raw_directory(diagnostic_manifest, raw)
    _, _, decisions = build_hybrid(diagnostic_manifest)
    parse_path = diagnostic_root / "outputs/unlimited_ocr/parse_diagnostics.jsonl"
    parse_diagnostics = [
        json.loads(line) for line in parse_path.read_text(encoding="utf-8").splitlines() if line.strip()
    ]
    expected_codes = {"raw_page_not_selected", "raw_page_missing"}
    assert expected_codes <= {row["code"] for row in parse_diagnostics}
    assert any(obj.metadata["location_status"] == "unsupported" for obj in parsed)
    assert any(obj.metadata["table_structure_status"] == "unsupported" for obj in parsed)
    assert all(row["action"] == "reject" for row in decisions)
    report = {
        "scope": "independent_harness_only",
        "main_directory_loader_executed": False,
        "network_requests": 0,
        "server_or_gpu_contacted": False,
        "model_requested": None,
        "main_pipeline_or_domain_sheets_executed": False,
        "diagnostic_fixtures": [
            "normalized coordinate above 999 is unsupported",
            "rowspan=0 is unsupported",
            "raw file page_9999 is outside the selected page sample",
            "selected page 2 has no OCR raw result",
        ],
        "parse_diagnostics": parse_diagnostics,
        "unsupported_objects": [obj.to_dict() for obj in parsed],
        "hybrid_decisions": decisions,
    }
    write_json(diagnostic_root / "diagnostics.json", report)
    return report


def run_demo(destination: Path) -> dict:
    started = time.perf_counter()
    destination.mkdir(parents=True, exist_ok=True)
    previous = None
    previous_path = destination / "demo_summary.json"
    if previous_path.is_file():
        try:
            previous = json.loads(previous_path.read_text(encoding="utf-8"))
        except (OSError, ValueError):
            previous = None
    source = destination / "offline_fixture.pdf"
    make_fixture_pdf(source)
    manifest = prepare_sample(source, destination / "run", pages=[1, 2], dpi=144, seed=42)
    _, baseline = extract_baseline(manifest)
    native_tables = {obj.page_number: obj for obj in baseline if obj.object_type == "table"}
    assert set(native_tables) == set(REFERENCE_ROWS), "fixture must contain one native table per page"
    for page, table in native_tables.items():
        assert len(table.rows) == 4 and len(table.rows[0]) == 4
        assert any(not cell for row in table.rows[1:] for cell in row[1:3]), "raster cells should be absent in native text"
    saved_raw = destination / "saved_raw"
    saved_raw.mkdir(parents=True, exist_ok=True)
    for page, table in native_tables.items():
        (saved_raw / f"page_{page:04d}.md").write_text(_raw_table(table, REFERENCE_ROWS[page]), encoding="utf-8")
    _, ocr = parse_raw_directory(manifest, saved_raw)
    _, hybrid, decisions = build_hybrid(manifest)
    assert sum(row["action"] == "replace" for row in decisions) == 2
    export_paths = {engine: str(export_project_objects(manifest, engine))
                    for engine in ("pymupdf", "unlimited_ocr", "hybrid")}
    golden_objects = []
    for page, table in native_tables.items():
        value = table.to_dict()
        value.update(object_id=f"fixture-golden-p{page}", engine="golden",
                     rows=REFERENCE_ROWS[page], include=True,
                     review_note="Synthetic truth defined by this generator; not a human-reviewed real document benchmark.")
        golden_objects.append(value)
    golden = write_jsonl(destination / "fixture_golden_v1.jsonl", golden_objects)
    summaries, matches = evaluate_run(manifest, golden)
    reports = write_reports(manifest, summaries, matches)
    metrics = {row["engine"]: row for row in summaries}
    assert metrics["pymupdf"]["total_golden_cell_restoration"] < 1
    assert metrics["unlimited_ocr"]["total_golden_cell_restoration"] == 1
    assert metrics["hybrid"]["total_golden_cell_restoration"] == 1
    diagnostic_report = diagnostics_demo(destination, manifest, baseline, native_tables)
    workbook = load_workbook(reports[-1], read_only=True, data_only=True)
    try:
        expected_sheets = {"summary", "pymupdf_objects", "uocr_objects", "hybrid_objects",
                           "merge_decisions", "golden_matches", "table_row_metrics", "measurements"}
        assert expected_sheets <= set(workbook.sheetnames)
        verified_sheets = workbook.sheetnames
    finally:
        workbook.close()
    report = {
        "purpose": "offline wiring/evaluation fixture demo; not a real OCR accuracy benchmark",
        "scope": "independent_harness_only",
        "pdf_sha256": sha256_file(source),
        "golden_sha256": sha256_file(golden),
        "selected_pages": [1, 2],
        "fixture_generator": str(Path(__file__).resolve()),
        "synthetic_golden_version": "fixture-v1",
        "ocr_source": "hand-authored saved fixture raw; no OCR model executed",
        "server_or_gpu_contacted": False,
        "network_requests": 0,
        "model_changes": False,
        "main_pipeline_or_domain_sheets_executed": False,
        "elapsed_local_demo_seconds": round(time.perf_counter()-started, 3),
        "request_elapsed_seconds": None,
        "vram_peak_gb": None,
        "paths": {
            "pdf": str(source), "manifest": str(manifest), "golden": str(golden),
            "saved_raw": str(saved_raw), "project_document_objects": export_paths,
            "evaluation_json": str(reports[0]), "evaluation_markdown": str(reports[1]),
            "evaluation_xlsx": str(reports[2]),
            "diagnostics": str(destination / "diagnostic_cases/diagnostics.json"),
        },
        "summaries": summaries,
        "hybrid_decisions": decisions,
        "verified_excel_sheets": verified_sheets,
        "diagnostic_summary": {
            "unsupported_object_count": len(diagnostic_report["unsupported_objects"]),
            "rejected_object_count": sum(row["action"] == "reject" for row in diagnostic_report["hybrid_decisions"]),
            "parse_diagnostics": diagnostic_report["parse_diagnostics"],
        },
    }
    report["repeatability"] = {
        "previous_run_present": isinstance(previous, dict),
        "same_pdf_sha256": previous.get("pdf_sha256") == report["pdf_sha256"] if isinstance(previous, dict) else None,
        "same_golden_sha256": previous.get("golden_sha256") == report["golden_sha256"] if isinstance(previous, dict) else None,
        "same_metrics": previous.get("summaries") == report["summaries"] if isinstance(previous, dict) else None,
    }
    summary_path = write_json(destination / "demo_summary.json", report)
    guide = [
        "# 오프라인 OCR 하네스 재실행 예시",
        "",
        "이 결과는 생성한 동일 PDF와 동일 fixture-v1 정답으로 하네스 연결·평가를 검증한 예시입니다.",
        "숫자 셀은 PDF에서 실제 보이는 래스터 이미지로 구성되어 native text에는 없습니다.",
        "saved_raw는 생성 스크립트가 정답에서 작성한 저장 응답입니다. 실제 OCR 모델의 성능 결과가 아닙니다.",
        "서버/GPU/외부 모델 요청은 0회이며 본체의 업무 데이터 시트는 이번 예시에서 실행하지 않습니다.",
        "",
        "저장소 루트에서 프로젝트 의존성이 설치된 Python으로 재실행:",
        "",
        "```powershell",
        "python scripts/run_ocr_fixture_demo.py",
        "```",
        "",
        "실제 이번 실행 인터프리터:",
        "",
        "```powershell",
        "$env:PYTHON_DOTENV_DISABLED = '1'",
        *(
            ["$env:PYTHONPATH = '" + os.environ["PYTHONPATH"].replace("'", "''") + "'"]
            if os.environ.get("PYTHONPATH") else []
        ),
        "& '" + sys.executable.replace("'", "''") + "' scripts/run_ocr_fixture_demo.py",
        "```",
        "",
        "기본 산출물은 output/ocr_fixture_demo_20261010 하위의 정해진 예시 파일에 저장됩니다.",
        "기존 main의 실행 결과·원본 문서는 수정하지 않으며 다른 파일을 삭제하지 않습니다.",
        "",
        "| 엔진 | 객체 리콜 | 객체 정밀도 | 매칭 표 셀 정확도 | 전체 정답 셀 복원율 | 표 행 리콜 |",
        "|---|---:|---:|---:|---:|---:|",
    ]
    for row in summaries:
        guide.append(f"| {row['engine']} | {row['object_recall']} | {row['object_precision']} | "
                     f"{row['matched_table_cell_accuracy']} | {row['total_golden_cell_restoration']} | "
                     f"{row['table_row_recall']} |")
    guide.extend([
        "",
        "각 엔진의 PDF 해시·정답 해시·표본 페이지·골든 버전은 동일합니다.",
        "입력 순서나 반복 실행으로 중복 객체를 추가하지 않습니다. 두 번째 페이지는 90도 회전 PDF입니다.",
        "전체 JSON/Markdown/XLSX는 run/reports, 엔진별 객체 JSONL은 run/outputs 아래에 있습니다.",
        "DocumentObject export에는 native engine과 OCR engine을 각각 보존합니다.",
        "",
        "diagnostic_cases에는 범위를 벗어난 0~999 좌표, 잘못된 rowspan, 표본 밖 페이지,",
        "선택 페이지 OCR 결과 누락을 실제 파싱·병합한 기록이 있습니다.",
        "지원하지 않는 객체를 hybrid에서 보류/거절한 결과와 요청 범위를 diagnostics.json에서 확인할 수 있습니다.",
        "",
        "로컬 예시 실행 시간만 demo_summary.json에 측정합니다. 서버 요청 시간과 최대 VRAM은 측정하지 않습니다.",
        "실제 문서의 사람 검수 골든과 실제 서버 정확도·처리 시간 검증은 별도로 남아 있습니다.",
        "",
        "주요 산출물:",
        "",
        "- offline_fixture.pdf",
        "- fixture_golden_v1.jsonl",
        "- demo_summary.json",
        "- run/reports/evaluation.json",
        "- run/reports/evaluation.md",
        "- run/reports/evaluation.xlsx",
        "- diagnostic_cases/diagnostics.json",
    ])
    (destination / "README_demo.md").write_text("\n".join(guide), encoding="utf-8")
    print(json.dumps({"summary": str(summary_path), "reports": [str(p) for p in reports],
                      "network_requests": 0, "pdf_sha256": report["pdf_sha256"]}, ensure_ascii=False))
    return report


def main() -> int:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--output", default=str(PROJECT_ROOT / "output/ocr_fixture_demo_20261010"))
    args = parser.parse_args()
    destination = Path(args.output).resolve()
    allowed = (PROJECT_ROOT / "output").resolve()
    if destination == allowed or not destination.is_relative_to(allowed):
        parser.error("--output must be a subdirectory of this checkout's output directory")
    run_demo(destination)
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
