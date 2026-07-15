"""Validate generated Excel against source PDF/HWP at a content-sanity level.

The verifier is intentionally conservative: it checks sheet structure, non-empty core
outputs, source-token support for textual fields, numeric support where text extraction
can expose it, and image-observation coverage when visual analysis is expected.
"""

from __future__ import annotations

import argparse
import json
import re
from pathlib import Path
from typing import Iterable

import fitz
from openpyxl import load_workbook

EXPECTED_SHEETS = [
    "용도별 자동차(현황)",
    "용도별 에너지(현황)",
    "온실가스(현황전망목표)",
    "감축전략(계획실적)",
    "감축전략(정성사업)",
    "이미지·그래프 판독결과",
    "지자체별 요약카드",
]
CORE_NONEMPTY = [
    "온실가스(현황전망목표)",
    "감축전략(계획실적)",
    "지자체별 요약카드",
]


def _pdf_text(pdf_path: Path) -> tuple[str, dict[int, str]]:
    doc = fitz.open(str(pdf_path))
    page_text: dict[int, str] = {}
    all_parts: list[str] = []
    for idx, page in enumerate(doc, start=1):
        text = page.get_text("text", sort=True)
        page_text[idx] = text
        all_parts.append(text)
    doc.close()
    return "\n".join(all_parts), page_text


def _norm_text(text: str) -> str:
    return re.sub(r"\s+", "", str(text or "")).casefold()


def _tokens(text: str) -> list[str]:
    # Keep meaningful Korean/ASCII tokens; drop generic administrative words.
    stop = {"서울특별시", "계획", "현황", "목표", "부문", "사업", "기타", "합계", "직접배출", "간접배출"}
    found = re.findall(r"[가-힣A-Za-z0-9]{2,}", str(text or ""))
    return [t for t in found if t not in stop and not t.isdigit()]


def _row_dicts(ws) -> list[dict]:
    headers = [cell.value for cell in ws[1]]
    rows = []
    for values in ws.iter_rows(min_row=2, values_only=True):
        if not any(v is not None and v != "" for v in values):
            continue
        rows.append({headers[idx]: value for idx, value in enumerate(values)})
    return rows


def _number_variants(value) -> set[str]:
    if value is None or value == "":
        return set()
    try:
        number = float(str(value).replace(",", ""))
    except ValueError:
        return set()
    variants = {str(value).replace(",", "").replace(".0", "")}
    variants.add(str(int(round(number))))
    if abs(number) >= 1000:
        variants.add(f"{int(round(number)):,}")
    # Source sometimes uses thousand-ton units while Excel may preserve visible value.
    if abs(number) >= 1000 and number % 1000 == 0:
        variants.add(str(int(round(number / 1000))))
        variants.add(f"{int(round(number / 1000)):,}")
    return {re.sub(r"\s+", "", v) for v in variants if v}


def _any_token_supported(values: Iterable[str], source_norm: str, min_hits: int = 1) -> tuple[bool, list[str]]:
    tokens: list[str] = []
    for value in values:
        tokens.extend(_tokens(value))
    unique = []
    for token in tokens:
        if token not in unique:
            unique.append(token)
    hits = [token for token in unique if _norm_text(token) in source_norm]
    return len(hits) >= min_hits, hits


def verify(excel_path: Path, pdf_path: Path, *, require_images: bool = True) -> dict:
    source_text, page_text = _pdf_text(pdf_path)
    source_norm = _norm_text(source_text)
    source_digits = re.sub(r"\D+", "", source_text)

    wb = load_workbook(excel_path, data_only=True, read_only=True)
    failures: list[str] = []
    warnings: list[str] = []
    details: dict = {"sheets": {}}

    missing = [sheet for sheet in EXPECTED_SHEETS if sheet not in wb.sheetnames]
    if missing:
        failures.append(f"missing sheets: {missing}")

    for sheet in EXPECTED_SHEETS:
        if sheet not in wb.sheetnames:
            continue
        ws = wb[sheet]
        rows = _row_dicts(ws)
        details["sheets"][sheet] = {"rows": len(rows), "cols": ws.max_column}
        if sheet in CORE_NONEMPTY and not rows:
            failures.append(f"core sheet is empty: {sheet}")

    if "서울특별시" not in source_text:
        failures.append("source PDF text does not contain municipality 서울특별시")

    # Summary rows must be grounded by source terms and have evidence text.
    if "지자체별 요약카드" in wb.sheetnames:
        for row in _row_dicts(wb["지자체별 요약카드"]):
            item = row.get("항목")
            content = row.get("내용") or ""
            evidence = row.get("근거") or ""
            if not content:
                failures.append(f"summary item has empty content: {item}")
                continue
            supported, hits = _any_token_supported([content], source_norm, min_hits=1)
            if not supported:
                failures.append(f"summary content lacks source-token support: {item}")
            if not evidence:
                warnings.append(f"summary item has no evidence pointer: {item}")
            details.setdefault("summary_support", {})[item] = hits[:5]

    # Strategy names should not be hallucinated: at least one meaningful token should appear in source.
    if "감축전략(계획실적)" in wb.sheetnames:
        unsupported = []
        for row in _row_dicts(wb["감축전략(계획실적)"]):
            name = row.get("감축사업명") or ""
            detail = row.get("감축사업명_세부") or ""
            supported, hits = _any_token_supported([name, detail], source_norm, min_hits=1)
            if not supported:
                unsupported.append(str(name)[:60])
        if unsupported:
            failures.append(f"strategy names without source-token support: {unsupported[:10]}")

    # Numeric check is warning-only because graph/image values may not exist in extracted text.
    numeric_warnings = 0
    for sheet in ["온실가스(현황전망목표)", "감축전략(계획실적)"]:
        if sheet not in wb.sheetnames:
            continue
        ws = wb[sheet]
        headers = [cell.value for cell in ws[1]]
        year_cols = [h for h in headers if isinstance(h, int)]
        checked = 0
        supported = 0
        for row in _row_dicts(ws):
            for year in year_cols:
                value = row.get(year)
                variants = _number_variants(value)
                if not variants:
                    continue
                checked += 1
                if any(re.sub(r"\D+", "", variant) in source_digits for variant in variants):
                    supported += 1
        if checked and supported == 0:
            numeric_warnings += 1
            warnings.append(f"numeric values in {sheet} were not text-matched; may be table/image-derived or need manual review")
        elif checked and supported / checked < 0.5:
            warnings.append(
                f"numeric text-support rate in {sheet} is low: {supported}/{checked}; "
                "some values may be image/table-derived or need manual review"
            )
        details.setdefault("numeric_support", {})[sheet] = {"checked": checked, "text_supported": supported}

    if require_images and "이미지·그래프 판독결과" in wb.sheetnames:
        rows = _row_dicts(wb["이미지·그래프 판독결과"])
        if not rows:
            failures.append("image/graph observation sheet is empty while image validation is required")
        bad_pages = [r.get("페이지") for r in rows if not isinstance(r.get("페이지"), int) or r.get("페이지") not in page_text]
        if bad_pages:
            failures.append(f"image observations contain invalid page refs: {bad_pages[:5]}")
        details["image_observations"] = len(rows)

    return {
        "ok": not failures,
        "failures": failures,
        "warnings": warnings,
        "details": details,
    }


def main() -> int:
    parser = argparse.ArgumentParser()
    parser.add_argument("excel")
    parser.add_argument("pdf")
    parser.add_argument("--no-require-images", action="store_true")
    parser.add_argument("--output", default="")
    args = parser.parse_args()

    report = verify(Path(args.excel), Path(args.pdf), require_images=not args.no_require_images)
    text = json.dumps(report, ensure_ascii=False, indent=2)
    if args.output:
        Path(args.output).write_text(text, encoding="utf-8")
    print(text)
    return 0 if report["ok"] else 1


if __name__ == "__main__":
    raise SystemExit(main())
