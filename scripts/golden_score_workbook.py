"""골든셋과 파이프라인 출력 워크북 로딩."""
from __future__ import annotations

from pathlib import Path
from typing import Any

import openpyxl

from scripts.golden_score_contract import (
    골든열,
    계약헤더,
    필수채점필드,
    데이터시트,
    문자열,
    시트자료,
    값있음,
    출처유형목록,
    행,
    normalize_provenance_pages,
)


def _제외행(row: dict[str, Any]) -> bool:
    value = row.get("골든_채점제외")
    if isinstance(value, bool):
        return value
    if isinstance(value, (int, float)):
        return int(value) == 1
    return str(value or "").strip().casefold() in {"1", "y", "yes", "true"}


def _제목행_추정(headers: list[str], sheet_name: str) -> bool:
    first = headers[0] if headers else ""
    return first.startswith(sheet_name[:2]) or first.startswith(f"{int(sheet_name[:2])}.") or "제목" in first


def _행목록(ws, headers: list[str], *, golden: bool, sheet_name: str) -> tuple[list[행], int, list[str], list[str]]:
    rows: list[행] = []
    excluded = 0
    errors: list[str] = []
    warnings: list[str] = []
    for row_number, values in enumerate(ws.iter_rows(min_row=2, values_only=True), start=2):
        if not any(값있음(value) for value in values):
            continue
        data = {header: values[idx] if idx < len(values) else None for idx, header in enumerate(headers) if header}
        if golden and _제외행(data):
            excluded += 1
            continue
        source_type = 문자열(data.get("골든_출처유형")) if golden else ""
        source_pages = normalize_provenance_pages(data.get("골든_출처페이지")) if golden else ""
        if golden and sheet_name != "00_문서메타" and source_type not in 출처유형목록:
            errors.append(f"{sheet_name} {row_number}행: 골든_출처유형 값이 허용 코드가 아닙니다: {source_type or '빈칸'}")
        if golden and sheet_name != "00_문서메타" and not source_pages:
            warnings.append(f"{sheet_name} {row_number}행: 골든_출처페이지가 비어 있습니다")
        rows.append(행(row_number, data, source_type, source_pages))
    return rows, excluded, errors, warnings


def _골든시트(ws, sheet_name: str) -> 시트자료:
    headers = [문자열(cell.value) for cell in ws[1]]
    errors: list[str] = []
    # 계약 검증은 config.EXCEL_HEADERS 순서 일치가 아니라 "채점 축이 읽는 필드가
    # 골든 헤더에 존재하는가"로 한다. 행 값은 헤더 이름으로 읽으므로 열 순서·추가
    # 열은 채점에 영향이 없고, config가 확장돼도(v8) 같은 골든을 같은 축으로 잰다.
    required = 필수채점필드(sheet_name)
    missing = [field for field in required if field not in headers]
    if missing:
        if _제목행_추정(headers, sheet_name):
            errors.append(f"{sheet_name}: 1행이 제목 행으로 추정됩니다. 골든셋은 1행이 헤더여야 합니다.")
        else:
            errors.append(f"{sheet_name}: 채점 필수 컬럼이 없습니다: {missing}")
    if sheet_name != "00_문서메타" and not errors:
        if "골든_출처유형" not in headers or "골든_출처페이지" not in headers:
            errors.append(f"{sheet_name}: 골든_출처유형·골든_출처페이지 컬럼이 계약 컬럼 뒤에 필요합니다.")
    rows, excluded, row_errors, row_warnings = _행목록(ws, headers, golden=True, sheet_name=sheet_name)
    errors.extend(row_errors)
    if errors:
        return 시트자료(sheet_name, headers, [], "형식 오류", errors, excluded, row_warnings)
    return 시트자료(sheet_name, headers, rows, 제외행수=excluded, 경고=row_warnings)


def _출력시트(ws, sheet_name: str) -> 시트자료:
    headers = [문자열(cell.value) for cell in ws[1]]
    rows, _, _, _ = _행목록(ws, headers, golden=False, sheet_name=sheet_name)
    return 시트자료(sheet_name, headers, rows)


def 워크북읽기(path: Path, *, golden: bool) -> dict[str, 시트자료]:
    wb = openpyxl.load_workbook(path, data_only=True, read_only=True)
    sheets: dict[str, 시트자료] = {}
    for sheet_name in 데이터시트:
        if sheet_name not in wb.sheetnames:
            sheets[sheet_name] = 시트자료(sheet_name, [], [], "골든 없음" if golden else "출력 없음")
            continue
        sheets[sheet_name] = _골든시트(wb[sheet_name], sheet_name) if golden else _출력시트(wb[sheet_name], sheet_name)
    return sheets
