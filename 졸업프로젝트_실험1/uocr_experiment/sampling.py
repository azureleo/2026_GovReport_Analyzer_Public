from __future__ import annotations

import random
import re
from collections import defaultdict
from pathlib import Path
from typing import Any, Iterable

import fitz
from openpyxl import load_workbook

from .io_utils import sha256_file, write_json
from .models import SamplePage


INVENTORY_SHEET = "21_원문객체인벤토리"
_TABLE_RE = re.compile(r"(?:^|\s|[\[(])표\s*\d+(?:[-–]\d+)?")
_FIGURE_RE = re.compile(r"(?:^|\s|[\[(])(?:그림|도표|그래프|Figure|Fig\.?)\s*\d+", re.I)


def _as_int(value: Any, default: int = 0) -> int:
    try:
        return int(float(value))
    except (TypeError, ValueError):
        return default


def read_inventory(path: str | Path | None) -> list[dict[str, Any]]:
    if not path:
        return []
    workbook_path = Path(path).resolve()
    workbook = load_workbook(workbook_path, read_only=True, data_only=True)
    try:
        if INVENTORY_SHEET not in workbook.sheetnames:
            raise ValueError(f"'{INVENTORY_SHEET}' 시트가 없습니다: {workbook_path}")
        sheet = workbook[INVENTORY_SHEET]
        values = sheet.iter_rows(values_only=True)
        first_row = next(values, None)
        if first_row is None:
            return []
        headers = [str(value or "").strip() for value in first_row]
        records: list[dict[str, Any]] = []
        for row in values:
            record = {header: value for header, value in zip(headers, row) if header}
            page = _as_int(record.get("출처페이지"))
            if page > 0:
                records.append(record)
        return records
    finally:
        workbook.close()


def infer_inventory_from_pdf(pdf_path: str | Path) -> list[dict[str, Any]]:
    """인벤토리 엑셀이 없을 때 표/그림 캡션과 이미지로 후보 페이지를 만든다."""
    records: list[dict[str, Any]] = []
    document = fitz.open(str(Path(pdf_path)))
    try:
        for index, page in enumerate(document, start=1):
            text = page.get_text("text") or ""
            table_matches = list(_TABLE_RE.finditer(text))
            figure_matches = list(_FIGURE_RE.finditer(text))
            for sequence, match in enumerate(table_matches, start=1):
                records.append({
                    "객체ID": f"heuristic-p{index}-table-{sequence}",
                    "객체유형": "table",
                    "출처페이지": index,
                    "캡션": match.group(0).strip(),
                    "연결상태": "미확인",
                    "행수": 0,
                    "열수": 0,
                })
            for sequence, match in enumerate(figure_matches, start=1):
                records.append({
                    "객체ID": f"heuristic-p{index}-figure-{sequence}",
                    "객체유형": "figure",
                    "출처페이지": index,
                    "캡션": match.group(0).strip(),
                    "연결상태": "미확인",
                    "행수": 0,
                    "열수": 0,
                })
            if not figure_matches and page.get_images(full=True):
                records.append({
                    "객체ID": f"heuristic-p{index}-image-1",
                    "객체유형": "image",
                    "출처페이지": index,
                    "캡션": "",
                    "연결상태": "미확인",
                    "행수": 0,
                    "열수": 0,
                })
    finally:
        document.close()
    return records


def _classify(record: dict[str, Any]) -> set[str]:
    status = str(record.get("연결상태") or "").strip()
    obj_type = str(record.get("객체유형") or "").strip().lower()
    rows = _as_int(record.get("행수"))
    columns = _as_int(record.get("열수"))
    labels: set[str] = set()
    if "미확인" in status or status in {"누락", "미연결"}:
        labels.add("unconfirmed")
    if "부분" in status or "약함" in status:
        labels.add("partial")
    if ("표" in obj_type or "table" in obj_type) and (rows >= 8 or columns >= 5):
        labels.add("complex_table")
    if any(token in obj_type for token in ("figure", "image", "graph", "chart", "그림", "그래프", "차트")):
        labels.add("visual")
    if status == "확인":
        labels.add("confirmed_control")
    if not labels:
        labels.add("other")
    return labels


def select_sample_pages(
    records: Iterable[dict[str, Any]],
    *,
    target_size: int = 50,
    seed: int = 42,
    manual_pages: Iterable[int] | None = None,
) -> tuple[list[int], dict[int, dict[str, Any]]]:
    """상태/유형별 층화 표본을 고정 시드로 선택한다."""
    page_meta: dict[int, dict[str, Any]] = {}
    groups: dict[str, list[int]] = defaultdict(list)
    for record in records:
        page = _as_int(record.get("출처페이지"))
        if page <= 0:
            continue
        meta = page_meta.setdefault(page, {
            "categories": set(),
            "source_statuses": set(),
            "expected_object_ids": [],
            "inventory_objects": [],
        })
        labels = _classify(record)
        meta["categories"].update(labels)
        status = str(record.get("연결상태") or "").strip()
        if status:
            meta["source_statuses"].add(status)
        object_id = str(record.get("객체ID") or "").strip()
        if object_id:
            meta["expected_object_ids"].append(object_id)
        meta["inventory_objects"].append({
            "object_id": object_id,
            "object_type": str(record.get("객체유형") or ""),
            "caption": str(record.get("캡션") or ""),
            "number": str(record.get("번호") or ""),
            "status": status,
            "rows": _as_int(record.get("행수")),
            "columns": _as_int(record.get("열수")),
        })

    for page, meta in page_meta.items():
        for label in meta["categories"]:
            groups[label].append(page)

    if manual_pages:
        selected = sorted({int(page) for page in manual_pages if int(page) > 0})
    else:
        rng = random.Random(seed)
        for pages in groups.values():
            rng.shuffle(pages)
        quotas = {
            "unconfirmed": round(target_size * 0.30),
            "partial": round(target_size * 0.20),
            "complex_table": round(target_size * 0.20),
            "visual": round(target_size * 0.20),
            "confirmed_control": max(1, target_size - round(target_size * 0.90)),
        }
        selected_set: set[int] = set()
        for label, quota in quotas.items():
            for page in groups.get(label, []):
                if page not in selected_set:
                    selected_set.add(page)
                if sum(1 for item in selected_set if label in page_meta[item]["categories"]) >= quota:
                    break
        remaining = list(page_meta)
        rng.shuffle(remaining)
        for page in remaining:
            if len(selected_set) >= target_size:
                break
            selected_set.add(page)
        selected = sorted(selected_set)[:target_size]

    for page in selected:
        page_meta.setdefault(page, {
            "categories": {"manual"},
            "source_statuses": set(),
            "expected_object_ids": [],
            "inventory_objects": [],
        })
    return selected, page_meta


def prepare_sample(
    pdf_path: str | Path,
    workspace: str | Path,
    *,
    inventory_path: str | Path | None = None,
    pages: Iterable[int] | None = None,
    sample_size: int = 50,
    seed: int = 42,
    dpi: int = 300,
) -> Path:
    pdf = Path(pdf_path).resolve()
    run_dir = Path(workspace).resolve()
    image_dir = run_dir / "pages"
    image_dir.mkdir(parents=True, exist_ok=True)

    records = read_inventory(inventory_path) if inventory_path else infer_inventory_from_pdf(pdf)
    selected, page_meta = select_sample_pages(
        records,
        target_size=max(1, sample_size),
        seed=seed,
        manual_pages=pages,
    )
    document = fitz.open(str(pdf))
    try:
        selected = [page for page in selected if 1 <= page <= document.page_count]
        if not selected:
            raise ValueError("선택된 표본 페이지가 없습니다. --pages 또는 인벤토리 파일을 확인하세요.")
        sample_pages: list[SamplePage] = []
        scale = max(72, dpi) / 72
        for page_number in selected:
            page = document[page_number - 1]
            pixmap = page.get_pixmap(matrix=fitz.Matrix(scale, scale), alpha=False)
            image_path = image_dir / f"page_{page_number:04d}.png"
            pixmap.save(image_path)
            meta = page_meta[page_number]
            sample_pages.append(SamplePage(
                page_number=page_number,
                image_path=str(image_path),
                categories=sorted(meta["categories"]),
                source_statuses=sorted(meta["source_statuses"]),
                expected_object_ids=list(dict.fromkeys(meta["expected_object_ids"])),
                pdf_width=float(page.rect.width),
                pdf_height=float(page.rect.height),
                image_width=int(pixmap.width),
                image_height=int(pixmap.height),
            ))
    finally:
        document.close()

    inventory_by_page = {
        str(page): page_meta[page]["inventory_objects"]
        for page in selected
    }
    manifest = {
        "schema_version": 1,
        "experiment": "Unlimited-OCR integration benchmark",
        "pdf_path": str(pdf),
        "pdf_sha256": sha256_file(pdf),
        "inventory_path": str(Path(inventory_path).resolve()) if inventory_path else None,
        "inventory_sha256": sha256_file(inventory_path) if inventory_path else None,
        "selection": {"sample_size": sample_size, "seed": seed, "dpi": dpi},
        "pages": [item.to_dict() for item in sample_pages],
        "inventory_by_page": inventory_by_page,
    }
    manifest_path = run_dir / "sample_manifest.json"
    write_json(manifest_path, manifest)
    return manifest_path
