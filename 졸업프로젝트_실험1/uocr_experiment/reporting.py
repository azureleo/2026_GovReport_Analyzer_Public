from __future__ import annotations

import json
from pathlib import Path
from typing import Any, Iterable

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from .io_utils import load_manifest, load_objects, read_jsonl, write_json


def _cell(value: Any) -> Any:
    if isinstance(value, (dict, list, tuple, set)):
        value = json.dumps(value, ensure_ascii=False, default=str)
    if value is None or isinstance(value, (str, int, float, bool)):
        return value
    return str(value)


def _write_sheet(workbook: Workbook, name: str, rows: Iterable[dict[str, Any]]) -> None:
    sheet = workbook.create_sheet(name[:31])
    values = list(rows)
    if not values:
        sheet.append(["데이터 없음"])
        return
    headers: list[str] = []
    for row in values:
        for key in row:
            if key not in headers:
                headers.append(key)
    sheet.append(headers)
    for row in values:
        sheet.append([_cell(row.get(key)) for key in headers])
    for cell in sheet[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="2F5597")
        cell.alignment = Alignment(horizontal="center", vertical="center")
    sheet.freeze_panes = "A2"
    sheet.auto_filter.ref = sheet.dimensions
    for column_index, header in enumerate(headers, start=1):
        width = min(50, max(10, len(str(header)) + 2))
        for cell in sheet.iter_cols(min_col=column_index, max_col=column_index, min_row=2, max_row=min(sheet.max_row, 100)):
            for item in cell:
                width = min(50, max(width, len(str(item.value or "")) + 2))
                item.alignment = Alignment(vertical="top", wrap_text=True)
        sheet.column_dimensions[get_column_letter(column_index)].width = width


def write_reports(
    manifest_path: str | Path,
    summaries: list[dict[str, Any]],
    matches: list[dict[str, Any]],
) -> tuple[Path, Path, Path]:
    manifest = load_manifest(manifest_path)
    run_dir = Path(manifest_path).resolve().parent
    report_dir = run_dir / "reports"
    report_dir.mkdir(parents=True, exist_ok=True)
    decisions = read_jsonl(run_dir / "outputs" / "hybrid" / "decisions.jsonl")

    json_path = write_json(report_dir / "evaluation.json", {
        "manifest": str(Path(manifest_path).resolve()),
        "pdf_sha256": manifest.get("pdf_sha256"),
        "summaries": summaries,
        "matches": matches,
    })

    lines = [
        "# Unlimited-OCR 비교 실험 결과",
        "",
        f"- PDF: `{manifest.get('pdf_path')}`",
        f"- 표본 페이지: {len(manifest.get('pages', []))}개",
        f"- 평가 모드: `{summaries[0].get('evaluation_mode') if summaries else 'unknown'}`",
        "",
        "## 엔진별 요약",
        "",
        "| 엔진 | 객체 수 | 객체 리콜 | 객체 정밀도 | 문자 정확도 | 표 셀 정확도 |",
        "|---|---:|---:|---:|---:|---:|",
    ]
    for item in summaries:
        lines.append(
            "| {engine} | {count} | {recall} | {precision} | {char} | {cell} |".format(
                engine=item.get("engine", ""),
                count=item.get("object_count", 0),
                recall=item.get("object_recall", "-") if item.get("object_recall") is not None else "-",
                precision=item.get("object_precision", "-") if item.get("object_precision") is not None else "-",
                char=item.get("character_accuracy", "-") if item.get("character_accuracy") is not None else "-",
                cell=item.get("table_cell_accuracy", "-") if item.get("table_cell_accuracy") is not None else "-",
            )
        )
    lines.extend([
        "",
        "## 해석 주의",
        "",
        "골든셋 없이 실행한 결과는 객체 수와 페이지 커버리지를 보여주는 진단값이며 정확도가 아닙니다. "
        "`golden_template.jsonl`을 사람이 검수한 뒤 다시 평가해야 모델별 정확도를 비교할 수 있습니다.",
    ])
    markdown_path = report_dir / "evaluation.md"
    markdown_path.write_text("\n".join(lines), encoding="utf-8")

    workbook = Workbook()
    active_sheet = workbook.active
    if active_sheet is not None:
        workbook.remove(active_sheet)
    _write_sheet(workbook, "summary", summaries)
    _write_sheet(workbook, "sample_pages", manifest.get("pages", []))
    for engine, sheet_name in (
        ("pymupdf", "pymupdf_objects"),
        ("unlimited_ocr", "uocr_objects"),
        ("hybrid", "hybrid_objects"),
    ):
        objects = load_objects(run_dir / "outputs" / engine / "objects.jsonl")
        _write_sheet(workbook, sheet_name, [item.to_dict() for item in objects])
    _write_sheet(workbook, "merge_decisions", decisions)
    _write_sheet(workbook, "golden_matches", matches)
    xlsx_path = report_dir / "evaluation.xlsx"
    workbook.save(xlsx_path)
    return Path(json_path), markdown_path, xlsx_path
